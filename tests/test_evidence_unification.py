"""Regressions for visible evidence, physical source identity and safe projection."""

import json

import pytest
from docx import Document
from docx.oxml import OxmlElement

from bridge_evidence.contest.long_document import (
    ChunkPlan,
    ChunkTask,
    DocumentSnapshot,
    SourceSlice,
    merge_fragments,
    parse_fragment,
)


def _task(snapshot, owned=(0,), context=()):
    def source(i):
        r = snapshot.regions[i]
        return SourceSlice(r.source_id, 0, len(r.text), r.text)

    return ChunkTask("test", [source(i) for i in owned], [source(i) for i in context])


def _payload(task, **extra):
    return {
        "covered": task.owned_keys,
        "基本信息": {},
        "basic_sources": {},
        "维修建议": [],
        "病害明细": [],
        **extra,
    }


def _history_row(task):
    identity = {
        "bridge": "测试桥",
        "span": "左幅",
        "component": "梁",
        "position": "1跨",
        "year": 2018,
    }
    return {
        "source_id": "p0",
        "source_key": task.owned_keys[0],
        "instance_id": "defect-1",
        "fields": {"病害描述": "裂缝"},
        "identity": identity,
        "history_identity": {**identity, "year": 2017},
        "temporal": "已有",
        "history_sources": ["p1"],
    }


def test_unseen_history_and_model_only_identity_cannot_establish_relationship():
    snapshot = DocumentSnapshot.from_text("2018 测试桥左幅1跨梁裂缝\n未提供的历史段落")
    task = _task(snapshot)
    row = _history_row(task)
    fragment = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    )
    assert fragment.diseases[0]["temporal"] == "证据不足"
    assert fragment.diseases[0]["fields"]["是否新增"] is None


def test_full_source_cannot_be_used_to_escape_visible_context_slice():
    snapshot = DocumentSnapshot.from_text(
        "2018 测试桥左幅1跨梁裂缝\n2017 测试桥左幅1跨梁裂缝隐藏尾部"
    )
    task = _task(snapshot)
    task.context = [SourceSlice("p1", 0, 4, "2017")]
    row = _history_row(task)
    row["history_evidence"] = [
        {"source_key": "p1:0:4", "quote": "2017 测试桥左幅1跨梁裂缝"}
    ]
    fragment = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    )
    assert fragment.diseases[0]["temporal"] == "证据不足"


def test_explicit_visible_history_relationship_keeps_auditable_evidence():
    snapshot = DocumentSnapshot.from_text(
        "2018 测试桥左幅1跨梁裂缝\n2017 测试桥左幅1跨梁裂缝\n2018与2017为测试桥左幅1跨梁同一裂缝"
    )
    task = _task(snapshot, context=(1, 2))
    row = _history_row(task)
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text.strip()}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text.strip()}
    ]
    row["relationship_evidence"] = [
        {"source_key": task.context[1].key, "quote": task.context[1].text}
    ]
    fragment = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot, "call"
    )
    assert fragment.diseases[0]["temporal"] == "已有"
    # All three sources are responsibility regions, although this fragment sees
    # the other two as context. Merge a complete one-task plan for provenance.
    task.owned.extend(task.context)
    task.context = []
    fragment.covered = task.owned_keys
    plan = ChunkPlan(snapshot.digest, [task], 8192, 4096)
    merged = merge_fragments(snapshot, plan, [fragment])
    provenance = merged["provenance"][0]
    assert provenance["source_key"] == row["source_key"]
    assert provenance["instance_id"] == "defect-1"
    assert provenance["history_evidence"][0]["source_key"].startswith("p1:")
    assert provenance["source_ranges"][0]["start"] == 0


def test_exact_basic_quotes_are_checked_against_owned_slice():
    snapshot = DocumentSnapshot.from_text("测试桥评分82")
    task = _task(snapshot)
    payload = _payload(
        task,
        基本信息={"总体评分": "82"},
        basic_sources={"总体评分": "p0"},
        basic_quotes={"总体评分": "93"},
    )
    # The forged quote is never evidence; the value is re-located verbatim instead.
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic["总体评分"] == "82"
    assert fragment.basic_evidence["总体评分"]["quote"] == "82"
    payload["基本信息"]["总体评分"] = "93"
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic["总体评分"] is None and fragment.basic_evidence == {}


def test_quote_object_cannot_smuggle_fields_or_non_scalar_values():
    snapshot = DocumentSnapshot.from_text("裂缝")
    task = _task(snapshot)
    row = {"source_id": "p0", "fields": {"病害描述": "裂缝", "嵌套": {"x": 1}}, "quotes": ["裂缝"]}
    fragment = parse_fragment(json.dumps(_payload(task, 病害明细=[row])), task, snapshot)
    assert fragment.diseases[0]["quotes"] == {"病害描述": "裂缝"}  # verified field, not the list
    assert fragment.diseases[0]["fields"]["病害描述"] == "裂缝"
    assert "嵌套" not in fragment.diseases[0]["fields"]


def test_snapshot_uses_ordered_nested_wrapped_and_supplementary_sources(tmp_path):
    document = Document()
    document.add_paragraph("开头")
    wrapped = document.add_paragraph("包装段落")._p
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    wrapped.addprevious(control)
    control.append(content)
    content.append(wrapped)
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).paragraphs[0].text = "之前"
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "内表裂缝"
    outer.cell(0, 0).paragraphs[-1].text = "之后"
    document.sections[0].header.paragraphs[0].text = "页眉桥名"
    document.sections[0].footer.paragraphs[0].text = "页脚年份"
    path = tmp_path / "nested.docx"
    document.save(path)
    snapshot = DocumentSnapshot.from_docx(path)
    texts = [region.text for region in snapshot.regions]
    assert texts[:5] == ["开头", "包装段落", "之前", "内表裂缝", "之后"]
    assert {"页眉桥名", "页脚年份"} <= set(texts)
    assert len({region.source_id for region in snapshot.regions}) == len(
        snapshot.regions
    )
    assert all(region.location.get("source_path") for region in snapshot.regions)


def test_task_coverage_does_not_claim_opaque_document_content_verified(tmp_path):
    document = Document()
    document.add_paragraph("正文")._p.append(OxmlElement("w:object"))
    path = tmp_path / "opaque.docx"
    document.save(path)
    snapshot = DocumentSnapshot.from_docx(path)
    task = _task(snapshot, owned=tuple(range(len(snapshot.regions))))
    coverage = ChunkPlan(snapshot.digest, [task], 8192, 4096).coverage(snapshot)
    assert coverage["complete"]
    assert not coverage["document_complete"]
    assert coverage["unverified_content"]


def test_different_score_objects_are_preserved_without_false_conflict_or_selection():
    snapshot = DocumentSnapshot.from_text("2017 左幅评分82\n2018 右幅评分93")
    tasks = [_task(snapshot, owned=(i,)) for i in range(2)]
    for i, task in enumerate(tasks):
        task.task_id = f"task-{i}"
    fragments = []
    for i, (year, side, score) in enumerate(
        ((2017, "左幅", "82"), (2018, "右幅", "93"))
    ):
        value = _payload(
            tasks[i],
            基本信息={"总体评分": score},
            basic_sources={"总体评分": f"p{i}"},
            basic_quotes={"总体评分": tasks[i].owned[0].text.strip()},
            basic_scopes={"总体评分": {"year": year, "span": side}},
        )
        fragments.append(
            parse_fragment(json.dumps(value), tasks[i], snapshot, f"call-{i}")
        )
    result = merge_fragments(
        snapshot, ChunkPlan(snapshot.digest, tasks, 8192, 4096), fragments
    )
    assert result["conflicts"] == []
    assert len(result["basic_observations"]) == 2
    assert len(result["basic_ambiguities"]) == 1
    assert result["data"]["基本信息"]["总体评分"] is None
    assert {row["scope"]["span"] for row in result["provenance"]} == {"左幅", "右幅"}
    assert all(row["source_ranges"] and row["quotes"] for row in result["provenance"])


def test_equal_score_scope_conflict_remains_explicit():
    snapshot = DocumentSnapshot.from_text("2018 左幅评分82\n2018 左幅评分93")
    tasks = [_task(snapshot, owned=(i,)) for i in range(2)]
    fragments = []
    for i, task in enumerate(tasks):
        task.task_id = f"task-{i}"
        payload = _payload(
            task,
            基本信息={"总体评分": str(82 + 11 * i)},
            basic_sources={"总体评分": f"p{i}"},
            basic_scopes={"总体评分": {"year": 2018, "span": "左幅"}},
            basic_quotes={"总体评分": task.owned[0].text.strip()},
        )
        fragments.append(parse_fragment(json.dumps(payload), task, snapshot))
    result = merge_fragments(
        snapshot, ChunkPlan(snapshot.digest, tasks, 8192, 4096), fragments
    )
    assert len(result["conflicts"]) == 1
    assert result["basic_ambiguities"] == []
    assert result["data"]["基本信息"]["总体评分"] is None


def test_basic_source_on_split_region_requires_explicit_interval():
    snapshot = DocumentSnapshot.from_text("前文评分82后文评分93")
    region = snapshot.regions[0]
    task = ChunkTask(
        "test",
        [
            SourceSlice("p0", 0, 6, region.text[:6]),
            SourceSlice("p0", 6, len(region.text), region.text[6:]),
        ],
        [],
    )
    payload = _payload(
        task, 基本信息={"总体评分": "82"}, basic_sources={"总体评分": "p0"}
    )
    # An ambiguous split source is resolved by the slice that contains the value.
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic_evidence["总体评分"]["quote_start"] == 4
    payload["basic_sources"]["总体评分"] = {
        "source_id": "p0",
        "source_key": task.owned[0].key,
        "quote": "82",
    }
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic_evidence["总体评分"]["quote_start"] == 4


@pytest.mark.parametrize(
    "negation", ["不是同一", "尚不能证明属于同一", "是否同一", "可能为同一"]
)
def test_uncertain_or_negative_relationship_cannot_be_published_as_existing(negation):
    snapshot = DocumentSnapshot.from_text(
        f"2018 测试桥左幅1跨梁裂缝\n2017 测试桥左幅1跨梁裂缝\n2018与2017测试桥左幅1跨梁{negation}裂缝"
    )
    task = _task(snapshot, context=(1, 2))
    row = _history_row(task)
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text}
    ]
    row["relationship_evidence"] = [
        {"source_key": task.context[1].key, "quote": task.context[1].text}
    ]
    fragment = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    )
    assert fragment.diseases[0]["temporal"] == "证据不足"


def test_exact_token_cache_does_not_change_chunks_or_recount_same_rendering():
    from bridge_evidence.contest.evidence_manifest import (
        make_chunks,
        manifest_from_text,
    )

    manifest = manifest_from_text("桥名\n" + "长行裂缝" * 40 + "\n尾部")
    baseline = make_chunks(manifest, len, 768)
    observed = []

    def count(text):
        observed.append(text)
        return len(text)

    result = make_chunks(manifest, count, 768)
    assert result == baseline
    assert len(observed) == len(set(observed))


def test_legacy_snapshot_digest_retains_previous_version_identity():
    from dataclasses import asdict

    from bridge_evidence.contest.run_state import json_digest

    snapshot = DocumentSnapshot.from_text("桥")
    old = {
        "document_sha256": snapshot.document_sha256,
        "regions": [asdict(region) for region in snapshot.regions],
        "parser": "bound-text-v1",
    }
    assert snapshot.digest == json_digest(old)


def test_identity_number_is_not_accepted_as_substring_of_another_component():
    snapshot = DocumentSnapshot.from_text(
        "2018 测试桥左幅11跨梁裂缝\n2017 测试桥左幅11跨梁裂缝\n2018与2017测试桥左幅11跨梁同一裂缝"
    )
    task = _task(snapshot, context=(1, 2))
    row = _history_row(task)  # claims 1跨 although every quote says 11跨
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text}
    ]
    row["relationship_evidence"] = [
        {"source_key": task.context[1].key, "quote": task.context[1].text}
    ]
    fragment = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    )
    assert fragment.diseases[0]["temporal"] == "证据不足"


def test_chart_without_raster_image_remains_unverified(tmp_path):
    document = Document()
    document.add_paragraph("图表说明")._p.append(OxmlElement("c:chart"))
    path = tmp_path / "chart.docx"
    document.save(path)
    snapshot = DocumentSnapshot.from_docx(path)
    task = _task(snapshot)
    coverage = ChunkPlan(snapshot.digest, [task], 8192, 4096).coverage(snapshot)
    assert not coverage["document_complete"]
    assert any(item["kind"] == "chart" for item in coverage["unverified_content"])


@pytest.mark.parametrize("value", ["8", "82", 82])
def test_basic_quote_must_support_the_complete_numeric_scalar(value):
    snapshot = DocumentSnapshot.from_text("总体评分82.5")
    task = _task(snapshot)
    payload = _payload(
        task,
        基本信息={"总体评分": value},
        basic_sources={"总体评分": "p0"},
        basic_quotes={"总体评分": "82.5"},
    )
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic["总体评分"] is None and fragment.basic_evidence == {}


def test_multiple_scoped_scores_in_one_slice_survive_as_separate_observations():
    snapshot = DocumentSnapshot.from_text("2017 左幅评分82；2018 右幅评分93")
    task = _task(snapshot)
    payload = _payload(
        task,
        basic_observations=[
            {
                "field": "总体评分",
                "value": score,
                "scope": {"year": year, "span": span},
                "source_id": "p0",
                "source_key": task.owned[0].key,
                "quote": quote,
            }
            for year, span, score, quote in (
                (2017, "左幅", "82", "2017 左幅评分82"),
                (2018, "右幅", "93", "2018 右幅评分93"),
            )
        ],
    )
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    merged = merge_fragments(
        snapshot, ChunkPlan(snapshot.digest, [task], 8192, 4096), [fragment]
    )
    assert [item["value"] for item in merged["basic_observations"]] == ["82", "93"]
    assert merged["data"]["基本信息"]["总体评分"] is None
    assert len(merged["basic_ambiguities"]) == 1
    assert len(merged["provenance"]) == 2
    assert merged["conflicts"] == []


def test_scoped_observation_cannot_borrow_year_from_another_quote_in_slice():
    snapshot = DocumentSnapshot.from_text("2017 左幅评分82；2018 右幅评分93")
    task = _task(snapshot)
    payload = _payload(
        task,
        basic_observations=[
            {
                "field": "总体评分",
                "value": "82",
                "scope": {"year": 2018, "span": "右幅"},
                "source_id": "p0",
                "source_key": task.owned[0].key,
                "quote": "2017 左幅评分82",
            }
        ],
    )
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic_observations == [] and fragment.basic_evidence == {}


def test_legacy_full_basic_schema_null_placeholders_are_not_unsourced_claims():
    from bridge_evidence.contest.inference_core import BASIC_FIELDS

    snapshot = DocumentSnapshot.from_text("测试桥")
    task = _task(snapshot)
    basic = dict.fromkeys(BASIC_FIELDS)
    basic["桥梁名称"] = "测试桥"
    payload = _payload(task, 基本信息=basic, basic_sources={"桥梁名称": "p0"})
    fragment = parse_fragment(json.dumps(payload), task, snapshot, "call-1")
    assert fragment.basic == basic
    assert set(fragment.basic_sources) == {"桥梁名称"}
    assert [item["field"] for item in fragment.basic_observations] == ["桥梁名称"]
    merged = merge_fragments(
        snapshot, ChunkPlan(snapshot.digest, [task], 8192, 4096), [fragment]
    )
    assert merged["data"]["基本信息"] == basic
    assert len(merged["provenance"]) == 1
    assert merged["conflicts"] == []


@pytest.mark.parametrize("missing", [None, "", " \t\n"])
def test_empty_basic_slots_remain_empty_without_inventing_source_evidence(missing):
    snapshot = DocumentSnapshot.from_text("病害记录")
    task = _task(snapshot)
    payload = _payload(
        task,
        基本信息={"总体评分": missing},
        basic_observations=[{"field": "上一次评分", "value": missing}],
    )
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic == {"总体评分": None, "上一次评分": None}
    assert fragment.basic_sources == {}
    assert fragment.basic_evidence == {}
    assert fragment.basic_observations == []
    merged = merge_fragments(
        snapshot, ChunkPlan(snapshot.digest, [task], 8192, 4096), [fragment]
    )
    assert merged["data"]["基本信息"] == fragment.basic
    assert merged["provenance"] == []
    assert merged["basic_observations"] == []


@pytest.mark.parametrize("value", [0, False, "0", "无", "null", [], {}, "桥"])
def test_unsupported_placeholder_or_invalid_basic_values_are_dropped(value):
    snapshot = DocumentSnapshot.from_text("桥梁记录无")
    task = _task(snapshot)
    fragment = parse_fragment(
        json.dumps(_payload(task, 基本信息={"总体评分": value})), task, snapshot
    )
    assert fragment.basic["总体评分"] is None
    assert fragment.basic_evidence == {} and fragment.basic_observations == []


@pytest.mark.parametrize(
    "relation",
    [
        "2018与2017测试桥左幅1跨梁在同一位置出现不同裂缝，2018裂缝为新增",
        "2018与2017测试桥左幅1跨梁在同一位置均有裂缝",
        "2018与2017测试桥左幅1跨梁同一类型裂缝",
        "2018与2017测试桥左幅1跨梁同一构件的裂缝",
        "2018与2017测试桥左幅1跨梁同一条裂缝，但2018为另一条新增裂缝",
        "2018与2017测试桥左幅1跨梁裂缝，并非同一处病害",
    ],
)
def test_shared_location_or_conflicting_defect_statement_never_proves_continuity(
    relation,
):
    snapshot = DocumentSnapshot.from_text(
        f"2018 测试桥左幅1跨梁裂缝\n2017 测试桥左幅1跨梁裂缝\n{relation}"
    )
    task = _task(snapshot, context=(1, 2))
    row = _history_row(task)
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text}
    ]
    row["relationship_evidence"] = [
        {"source_key": task.context[1].key, "quote": task.context[1].text}
    ]
    disease = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    ).diseases[0]
    assert disease["temporal"] == "证据不足"
    assert disease["temporal_verification"] == "insufficient_evidence"
    assert disease["fields"]["是否新增"] is None


def test_same_defect_predicate_must_refer_to_the_extracted_disease():
    snapshot = DocumentSnapshot.from_text(
        "2018 测试桥左幅1跨梁裂缝及剥落\n2017 测试桥左幅1跨梁裂缝及剥落\n"
        "2018与2017测试桥左幅1跨梁同一裂缝"
    )
    task = _task(snapshot, context=(1, 2))
    row = _history_row(task)
    row["fields"] = {"病害类型": "剥落", "病害描述": "混凝土剥落"}
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text}
    ]
    row["relationship_evidence"] = [
        {"source_key": task.context[1].key, "quote": task.context[1].text}
    ]
    disease = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    ).diseases[0]
    assert disease["temporal"] == "证据不足"


def test_relationship_years_and_same_defect_predicate_cannot_be_spliced_from_different_quotes():
    snapshot = DocumentSnapshot.from_text(
        "2018 测试桥左幅1跨梁裂缝\n2017 测试桥左幅1跨梁裂缝\n"
        "2018测试桥左幅1跨梁同一条裂缝\n2017报告提及测试桥左幅1跨梁"
    )
    task = _task(snapshot, context=(1, 2, 3))
    row = _history_row(task)
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text}
    ]
    row["relationship_evidence"] = [
        {"source_key": part.key, "quote": part.text} for part in task.context[1:]
    ]
    disease = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    ).diseases[0]
    assert disease["temporal"] == "证据不足"


@pytest.mark.parametrize("predicate", ["同一裂缝", "同一条裂缝", "同一处裂缝"])
def test_explicit_same_defect_predicate_with_both_years_retains_existing_status(
    predicate,
):
    snapshot = DocumentSnapshot.from_text(
        f"2018 测试桥左幅1跨梁裂缝\n2017 测试桥左幅1跨梁裂缝\n2018与2017测试桥左幅1跨梁{predicate}"
    )
    task = _task(snapshot, context=(1, 2))
    row = _history_row(task)
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text}
    ]
    row["relationship_evidence"] = [
        {"source_key": task.context[1].key, "quote": task.context[1].text}
    ]
    disease = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    ).diseases[0]
    assert disease["temporal"] == "已有"
    assert disease["temporal_verification"] == "explicit_visible_evidence"


def test_relationship_quote_cannot_hide_conflicting_text_in_visible_source():
    prefix = "2018与2017测试桥左幅1跨梁同一裂缝"
    snapshot = DocumentSnapshot.from_text(
        f"2018 测试桥左幅1跨梁裂缝\n2017 测试桥左幅1跨梁裂缝\n{prefix}，但2018实为另一条新增裂缝"
    )
    task = _task(snapshot, context=(1, 2))
    row = _history_row(task)
    row["current_evidence"] = [
        {"source_key": task.owned[0].key, "quote": task.owned[0].text}
    ]
    row["history_evidence"] = [
        {"source_key": task.context[0].key, "quote": task.context[0].text}
    ]
    row["relationship_evidence"] = [
        {"source_key": task.context[1].key, "quote": prefix}
    ]
    disease = parse_fragment(
        json.dumps(_payload(task, 病害明细=[row])), task, snapshot
    ).diseases[0]
    assert disease["temporal"] == "证据不足"


def test_basic_source_accepts_exact_owned_key_without_losing_interval_binding():
    snapshot = DocumentSnapshot.from_text(
        "桥梁名称：运行检查桥；报告日期：2026年10月10日"
    )
    task = _task(snapshot)
    payload = _payload(
        task,
        基本信息={"桥梁名称": "运行检查桥", "报告日期": "2026年10月10日"},
        basic_sources={"桥梁名称": task.owned[0].key, "报告日期": task.owned[0].key},
        basic_quotes={
            "桥梁名称": "桥梁名称：运行检查桥",
            "报告日期": "报告日期：2026年10月10日",
        },
        basic_scopes={"桥梁名称": {}, "报告日期": {}},
    )
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic_sources == {"桥梁名称": "p0", "报告日期": "p0"}
    assert all(
        item["source_key"] == task.owned[0].key
        for item in fragment.basic_evidence.values()
    )


@pytest.mark.parametrize("ref", ["p1:0:6", "p0:0:999", "p0:0:2"])
def test_basic_source_key_compatibility_does_not_accept_context_or_invented_intervals(
    ref,
):
    snapshot = DocumentSnapshot.from_text("本次桥名\n只读历史内容")
    task = _task(snapshot, context=(1,))
    payload = _payload(
        task, 基本信息={"桥梁名称": "本次桥名"}, basic_sources={"桥梁名称": ref}
    )
    # Context or invented intervals are never used; the owned verbatim text is.
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic_sources == {"桥梁名称": "p0"}
    assert fragment.basic_evidence["桥梁名称"]["source_key"] == task.owned[0].key


def test_unsupported_basic_scope_falls_back_to_the_safe_empty_scope():
    snapshot = DocumentSnapshot.from_text(
        "桥梁名称：运行检查桥；报告日期：2026年10月10日"
    )
    task = _task(snapshot)
    payload = _payload(
        task,
        基本信息={"桥梁名称": "运行检查桥"},
        basic_sources={"桥梁名称": task.owned[0].key},
        basic_quotes={"桥梁名称": "桥梁名称：运行检查桥"},
        basic_scopes={"桥梁名称": {"year": "2026", "span": "全桥"}},
    )
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.basic["桥梁名称"] == "运行检查桥"
    assert fragment.basic_observations[0]["scope"] == {}


def test_advice_alias_is_mapped_to_the_canonical_field_instead_of_omitted():
    snapshot = DocumentSnapshot.from_text("桥面伸缩缝有杂物，建议清理")
    task = _task(snapshot)
    payload = _payload(
        task,
        维修建议=[
            {
                "source_id": "p0",
                "source_key": task.owned[0].key,
                "fields": {"建议": "桥面伸缩缝有杂物，建议清理"},
            }
        ],
    )
    fragment = parse_fragment(json.dumps(payload), task, snapshot)
    assert fragment.advice[0]["fields"] == {"建议内容": "桥面伸缩缝有杂物，建议清理"}


def test_fragment_prompt_defines_canonical_schema_and_noninvented_scopes():
    from bridge_evidence.contest.docx_writer import ADVICE_FIELDS, DEFECT_FIELDS
    from bridge_evidence.contest.inference_core import BASIC_FIELDS
    from bridge_evidence.contest.long_document import FRAGMENT_INSTRUCTION

    assert all(
        field in FRAGMENT_INSTRUCTION
        for field in BASIC_FIELDS | set(ADVICE_FIELDS) | set(DEFECT_FIELDS)
    )
    assert "现有建议字段" not in FRAGMENT_INSTRUCTION
    assert "现有病害字段" not in FRAGMENT_INSTRUCTION
    assert "basic_scopes" in FRAGMENT_INSTRUCTION and "{}" in FRAGMENT_INSTRUCTION
    assert "全桥" in FRAGMENT_INSTRUCTION and "basic_quotes" in FRAGMENT_INSTRUCTION
