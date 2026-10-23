from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from rapidfuzz.fuzz import WRatio
from scipy.optimize import linear_sum_assignment

from bridge_evidence.adapters import FormatAdapter


def base_stem(path: Path) -> str:
    value = path.stem
    value = re.sub(r"[-—_ ]*(?:含有?|有|无)?对比年度的?信息提取报告.*$", "", value)
    value = re.sub(r"[-—_ ]*信息提取报告.*$", "", value)
    value = value.replace(".docx", "").replace(".doc", "")
    return re.sub(r"[\s_\-—（）()·.,，。]+", "", value).lower()


def logical_row(row) -> list[str]:
    seen: set[int] = set()
    values: list[str] = []
    for cell in row.cells:
        identity = id(cell._tc)
        if identity in seen:
            continue
        seen.add(identity)
        values.append(" ".join(cell.text.split()))
    return values


def canonical_header(value: str) -> str:
    """Normalize harmless label-template header variants.

    Released labels use the same semantic schema but a handful of documents
    retain a year suffix or the synonym ``位置``.  Canonicalizing headers is
    necessary for a faithful holdout diagnostic; it does not change extractor
    output and cannot expose any row value.
    """

    value = " ".join(value.split())
    value = re.sub(r"（含照片编号）|\(含照片编号\)", "", value)
    value = re.sub(r"（20\d{2}年）|\(20\d{2}年\)", "", value)
    aliases = {
        "病害位置": "病害部位",
        "上次定检状态": "上一次定检状态",
        "是否新病害": "是否新增",
        "发展情况": "发展程度",
    }
    return aliases.get(value, value)


@dataclass(frozen=True)
class LabelReport:
    path: Path
    summary: dict[str, str]
    recommendations: list[dict[str, str]]
    diseases: list[dict[str, str]]


def parse_label(path: Path, adapter: FormatAdapter | None = None) -> LabelReport:
    if path.suffix.lower() == ".doc":
        if adapter is None:
            raise ValueError("A FormatAdapter is required to parse a binary .doc label")
        document = Document(adapter.adapt(path).docx_path)
    else:
        document = Document(path)
    if len(document.tables) < 3:
        raise ValueError(f"Label does not contain three tables: {path}")
    summary: dict[str, str] = {}
    for row in document.tables[0].rows[1:]:
        values = logical_row(row)
        if len(values) >= 2:
            summary[values[0]] = values[1]
    rec_headers = [canonical_header(value) for value in logical_row(document.tables[1].rows[0])]
    recommendations = []
    for row in document.tables[1].rows[1:]:
        values = logical_row(row)
        if values:
            recommendations.append(dict(zip(rec_headers, values)))
    disease_headers = [canonical_header(value) for value in logical_row(document.tables[2].rows[0])]
    diseases = []
    for row in document.tables[2].rows[1:]:
        values = logical_row(row)
        if values:
            diseases.append(dict(zip(disease_headers, values)))
    return LabelReport(path, summary, recommendations, diseases)


def pair_files(input_dir: Path, label_dir: Path) -> list[tuple[Path, Path]]:
    inputs = sorted(path for path in input_dir.iterdir() if path.suffix.lower() in {".doc", ".docx"})
    labels = sorted(path for path in label_dir.iterdir() if path.suffix.lower() in {".doc", ".docx"})
    if not inputs or not labels:
        return []

    input_names = [base_stem(path) for path in inputs]
    label_names = [base_stem(path) for path in labels]
    scores: list[list[float]] = []
    for label_name in label_names:
        row: list[float] = []
        for input_name in input_names:
            if label_name == input_name:
                score = 100.0
            elif label_name in input_name or input_name in label_name:
                shorter = min(len(label_name), len(input_name))
                longer = max(len(label_name), len(input_name))
                score = 96.0 + (shorter / longer if longer else 0.0)
            else:
                score = float(WRatio(label_name, input_name))
            row.append(score)
        scores.append(row)

    label_indices, input_indices = linear_sum_assignment(scores, maximize=True)
    pairs: list[tuple[Path, Path]] = []
    for label_index, input_index in sorted(zip(label_indices, input_indices)):
        score = scores[label_index][input_index]
        if score < 55.0:
            continue
        pairs.append((inputs[input_index], labels[label_index]))
    return pairs
