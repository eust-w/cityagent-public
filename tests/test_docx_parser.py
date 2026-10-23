from __future__ import annotations

from docx import Document
from PIL import Image

from bridge_evidence.parsers.docx import DocxStructureParser


def test_parser_reads_paragraphs_and_tables_embedded_in_cells(tmp_path) -> None:
    path = tmp_path / "embedded.docx"
    document = Document()
    document.add_paragraph("顶层封面")
    outer = document.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "5.4 全桥BCI值"
    cell.add_paragraph("桥梁BCI=95.35，整体技术状况等级评定为A级。")
    nested = cell.add_table(rows=2, cols=3)
    nested.rows[0].cells[0].text = "部位名称"
    nested.rows[0].cells[1].text = "技术状况指数"
    nested.rows[0].cells[2].text = "BCI"
    nested.rows[1].cells[0].text = "桥面系"
    nested.rows[1].cells[1].text = "69.00"
    nested.rows[1].cells[2].text = "95.35"
    document.save(path)

    parsed = DocxStructureParser().parse(path)

    assert "桥梁BCI=95.35" in parsed.all_text
    depth_by_text = {paragraph.text: paragraph.table_depth for paragraph in parsed.paragraphs if paragraph.text}
    assert depth_by_text["顶层封面"] == 0
    assert depth_by_text["桥梁BCI=95.35，整体技术状况等级评定为A级。"] == 1
    assert depth_by_text["桥面系"] == 2
    assert any(table.header[:3] == ["部位名称", "技术状况指数", "BCI"] for table in parsed.tables)
    rating_table = next(table for table in parsed.tables if table.header and table.header[0] == "部位名称")
    assert any("全桥BCI值" in item for item in rating_table.context)


def test_parser_reads_header_footer_and_image_occurrences(tmp_path) -> None:
    image_path = tmp_path / "defect.png"
    Image.new("RGB", (24, 24), "orange").save(image_path)
    path = tmp_path / "stories.docx"
    document = Document()
    document.add_paragraph("正文结论")
    document.sections[0].header.paragraphs[0].text = "页眉桥梁名称"
    document.sections[0].footer.paragraphs[0].text = "页脚报告编号"
    document.add_picture(str(image_path))
    document.add_picture(str(image_path))
    document.save(path)

    parsed = DocxStructureParser().parse(path)

    stories = {(paragraph.story, paragraph.text) for paragraph in parsed.paragraphs}
    assert ("body", "正文结论") in stories
    assert ("header", "页眉桥梁名称") in stories
    assert ("footer", "页脚报告编号") in stories
    assert len(parsed.images) == 2
    assert {image.story for image in parsed.images} == {"body"}
    assert len({image.target for image in parsed.images}) == 1
