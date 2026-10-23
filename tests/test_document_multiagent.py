from copy import deepcopy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
import time

from docx import Document
import pytest

from bridge_evidence.contest import document_multiagent as pipeline
from bridge_evidence.contest import document_multiagent_worker as worker
from bridge_evidence.contest.document_entries import EntryStore
from bridge_evidence.contest.document_multiagent_cli import main
from bridge_evidence.contest.document_workspace import DocumentWorkspace
from bridge_evidence.contest.evidence_manifest import manifest_from_text
from bridge_evidence.contest.run_state import atomic_json, json_digest


def candidate(tmp_path, manifest, task="part1", field="桥梁名称", value="合成桥", rows=False):
    store = EntryStore(manifest, task, tmp_path / task)
    unit = manifest["units"][0]
    store.record_read(unit["id"], 0, len(unit["text"]))
    entry = {"id": "same-local-id", "section": "基本信息", "field": field, "value": value,
             "source_refs": [{"evidence_id": unit["id"]}], "status": "explicit"}
    if rows:
        entry.pop("field"); entry.pop("value")
        entry.update(section="病害明细", row={"病害部位": "左侧", "病害描述": "裂缝", "发展程度": None})
    store.save_entries([entry])
    return {"task_id": task, "entries": store.snapshot(), "handoff": {"unresolved": ["年份待核"]},
            "manifest_sha256": json_digest(manifest), "status": "failed", "generation_attempts_settled": True,
            "model_calls": 1}


def test_same_local_ids_are_namespaced_and_conflicting_basic_values_not_chosen(tmp_path):
    manifest = manifest_from_text("合成桥在历史段和当前段重名。")
    workspace = DocumentWorkspace(manifest, tmp_path / "merge", {"model": "synthetic"})
    results = [candidate(tmp_path, manifest, "a"), candidate(tmp_path, manifest, "b")]
    receipt = pipeline.import_candidates(workspace, results)
    state = workspace.snapshot()
    assert receipt["semantic_approval"] is False
    assert len(state["facts"]) == 2
    assert "桥梁名称" not in state["draft"]["基本信息"]
    assert len(state["questions"]) == 3
    assert all(fact["source_refs"] for fact in state["facts"].values())
    assert workspace.delivery_status()["ready"] is False


def test_rows_are_directly_deliverable_and_duplicates_left_to_model(tmp_path):
    manifest = manifest_from_text("左侧裂缝，去年记录不明确。")
    workspace = DocumentWorkspace(manifest, tmp_path / "merge", {"model": "synthetic"})
    results = [candidate(tmp_path, manifest, task, rows=True) for task in ("a", "b")]
    pipeline.import_candidates(workspace, results)
    data = workspace.snapshot()["draft"]
    assert len(data["病害明细"]) == 2
    assert all(row["发展程度"] is None for row in data["病害明细"])
    path = pipeline.write_partial(tmp_path, results, "合成验证")
    document = Document(path)
    assert len(document.tables[-1].rows) == 3
    assert any("未完成交付" in paragraph.text for paragraph in document.paragraphs)


@pytest.mark.parametrize("override", [{"report_timeout": float("nan")}, {"report_timeout": True},
    {"context": 65536}, {"workers": 0}, {"max_model_calls": 3}, {"compression_trigger_tokens": 262144}])
def test_settings_reject_invalid_shared_budgets(override):
    with pytest.raises(ValueError):
        pipeline.MultiAgentSettings(**override)


def test_hash_pinned_plan_only_does_not_create_clients(tmp_path, monkeypatch):
    manifest = manifest_from_text("1 结构检查\n合成原文")
    source = tmp_path / "source.json"
    atomic_json(source, manifest)
    monkeypatch.setattr(pipeline, "make_client", lambda *a, **k: pytest.fail("No inference during planning"))
    args = ["--manifest", str(source), "--manifest-sha256", hashlib.sha256(source.read_bytes()).hexdigest(),
            "--scope", "official-training", "--output-dir", str(tmp_path / "plan")]
    assert main(args) == 0
    assert (tmp_path / "plan" / "plan.json").is_file()
    args[3] = "0" * 64
    args[-1] = str(tmp_path / "wrong")
    with pytest.raises(SystemExit):
        main(args)
    assert not (tmp_path / "wrong").exists()


@pytest.mark.parametrize("batch_mode", [False, True])
def test_unsettled_child_blocks_parent_and_keeps_saved_word(tmp_path, monkeypatch, batch_mode):
    manifest = manifest_from_text("合成桥")
    def dispatch(tasks, **kwargs):
        item = candidate(tmp_path, manifest, tasks[0]["id"])
        item["generation_attempts_settled"] = False
        path = tmp_path / "child.json"
        atomic_json(path, item)
        return {"status": "unknown", "results": [{"task_id": item["task_id"], "result_path": str(path),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "settled": False}]}
    monkeypatch.setattr(pipeline, "make_client", lambda *a, **k: pytest.fail("Unknown generation must block parent"))
    config = {"model": "synthetic"}
    observed = []
    if batch_mode:
        from bridge_evidence.contest import document_batch_settlement, document_generation_settlement
        config["batch_admission"] = {"batch_id": "batch", "document_id": "this-document", "deadline": time.monotonic() + 60}
        def shared_observer(actual_dispatch, **kwargs):
            observed.append((actual_dispatch, kwargs))
            return {"may_continue_aggregation": False, "error_code": "SYNTHETIC_UNKNOWN"}
        monkeypatch.setattr(document_batch_settlement, "observe_batch_child_settlement", shared_observer)
        monkeypatch.setattr(document_generation_settlement, "observe_generation_quiescence",
            lambda *a, **k: pytest.fail("Batch must not wait for other documents' service requests"))
    result = pipeline.run_multiagent(manifest, tmp_path / "run", config, dispatcher=dispatch)
    if batch_mode:
        assert len(observed) == 1
        assert observed[0][1]["batch_config"] == config["batch_admission"]
        assert observed[0][1]["manifest_sha256"] == json_digest(manifest)
        assert observed[0][1]["run_directory"] == tmp_path / "run" / "children"
    assert result["complete"] is False
    assert "Unsettled" in result["error"]["message"]
    assert result["child_results"][0]["entry_count"] == 1
    assert "合成桥" in " ".join(cell.text for row in Document(result["partial_docx"]).tables[0].rows for cell in row.cells)


def test_settled_failed_child_candidates_reach_merger_without_budget_expansion(tmp_path, monkeypatch):
    manifest = manifest_from_text("合成桥")
    def dispatch(tasks, **kwargs):
        assert kwargs["max_workers"] == 2
        task = tasks[0]
        taskdir = tmp_path / "child"
        taskdir.mkdir()
        command = kwargs["worker_command"](task, taskdir)
        request = json.loads((taskdir / "request.json").read_bytes())
        assert request["max_model_calls"] == 128
        assert "document_multiagent_worker" in command[-3]
        item = candidate(tmp_path, manifest, task["id"])
        path = taskdir / "result.json"
        atomic_json(path, item)
        return {"status": "failed", "results": [{"task_id": task["id"], "result_path": str(path),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "settled": True}]}
    fake_client = SimpleNamespace(base_url="http://unused", model="synthetic", api_key="synthetic")
    monkeypatch.setattr(pipeline, "make_client", lambda *a, **k: fake_client)
    class Merger:
        artifact_delivery = None
        def __init__(self, source, client, config, directory, **kwargs):
            assert kwargs["settings"].max_model_calls == 128
            assert kwargs["child_summary"][0]["status"] == "failed"
            assert kwargs["child_summary"][0]["primary_evidence_ids"]
            self.workspace = DocumentWorkspace(source, directory, {"model": "synthetic"})
        def run(self):
            assert self.workspace.snapshot()["draft"]["基本信息"]["桥梁名称"] == "合成桥"
            raise RuntimeError("Synthetic Word self-check failure")
        def _model_calls(self):
            return []
    result = pipeline.run_multiagent(manifest, tmp_path / "run", {"model": "synthetic"},
                                    dispatcher=dispatch, workflow_type=Merger)
    assert not result["complete"]
    assert result["error"]["message"] == "Synthetic Word self-check failure"
    assert Path(result["partial_docx"]).is_file()


class FakeGateway:
    instance = None
    def __init__(self, client, **kwargs):
        self.receipts = []
        self.url = "http://unused"
        FakeGateway.instance = self
    def __enter__(self):
        return self
    def __exit__(self, *args):
        pass


@pytest.mark.parametrize("mode", ["success", "missing_read", "empty_final", "stale_handoff", "unknown_generation"])
def test_child_handoff_requires_actual_reading_and_current_saved_state(tmp_path, monkeypatch, mode):
    manifest = manifest_from_text("合成桥。")
    source = tmp_path / "source.json"
    atomic_json(source, manifest)
    unit = manifest["units"][0]
    request = {"manifest_path": str(source), "manifest_sha256": json_digest(manifest),
        "task": {"id": "part1", "primary_evidence_ids": [unit["id"]]},
        "model_config": {}, "deadline_monotonic": time.monotonic() + 1000, "max_model_calls": 8}
    client = SimpleNamespace(calls=[], model="synthetic", api_key="synthetic", count_tokens=lambda x: len(str(x)))
    monkeypatch.setattr(worker, "AuditedModelGateway", FakeGateway)
    def react(prompt, handlers, config):
        assert config["profile"] == "section"
        assert config["timeout"] <= 945
        gateway = FakeGateway.instance
        gateway.on_receipt({"status": "success", "call_id": "synthetic1"})
        if mode != "missing_read":
            payload = {"messages": [{"role": "user", "content": prompt}]}
            audit = {"request_id": "synthetic-request", "call_id": "synthetic1",
                     "input_delivery_confirmed": True, "http_status": 200}
            gateway.on_prepared_request(payload, audit)
            gateway.on_input_observed(payload, audit)
        receipt = handlers["complete_section"]("完成分工阅读", [])
        if mode == "missing_read":
            assert receipt["status"] == "incomplete_reading"
        else:
            assert receipt["status"] == "handed_off"
        if mode == "stale_handoff":
            result = handlers["save_entries"]([{"id": "name", "section": "基本信息", "field": "桥梁名称",
                "value": "合成桥", "source_refs": [{"evidence_id": unit["id"]}], "status": "explicit"}])
            assert result["ok"]
        if mode == "unknown_generation":
            gateway.receipts.append({"generation_attempted": True, "request_id": "no-client-receipt"})
        return {"status": "completed", "native_final_event": mode != "empty_final"}
    result = worker.run_section(request, tmp_path / "section", client=client, react_runner=react)
    assert (tmp_path / "section" / "model-source-map.json").is_file()
    exposure_plan = json.loads((tmp_path / "section" / "initial-exposure-plan.json").read_bytes())
    assert exposure_plan["read_receipt"] is False
    if mode == "missing_read":
        assert not result["entries"]["reads"]
    assert result["document_delivery_complete"] is False
    if mode == "unknown_generation":
        assert not result["generation_attempts_settled"]
    else:
        assert result["generation_attempts_settled"]
        assert (result["status"] == "completed") == (mode == "success")


def test_initial_prompt_has_complete_units_only_and_preserves_ranges():
    manifest = manifest_from_text("短段\n" + "长" * 1000)
    task = {"id": "t", "primary_evidence_ids": [u["id"] for u in manifest["units"]]}
    prompt, exposed = worker.section_prompt(manifest, task, max_initial_characters=1300)
    assert len(prompt) <= 1300
    assert all(span["end"] == len(next(u["text"] for u in manifest["units"] if u["id"] == span["evidence_id"])) for span in exposed)
    reads = {span["evidence_id"]: [[span["start"], span["end"]]] for span in exposed}
    assert worker.unread_primary(manifest, task, reads)
    assert "未内嵌" in prompt


def test_initial_prompt_distinguishes_merged_cell_from_actual_blank():
    manifest = manifest_from_text(" | 裂缝1条")
    unit = manifest["units"][0]
    unit.update(kind="table_row", table_id="T1", row_index=2, segment_index=0,
        cells=[{"column_index": 0, "column_span": 1, "vertical_merge": "continue",
                "merged_context_text": "左幅1号梁", "merged_from_source": "cell-above"}])
    manifest["tables"] = [{"id": "T1", "parent_table_id": None, "header_evidence_ids": []}]
    task = {"id": "t", "primary_evidence_ids": [unit["id"]]}
    merged, _ = worker.section_prompt(manifest, task)
    blank = deepcopy(manifest)
    blank["units"][0]["cells"] = [{"column_index": 0, "text": ""}]
    actual_blank, _ = worker.section_prompt(blank, task)
    assert merged != actual_blank
    assert "左幅1号梁" in merged
    assert "cell-above" not in merged  # Exact paths remain in the program's source map.
