"""Material organization only: CPU evidence and canonical input/coverage tests."""
import copy
import json

import pytest
from docx import Document

from bridge_evidence.contest.document_section_input import compact_section_prompt, SectionInputBudgetError
from bridge_evidence.contest.document_task_plan import build_document_navigation, build_document_task_plan
from bridge_evidence.contest.document_source_delivery import SourceDeliveryLedger
from bridge_evidence.contest.evidence_manifest import manifest_from_text, build_manifest


class MaterialIndex:
    """Small exact-mapping double; real XML bindings are exercised separately."""
    def __init__(self, manifest, packets=None, headings=None):
        self.manifest=copy.deepcopy(manifest)
        self.sections=[]
        for ordinal, (position,level) in enumerate(headings or []):
            unit=manifest['units'][position]
            self.sections.append({'id':f'#/sections/S{ordinal}', 'parent':None, 'region':'content',
                'title':unit['text'],'level':level,'heading_spans':[self.span(unit['id'])],
                'unit_start':position,'unit_end':position+1})
        self.packets=[];self.values={}
        for ordinal, item in enumerate(packets or []):
            identifier=f'#/packets/P{ordinal}'
            self.packets.append({'id':identifier,'section_id':'#/sections/S0'})
            self.values[identifier]={'spans':[self.span(e) for e in item.get('primary',[])],
                'context':[self.span(e) for e in item.get('context',[])],
                'unresolved':copy.deepcopy(item.get('unresolved',[]))}
    def span(self,eid):
        unit=next(u for u in self.manifest['units'] if u['id']==eid)
        return {'evidence_id':eid,'start':0,'end':len(unit['text']),'text':unit['text']}
    def validate_manifest(self,manifest):
        if manifest!=self.manifest:raise ValueError('Material index belongs to another manifest')
    def packet(self,ref):return copy.deepcopy(self.values[ref])


def task(manifest,primary=None,context=None):
    return {'id':'section-task','primary_evidence_ids':primary or [u['id'] for u in manifest['units']],
        'context_evidence_ids':context or [],'headings':[],'table_ids':[],
        'material_section_refs':[]}


def data(prompt):return json.loads(prompt.split('\n',1)[1])
def rows(prompt):return [row for group in data(prompt)['materials'] for row in group['units']]


def test_material_units_present_once_with_exact_canonical_text_and_ledger_registration(tmp_path):
    manifest=manifest_from_text('桥梁名称\n2018年与2014年比较\n  支座裂缝😀\t位置  \n未发生明显突变')
    ids=[u['id'] for u in manifest['units']]
    index=MaterialIndex(manifest,[{'primary':ids[1:3],'context':[ids[0]]},{'primary':ids[2:]}])
    current=task(manifest,primary=ids[1:],context=[ids[0]])
    prompt,exposed=compact_section_prompt(manifest,current,material_index=index)
    actual=rows(prompt)
    assert len(actual)==len(set(r['evidence_id'] for r in actual))==4
    assert {r['evidence_id']:r['text'] for r in actual}=={u['id']:u['text'] for u in manifest['units']}
    assert prompt.count(manifest['document_sha256'])==1
    ledger=SourceDeliveryLedger(manifest,tmp_path/'ledger',run_id='run',agent_id='task')
    ledger.register_initial_message(prompt,exposed)
    assert ledger.confirmed_ranges()=={}
    assert all(e=={'evidence_id':u['id'],'start':0,'end':len(u['text'])} for e,u in zip(exposed,manifest['units']))


def test_unmapped_primary_remains_canonical_fallback_without_copying_mineru_metadata():
    manifest=manifest_from_text('原文一\n原文二\n未映射原文三')
    ids=[u['id'] for u in manifest['units']]
    index=MaterialIndex(manifest,[{'primary':[ids[0]],'unresolved':[{'ref':'#/texts/9','code':'unmapped'}]}])
    index.packets[0]['unused']={'source_map':'XML-PATH'*1000,'context':'DO-NOT-COPY'*1000}
    prompt,exposed=compact_section_prompt(manifest,task(manifest),material_index=index)
    assert len(exposed)==3
    assert data(prompt)['unresolved_material_references']==[{'ref':'#/texts/9','code':'unmapped'}]
    assert next(g for g in data(prompt)['materials'] if g.get('group')=='canonical-fallback')['units'][0]['text']==manifest['units'][1]['text']
    assert 'DO-NOT-COPY' not in prompt and 'XML-PATH' not in prompt
    assert 'cell_fields' not in data(prompt) and 'source_map' not in data(prompt)


def test_material_initial_budget_is_full_message_and_does_not_truncate_whole_units():
    manifest=manifest_from_text('很长原文'*5000+'\n短原文')
    index=MaterialIndex(manifest)
    prompt,exposed=compact_section_prompt(manifest,task(manifest),3000,material_index=index)
    assert len(prompt)<=3000 and len(exposed)==1
    assert rows(prompt)[0]['text']=='短原文'
    assert len(data(prompt)['task']['primary_evidence_ids'])==2
    with pytest.raises(SectionInputBudgetError):compact_section_prompt(manifest,task(manifest),100,material_index=index)
    full,full_exposed=compact_section_prompt(manifest,task(manifest),50000,material_index=index)
    same,same_exposed=compact_section_prompt(manifest,task(manifest),len(full),material_index=index)
    assert same==full and same_exposed==full_exposed
    shorter,less=compact_section_prompt(manifest,task(manifest),len(full)-1,material_index=index)
    assert len(shorter)<=len(full)-1 and len(less)<len(full_exposed)


def test_material_binding_and_span_text_mismatch_fail_closed():
    manifest=manifest_from_text('原文')
    eid=manifest['units'][0]['id'];index=MaterialIndex(manifest,[{'primary':[eid]}])
    changed=copy.deepcopy(manifest);changed['units'][0]['text']='改变'
    with pytest.raises(ValueError):compact_section_prompt(changed,task(changed),material_index=index)
    index.values['#/packets/P0']['spans'][0]['text']='虚构'
    with pytest.raises(ValueError,match='canonical source'):compact_section_prompt(manifest,task(manifest),material_index=index)


def test_repaired_material_heading_replaces_regex_false_positive_and_uses_same_e_ids():
    manifest=manifest_from_text('封面\n1 工程概述\n介绍\n2015 年首次检测状态\n正文\n2 检测结果\n病害\n3 建议\n建议正文')
    index=MaterialIndex(manifest,headings=[(1,1),(5,1),(7,1)])
    navigation=build_document_navigation(manifest,material_index=index)
    ids=[u['id'] for u in manifest['units']]
    assert [h['evidence_id'] for h in navigation['headings']]==[ids[1],ids[5],ids[7]]
    assert ids[3] not in navigation['section_evidence_ids']
    plan=build_document_task_plan(manifest,target_tasks=3,material_index=index)
    assert plan['partition_basis']=='mineru_heading_boundaries'
    assert [e for t in plan['tasks'] for e in t['primary_evidence_ids']]==ids
    assert plan['coverage']['all_units_assigned_once']
    assert all(t['material_section_refs'] for t in plan['tasks'])
    assert all(h['title']==manifest['units'][h['unit_index']]['text'] for h in navigation['headings'])


def test_no_exact_material_heading_preserves_legacy_navigation():
    manifest=manifest_from_text('1 工程概述\n正文\n2 建议\n建议正文')
    index=MaterialIndex(manifest)
    assert build_document_navigation(manifest,material_index=index)==build_document_navigation(manifest)
    assert compact_section_prompt(manifest,task(manifest),material_index=None)==compact_section_prompt(manifest,task(manifest))


def test_merged_nested_table_layout_and_candidate_header_remain_distinct(tmp_path):
    path=tmp_path/'merged.docx';doc=Document();doc.add_paragraph('1 工程概述')
    table=doc.add_table(rows=3,cols=3)
    table.cell(0,0).text='部件';table.cell(0,1).text='分数';table.cell(0,2).text='等级'
    table.cell(1,0).merge(table.cell(1,1)).text='综合评分';table.cell(1,2).text='94.57/A'
    table.cell(2,0).text='父表';child=table.cell(2,1).add_table(rows=2,cols=1)
    child.cell(0,0).text='子表头';child.cell(1,0).text='子表原文'
    doc.add_paragraph('2 建议');doc.add_paragraph('维修建议');doc.save(path)
    manifest=build_manifest(path);index=MaterialIndex(manifest,headings=[(0,1)])
    prompt,exposed=compact_section_prompt(manifest,task(manifest),material_index=index)
    table_metadata=data(prompt)['tables']
    assert len(table_metadata)==2 and table_metadata[1]['parent']==table_metadata[0]['id']
    assert table_metadata[0]['header_kind']=='first_row_candidate'
    merged=next(r for r in rows(prompt) if '综合评分' in r['text'])
    assert merged['cells'][0][:2]==[0,2]
    assert merged['text'][merged['cells'][0][2]:merged['cells'][0][3]]=='综合评分'
    plan=build_document_task_plan(manifest,material_index=index)
    assert plan['coverage']['table_families_unbroken']
    assert len(exposed)==len(manifest['units'])
