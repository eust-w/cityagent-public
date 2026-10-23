from __future__ import annotations

import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from PIL import Image

from bridge_evidence.contest.evidence_manifest import build_manifest
from bridge_evidence.contest.visual_inventory import inventory_docx

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _mutate(path, callback):
    with zipfile.ZipFile(path) as archive:
        members = {info.filename: archive.read(info) for info in archive.infolist()}
    callback(members)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)


def _report(tmp_path):
    picture = tmp_path / "image.png"
    Image.new("RGB", (8, 8), "green").save(picture)
    document = Document()
    document.add_paragraph("original body")
    document.add_picture(str(picture))
    for label, story in (
        ("header", document.sections[0].header),
        ("footer", document.sections[0].footer),
    ):
        paragraph = story.paragraphs[0]
        paragraph.text = "original " + label
        paragraph.add_run().add_picture(str(picture))
        paragraph.add_run()._r.append(OxmlElement("w:sym"))
    path = tmp_path / "report.docx"
    document.save(path)
    return path


def test_custom_named_stories_preserve_text_images_symbols_and_exact_paths(tmp_path):
    path = _report(tmp_path)
    original = build_manifest(path)

    def rename(members):
        for old, new in (("header1", "customHeader"), ("footer1", "customFooter")):
            members[f"word/{new}.xml"] = members.pop(f"word/{old}.xml")
            members[f"word/_rels/{new}.xml.rels"] = members.pop(
                f"word/_rels/{old}.xml.rels"
            )
            members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
                f"/{old}.xml".encode(), f"/{new}.xml".encode()
            )
            members["word/_rels/document.xml.rels"] = members[
                "word/_rels/document.xml.rels"
            ].replace(f'"{old}.xml"'.encode(), f'"{new}.xml"'.encode())

    _mutate(path, rename)
    assert Document(path).sections[0].header.paragraphs[0].text == "original header"
    updated = build_manifest(path)
    assert (
        updated["parser_version"]
        == original["parser_version"]
        == "evidence-manifest-v2"
    )
    assert sorted(u["text"] for u in updated["units"]) == sorted(
        u["text"] for u in original["units"]
    )
    assert len(updated["images"]) == len(original["images"]) == 3
    assert (
        len(updated["unverified_content"]) == len(original["unverified_content"]) == 2
    )
    assert {r["source_path"] for r in updated["images"]} == {
        r["source_path"]
        for r in inventory_docx(path)["occurrences"]
        if r["kind"] == "image"
    }
    assert {u["context"]["story"] for u in updated["units"]} == {
        "body",
        "header",
        "footer",
    }
    assert [u for u in updated["units"] if u["context"]["story"] == "body"] == [
        u for u in original["units"] if u["context"]["story"] == "body"
    ]
    assert any(
        r["source_path"].startswith("word/customHeader.xml#")
        for r in updated["unverified_content"]
    )


@pytest.mark.parametrize("kind", ["footnotes", "endnotes", "comments"])
def test_custom_notes_and_comments_are_discovered_by_opc_identity(tmp_path, kind):
    path = _report(tmp_path)

    def insert(members):
        part = "custom-" + kind + ".xml"
        root = etree.Element(f"{{{W}}}{kind}", nsmap={"w": W})
        node = etree.SubElement(root, f"{{{W}}}p")
        run = etree.SubElement(node, f"{{{W}}}r")
        etree.SubElement(run, f"{{{W}}}t").text = kind + " source text"
        etree.SubElement(run, f"{{{W}}}sym")
        members["word/" + part] = etree.tostring(root)
        content_types = etree.fromstring(members["[Content_Types].xml"])
        etree.SubElement(
            content_types,
            f"{{{CT}}}Override",
            PartName="/word/" + part,
            ContentType=f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}+xml",
        )
        members["[Content_Types].xml"] = etree.tostring(content_types)
        rels = etree.fromstring(members["word/_rels/document.xml.rels"])
        etree.SubElement(
            rels,
            f"{{{RELS}}}Relationship",
            Id="rIdCustomStory",
            Target=part,
            Type=R + "/" + kind,
        )
        members["word/_rels/document.xml.rels"] = etree.tostring(rels)

    _mutate(path, insert)
    result = build_manifest(path)
    assert any(
        u["text"] == kind + " source text" and u["context"]["story"] == kind
        for u in result["units"]
    )
    assert any(
        o["source_path"].startswith("word/custom-" + kind + ".xml#")
        for o in result["unverified_content"]
    )


@pytest.mark.parametrize("missing", [False, True])
def test_orphan_or_missing_declared_story_fails_closed(tmp_path, missing):
    path = _report(tmp_path)

    def insert(members):
        if not missing:
            members["word/unreachable.xml"] = (
                f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>must not disappear</w:t></w:r></w:p></w:hdr>'.encode()
            )
        types = etree.fromstring(members["[Content_Types].xml"])
        etree.SubElement(
            types,
            f"{{{CT}}}Override",
            PartName="/word/unreachable.xml",
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
        )
        members["[Content_Types].xml"] = etree.tostring(types)

    _mutate(path, insert)
    with pytest.raises(ValueError, match="missing or unreachable"):
        build_manifest(path)


def test_relationship_discovers_custom_story_with_generic_content_type(tmp_path):
    path = _report(tmp_path)

    def mutate(members):
        members["word/unusual.xml"] = members.pop("word/header1.xml")
        members["word/_rels/unusual.xml.rels"] = members.pop(
            "word/_rels/header1.xml.rels"
        )
        members["word/_rels/document.xml.rels"] = members[
            "word/_rels/document.xml.rels"
        ].replace(b'"header1.xml"', b'"unusual.xml"')
        types = etree.fromstring(members["[Content_Types].xml"])
        for item in types:
            if item.get("PartName") == "/word/header1.xml":
                item.set("PartName", "/word/unusual.xml")
                item.set("ContentType", "application/xml")
        members["[Content_Types].xml"] = etree.tostring(types)

    _mutate(path, mutate)
    value = build_manifest(path)
    assert any(
        unit["text"] == "original header"
        and unit["source_path"].startswith("word/unusual.xml#")
        for unit in value["units"]
    )
    assert len(value["images"]) == 3


def test_story_identity_conflict_and_external_story_are_rejected(tmp_path):
    path = _report(tmp_path)

    def conflict(members):
        members["word/_rels/document.xml.rels"] = members[
            "word/_rels/document.xml.rels"
        ].replace((R + "/header").encode(), (R + "/footer").encode())

    _mutate(path, conflict)
    with pytest.raises(ValueError, match="Conflicting"):
        build_manifest(path)
    path = _report(tmp_path)

    def external(members):
        rels = etree.fromstring(members["word/_rels/document.xml.rels"])
        for item in rels:
            if item.get("Type") == R + "/header":
                item.set("TargetMode", "External")
                item.set("Target", "https://example.invalid/story.xml")
        members["word/_rels/document.xml.rels"] = etree.tostring(rels)

    _mutate(path, external)
    with pytest.raises(ValueError, match="External"):
        build_manifest(path)
