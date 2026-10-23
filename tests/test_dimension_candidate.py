from __future__ import annotations

import json
import importlib.util
from pathlib import Path

from docx import Document

SCRIPT_PATH = Path(__file__).parents[1] / "scripts" / "compose_dimension_candidate.py"
SPEC = importlib.util.spec_from_file_location("compose_dimension_candidate", SCRIPT_PATH)
assert SPEC is not None and SPEC.loader is not None
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)
compose = MODULE.compose


def _write_result(path: Path, *, prefix: str, temporal: tuple[str, str, str]) -> None:
    document = Document()
    summary = document.add_table(rows=17, cols=3)
    for index, row in enumerate(summary.rows):
        row.cells[0].text = f"field-{index}"
        row.cells[1].text = f"{prefix}-summary-{index}"
        row.cells[2].text = "note"
    recommendations = document.add_table(rows=2, cols=4)
    recommendations.rows[0].cells[0].text = "序号"
    for index, value in enumerate(("1", f"{prefix}-category", f"{prefix}-content", f"{prefix}-component")):
        recommendations.rows[1].cells[index].text = value
    details = document.add_table(rows=2, cols=7)
    for index, value in enumerate(("1", f"{prefix}-location", f"{prefix}-type", f"{prefix}-description", *temporal)):
        details.rows[1].cells[index].text = value
    document.save(path)


def _write_state(path: Path) -> None:
    path.write_text(
        json.dumps({"diseases": [{"evidence": ["T2R1"]}]}, ensure_ascii=False),
        encoding="utf-8",
    )


def test_compose_preserves_candidate_summary_and_temporal_fields(tmp_path: Path) -> None:
    baseline_result = tmp_path / "baseline-result"
    baseline_state = tmp_path / "baseline-state"
    candidate_result = tmp_path / "candidate-result"
    candidate_state = tmp_path / "candidate-state"
    output_result = tmp_path / "output"
    for directory in (baseline_result, baseline_state, candidate_result, candidate_state):
        directory.mkdir()
    name = "sample.docx"
    _write_result(baseline_result / name, prefix="baseline", temporal=("未明确", "未明确", "未明确"))
    _write_result(candidate_result / name, prefix="candidate", temporal=("是", "未发现", "无"))
    _write_state(baseline_state / "sample.json")
    _write_state(candidate_state / "sample.json")

    report = compose(
        baseline_result=baseline_result,
        baseline_state=baseline_state,
        candidate_result=candidate_result,
        candidate_state=candidate_state,
        output_result=output_result,
    )

    assert report["passed"] is True
    document = Document(output_result / name)
    # Candidate summary survives except for the category-count row, which
    # follows the restored recommendation ledger.
    assert document.tables[0].rows[13].cells[1].text == "candidate-summary-13"
    assert document.tables[0].rows[16].cells[1].text == "baseline-summary-16"
    assert document.tables[1].rows[1].cells[2].text == "baseline-content"
    assert [cell.text for cell in document.tables[2].rows[1].cells[1:4]] == [
        "baseline-location",
        "baseline-type",
        "baseline-description",
    ]
    assert [cell.text for cell in document.tables[2].rows[1].cells[4:7]] == [
        "是",
        "未发现",
        "无",
    ]
