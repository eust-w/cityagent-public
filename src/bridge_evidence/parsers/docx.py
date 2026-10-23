from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import parse_xml
from docx.oxml.ns import qn


def compact(text: str) -> str:
    return " ".join(text.replace("\u3000", " ").replace("\xa0", " ").split())


@dataclass(frozen=True)
class ParagraphBlock:
    index: int
    text: str
    style: str
    # Native DOCX narrative paragraphs normally live at depth 0 and table
    # cells at depth 1.  Legacy binary Word conversion often wraps the whole
    # report in one outer table, shifting those depths to 1 and 2.  Retaining
    # the relative depth lets evidence tools distinguish section summaries
    # from photo-ledger cells without relying on a report year or filename.
    table_depth: int = 0
    story: str = "body"


@dataclass(frozen=True)
class TableBlock:
    index: int
    rows: list[list[str]]
    context: list[str] = field(default_factory=list)
    story: str = "body"

    @property
    def header(self) -> list[str]:
        return self.rows[0] if self.rows else []

    def contains(self, *terms: str) -> bool:
        haystack = "|".join(self.header)
        return all(term in haystack for term in terms)

    def is_disease_table(self) -> bool:
        header = "|".join("".join(cell.split()) for cell in self.header)
        has_location = any(term in header for term in ("病害位置", "构件位置", "位置", "桥幅", "桥跨", "构件部位"))
        has_type = any(term in header for term in ("病害类型", "病害种类", "损坏类型"))
        has_description = any(term in header for term in ("病害描述", "病害情况", "具体情况", "具体位置", "检查结果", "检测结果"))
        # Several official templates put the full location directly in the
        # description column and therefore omit a dedicated location column.
        return has_type and has_description and (has_location or "序号" in header)


@dataclass(frozen=True)
class ImageBlock:
    index: int
    relationship_id: str
    target: str
    alt_text: str = ""
    story: str = "body"


@dataclass(frozen=True)
class ParsedDocument:
    path: Path
    paragraphs: list[ParagraphBlock]
    tables: list[TableBlock]
    images: list[ImageBlock] = field(default_factory=list)

    @property
    def all_text(self) -> str:
        return "\n".join(p.text for p in self.paragraphs if p.text)

    @property
    def full_text(self) -> str:
        table_text = (" | ".join(row) for table in self.tables for row in table.rows)
        return "\n".join([self.all_text, *table_text])


class DocxStructureParser:
    def parse(self, path: Path) -> ParsedDocument:
        document: DocumentObject = Document(path)
        # ``Document.paragraphs`` and ``Document.tables`` only expose blocks
        # that are direct children of the Word body.  A sizeable family of
        # legacy binary-Word reports is converted into one outer table, while
        # its real paragraphs and rating/disease tables live inside cells,
        # structured-document tags, or text boxes.  Reading only the direct
        # children therefore makes an otherwise visible ``BCI=...`` sentence
        # completely invisible to the evidence layer.
        #
        # Walk the OOXML body in document order instead.  Each XML block is
        # visited exactly once, so this remains independent of filenames,
        # report years, and bridge identities while recovering every embedded
        # paragraph/table that Word itself renders.
        paragraphs: list[ParagraphBlock] = []
        tables: list[TableBlock] = []
        images: list[ImageBlock] = []
        paragraph_tag = qn("w:p")
        table_tag = qn("w:tbl")
        style_names = {style.style_id: style.name for style in document.styles}
        story_parts: list[tuple[str, Any, Any]] = [
            ("body", document.element.body, document.part)
        ]
        for part in document.part.package.parts:
            part_name = str(part.partname).lstrip("/")
            story = self._story_name(part_name)
            if story is None or part_name == "word/document.xml":
                continue
            story_parts.append((story, parse_xml(part.blob), part))

        for story, root, part in story_parts:
            recent: list[str] = []
            for element in root.iter():
                if element.tag == paragraph_tag:
                    text = compact(self._xml_text(element))
                    properties = element.find(qn("w:pPr"))
                    style_node = properties.find(qn("w:pStyle")) if properties is not None else None
                    style_id = style_node.get(qn("w:val"), "") if style_node is not None else ""
                    paragraphs.append(
                        ParagraphBlock(
                            len(paragraphs),
                            text,
                            style_names.get(style_id, style_id),
                            self._table_depth(element),
                            story,
                        )
                    )
                    if text:
                        recent.append(text)
                        # Keep enough heading history for long narrative sections:
                        # some legacy reports place 15-20 summary paragraphs between
                        # a main/approach bridge heading and the evidence table.
                        recent = recent[-30:]
                elif element.tag == table_tag:
                    tables.append(
                        TableBlock(
                            len(tables),
                            self._logical_rows(element),
                            list(recent),
                            story,
                        )
                    )
            images.extend(self._images(root, part, story, len(images)))
        return ParsedDocument(
            path=path,
            paragraphs=paragraphs,
            tables=tables,
            images=images,
        )

    @staticmethod
    def _story_name(part_name: str) -> str | None:
        if re.fullmatch(r"word/header\d+\.xml", part_name):
            return "header"
        if re.fullmatch(r"word/footer\d+\.xml", part_name):
            return "footer"
        for story in ("footnotes", "endnotes", "comments"):
            if part_name == f"word/{story}.xml":
                return story
        return None

    @staticmethod
    def _images(root: Any, part: Any, story: str, start: int) -> list[ImageBlock]:
        image_tags = {
            qn("a:blip"),
            "{urn:schemas-microsoft-com:vml}imagedata",
        }
        relationship_attributes = (qn("r:embed"), qn("r:id"), qn("r:link"))
        doc_properties_tag = qn("wp:docPr")
        images: list[ImageBlock] = []
        for element in root.iter():
            if element.tag not in image_tags:
                continue
            relationship_id = next(
                (
                    element.get(attribute)
                    for attribute in relationship_attributes
                    if element.get(attribute)
                ),
                "",
            )
            relation = part.rels.get(relationship_id) if relationship_id else None
            if relation is None:
                target = ""
            elif relation.is_external:
                target = str(relation.target_ref)
            else:
                target = str(relation.target_part.partname).lstrip("/")

            alt_text = ""
            ancestor = element.getparent()
            while ancestor is not None and not alt_text:
                properties = next(
                    (node for node in ancestor.iter() if node.tag == doc_properties_tag),
                    None,
                )
                if properties is not None:
                    alt_text = properties.get("descr") or properties.get("title") or ""
                ancestor = ancestor.getparent()
            images.append(
                ImageBlock(
                    index=start + len(images),
                    relationship_id=relationship_id,
                    target=target,
                    alt_text=compact(alt_text),
                    story=story,
                )
            )
        return images

    @staticmethod
    def _table_depth(element: Any) -> int:
        """Count enclosing OOXML tables for a paragraph."""
        table_tag = qn("w:tbl")
        depth = 0
        parent = element.getparent()
        while parent is not None:
            if parent.tag == table_tag:
                depth += 1
            parent = parent.getparent()
        return depth

    @staticmethod
    def _xml_text(element: Any) -> str:
        """Return visible OOXML text without repeated python-docx XPath calls."""
        text_tag = qn("w:t")
        tab_tag = qn("w:tab")
        break_tags = {qn("w:br"), qn("w:cr")}
        parts: list[str] = []
        for node in element.iter():
            if node.tag == text_tag:
                parts.append(node.text or "")
            elif node.tag == tab_tag:
                parts.append(" ")
            elif node.tag in break_tags:
                parts.append("\n")
        return "".join(parts)

    @classmethod
    def _cell_text_without_nested_tables(cls, cell_element: Any) -> str:
        """Read direct cell content; descendant tables are parsed separately."""
        paragraph_tag = qn("w:p")
        table_tag = qn("w:tbl")
        paragraphs: list[str] = []

        def visit(element: Any) -> None:
            for child in element:
                if child.tag == table_tag:
                    continue
                if child.tag == paragraph_tag:
                    paragraphs.append(cls._xml_text(child))
                    continue
                visit(child)

        visit(cell_element)
        return "\n".join(paragraphs)

    @classmethod
    def _logical_rows(cls, table_element: Any) -> list[list[str]]:
        """Extract a table in one OOXML pass, preserving vertical merges.

        Using ``Table.rows -> row.cells -> cell.text`` repeatedly evaluates
        XPath against the full document.  On long-span reports containing
        thousands of cells and nested layout tables that becomes effectively
        quadratic.  Direct-child OOXML traversal is linear and also prevents
        an outer layout cell from duplicating every descendant table.
        """
        rows: list[list[str]] = []
        previous_vertical: dict[int, str] = {}
        for row in table_element.findall(qn("w:tr")):
            values: list[str] = []
            column = 0
            for cell in row.findall(qn("w:tc")):
                properties = cell.find(qn("w:tcPr"))
                span_node = properties.find(qn("w:gridSpan")) if properties is not None else None
                try:
                    span = max(1, int(span_node.get(qn("w:val"), "1"))) if span_node is not None else 1
                except ValueError:
                    span = 1

                value = compact(cls._cell_text_without_nested_tables(cell))
                merge = properties.find(qn("w:vMerge")) if properties is not None else None
                merge_value = merge.get(qn("w:val")) if merge is not None else None
                if merge is not None and merge_value != "restart" and not value:
                    value = previous_vertical.get(column, "")

                if merge is not None:
                    for offset in range(span):
                        previous_vertical[column + offset] = value
                else:
                    for offset in range(span):
                        previous_vertical.pop(column + offset, None)

                values.append(value)
                column += span
            rows.append(values)
        return rows
