from pathlib import Path

import pytest
from docx import Document
from test_chunked_integration import CharTokenizer

import run
from bridge_evidence.contest.candidate import validate_config
from bridge_evidence.contest.run_state import file_digest


def arguments(tmp_path, workflow='chunked'):
    args = run.build_parser().parse_args(['--engine', 'direct', '--workflow', workflow,
                                         '--work-dir', str(tmp_path / 'work')])
    args.candidate = validate_config({'workflow': workflow, 'input_budget': 3000,
                                      'output_budget': 500})
    return args


def test_chunk_preparation_does_not_extract_fulltext_again(tmp_path, monkeypatch):
    path = tmp_path / 'report.docx'
    doc = Document()
    doc.add_paragraph('测试桥')
    doc.save(path)
    def unexpected(*args):
        pytest.fail('chunk preparation should reuse snapshot text')
    monkeypatch.setattr(run, 'extract_docx_text', unexpected)
    prepared = run.prepare_document(arguments(tmp_path), path, CharTokenizer())
    assert prepared['document_sha256'] == file_digest(path)
    assert prepared['snapshot'].regions[0].text == '测试桥'
    assert prepared['chunk_plan'].coverage(prepared['snapshot'])['complete']


def test_doc_chunk_preparation_only_converts_once(tmp_path, monkeypatch):
    path = tmp_path / 'old.doc'
    path.write_bytes(b'synthetic DOC fixture')
    converted = tmp_path / 'old.docx'
    doc = Document()
    doc.add_paragraph('测试桥')
    doc.save(converted)
    from bridge_evidence.contest.run_state import atomic_json
    atomic_json(tmp_path / 'conversion.json', {
        'source_sha256': file_digest(path), 'output_sha256': file_digest(converted),
        'converter': 'LibreOffice test', 'conversion': 'docx',
    })
    monkeypatch.setattr(run, 'convert_doc_to_docx', lambda *args: converted)
    monkeypatch.setattr(run, 'convert_doc_to_text', lambda *args: pytest.fail('second conversion'))
    prepared = run.prepare_document(arguments(tmp_path), path, CharTokenizer())
    assert prepared['target_docx'] == str(converted)


def test_whole_preparation_preserves_existing_text_contract(tmp_path, monkeypatch):
    path = tmp_path / 'old.doc'
    path.write_bytes(b'synthetic DOC fixture')
    target = tmp_path / 'old.docx'
    doc = Document()
    doc.add_paragraph('legacy text')
    doc.save(target)
    from bridge_evidence.contest.run_state import atomic_json
    atomic_json(tmp_path / 'conversion.json', {
        'source_sha256': file_digest(path), 'output_sha256': file_digest(target),
        'converter': 'LibreOffice test', 'conversion': 'docx',
    })
    monkeypatch.setattr(run, 'convert_doc_to_docx', lambda *args: target)
    monkeypatch.setattr(run, 'convert_doc_to_text', lambda *args: 'legacy text')
    prepared = run.prepare_document(arguments(tmp_path, 'whole'), path, CharTokenizer())
    assert prepared['text'] == 'legacy text'
    assert 'snapshot' not in prepared
    assert prepared['source_evidence']['source_sha256'] == file_digest(path)


def test_prepared_model_cli_requires_pinned_receipt_pair():
    args = run.build_parser().parse_args(['--prepared-model-identity', '/tmp/model.json',
                                         '--prepared-model-identity-sha256', 'a' * 64])
    assert args.prepared_model_identity == Path('/tmp/model.json')


def test_partial_resume_defers_plan_to_checkpoint_reader(tmp_path, monkeypatch):
    from bridge_evidence.contest import chunked_inference
    path = tmp_path / 'report.docx'
    doc = Document()
    doc.add_paragraph('测试桥')
    doc.save(path)
    args = arguments(tmp_path)
    args.resume = True
    monkeypatch.setattr(chunked_inference, 'prepare_chunk_plan',
                        lambda *args: pytest.fail('resume rebuilt plan before checkpoint'))
    prepared = run.prepare_document(args, path, CharTokenizer())
    assert prepared['chunk_plan'] is None


def test_chunk_preparation_binds_exact_source_and_prepared_document(tmp_path):
    path = tmp_path / 'report.docx'
    doc = Document()
    doc.add_paragraph('桥梁来源绑定')
    doc.save(path)
    prepared = run.prepare_document(arguments(tmp_path), path, CharTokenizer())
    assert prepared['source_evidence'] == {
        'schema_version': 1,
        'source_sha256': file_digest(path),
        'prepared_docx_path': str(path.resolve()),
        'prepared_docx_sha256': file_digest(path),
        'conversion': None,
    }


def test_source_evidence_rejects_unbound_doc_conversion(tmp_path):
    path, target = tmp_path / 'report.doc', tmp_path / 'report.docx'
    path.write_bytes(b'original doc')
    target.write_bytes(b'converted docx')
    from bridge_evidence.contest.run_state import atomic_json
    metadata = {
        'source_sha256': '0' * 64, 'output_sha256': file_digest(target),
        'converter': 'LibreOffice test', 'conversion': 'docx',
    }
    atomic_json(tmp_path / 'conversion.json', metadata)
    with pytest.raises(ValueError, match='conversion'):
        run.bind_source_evidence(path, target)
    metadata['source_sha256'] = file_digest(path)
    atomic_json(tmp_path / 'conversion.json', metadata)
    assert run.bind_source_evidence(path, target)['conversion'] == metadata
