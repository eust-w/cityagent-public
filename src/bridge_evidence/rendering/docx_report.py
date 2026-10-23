from __future__ import annotations

from collections import Counter
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from bridge_evidence.evidence.disease_graph import published_diseases
from bridge_evidence.models.schema import DocumentState


# Use a system CJK family for the submission result.  The previous user-local
# LXGW font disappeared when LibreOffice rendered with an isolated HOME,
# leaving Chinese text as tofu boxes.  Arial Unicode MS is installed as a
# system-wide TTF on the competition workstation; unlike macOS TTC families it
# is also rendered reliably by LibreOffice during the package visual audit.
FONT = "Arial Unicode MS"
TABLE_WIDTH = 9360
TABLE_INDENT = 120
HEADER_FILL = "E8EEF5"
BORDER_COLOR = "AAB4C0"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)


def _set_run(run, *, size: float = 11, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _set_cell_margins(cell, top: int = 80, left: int = 120, bottom: int = 80, right: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"Column widths must sum to {TABLE_WIDTH}, got {sum(widths)}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _set_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), BORDER_COLOR)


def _cell_text(cell, text: str, *, header: bool = False, center: bool = False, size: float = 9.2) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    _set_run(paragraph.add_run(str(text or "")), size=size, bold=header)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        _shade_cell(cell, HEADER_FILL)


def _add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        _cell_text(table.rows[0].cells[index], header, header=True, center=True, size=9.4)
    _set_repeat_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _cell_text(cells[index], value, center=index == 0 or len(str(value)) < 10)
    _set_table_geometry(table, widths)
    _set_borders(table)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


class DocxReportRenderer:
    """Render the canonical state using compact_reference_guide + memo_masthead."""

    def render(self, state: DocumentState, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        self._configure_document(document, state)
        self._add_masthead(document, state)
        document.add_heading("1、简要信息", level=1)
        _add_table(document, ["字段", "内容", "说明"], self._summary_rows(state), [1800, 4300, 3260])

        document.add_heading("2、详细信息", level=1)
        lead = document.add_paragraph("在简要信息基础上，进一步提取结构化明细与证据约束结论。")
        lead.paragraph_format.space_after = Pt(8)
        document.add_heading("（1）详细结论", level=2)
        for paragraph_text in self._detailed_conclusion(state):
            paragraph = document.add_paragraph(paragraph_text)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.25

        document.add_heading("（2）建议明细", level=2)
        # The source reports often restart their local sequence number for each
        # component or carriageway.  The submitted table is one merged ledger,
        # so its display sequence must be globally continuous.
        recommendation_rows = []
        for index, r in enumerate(state.recommendations, start=1):
            category = r.category if r.category in {"立即处置", "尽快维修", "预防性养护"} else "预防性养护"
            content = (r.content or "加强日常巡检与养护观测").strip()
            raw_comp = (r.component or "桥梁总体").strip()
            clean_comp = re.sub(r"^\d+[\.、\s]+", "", raw_comp).strip() or raw_comp
            recommendation_rows.append([str(index), category, content, clean_comp])
        _add_table(document, ["序号", "建议类别", "建议内容", "病害部位"], recommendation_rows, [600, 1500, 4960, 2300])

        document.add_heading("病害列表", level=2)
        disease_rows = []
        for index, disease in enumerate(published_diseases(state), start=1):
            raw_location = (disease.location or disease.component or "未明确").strip()
            clean_location = re.sub(r"^\d+[\.、\s]+", "", raw_location).strip() or raw_location
            clean_type = (disease.type or "未明确").strip()
            clean_desc = (disease.description or "未明确").strip()

            if disease.is_new is True:
                is_new_str = "是"
                history_str = "上一次定检未发现"
                dev = str(disease.development or "").strip()
                development_str = dev if dev and dev != "未明确" else "无"
            else:
                is_new_str = "否"
                prev = str(disease.previous_state or "").strip()
                if not prev or prev == "未明确":
                    history_str = "上一次定检未发现"
                elif prev == "未发现":
                    history_str = "上一次定检未发现"
                elif prev == "无":
                    history_str = "无"
                elif prev.startswith("上一次") or prev.startswith("上次"):
                    history_str = prev
                elif "已发现" in prev:
                    history_str = "上一次定检已发现，本次基本维持。" if prev == "已发现" else f"上一次定检{prev}"
                else:
                    history_str = f"上一次定检已发现，{prev}"

                dev = str(disease.development or "").strip()
                development_str = dev if dev and dev != "未明确" else "无"

            disease_rows.append([
                str(index),
                clean_location,
                clean_type,
                clean_desc,
                is_new_str,
                history_str,
                development_str,
            ])
        _add_table(document, ["序号", "病害部位", "病害类型", "病害描述", "是否新增", "上一次定检状态", "发展程度"], disease_rows, [500, 1100, 1050, 3450, 800, 1260, 1200])

        self._add_narrative_section(document, "病害成因", state.narrative.causes)
        self._add_narrative_section(document, "处置建议", state.narrative.dispositions or [r.content for r in state.recommendations])
        self._add_narrative_section(document, "安全影响", state.narrative.safety_impacts)
        document.save(path)
        self.audit_semantics(state, path)
        return path

    @classmethod
    def audit_semantics(cls, state: DocumentState, path: Path) -> None:
        """Fail closed when rendered table values diverge from canonical state."""

        document = Document(path)
        if len(document.tables) != 3:
            raise RuntimeError(f"rendered report must contain exactly three tables, got {len(document.tables)}")
        summary = {
            row.cells[0].text.strip(): row.cells[1].text.strip()
            for row in document.tables[0].rows[1:]
            if len(row.cells) >= 2
        }
        expected = {row[0]: row[1] for row in cls._summary_rows(state)}
        mismatches = [
            field
            for field, value in expected.items()
            if summary.get(field, "") != value
        ]
        if mismatches:
            raise RuntimeError("rendered summary semantic mismatch: " + ", ".join(mismatches[:8]))
        recommendation_count = len(document.tables[1].rows) - 1
        disease_count = len(document.tables[2].rows) - 1
        if recommendation_count != len(state.recommendations):
            raise RuntimeError("rendered recommendation row count mismatch")
        if disease_count != len(published_diseases(state)):
            raise RuntimeError("rendered disease row count mismatch")

    @staticmethod
    def _configure_document(document: Document, state: DocumentState) -> None:
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        normal = document.styles["Normal"]
        normal.font.name = FONT
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.font.size = Pt(11)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        for style_name, size, color, before, after in (
            ("Heading 1", 16, BLUE, 18, 10),
            ("Heading 2", 13, BLUE, 14, 7),
            ("Heading 3", 12, DARK_BLUE, 10, 5),
        ):
            style = document.styles[style_name]
            style.font.name = FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            style.font.size = Pt(size)
            style.font.color.rgb = color
            style.font.bold = True
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.space_after = Pt(0)
        _set_run(header.add_run("BridgeEvidence｜城市桥梁定检信息提取"), size=8.5, color=RGBColor(100, 110, 120))
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_run(footer.add_run("第 "), size=8.5, color=RGBColor(100, 110, 120))
        run = footer.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for node in (begin, instruction, separate, text, end):
            run._r.append(node)
        _set_run(footer.add_run(" 页"), size=8.5, color=RGBColor(100, 110, 120))

    @staticmethod
    def _add_masthead(document: Document, state: DocumentState) -> None:
        title = document.add_paragraph()
        title.paragraph_format.space_before = Pt(8)
        title.paragraph_format.space_after = Pt(4)
        _set_run(title.add_run("城市桥梁定检信息提取报告"), size=23, bold=True, color=RGBColor(0, 0, 0))
        subtitle = document.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(14)
        _set_run(subtitle.add_run(str(state.bridge.name.value or "未明确")), size=14, bold=True, color=RGBColor(55, 55, 55))
        for label, value in (
            ("报告日期", state.bridge.report_date.value),
            ("生成系统", "桥鉴（BridgeEvidence）证据闭环智能体"),
            ("证据覆盖", f"{state.validation.evidence_coverage:.1%}"),
        ):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            _set_run(paragraph.add_run(f"{label}："), size=10.5, bold=True)
            _set_run(paragraph.add_run(str(value or "未明确")), size=10.5)

    @staticmethod
    def _summary_rows(state: DocumentState) -> list[list[str]]:
        immediate = sum(r.category == "立即处置" for r in state.recommendations)
        repair = sum(r.category == "尽快维修" for r in state.recommendations)
        preventive = sum(r.category == "预防性养护" for r in state.recommendations)
        values = [
            ("桥梁名称", state.bridge.name.value, "桥梁全称"),
            ("报告日期", state.bridge.report_date.value, "报告出具日期"),
            ("总体评分", state.scores.overall.score.value, "总体技术状况评分"),
            ("总体等级", state.scores.overall.grade.value, "总体技术状况等级"),
            ("上部结构评分", state.scores.superstructure.score.value, "上部结构技术状况评分"),
            ("上部结构等级", state.scores.superstructure.grade.value, "上部结构技术状况等级"),
            ("下部结构评分", state.scores.substructure.score.value, "下部结构技术状况评分"),
            ("下部结构等级", state.scores.substructure.grade.value, "下部结构技术状况等级"),
            ("桥面系评分", state.scores.deck.score.value, "桥面系技术状况评分"),
            ("桥面系等级", state.scores.deck.grade.value, "桥面系技术状况等级"),
            ("上一次总体评分", state.scores.previous.score.value, "上一次定检总体技术状况评分"),
            ("上一次总体等级", state.scores.previous.grade.value, "上一次定检总体技术状况等级"),
            ("病害发展趋势与具体说明", state.narrative.development_trend, "与上一次定检相比的发展趋势"),
            ("总体结论", state.narrative.overall_conclusion, "总体结论的简要说明"),
            ("主要风险点", state.narrative.key_risks, "比较严重或突出的病害描述"),
            ("建议", f"{immediate}条立即处置、{repair}条尽快维修、{preventive}条预防性养护", "建议分类统计"),
        ]
        return [[str(label), "未明确" if value in (None, "") else str(value), note] for label, value, note in values]

    @staticmethod
    def _detailed_conclusion(state: DocumentState) -> list[str]:
        scores = state.scores
        score_text = (
            f"经综合评定，该桥总体技术状况评分 {scores.overall.score.value or '未明确'} 分，"
            f"总体技术状况等级为 {scores.overall.grade.value or '未明确'}。其中，上部结构评分 "
            f"{scores.superstructure.score.value or '未明确'} 分（{scores.superstructure.grade.value or '未明确'}），"
            f"下部结构评分 {scores.substructure.score.value or '未明确'} 分（{scores.substructure.grade.value or '未明确'}），"
            f"桥面系评分 {scores.deck.score.value or '未明确'} 分（{scores.deck.grade.value or '未明确'}）。"
        )
        disease_counts = Counter((d.component, d.type) for d in state.diseases)
        disease_text = "；".join(f"{component}{disease_type}{count}处" for (component, disease_type), count in disease_counts.most_common(18))
        result = [score_text]
        if state.narrative.development_trend:
            result.append(state.narrative.development_trend)
        if disease_text:
            result.append(f"本次共提取 {len(state.diseases)} 条病害记录，主要包括：{disease_text}。")
        if state.narrative.overall_conclusion:
            result.append(state.narrative.overall_conclusion)
        return result

    @staticmethod
    def _add_narrative_section(document: Document, title: str, paragraphs: list[str]) -> None:
        document.add_heading(title, level=2)
        for text in paragraphs or ["未明确。"]:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.25
