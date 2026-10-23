from copy import deepcopy
from dataclasses import replace
import hashlib
import zipfile

from docx import Document
import pytest

from bridge_evidence.contest.evidence_manifest import build_manifest
from bridge_evidence.contest.source_object_dispositions import occurrence_id as ledger_id
from bridge_evidence.contest.source_object_catalog import CatalogError, CatalogLimits, build_source_object_catalog
from bridge_evidence.contest.run_state import json_digest

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = (f'xmlns:w="{W}" xmlns:r="{R}" xmlns:v="urn:schemas-microsoft-com:vml" '
      'xmlns:o="urn:schemas-microsoft-com:office:office"')
OBJECT = ('<w:object><v:shape id="shape1"><v:imagedata r:id="rImg"/></v:shape>'
          '<o:OLEObject Type="Embed" ProgID="Equation.DSMT4" ShapeID="shape1" r:id="rOle"/></w:object>')
BIN = b"original OLE bytes: not decoded or interpreted"
CACHE = b"separate cached image bytes"


def sha(data):
    return hashlib.sha256(data).hexdigest()


def rewrite(path, updates=None, extras=None):
    with zipfile.ZipFile(path) as z:
        files = [(x.filename, z.read(x)) for x in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            value = (updates or {}).get(name, data)
            if value is not None:
                z.writestr(name, value)
        for name, data in (extras or []):
            z.writestr(name, data)


def make(tmp_path, body=None):
    path = tmp_path / "synthetic.docx"
    Document().save(path)
    body = body if body is not None else '<w:p><w:r><w:t>2018 左幅 原文上下文</w:t>' + OBJECT + '</w:r></w:p>'
    with zipfile.ZipFile(path) as z:
        content_types = z.read("[Content_Types].xml").decode()
    content_types = content_types.replace('</Types>',
        '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.oleObject"/>'
        '<Default Extension="png" ContentType="image/png"/></Types>')
    xml = f'<?xml version="1.0" encoding="UTF-8"?><w:document {NS}><w:body>{body}</w:body></w:document>'.encode()
    rels = (f'<Relationships xmlns="{REL}"><Relationship Id="rOle" Type="{R}/oleObject" Target="embeddings/native.bin"/>'
            f'<Relationship Id="rImg" Type="{R}/image" Target="media/cache.png"/></Relationships>').encode()
    rewrite(path, {"word/document.xml": xml, "word/_rels/document.xml.rels": rels,
                   "[Content_Types].xml": content_types.encode()},
            [("word/embeddings/native.bin", BIN), ("word/media/cache.png", CACHE)])
    return path, build_manifest(path)


def mutated(path, manifest, updates=None, extras=None):
    rewrite(path, updates, extras)
    manifest = deepcopy(manifest)
    manifest["document_sha256"] = sha(path.read_bytes())
    return manifest


def rel_replace(path, manifest, old, new):
    with zipfile.ZipFile(path) as z:
        data = z.read("word/_rels/document.xml.rels")
    return mutated(path, manifest, {"word/_rels/document.xml.rels": data.replace(old.encode(), new.encode())})


def test_same_bytes_different_positions_separate_identity_and_ledger_compatibility(tmp_path):
    p = '<w:p><w:r><w:t>正文依据</w:t>' + OBJECT + '</w:r></w:p>'
    path, manifest = make(tmp_path, p + p)
    before = deepcopy(manifest)
    c = build_source_object_catalog(path, manifest)
    assert c["ledger_ready"] and len(c["entries"]) == len(c["occurrences"]) == 4
    # V2 independently inventories each container and its nested OLE instance.
    objects = [entry for entry in c["entries"] if entry["inventory_record"]["kind"] == "object"]
    a, b = [entry["occurrence"] for entry in objects]
    assert a["locator"] != b["locator"] and ledger_id(a) != ledger_id(b)
    assert a["raw_sha256"] == b["raw_sha256"] == sha(OBJECT.encode())
    assert a["native_sha256"] == b["native_sha256"] == sha(BIN)
    assert a["cache_sha256"] == b["cache_sha256"] == sha(CACHE)
    assert [x["object_id"] for x in objects] == [ledger_id(a), ledger_id(b)]
    assert len({x["object_id"] for x in c["entries"]}) == 4
    with zipfile.ZipFile(path) as archive:
        for entry in objects:
            span = entry["raw_xml_span"]
            assert archive.read(span["member"])[span["start"]:span["end"]] == OBJECT.encode()
    assert manifest == before and all(e["representation_state"] == "unread" for e in c["entries"])
    assert not c["all_source_semantics_verified"]


def test_table_object_full_row_and_header_context_not_cropped(tmp_path):
    header = '构件名称及检测单位mm' * 200
    body = ('<w:tbl><w:tr><w:trPr><w:tblHeader/></w:trPr><w:tc><w:p><w:r><w:t>' + header
            + '</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:p><w:r><w:t>右幅桥墩2019</w:t>'
            + OBJECT + '</w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>相邻单元原文</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
    path, manifest = make(tmp_path, body)
    c = build_source_object_catalog(path, manifest)
    refs = c["occurrences"][0]["context_refs"]
    units = {u["id"]: u["text"] for u in manifest["units"]}
    assert c["ledger_ready"] and any(units[r["evidence_id"]] == header for r in refs)
    assert any('右幅桥墩2019 | 相邻单元原文' == units[r["evidence_id"]] for r in refs)
    assert all(r["start"] == 0 and r["end"] == len(units[r["evidence_id"]]) for r in refs)
    assert c["entries"][0]["task_context_sufficient"] is None


def test_unsupported_symbol_preserves_exact_original_bytes(tmp_path):
    symbol = '<w:sym w:font="Symbol" w:char="F061" />'
    body = '<w:p><w:r><w:t>变量语境</w:t>' + symbol + '</w:r></w:p>'
    path, manifest = make(tmp_path, body)
    c = build_source_object_catalog(path, manifest)
    occurrence = c["occurrences"][0]
    assert occurrence["kind"] == "unsupported" and occurrence["raw_sha256"] == sha(symbol.encode())
    assert occurrence["native_sha256"] is None and occurrence["cache_sha256"] is None
    assert c["entries"][0]["inventory_record"] == manifest["unverified_content"][0]
    assert c["entries"][0]["representation_state"] == "unread"


def test_empty_context_is_explicit_blocker_even_when_bytes_resolve(tmp_path):
    path, manifest = make(tmp_path, '<w:p><w:r>' + OBJECT + '</w:r></w:p>')
    c = build_source_object_catalog(path, manifest)
    assert c["occurrences"][0]["context_refs"] == []
    assert c["occurrences"][0]["native_sha256"] == sha(BIN)
    assert "EMPTY_CONTEXT" in c["entries"][0]["diagnostics"] and not c["ledger_ready"]


def test_cache_shape_mismatch_is_not_a_guessed_image_binding(tmp_path):
    path, manifest = make(tmp_path, '<w:p><w:r><w:t>说明</w:t>' + OBJECT.replace('ShapeID="shape1"', 'ShapeID="other"') + '</w:r></w:p>')
    c = build_source_object_catalog(path, manifest)
    assert "CACHE_SHAPE_MISMATCH" in c["entries"][0]["diagnostics"]
    assert c["occurrences"][0]["cache_sha256"] is None and not c["ledger_ready"]


def test_relationship_type_mismatch_does_not_relabel_cached_pixels_as_native(tmp_path):
    path, manifest = make(tmp_path)
    manifest = rel_replace(path, manifest, R + '/oleObject', R + '/image')
    c = build_source_object_catalog(path, manifest)
    assert "RELATIONSHIP_TYPE_MISMATCH" in c["entries"][0]["diagnostics"]
    assert c["occurrences"][0]["native_sha256"] is None


def test_missing_member_and_missing_relation_preserve_each_inventory_entry(tmp_path):
    path, manifest = make(tmp_path)
    manifest = mutated(path, manifest, {"word/embeddings/native.bin": None})
    c = build_source_object_catalog(path, manifest)
    assert c["inventory_count"] == 2 and "MISSING_MEMBER" in c["entries"][0]["diagnostics"]
    manifest = rel_replace(path, manifest, 'Id="rOle"', 'Id="different"')
    c = build_source_object_catalog(path, manifest)
    assert c["inventory_count"] == 2 and "MISSING_RELATIONSHIP" in c["entries"][0]["diagnostics"]


def test_external_ole_link_stays_blocking_without_fetching(tmp_path):
    path, manifest = make(tmp_path)
    manifest = rel_replace(path, manifest, 'Target="embeddings/native.bin"', 'Target="https://example.invalid/private" TargetMode="External"')
    c = build_source_object_catalog(path, manifest)
    assert "EXTERNAL_RELATIONSHIP" in c["entries"][0]["diagnostics"]
    assert not c["entries"][0]["native_parts"] and c["occurrences"][0]["native_sha256"] is None


def test_unsafe_zip_members_duplicate_and_large_packages_fail_closed(tmp_path):
    for name, expected in [('word/../escape.bin', 'UNSAFE_ZIP_MEMBER'), ('word/document.xml', 'DUPLICATE_ZIP_MEMBER')]:
        path, manifest = make(tmp_path)
        with pytest.warns(UserWarning) if name == 'word/document.xml' else __import__('contextlib').nullcontext():
            manifest = mutated(path, manifest, extras=[(name, b'x')])
        with pytest.raises(CatalogError, match=expected):
            build_source_object_catalog(path, manifest)
    path, manifest = make(tmp_path)
    with pytest.raises(CatalogError, match='PACKAGE_LIMIT'):
        build_source_object_catalog(path, manifest, limits=replace(CatalogLimits(), package_bytes=100))
    with pytest.raises(CatalogError, match='MEMBER_LIMIT'):
        build_source_object_catalog(path, manifest, limits=replace(CatalogLimits(), member_bytes=100))


def test_relationship_traversal_and_encoded_external_targets_are_rejected(tmp_path):
    for target in ('../../outside.bin', '%2e%2e/%2e%2e/out.bin', 'file%3A///secret', '/etc/passwd', 'a\\b'):
        path, manifest = make(tmp_path)
        manifest = rel_replace(path, manifest, 'embeddings/native.bin', target)
        c = build_source_object_catalog(path, manifest)
        assert 'UNSAFE_RELATIONSHIP_TARGET' in c['entries'][0]['diagnostics']
        assert not c['ledger_ready']


def test_source_hash_change_never_reuses_old_inventory(tmp_path):
    path, manifest = make(tmp_path)
    rewrite(path, extras=[('changed.txt', b'changed')])
    with pytest.raises(CatalogError, match='SOURCE_HASH_MISMATCH'):
        build_source_object_catalog(path, manifest)


def test_unknown_schema_and_bad_locator_are_explicit_and_not_silently_dropped(tmp_path):
    path, manifest = make(tmp_path)
    wrong = deepcopy(manifest); wrong['parser_version'] = 'future-v99'
    with pytest.raises(CatalogError, match='MANIFEST_SCHEMA_UNSUPPORTED'):
        build_source_object_catalog(path, wrong)
    wrong = deepcopy(manifest); wrong['units'][0]['text'] = 'not in the source'
    with pytest.raises(CatalogError, match='MANIFEST_TEXT_INVENTORY_MISMATCH'):
        build_source_object_catalog(path, wrong)
    wrong = deepcopy(manifest); wrong['unverified_content'][0]['source_path'] = 'word/document.xml#//*[1]'
    c = build_source_object_catalog(path, wrong)
    assert c['inventory_count'] == len(c['entries']) == 2 and c['entries'][0]['occurrence'] is None
    assert 'LOCATOR_NOT_FOUND' in c['entries'][0]['diagnostics'] and c['unlisted_xml_occurrences']
    wrong = deepcopy(manifest); wrong['unverified_content'][0]['future_field'] = 'unknown'
    c = build_source_object_catalog(path, wrong)
    assert c['entries'][0]['inventory_record']['future_field'] == 'unknown'
    assert 'INVENTORY_SCHEMA_UNSUPPORTED' in c['entries'][0]['diagnostics']


def test_omitted_inventory_and_duplicate_inventory_cannot_be_ledger_ready(tmp_path):
    path, manifest = make(tmp_path)
    missing = deepcopy(manifest); missing['unverified_content'] = []
    c = build_source_object_catalog(path, missing)
    assert len(c['unlisted_xml_occurrences']) == 2 and not c['ledger_ready']
    duplicate = deepcopy(manifest); duplicate['unverified_content'] *= 2
    c = build_source_object_catalog(path, duplicate)
    assert len(c['entries']) == 4 and not c['ledger_ready']
    assert 'DUPLICATE_INVENTORY_POSITION' in c['entries'][2]['diagnostics']


def test_body_level_altchunk_uses_adjacent_context_and_internal_relation(tmp_path):
    body = '<w:p><w:r><w:t>附件说明</w:t></w:r></w:p><w:altChunk r:id="rOle"/><w:p><w:r><w:t>后续说明</w:t></w:r></w:p>'
    path, manifest = make(tmp_path, body)
    manifest = rel_replace(path, manifest, R + '/oleObject', R + '/aFChunk')
    c = build_source_object_catalog(path, manifest)
    assert c['ledger_ready'] and c['occurrences'][0]['kind'] == 'unsupported'
    assert len(c['occurrences'][0]['context_refs']) == 2
    assert c['occurrences'][0]['native_sha256'] == sha(BIN)


def test_xml_entities_and_ambiguous_relationship_ids_are_not_accepted(tmp_path):
    path, manifest = make(tmp_path)
    manifest = rel_replace(path, manifest, 'Id="rImg"', 'Id="rOle"')
    c = build_source_object_catalog(path, manifest)
    assert 'AMBIGUOUS_RELATIONSHIP' in c['entries'][0]['diagnostics'] and not c['ledger_ready']
    path, manifest = make(tmp_path)
    with zipfile.ZipFile(path) as z:
        xml = z.read('word/document.xml')
    xml = xml.replace(b'?>', b'?><!DOCTYPE w:document [<!ENTITY x "entity">]>', 1)
    manifest = mutated(path, manifest, {'word/document.xml': xml})
    with pytest.raises(CatalogError, match='XML_DTD_FORBIDDEN'):
        build_source_object_catalog(path, manifest)


def test_v1_legacy_manifest_remains_readable_without_rewriting_it(tmp_path):
    path, current = make(tmp_path, '<w:p><w:r><w:t>Legacy context</w:t><w:sym/></w:r></w:p>')
    # A complete historical V1 inventory has no V2 legacy aliases.
    legacy = deepcopy(current)
    legacy['parser_version'] = 'evidence-manifest-v1'
    for unit in legacy['units']:
        unit.pop('legacy_source_id', None)
    for row in legacy['unverified_content']:
        row['reason'] = 'Embedded/external object or font symbol needs separate interpretation'
    original = deepcopy(legacy)
    catalog = build_source_object_catalog(path, legacy)
    assert catalog['ledger_ready'] and catalog['inventory_count'] == 1
    assert catalog['manifest_parser_version'] == 'evidence-manifest-v1'
    assert catalog['manifest_sha256'] == json_digest(original)
    assert legacy == original


def test_v1_unlisted_nested_ole_blocks_without_upgrading_original_inventory(tmp_path):
    path, current = make(tmp_path)
    legacy = deepcopy(current)
    legacy['parser_version'] = 'evidence-manifest-v1'
    for unit in legacy['units']:
        unit.pop('legacy_source_id', None)
    legacy['unverified_content'] = [row for row in legacy['unverified_content']
                                    if row['kind'] == 'object']
    original = deepcopy(legacy)
    catalog = build_source_object_catalog(path, legacy)
    missing = [row['source_path'] for row in current['unverified_content']
               if row['kind'] == 'OLEObject']
    assert catalog['unlisted_xml_occurrences'] == missing
    assert not catalog['ledger_ready'] and catalog['inventory_count'] == 1
    assert [entry['inventory_record'] for entry in catalog['entries']] == original['unverified_content']
    assert catalog['manifest_parser_version'] == 'evidence-manifest-v1'
    assert catalog['manifest_sha256'] == json_digest(original)
    assert legacy == original


def test_manifest_version_downgrade_cannot_hide_formula_occurrences(tmp_path):
    formula = '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>'
    path, current = make(tmp_path, '<w:p><w:r><w:t>Source context</w:t>' + formula + '</w:r></w:p>')
    forged = deepcopy(current)
    forged['parser_version'] = 'evidence-manifest-v1'
    for unit in forged['units']:
        unit.pop('legacy_source_id', None)
    forged['unverified_content'] = []
    original = deepcopy(forged)
    catalog = build_source_object_catalog(path, forged)
    assert len(current['unverified_content']) == 1
    assert catalog['unlisted_xml_occurrences'] == [current['unverified_content'][0]['source_path']]
    assert not catalog['ledger_ready'] and catalog['inventory_count'] == 0
    assert catalog['entries'] == []
    assert catalog['manifest_parser_version'] == 'evidence-manifest-v1'
    assert catalog['manifest_sha256'] == json_digest(original)
    assert forged == original


def test_v2_keeps_complete_manifest_image_hash_and_nested_object_instances(tmp_path):
    path, manifest = make(tmp_path)
    before = deepcopy(manifest)
    catalog = build_source_object_catalog(path, manifest)
    assert catalog['manifest_parser_version'] == 'evidence-manifest-v2'
    assert catalog['manifest_sha256'] == json_digest(before)
    assert catalog['source_sha256'] == sha(path.read_bytes())
    assert catalog['image_inventory_sha256'] == json_digest(before['images'])
    assert catalog['image_inventory_count'] == len(before['images']) == 1
    assert [entry['inventory_record'] for entry in catalog['entries']] == before['unverified_content']
    assert [entry['inventory_record']['kind'] for entry in catalog['entries']] == ['object', 'OLEObject']
    assert len({entry['object_id'] for entry in catalog['entries']}) == 2
    assert manifest == before and catalog['ledger_ready']
    forged = deepcopy(manifest)
    forged['units'][0]['legacy_source_id'] = 'not-the-original-position'
    with pytest.raises(CatalogError, match='MANIFEST_TEXT_INVENTORY_MISMATCH'):
        build_source_object_catalog(path, forged)


@pytest.mark.parametrize('generic_content_type', [False, True])
def test_v2_nonstandard_story_parts_are_replayed_and_objects_not_omitted(tmp_path, generic_content_type):
    from docx.oxml import OxmlElement

    path = tmp_path / 'custom-story.docx'
    document = Document()
    document.add_paragraph('Body context')
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.text = 'Header context'
    paragraph.add_run()._r.append(OxmlElement('w:sym'))
    document.save(path)
    with zipfile.ZipFile(path) as archive:
        header = archive.read('word/header1.xml')
        types = archive.read('[Content_Types].xml').replace(b'/word/header1.xml', b'/word/unusual.xml')
        relations = archive.read('word/_rels/document.xml.rels').replace(b'"header1.xml"', b'"unusual.xml"')
    if generic_content_type:
        types = types.replace(b'application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml',
                              b'application/xml')
    rewrite(path, {'word/header1.xml': None, '[Content_Types].xml': types,
                   'word/_rels/document.xml.rels': relations}, [('word/unusual.xml', header)])
    manifest = build_manifest(path)
    catalog = build_source_object_catalog(path, manifest)
    assert catalog['ledger_ready'] and catalog['inventory_count'] == 1
    assert catalog['occurrences'][0]['locator'].startswith('word/unusual.xml#')
    assert catalog['manifest_sha256'] == json_digest(manifest)
    forged = deepcopy(manifest)
    forged['units'] = [unit for unit in forged['units'] if unit['context']['story'] != 'header']
    with pytest.raises(CatalogError, match='MANIFEST_TEXT_INVENTORY_MISMATCH'):
        build_source_object_catalog(path, forged)

    # Switching the caller's schema version must not narrow the independent
    # source scan to conventional filenames, even with generic content types.
    forged['parser_version'] = 'evidence-manifest-v1'
    for unit in forged['units']:
        unit.pop('legacy_source_id', None)
    forged['unverified_content'] = []
    original = deepcopy(forged)
    with pytest.raises(CatalogError, match='MANIFEST_TEXT_INVENTORY_MISMATCH'):
        build_source_object_catalog(path, forged)
    assert forged == original


def test_v2_formula_chart_and_diagram_inventory_is_not_dropped(tmp_path):
    formula = '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>'
    chart = '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" r:id="rMissing"/>'
    diagram = '<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" r:dm="rMissing"/>'
    path, manifest = make(tmp_path, '<w:p><w:r><w:t>Source context</w:t>' + formula + chart + diagram + '</w:r></w:p>')
    catalog = build_source_object_catalog(path, manifest)
    assert [entry['inventory_record']['kind'] for entry in catalog['entries']] == ['oMath', 'chart', 'relIds']
    assert catalog['occurrences'][0]['kind'] == 'formula'
    assert len(catalog['occurrences']) == 3
    assert not catalog['ledger_ready'] and not catalog['unlisted_xml_occurrences']
    missing = deepcopy(manifest)
    missing['unverified_content'] = missing['unverified_content'][1:]
    missing_catalog = build_source_object_catalog(path, missing)
    assert len(missing_catalog['unlisted_xml_occurrences']) == 1
    assert not missing_catalog['ledger_ready']
