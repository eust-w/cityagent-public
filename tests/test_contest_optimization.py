
import pytest

from bridge_evidence.contest.candidate import select_candidate
from bridge_evidence.contest.inference_core import build_prompt
from bridge_evidence.contest.long_document import (
    FRAGMENT_INSTRUCTION,
    ChunkTask,
    DocumentSnapshot,
    ExtractionFragment,
    SourceSlice,
    merge_fragments,
    plan_chunks,
    source_metadata,
)
from bridge_evidence.contest.timing import TimingLedger


class Tokenizer:
    def apply_chat_template(self, messages, **kw):
        return "\n".join(m["content"] for m in messages)

    def __call__(self, text, **kw):
        return {"input_ids": list(range(len(text)))}


def chunk_input_budget(snapshot, source_tokens=512):
    """Use the actual prompt/template/metadata overhead plus a small source slot."""
    tokenizer = Tokenizer()
    overhead = []
    for region in snapshot.regions:
        # End offsets use the longest possible digits for this source, while
        # empty text measures overhead independently of its intentionally long body.
        empty = SourceSlice(region.source_id, len(region.text), len(region.text), "", source_metadata(region))
        task = ChunkTask("budget", [empty], [])
        overhead.append(len(build_prompt(
            tokenizer, task.text(), instruction=FRAGMENT_INSTRUCTION,
            max_model_len=1_000_000, max_new_tokens=100,
        )))
    return max(overhead) + source_tokens


def test_oversize_rows_coverage_and_duplicate_instances():
    snapshot = DocumentSnapshot.from_text("重复原文\n重复原文\n" + "尾" * 16000)
    plan = plan_chunks(
        snapshot, Tokenizer(), input_budget=chunk_input_budget(snapshot), output_budget=100, overlap=0
    )
    assert len(plan.tasks) > 10
    assert plan.coverage(snapshot)["complete"]
    assert "".join(s.text for task in plan.tasks for s in task.owned) == "".join(
        r.text for r in snapshot.regions
    )
    assert snapshot.regions[0].source_id != snapshot.regions[1].source_id
    assert all(task.prompt_tokens <= plan.input_budget for task in plan.tasks)


def test_merge_keeps_equal_text_from_distinct_rows_and_rejects_partial():
    snapshot = DocumentSnapshot.from_text("重复\n重复")
    plan = plan_chunks(
        snapshot, Tokenizer(), input_budget=chunk_input_budget(snapshot), output_budget=100, overlap=0
    )
    fragments = []
    for task in plan.tasks:
        rows = [
            {"source_id": r.source_id, "fields": {"病害描述": "重复", "病害部位": "梁"}}
            for r in task.owned
        ]
        fragments.append(ExtractionFragment(
            task.task_id, task.owned_keys, {"桥梁名称": "桥"}, [], rows
        ))
    assert len(merge_fragments(snapshot, plan, fragments)["data"]["病害明细"]) == 2
    with pytest.raises(ValueError, match="coverage"):
        merge_fragments(snapshot, plan, [])


def test_timing_unique_calls_failure_resume(tmp_path):
    ledger = TimingLedger(tmp_path / "ledger.sqlite")
    ledger.start("a", role="extraction")
    ledger.finish("a", 3.0, status="failed")
    ledger.finish("a", 3.0, status="failed")
    assert TimingLedger(tmp_path / "ledger.sqlite").summary()["generation_seconds"] == 3
    with pytest.raises(ValueError):
        ledger.finish("a", 4.0, status="success")


def test_candidate_requires_all_quality_speed_failure_and_repeat_gates():
    baseline = {
        "f1": [0.6, 0.3, 0.25],
        "wall_seconds": 100,
        "failures": 2,
        "conditions": "same",
    }
    good = dict(baseline, wall_seconds=79, f1=[0.6, 0.3, 0.25], repeated=True)
    assert select_candidate(baseline, [good]) == good
    assert select_candidate(baseline, [dict(good, f1=[0.59999, 0.3, 0.25])]) == baseline
    assert select_candidate(baseline, [dict(good, repeated=False)]) == baseline
    assert select_candidate(baseline, [dict(good, failures=3)]) == baseline
    assert select_candidate(baseline, [dict(good, f1=[0.58, 0.3, 0.25])]) == baseline


@pytest.mark.parametrize(
    "config",
    [
        {"batch_size": 3},
        {"input_budget": 0},
        {"workflow": "judge"},
        {"unknown": True},
        {"engine_parameters": {"speculative": True}},
        {"engine_parameters": {"enforce_eager": "false"}},
        {"engine_parameters": {"gpu_memory_utilization": 1.0}},
    ],
)
def test_config_boundaries(config):
    from bridge_evidence.contest.candidate import validate_config

    with pytest.raises(ValueError):
        validate_config(config)


def test_snapshot_retains_merged_cells_headers_and_order(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("桥 2018年")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "构件"
    table.cell(0, 1).text = "病害"
    table.cell(1, 0).text = "梁"
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 1).text = "裂缝"
    table.cell(2, 1).text = "裂缝"
    doc.add_paragraph("尾部 2017年")
    path = tmp_path / "bridge.docx"
    doc.save(path)
    snapshot = DocumentSnapshot.from_docx(path)
    assert snapshot.regions[0].kind == "paragraph"
    assert snapshot.regions[-1].text == "尾部 2017年"
    assert snapshot.sources()["t0r1"].cells[0]["v_merge"] == "restart"
    assert snapshot.sources()["t0r2"].cells[0]["v_merge"] == "continue"
    assert snapshot.sources()["t0r2"].header == ("构件", "病害")


def test_long_header_does_not_destroy_coverage(tmp_path):
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "超长表头" * 2000
    path = tmp_path / "long.docx"
    doc.save(path)
    snapshot = DocumentSnapshot.from_docx(path)
    plan = plan_chunks(snapshot, Tokenizer(), input_budget=chunk_input_budget(snapshot), output_budget=100)
    assert len(plan.tasks) > 1
    assert plan.coverage(snapshot)["complete"]
    assert "".join(part.text for task in plan.tasks for part in task.owned) == "超长表头" * 2000
    assert all(task.prompt_tokens <= plan.input_budget for task in plan.tasks)


def test_interrupted_timing_and_span_idempotency(tmp_path):
    ledger = TimingLedger(tmp_path / "ledger.sqlite")
    ledger.start("lost", role="extraction")
    ledger.recover()
    assert ledger.summary()["failures"] == 1
    ledger.span("load", "model_load", 2)
    ledger.span("load", "model_load", 2)
    assert ledger.summary()["spans"]["model_load"] == 2
    with pytest.raises(ValueError):
        ledger.span("load", "model_load", 3)


@pytest.mark.parametrize("seconds", [float("nan"), float("inf"), -1, 0])
def test_invalid_candidate_time_never_promotes(seconds):
    baseline = {
        "f1": [0.6, 0.3, 0.25],
        "wall_seconds": 100,
        "failures": 2,
        "conditions": "same",
    }
    candidate = {**baseline, "wall_seconds": seconds, "repeated": True}
    assert select_candidate(baseline, [candidate]) == baseline


def test_same_source_instance_with_conflicting_fields_cannot_silently_disappear():
    snapshot = DocumentSnapshot.from_text("裂缝")
    plan = plan_chunks(
        snapshot, Tokenizer(), input_budget=chunk_input_budget(snapshot), output_budget=100, overlap=0
    )
    task = plan.tasks[0]
    rows = [
        {
            "source_id": "p0",
            "source_key": task.owned_keys[0],
            "instance_id": "1",
            "fields": {"病害类型": value},
        }
        for value in ("裂缝", "剥落")
    ]
    fragment = ExtractionFragment(task.task_id, task.owned_keys, {}, [], rows)
    with pytest.raises(ValueError, match="instance"):
        merge_fragments(snapshot, plan, [fragment])
