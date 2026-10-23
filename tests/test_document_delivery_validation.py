"""Synthetic JSON and real local DOCX only; no reports, SDK or model calls."""
from __future__ import annotations

import copy
import hashlib
from pathlib import Path
import sys

from docx import Document
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from bridge_evidence.contest.document_delivery_validation import (
    validate_delivery_data, validate_delivery_docx,
)
from bridge_evidence.contest.docx_writer import (
    ADVICE_FIELDS, BASIC_FIELDS, DEFECT_FIELDS,
    write_numbered_extraction_docx as write_extraction_docx,
)
from bridge_evidence.contest.run_state import json_digest


def report():
    basic = {field: f"合成{field}" for field, _ in BASIC_FIELDS}
    basic.update({"桥梁名称": "合成桥·左幅", "总体评分": 0,
                  "上部结构评分": 90.5, "总体结论": "第一行\n第二行\t保留空白 ",
                  "上一次定检结果": "未知", "上一次总体评分": None, "上一次总体等级": ""})
    advice = [{field: f"合成{field}甲" for field in ADVICE_FIELDS},
              {field: f"合成{field}乙" for field in ADVICE_FIELDS}]
    advice[0]["序号"] = 0
    advice.append(copy.deepcopy(advice[0]))  # Repeated rows must not be deduplicated.
    defects = [{field: f"合成{field}甲" for field in DEFECT_FIELDS},
               {field: f"合成{field}乙" for field in DEFECT_FIELDS}]
    defects[0].update({"是否新增": "未知", "上一次定检状态": None, "发展程度": ""})
    return {"基本信息": basic, "维修建议": advice, "病害明细": defects}


def export(tmp_path, data=None):
    data = report() if data is None else data
    path = tmp_path / "synthetic.docx"
    write_extraction_docx(data, str(path))
    return data, path


def test_complete_readback_checks_values_layout_and_exact_bytes(tmp_path):
    data, path = export(tmp_path)
    before = copy.deepcopy(data)
    audit = validate_delivery_docx(data, path)
    assert data == before
    assert audit == {"schema_version": 1, "validator": "document_delivery_validation_v2_program_row_numbers",
                     "report_sha256": json_digest(data),
                     "docx_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                     "table_data_rows": [len(BASIC_FIELDS) - 2, 3, 2],
                     "cells_checked": (len(BASIC_FIELDS) - 1) * 3 + 4 * 4 + 3 * 7,
                     "paragraphs_checked": 11}
    assert "合成" not in str(audit)


@pytest.mark.parametrize("table_index,column_index", [
    (table, column) for table, width in ((0, 3), (1, len(ADVICE_FIELDS)),
                                        (2, len(DEFECT_FIELDS))) for column in range(width)
])
def test_same_shape_wrong_cell_is_rejected_in_every_column(tmp_path, table_index, column_index):
    data, path = export(tmp_path)
    doc = Document(path)
    doc.tables[table_index].rows[1].cells[column_index].text = "悄悄改写"
    doc.save(path)
    with pytest.raises(ValueError, match=f"table {table_index} row 1 column {column_index}"):
        validate_delivery_docx(data, path)


@pytest.mark.parametrize("field", [name for name, _ in BASIC_FIELDS])
def test_each_supported_basic_field_must_be_present_in_word(tmp_path, field):
    data = report()
    data["基本信息"][field] = "独立合成值"
    data, path = export(tmp_path, data)
    doc = Document(path)
    row = next(row for row in doc.tables[0].rows if row.cells[0].text == field)
    row.cells[1].text = ""
    doc.save(path)
    with pytest.raises(ValueError, match="differs from reviewed delivery"):
        validate_delivery_docx(data, path)


@pytest.mark.parametrize("table_index", [0, 1, 2])
def test_column_header_is_checked(tmp_path, table_index):
    data, path = export(tmp_path)
    doc = Document(path)
    doc.tables[table_index].rows[0].cells[1].text = "错误列名"
    doc.save(path)
    with pytest.raises(ValueError, match="row 0 column 1"):
        validate_delivery_docx(data, path)


@pytest.mark.parametrize("mutation", ["drop", "append", "reorder", "deduplicate"])
def test_row_count_order_and_duplicates_are_checked(tmp_path, mutation):
    data, path = export(tmp_path)
    doc = Document(path)
    table = doc.tables[1]
    if mutation == "drop":
        table._tbl.remove(table.rows[1]._tr)
    elif mutation == "deduplicate":
        table._tbl.remove(table.rows[3]._tr)
    elif mutation == "append":
        table.add_row()
    else:
        first, second = table.rows[1]._tr, table.rows[2]._tr
        first.addprevious(second)
    doc.save(path)
    with pytest.raises(ValueError, match="DOCX table 1"):
        validate_delivery_docx(data, path)


@pytest.mark.parametrize("mutation", ["title", "subtitle", "narrative", "heading", "missing_narrative",
                                     "table_position", "extra_table", "merged_cells", "nested_table"])
def test_content_bearing_layout_and_cell_topology_are_checked(tmp_path, mutation):
    data, path = export(tmp_path)
    doc = Document(path)
    if mutation in {"title", "subtitle", "narrative", "heading"}:
        index = {"title": 0, "subtitle": 1, "narrative": 6, "heading": 2}[mutation]
        doc.paragraphs[index].text = "同样表格行数却丢了审核内容"
    elif mutation == "missing_narrative":
        paragraph = doc.paragraphs[6]._p
        paragraph.getparent().remove(paragraph)
    elif mutation == "table_position":
        doc.paragraphs[3]._p.addprevious(doc.tables[0]._tbl)
    elif mutation == "extra_table":
        doc.add_table(rows=1, cols=1)
    elif mutation == "merged_cells":
        doc.tables[1].cell(1, 0).merge(doc.tables[1].cell(1, 1))
    else:
        doc.tables[2].cell(1, 0).add_table(rows=1, cols=1)
    doc.save(path)
    with pytest.raises(ValueError, match="DOCX"):
        validate_delivery_docx(data, path)


@pytest.mark.parametrize("section", ["基本信息", "维修建议", "病害明细"])
@pytest.mark.parametrize("unknown_value", [None, "", "原文没有提供的字段"])
def test_unexportable_field_is_rejected_even_if_empty(tmp_path, section, unknown_value):
    data = report()
    row = data[section] if section == "基本信息" else data[section][0]
    row["未知列"] = unknown_value
    with pytest.raises(ValueError, match="unsupported fields"):
        validate_delivery_data(data)
    # Schema rejection must happen before even trying to read an output file.
    with pytest.raises(ValueError, match="unsupported fields"):
        validate_delivery_docx(data, tmp_path / "does-not-exist.docx")


@pytest.mark.parametrize("bad_value", [False, True, float("nan"), float("inf"), [], {}, "含\x00空字节", "坏\ud800"])
def test_values_that_would_be_stringified_or_fail_xml_are_rejected(bad_value):
    data = report()
    data["病害明细"][0]["病害描述"] = bad_value
    with pytest.raises(ValueError, match="not an exportable scalar|unexportable XML"):
        validate_delivery_data(data)


@pytest.mark.parametrize("mutation", ["extra_section", "missing_section", "basic_null", "rows_null", "row_scalar"])
def test_silently_skipped_or_defaulted_shapes_are_rejected(mutation):
    data = report()
    if mutation == "extra_section":
        data["知识推断"] = {}
    elif mutation == "missing_section":
        del data["病害明细"]
    elif mutation == "basic_null":
        data["基本信息"] = None
    elif mutation == "rows_null":
        data["维修建议"] = None
    else:
        data["病害明细"].append("writer would silently skip this row")
    with pytest.raises(ValueError):
        validate_delivery_data(data)


@pytest.mark.parametrize("history,subtitle", [
    ({}, "（对比年度信息待确认）"),
    ({"上一次定检结果": "未知"}, "（对比年度信息待确认）"),
    ({"上一次定检结果": "未知", "上一次总体等级": "无"}, "（对比年度信息待确认）"),
    ({"上一次定检结果": "无"}, "（无对比年度）"),
    ({"上一次总体评分": 0}, "（有对比年度）"),
    ({"上一次总体评分": 80, "上一次定检结果": "未知"}, "（有对比年度）"),
])
def test_unknown_history_is_not_completed_and_numeric_zero_is_preserved(tmp_path, history, subtitle):
    data = {"基本信息": {"桥梁名称": "合成桥", **history}, "维修建议": [], "病害明细": []}
    data, path = export(tmp_path, data)
    doc = Document(path)
    assert doc.paragraphs[1].text == subtitle
    validate_delivery_docx(data, path)
    if history.get("上一次定检结果") == "未知":
        row = next(r for r in doc.tables[0].rows if r.cells[0].text == "上一次定检结果")
        row.cells[1].text = "无"
        doc.save(path)
        with pytest.raises(ValueError, match="differs from reviewed delivery"):
            validate_delivery_docx(data, path)


def test_only_line_break_representation_is_normalized(tmp_path):
    data = {"基本信息": {"桥梁名称": "合成桥", "总体结论": " a\r\nb\rc\nd\t \U0001f309"},
            "维修建议": [{"建议内容": "  1.00 mm\t  "}], "病害明细": []}
    data, path = export(tmp_path, data)
    validate_delivery_docx(data, path)
    doc = Document(path)
    assert doc.paragraphs[6].text == " a\n\nb\nc\nd\t \U0001f309"
    doc.tables[1].cell(1, 2).text = "1.00 mm"  # Whitespace is significant.
    doc.save(path)
    with pytest.raises(ValueError, match="table 1"):
        validate_delivery_docx(data, path)


def test_empty_unknown_values_do_not_create_facts_or_rows(tmp_path):
    data = {"基本信息": {"桥梁名称": None, "总体评分": "", "总体结论": None},
            "维修建议": [], "病害明细": []}
    data, path = export(tmp_path, data)
    assert validate_delivery_docx(data, path)["table_data_rows"] == [0, 0, 0]
    doc = Document(path)
    assert doc.paragraphs[0].text == "未知桥梁·信息提取报告"
    assert all(len(table.rows) == 1 for table in doc.tables)


def test_invalid_docx_is_an_explicit_validation_failure(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"synthetic invalid zip")
    with pytest.raises(ValueError, match="Cannot read exported DOCX"):
        validate_delivery_docx(report(), path)
