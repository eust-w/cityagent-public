"""Actual DOCX reconstruction with fixture model output, never real admission.

Only the model response and capability validator are fixtures. These tests exercise
the production parser, pixel inventory, durable ledger, merger, renderer and audit.
"""

import io
import json
from copy import deepcopy
from dataclasses import asdict

import pytest
from docx import Document
from lxml import etree
from PIL import Image
from test_chunked_integration import CharTokenizer, FragmentBackend
from test_contest_visual_evidence import mutate, response
from test_execution_audit import make_record

from bridge_evidence.contest.backends import AccountedBackend
from bridge_evidence.contest.candidate import validate_config
from bridge_evidence.contest.chunked_inference import infer_chunked, prepare_chunk_plan
from bridge_evidence.contest.docx_writer import write_extraction_docx
from bridge_evidence.contest.execution_audit import (
    audit_execution_records,
    audit_timing_records,
    rebuild_chunked_result,
    validate_execution_record,
)
from bridge_evidence.contest.facts import build_fact_graph
from bridge_evidence.contest.long_document import DocumentSnapshot
from bridge_evidence.contest.model_preserving_fusion import preserve_model_report
from bridge_evidence.contest.run_state import file_digest, json_digest
from bridge_evidence.contest.timing import TimingLedger
from bridge_evidence.contest.visual_evidence import (
    execute_visual_jobs,
    plan_visual_jobs,
    rebuild_visual_snapshot,
)
from bridge_evidence.contest.visual_inventory import A, W

MODEL = {"manifest": {"served_model": "fixed-test-model"}, "files": {"base": "a" * 64}}
RUN_IDENTITY = {"options": {"model_manifest": MODEL}}
FINGERPRINT = json_digest(RUN_IDENTITY)
MODEL_SHA = json_digest(MODEL)


@pytest.fixture
def visual_record(tmp_path, monkeypatch):
    from bridge_evidence.contest import vision_capability

    source = tmp_path / "inputs" / "bridge.docx"
    source.parent.mkdir()
    picture = tmp_path / "picture.png"
    Image.new("RGB", (20, 16), "green").save(picture)
    doc = Document()
    doc.add_paragraph("测试桥")
    for _ in range(2):
        doc.add_paragraph().add_run().add_picture(str(picture))
    doc.save(source)
    native = DocumentSnapshot.from_docx(source)
    config = validate_config({"input_budget": 8192, "output_budget": 500})
    native_plan = prepare_chunk_plan(native, CharTokenizer(), config)
    plan = plan_visual_jobs(source, native, MODEL_SHA, {"trial_id": FINGERPRINT})
    work = tmp_path / "work"
    work.mkdir()
    ledger = TimingLedger(work / "timing.sqlite")
    visual = execute_visual_jobs(
        plan, AccountedBackend(response, ledger), work / "visual"
    )
    effective = rebuild_visual_snapshot(
        native, plan, visual["evidence"], visual["calls"], visual["raw"]
    )
    backend = FragmentBackend()

    def extraction(prompt, maximum):
        result = backend(prompt, maximum)
        result["backend_request_id"] = "extract-" + str(backend.calls)
        return result

    accounted = AccountedBackend(extraction, ledger)
    result = infer_chunked(
        effective, CharTokenizer(), lambda p, n: accounted(p, n), config=config
    )
    proof = tmp_path / "capability.json"
    proof.write_text(json.dumps({"model_identity": MODEL}))
    capability = {
        "mode": "same-model",
        "model_identity_sha256": MODEL_SHA,
        "scope": "test fixture, no actual capability qualification",
    }
    monkeypatch.setattr(
        vision_capability, "validate_vision_capability", lambda path, model: capability
    )
    output = tmp_path / "result" / "bridge.docx"
    output.parent.mkdir()
    final, fusion = preserve_model_report(result["data"], label_conventions=True)
    write_extraction_docx(final, str(output))
    record = make_record(source, output, engine="direct")
    record.update(
        schema_version=2,
        run_fingerprint=FINGERPRINT,
        **{
            key: result[key]
            for key in (
                "snapshot",
                "snapshot_sha256",
                "chunk_plan",
                "plan_sha256",
                "fragments",
                "coverage",
                "provenance",
                "basic_observations",
                "basic_ambiguities",
                "conflicts",
                "contributing_call_ids",
            )
        },
    )
    calls = visual["calls"] + result["model_calls"]
    record.update(
        model_calls=calls,
        seconds=sum(call["seconds"] for call in calls),
        final_data=final,
        model_json_sha256=json_digest(result["data"]),
        final_json_sha256=json_digest(final),
        fusion={
            **fusion,
            "mode": "model_preserving",
            "result_sha256": json_digest(final),
        },
        source_evidence={
            "schema_version": 1,
            "source_sha256": file_digest(source),
            "prepared_docx_path": str(source),
            "prepared_docx_sha256": file_digest(source),
            "conversion": None,
        },
        generation_evidence={
            "schema_version": 1,
            "responses": {
                **{
                    item["response"]["call_id"]: item["response"]["text"]
                    for item in visual["raw"]
                },
                **result["generation_evidence"]["responses"],
            },
        },
        visual_evidence={
            "plan": plan,
            **visual,
            "capability": capability,
            "capability_path": str(proof),
            "model_identity_sha256": MODEL_SHA,
            "trial_id": FINGERPRINT,
            "native_chunk_plan": asdict(native_plan),
            "native_plan_sha256": native_plan.digest,
        },
        visual_contributing_call_ids=[call["call_id"] for call in visual["calls"]],
    )
    prefix = {
        "identity": json_digest(
            {
                "run": FINGERPRINT,
                "document": record["document"],
                "visual_plan": plan["plan_sha256"],
                "native_plan": native_plan.digest,
            }
        ),
        "initial_text_tasks": len(native_plan.tasks),
        "initial_visual_jobs": len(plan["jobs"]),
        "total_call_limit": 4 * (len(native_plan.tasks) + len(plan["jobs"])) + 2,
        "used_calls": len(visual["calls"]),
        "elapsed_seconds": visual["spent_seconds"],
    }
    record["shared_budget"] = {
        "prefix": prefix,
        "total_call_limit": prefix["total_call_limit"],
        "text_issued_calls": len(result["model_calls"]),
        "total_issued_calls": len(calls),
        "text_elapsed_seconds": result["wall_seconds"],
        "total_elapsed_seconds": prefix["elapsed_seconds"] + result["wall_seconds"],
    }
    record["workflow"]["stage_records"][1].update(
        model_call_ids=[call["call_id"] for call in calls],
        accepted_model_call_ids=[call["call_id"] for call in calls],
    )
    records = work / "records"
    records.mkdir()
    (work / "run_manifest.json").write_text(
        json.dumps({"identity": RUN_IDENTITY, "run_fingerprint": FINGERPRINT})
    )
    (output.parent / "infer_time.json").write_text(
        json.dumps({"infer_time": record["seconds"]})
    )
    return source, output, record, ledger, records


def check(fixture, *, strong=True, model_sha=MODEL_SHA):
    source, output, record, _, _ = fixture
    return validate_execution_record(
        record,
        expected_document=record["document"],
        expected_run_fingerprint=FINGERPRINT,
        output_path=output,
        profile="legacy",
        require_strong=strong,
        source_path=source,
        expected_model_identity_sha256=model_sha,
    )


def test_visual_sources_rebuild_and_repeated_occurrences_remain_distinct(visual_record):
    _, _, record, _, _ = visual_record
    assert check(visual_record) == []
    derived = [
        r for r in record["snapshot"]["regions"] if r["kind"] == "visual-transcription"
    ]
    assert len(derived) == 2 and len({r["source_id"] for r in derived}) == 2
    assert len(record["visual_contributing_call_ids"]) == 1
    assert rebuild_chunked_result(record)["data"]["病害明细"]
    graph = build_fact_graph(
        record, strict=True, allowed_sha256={record["document"]["sha256"]}
    ).to_dict()
    evidence = [item for item in graph["evidence"] if item.get("upstream_call_ids")]
    assert len(evidence) == 2
    for item in evidence:
        assert set(item["evidence_domains"].values()) == {"visual-transcription"}
        assert item["upstream_call_ids"] == record["visual_contributing_call_ids"]
        assert item["locations"][0]["bbox"] == [0, 0, 1, 1]


def test_empty_native_image_anchor_is_owned_without_inventing_text(visual_record):
    record = visual_record[2]
    empty = [r["source_id"] for r in record["snapshot"]["regions"] if r["text"] == ""]
    assert empty
    owned = [span for task in record["chunk_plan"]["tasks"] for span in task["owned"]]
    assert all(
        any(s["source_id"] == key and s["start"] == s["end"] == 0 for s in owned)
        for key in empty
    )
    assert check(visual_record) == []
    span = next(s for s in owned if s["text"])
    span.update(end=span["start"], text="")
    assert check(visual_record)


@pytest.mark.parametrize("mutation", ["media", "crop", "missing_instance"])
def test_actual_media_and_each_occurrence_are_independently_bound(
    visual_record, mutation
):
    source, _, record, _, _ = visual_record

    def change(values):
        if mutation == "media":
            stream = io.BytesIO()
            Image.new("RGB", (20, 16), "red").save(stream, format="PNG")
            name = next(name for name in values if name.startswith("word/media/"))
            values[name] = stream.getvalue()
        else:
            root = etree.fromstring(values["word/document.xml"])
            image = root.find(f".//{{{A}}}blip")
            if mutation == "crop":
                etree.SubElement(image.getparent(), f"{{{A}}}srcRect", l="1000")
            else:
                image.getparent().remove(image)
            values["word/document.xml"] = etree.tostring(root)

    mutate(source, change)
    # Rehashing the outer binding cannot substitute for source reconstruction.
    record["document"]["sha256"] = file_digest(source)
    record["source_evidence"].update(
        source_sha256=file_digest(source), prepared_docx_sha256=file_digest(source)
    )
    assert any("source" in error for error in check(visual_record))


@pytest.mark.parametrize(
    "mutation",
    ["raw", "call", "main_response", "contributor", "missing_evidence", "unbound_role"],
)
def test_visual_receipt_or_participation_tampering_fails(visual_record, mutation):
    record = visual_record[2]
    envelope = record["visual_evidence"]
    if mutation == "raw":
        envelope["raw"][0]["response"]["text"] += " "
    elif mutation == "call":
        record["model_calls"][0] = {
            **record["model_calls"][0],
            "prompt_sha256": "b" * 64,
        }
    elif mutation == "main_response":
        record["generation_evidence"]["responses"][
            record["visual_contributing_call_ids"][0]
        ] += " "
    elif mutation == "contributor":
        record["visual_contributing_call_ids"] = []
    elif mutation == "missing_evidence":
        envelope["evidence"] = []
    else:
        record.pop("visual_evidence")
    assert check(visual_record)


@pytest.mark.parametrize("model_sha", [None, "d" * 64])
def test_self_declared_model_identity_cannot_qualify(visual_record, model_sha):
    assert any(
        "fixed model identity" in error
        for error in check(visual_record, model_sha=model_sha)
    )


def test_capability_summary_without_actual_proof_is_rejected(visual_record):
    record = visual_record[2]
    record["visual_evidence"].pop("capability_path")
    assert check(visual_record)


def test_visual_role_cannot_downgrade_to_legacy_whole_receipt(visual_record):
    visual_record[2]["schema_version"] = 1
    visual_record[2].pop("visual_evidence")
    assert any("V2/V3 source-bound" in error for error in check(visual_record))


def test_actual_ledger_binds_visual_requests_and_batch_manifest(visual_record):
    source, output, record, ledger, records = visual_record
    (records / "bridge.json").write_text(json.dumps(record))
    result = audit_execution_records(
        source.parent,
        output.parent,
        records,
        FINGERPRINT,
        profile="legacy",
        require_strong=True,
    )
    assert result["passed"], result["errors"]
    with ledger.connect() as db:
        db.execute(
            "UPDATE calls SET receipt=? WHERE id=?",
            (json.dumps({"mock": True}), record["model_calls"][0]["call_id"]),
        )
    assert any(
        "mock ledger" in error
        for error in audit_timing_records(output.parent, records)["errors"]
    )


def test_empty_ocr_or_unreadable_image_cannot_clear_unverified_sources(visual_record):
    source, _, record, _, _ = visual_record
    envelope = record["visual_evidence"]
    # Independent parser rejects even a consistently rehashed empty response.
    raw = envelope["raw"][0]
    body = json.loads(raw["response"]["text"])
    body["text_blocks"] = []
    raw["response"]["text"] = json.dumps(body)
    raw["raw_sha256"] = json_digest(
        {key: value for key, value in raw.items() if key != "raw_sha256"}
    )
    native = DocumentSnapshot.from_docx(source)
    from bridge_evidence.contest.visual_evidence import _call

    calls = [_call(raw, envelope["plan"]["jobs"][0])]
    effective = rebuild_visual_snapshot(native, envelope["plan"], [], calls, [raw])
    assert effective.unverified_content
    assert not [r for r in effective.regions if r.kind.startswith("visual-")]


def test_ole_preview_never_resolves_its_embedded_object(visual_record, tmp_path):
    source, _, _, _, _ = visual_record

    def add_object(values):
        root = etree.fromstring(values["word/document.xml"])
        paragraph = root.find(f".//{{{W}}}p")
        obj = etree.SubElement(paragraph, f"{{{W}}}object")
        preview = root.find(f".//{{{W}}}drawing")
        preview.getparent().remove(preview)
        obj.append(preview)
        etree.SubElement(
            obj,
            "{urn:schemas-microsoft-com:office:office}OLEObject",
            ProgID="Equation.3",
        )
        values["word/document.xml"] = etree.tostring(root)

    mutate(source, add_object)
    native = DocumentSnapshot.from_docx(source)
    plan = plan_visual_jobs(source, native, MODEL_SHA, {"trial_id": FINGERPRINT})
    result = execute_visual_jobs(plan, response, tmp_path / "ole-cache")
    effective = rebuild_visual_snapshot(
        native, plan, result["evidence"], result["calls"], result["raw"]
    )
    assert any(
        item.get("kind") in {"object", "OLEObject"}
        for item in effective.unverified_content
    )
    assert not all(item.get("kind") == "image" for item in effective.unverified_content)


def test_visual_fact_domain_cannot_be_relabelled_native(visual_record):
    record = visual_record[2]
    visual = next(
        r for r in record["snapshot"]["regions"] if r["kind"].startswith("visual-")
    )
    visual["location"]["evidence_domain"] = "native-text"
    with pytest.raises(ValueError, match="typed source"):
        build_fact_graph(
            record, strict=True, allowed_sha256={record["document"]["sha256"]}
        )


def test_visual_photo_cannot_establish_literal_basic_fact(visual_record):
    record = deepcopy(visual_record[2])
    derived = next(
        r for r in record["snapshot"]["regions"] if r["kind"].startswith("visual-")
    )
    derived["kind"] = derived["location"]["evidence_domain"] = "visual-observation"
    source_id = derived["source_id"]
    provenance = next(
        p
        for p in record["provenance"]
        if p["section"] == "病害明细" and p["source_id"] == source_id
    )
    provenance.update(section="基本信息", field="总体评分")
    record["final_data"]["基本信息"]["总体评分"] = "98"
    with pytest.raises(ValueError, match="literal metadata"):
        build_fact_graph(
            record, strict=True, allowed_sha256={record["document"]["sha256"]}
        )


def test_photo_description_reference_preserves_its_observation_domain(visual_record):
    record = deepcopy(visual_record[2])
    derived = next(
        r for r in record["snapshot"]["regions"] if r["kind"].startswith("visual-")
    )
    derived["kind"] = derived["location"]["evidence_domain"] = "visual-observation"
    provenance = next(
        p
        for p in record["provenance"]
        if p["section"] == "病害明细" and p["source_id"] == derived["source_id"]
    )
    provenance["quotes"] = {"病害描述": derived["text"]}
    graph = build_fact_graph(
        record, strict=True, allowed_sha256={record["document"]["sha256"]}
    ).to_dict()
    evidence = next(
        item for item in graph["evidence"] if derived["source_id"] in item["source_ids"]
    )
    assert evidence["quotation_basis"] == "visual-observation"
    assert evidence["evidence_domains"] == {derived["source_id"]: "visual-observation"}
    provenance["history_evidence"] = [{"source_id": derived["source_id"]}]
    with pytest.raises(ValueError, match="temporal"):
        build_fact_graph(
            record, strict=True, allowed_sha256={record["document"]["sha256"]}
        )


@pytest.mark.parametrize(
    "mutation",
    [
        "missing",
        "identity",
        "native_hash",
        "visual_count",
        "limit",
        "call_count",
        "elapsed",
    ],
)
def test_visual_shared_budget_cannot_be_reset_or_expanded(visual_record, mutation):
    record = visual_record[2]
    shared = record["shared_budget"]
    if mutation == "missing":
        record.pop("shared_budget")
    elif mutation == "identity":
        shared["prefix"]["identity"] = "b" * 64
    elif mutation == "native_hash":
        record["visual_evidence"]["native_plan_sha256"] = "b" * 64
    elif mutation == "visual_count":
        shared["prefix"]["initial_visual_jobs"] += 1
        shared["prefix"]["total_call_limit"] += 4
        shared["total_call_limit"] += 4
    elif mutation == "limit":
        shared["total_call_limit"] += 1
    elif mutation == "call_count":
        shared["total_issued_calls"] -= 1
    else:
        shared["prefix"]["elapsed_seconds"] = 0
        shared["total_elapsed_seconds"] = shared["text_elapsed_seconds"]
    assert check(visual_record)


def test_frozen_strategy_text_counts_do_not_exclude_visual_budget(visual_record):
    from bridge_evidence.contest.rsi import StrategyBundle, base_action

    record = visual_record[2]
    shared = record["shared_budget"]
    bundle = StrategyBundle.baseline()
    state = {
        "phase": "schedule",
        "batch_size": 1,
        "remaining_calls": 1,
        "remaining_seconds": 1,
    }
    record.update(
        schema_version=3,
        rsi_bundle_sha256=bundle.digest,
        rsi={
            "bundle": bundle.as_dict(),
            "bundle_sha256": bundle.digest,
            "mode": "frozen",
            "initial_tasks": len(record["chunk_plan"]["tasks"]),
            "issued_calls": shared["text_issued_calls"],
            "spent_seconds": shared["text_elapsed_seconds"],
            "decisions": [
                {
                    "state": state,
                    "state_sha256": json_digest(state),
                    "bundle_sha256": bundle.digest,
                    "action": base_action(state),
                    "fallback": True,
                    "seconds": 0,
                }
            ],
        },
    )
    # A fixture fallback decision proves only the receipt contract, not RSI eligibility.
    assert check(visual_record) == []
    record["rsi"]["issued_calls"] = shared["total_issued_calls"]
    assert any("strategy budget" in error for error in check(visual_record))


def test_repair_audit_independently_rejects_photo_basic_score(monkeypatch):
    from dataclasses import replace

    from test_execution_reconstruction import repair_result

    from bridge_evidence.contest import chunked_inference
    from bridge_evidence.contest.execution_audit import (
        _rebuild_repair,
        _snapshot_and_plan,
    )
    from bridge_evidence.contest.long_document import (
        ExtractionFragment,
        merge_fragments,
    )

    value = repair_result()
    snapshot, plan = _snapshot_and_plan(value)
    merged = merge_fragments(
        snapshot, plan, [ExtractionFragment(**f) for f in value["fragments"]]
    )
    calls = {call["call_id"]: call for call in value["model_calls"]}
    original = snapshot.sources()["p1"]
    snapshot.regions[1] = replace(
        original,
        kind="visual-observation",
        location={**original.location, "evidence_domain": "visual-observation"},
    )
    # Even a compromised online tool filter must not bypass the offline audit.
    monkeypatch.setattr(
        chunked_inference,
        "read_evidence",
        lambda *args: [{"source_id": "p1", "text": original.text}],
    )
    with pytest.raises(ValueError, match="photo observation cannot repair"):
        _rebuild_repair(
            snapshot, merged, calls, value["generation_evidence"]["responses"], value
        )
