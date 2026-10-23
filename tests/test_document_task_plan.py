from __future__ import annotations

import copy
import hashlib
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement

from bridge_evidence.contest.document_task_plan import SCHEMA, build_document_task_plan
from bridge_evidence.contest.evidence_manifest import build_manifest, manifest_from_text


def _manifest(texts, *, styles=None):
    units = []
    for i, text in enumerate(texts):
        units.append({"id": f"original-{i}", "coverage_id": f"coverage-{i}",
                      "text": text, "kind": "paragraph", "style_id": (styles or {}).get(i, ""),
                      "source_path": f"source.xml#paragraph[{i}]",
                      "context": {"story": "body", "preceding_evidence_ids": []}})
    return {"document_sha256": "a" * 64, "units": units, "tables": [],
            "images": [], "unverified_content": [], "limitations": ["Synthetic parsed text only"]}


def _assert_lossless(manifest, plan):
    by_id = {unit["id"]: unit for unit in manifest["units"]}
    expected = list(by_id)
    assigned = [item for task in plan["tasks"] for item in task["primary_evidence_ids"]]
    assert assigned == expected
    assert len(set(assigned)) == len(assigned)
    for task in plan["tasks"]:
        assert task["primary_evidence_ids"] == expected[task["unit_index_start"]:task["unit_index_end"]]
        assert set(task["primary_evidence_ids"]).isdisjoint(task["context_evidence_ids"])
        assert task["grants_reading_coverage"] is False
        for span in task["primary_ranges"] + task["context_references"]:
            original = by_id[span["evidence_id"]]
            assert span["start"] == 0
            assert span["end"] == len(original["text"])
            assert span["source_path"] == original["source_path"]
            for key in ("source_start", "source_end", "table_id", "row_index", "segment_index"):
                if key in original:
                    assert span[key] == original[key]
    assert plan["coverage"]["all_units_assigned_once"] is True
    assert plan["coverage"]["table_families_unbroken"] is True
    assert plan["coverage"]["grants_reading_coverage"] is False
    assert plan["coverage"]["semantic_completeness_verified"] is False
    assert plan["coverage"]["visual_evidence_verified"] is False
    assert plan["global_source_access_required"] is True


def test_styled_chapters_are_contiguous_balanced_navigation_without_source_mutation():
    manifest = _manifest(["封面", "概况", "文字" * 80, "方法", "文字" * 80,
                          "记录", "文字" * 80, "结语", "文字" * 80],
                         styles={1: "Heading1", 3: "heading1", 5: "Heading1", 7: "Heading1"})
    before = copy.deepcopy(manifest)
    plan = build_document_task_plan(manifest)
    assert manifest == before
    assert plan["schema"] == SCHEMA
    assert len(plan["tasks"]) == 4
    assert plan["partition_basis"] == "heading_boundaries"
    assert all(heading["basis"] == "heading_style" for task in plan["tasks"] for heading in task["headings"])
    starts = {task["unit_index_start"] for task in plan["tasks"][1:]}
    assert starts <= {1, 3, 5, 7}
    _assert_lossless(manifest, plan)


def test_regex_chapters_subchapters_and_appendices_preserve_original_titles():
    texts = ["1 概述", "1.1 范围", "第一章 方法", "正文", "第二节 参数", "正文",
             "附录 A 原始记录", "正文", "Appendix B Original records", "正文"]
    manifest = _manifest(texts)
    plan = build_document_task_plan(manifest)
    headings = {h["evidence_id"]: h for task in plan["tasks"] for h in task["headings"]}
    assert headings["original-0"]["level"] == 1
    assert headings["original-1"]["level"] == 2
    assert headings["original-4"]["level"] == 2
    assert headings["original-6"]["title"] == texts[6]
    assert headings["original-8"]["basis"] == "regex_appendix_navigation_candidate"
    _assert_lossless(manifest, plan)


def test_toc_and_body_titles_with_same_words_are_not_coalesced():
    manifest = _manifest(["目  录", "\t1 概述\t1", "2 方法......4", "3 记录\t8",
                          "1 概述", "正文", "2 方法", "正文", "3 记录", "正文"],
                         styles={1: "Heading1", 3: "TOC1"})
    plan = build_document_task_plan(manifest, target_tasks=3)
    heading_ids = {h["evidence_id"] for task in plan["tasks"] for h in task["headings"]}
    assert heading_ids == {"original-4", "original-6", "original-8"}
    assert not {1, 2, 3}.intersection(task["unit_index_start"] for task in plan["tasks"][1:])
    _assert_lossless(manifest, plan)


@pytest.mark.parametrize("text", ["（1）列表项", "1、列表项", "2024-03-01", "1.5 mm", "2 0.05 m",
                                 "表 1 参数", "1 这是一个完整句子。", "不是标题\n跨行文字"])
def test_list_dates_measurements_and_sentences_do_not_become_regex_headings(text):
    plan = build_document_task_plan(_manifest([text]), target_tasks=1)
    assert plan["tasks"][0]["headings"] == []


def test_no_headings_falls_back_to_continuous_units_including_empty_paragraphs():
    manifest = _manifest(["", "散文甲😀", "", "散文乙", "", "散文丙", "", "末尾"])
    plan = build_document_task_plan(manifest)
    assert len(plan["tasks"]) == 4
    assert plan["partition_basis"] == "continuous_atomic_fallback"
    assert plan["coverage"]["empty_unit_count"] == 4
    _assert_lossless(manifest, plan)


def test_zero_units_and_empty_text_unit_are_distinct_scopes():
    manifest = _manifest([])
    plan = build_document_task_plan(manifest)
    assert plan["tasks"] == []
    assert plan["coverage"]["manifest_unit_count"] == 0
    _assert_lossless(manifest, plan)
    text_manifest = manifest_from_text("")
    text_plan = build_document_task_plan(text_manifest)
    assert len(text_plan["tasks"]) == 1
    assert text_plan["tasks"][0]["primary_ranges"][0]["end"] == 0
    _assert_lossless(text_manifest, text_plan)


def test_text_offsets_unicode_and_long_unit_remain_whole():
    manifest = manifest_from_text("  原始😀\r\n\r\n" + "很长😀" * 5000 + "\n最后\t尾部 ")
    plan = build_document_task_plan(manifest)
    _assert_lossless(manifest, plan)
    assert sum(task["original_character_count"] for task in plan["tasks"]) == sum(len(u["text"]) for u in manifest["units"])


def test_nested_table_family_and_row_fragments_are_one_task(tmp_path):
    path = tmp_path / "nested.docx"
    document = Document()
    document.add_heading("一段正文", 1)
    document.add_paragraph("开始")
    table = document.add_table(rows=2, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "嵌套前"
    nested = cell.add_table(rows=2, cols=2)
    nested.cell(0, 0).text = "栏目"
    nested.cell(1, 0).text = "记录"
    cell.paragraphs[-1].text = "嵌套后"
    table.cell(1, 0).text = "后续行"
    document.add_paragraph("单位：原文单位")
    document.add_paragraph("注：原文注释")
    document.add_heading("后续章节", 1)
    document.add_paragraph("末尾")
    document.save(path)
    manifest = build_manifest(path)
    plan = build_document_task_plan(manifest)
    ids = {unit["id"] for unit in manifest["units"] if unit.get("table_id")}
    owners = [task for task in plan["tasks"] if ids.intersection(task["primary_evidence_ids"])]
    assert len(owners) == 1
    assert ids <= set(owners[0]["primary_evidence_ids"])
    assert owners[0]["table_ids"] == [table["id"] for table in manifest["tables"]]
    retained = set(owners[0]["primary_evidence_ids"] + owners[0]["context_evidence_ids"])
    assert {u["id"] for u in manifest["units"] if u["text"].startswith(("单位：", "注："))} <= retained
    _assert_lossless(manifest, plan)


def test_merged_cells_headers_and_units_keep_exact_source_provenance(tmp_path):
    path = tmp_path / "merged.docx"
    document = Document()
    document.add_paragraph("表前文字")
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(0, 2).text = "单位"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "合并原值"
    table.cell(1, 1).text = "记录一"
    table.cell(2, 1).text = "记录二"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph("单位：mm")
    document.add_paragraph("注：必须保留")
    document.add_paragraph("结束")
    document.save(path)
    manifest = build_manifest(path)
    before = copy.deepcopy(manifest)
    plan = build_document_task_plan(manifest)
    task = next(task for task in plan["tasks"] if task["table_ids"])
    assert task["table_context"][0] == manifest["tables"][0]
    assert task["table_context"][0]["header_kind"] == "explicit"
    assert manifest == before
    merged = next(cell for u in manifest["units"] for cell in u.get("cells", []) if cell.get("merged_from_source"))
    assert merged["merged_context_text"] == "合并原值"
    assert any(cell["source_path"] == merged["merged_from_source"] for u in manifest["units"] for cell in u.get("cells", []))
    _assert_lossless(manifest, plan)


def test_single_giant_table_is_not_split_just_to_reach_four_tasks():
    manifest = _manifest(["巨表行" * 1000 for _ in range(30)])
    manifest["tables"] = [{"id": "original-table", "parent_table_id": None,
                            "header_evidence_ids": ["original-0"], "header_kind": "first_row_candidate"}]
    for i, unit in enumerate(manifest["units"]):
        unit.update(table_id="original-table", row_index=i)
    plan = build_document_task_plan(manifest)
    assert len(plan["tasks"]) == 1
    assert plan["tasks"][0]["table_context"][0]["header_kind"] == "first_row_candidate"
    _assert_lossless(manifest, plan)


@pytest.mark.parametrize("blank", ["", " \t\n"])
def test_table_context_crosses_blank_paragraphs_to_units_notes_but_not_next_chapter(blank):
    manifest = _manifest(["封面", "前言", "目录说明", "表格行" * 200, blank, "单位:mm", blank, "注:仅左侧",
                          blank, "2 下一章节", "正文" * 50])
    manifest["tables"] = [{"id": "original-table", "parent_table_id": None,
                            "header_evidence_ids": ["original-3"], "header_kind": "explicit"}]
    manifest["units"][3].update(table_id="original-table", row_index=0)
    plan = build_document_task_plan(manifest)
    owner = next(task for task in plan["tasks"] if "original-3" in task["primary_evidence_ids"])
    assert "original-5" not in owner["primary_evidence_ids"]  # Exercise cross-task context.
    retained = set(owner["primary_evidence_ids"] + owner["context_evidence_ids"])
    assert {f"original-{i}" for i in range(3, 9)} <= retained
    assert "original-9" not in owner["context_evidence_ids"]
    assert "original-10" not in owner["context_evidence_ids"]
    for context in owner["context_references"]:
        if context["evidence_id"] in {"original-4", "original-6", "original-8"}:
            assert context["end"] == len(blank)
    _assert_lossless(manifest, plan)


def test_table_context_preserves_blank_and_annotation_run_before_table():
    manifest = _manifest(["1 原章节", "注:原文限定", "", "单位:mm", "", "表格行" * 200,
                          "2 下一章节", "后续正文"])
    manifest["tables"] = [{"id": "original-table", "parent_table_id": None, "header_evidence_ids": []}]
    manifest["units"][5].update(table_id="original-table", row_index=0)
    plan = build_document_task_plan(manifest)
    owner = next(task for task in plan["tasks"] if "original-5" in task["primary_evidence_ids"])
    retained = set(owner["primary_evidence_ids"] + owner["context_evidence_ids"])
    assert {"original-1", "original-2", "original-3", "original-4"} <= retained
    assert "original-6" not in owner["context_evidence_ids"]
    _assert_lossless(manifest, plan)


def test_table_intervals_include_interleaved_non_table_units():
    manifest = _manifest(["表前", "行一", "夹在表族内的原文", "行二", "表后"])
    manifest["tables"] = [{"id": "outer", "parent_table_id": None, "header_evidence_ids": []},
                          {"id": "nested", "parent_table_id": "outer", "header_evidence_ids": []}]
    manifest["units"][1]["table_id"] = "outer"
    manifest["units"][3]["table_id"] = "nested"
    plan = build_document_task_plan(manifest)
    owner = next(task for task in plan["tasks"] if "original-1" in task["primary_evidence_ids"])
    assert {"original-1", "original-2", "original-3"} <= set(owner["primary_evidence_ids"])
    _assert_lossless(manifest, plan)


def test_original_context_links_and_supplementary_stories_are_preserved():
    manifest = _manifest(["开头", "1 方法", "方法原文", "2 记录", "记录原文", "页眉原文", "页脚原文"])
    manifest["units"][4]["context"].update(section_heading_id="original-3", preceding_evidence_ids=["original-2"])
    manifest["units"][5]["context"]["story"] = "header"
    manifest["units"][6]["context"]["story"] = "footer"
    plan = build_document_task_plan(manifest, target_tasks=3)
    owner = next(task for task in plan["tasks"] if "original-4" in task["primary_evidence_ids"])
    refs = owner["primary_evidence_ids"] + owner["context_evidence_ids"]
    assert {"original-2", "original-3"} <= set(refs)
    _assert_lossless(manifest, plan)


def test_heading_ancestors_are_navigation_even_in_continuous_fallback():
    manifest = _manifest(["1 概况", "1.1 范围", "段落" * 30, "1.2 参数", "段落" * 30,
                          "段落" * 30, "段落" * 30])
    plan = build_document_task_plan(manifest)
    owner = next(task for task in plan["tasks"] if "original-6" in task["primary_evidence_ids"])
    available = set(owner["primary_evidence_ids"] + owner["context_evidence_ids"])
    assert {"original-0", "original-3"} <= available
    _assert_lossless(manifest, plan)


def test_manifest_hash_and_ids_are_deterministic_and_source_config_bound():
    manifest = _manifest(["文字"] * 10)
    plan = build_document_task_plan(manifest)
    assert plan == build_document_task_plan(copy.deepcopy(manifest))
    canonical = json.dumps(manifest, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    assert plan["manifest_sha256"] == hashlib.sha256(canonical.encode()).hexdigest()
    other = copy.deepcopy(manifest)
    other["document_sha256"] = "b" * 64
    different = build_document_task_plan(other)
    assert {task["id"] for task in plan["tasks"]}.isdisjoint(task["id"] for task in different["tasks"])
    changed = build_document_task_plan(manifest, target_tasks=2)
    assert {task["id"] for task in plan["tasks"]}.isdisjoint(task["id"] for task in changed["tasks"])


def test_scope_limits_are_explicit_not_physical_document_completeness():
    manifest = _manifest(["已解析正文"])
    manifest["images"] = [{"id": "original-image", "verification": "unverified"}]
    manifest["unverified_content"] = [{"kind": "external"}]
    plan = build_document_task_plan(manifest)
    assert plan["coverage"]["image_count"] == 1
    assert plan["coverage"]["unverified_content_count"] == 1
    assert plan["source_limitations"] == manifest["limitations"]
    _assert_lossless(manifest, plan)


@pytest.mark.parametrize("target, maximum, expected", [(1, 4, 1), (8, 3, 3), (4, 1, 1), (3, 4, 3)])
def test_positive_task_limits(target, maximum, expected):
    manifest = _manifest(["原文"] * 12)
    plan = build_document_task_plan(manifest, target_tasks=target, max_tasks=maximum)
    assert len(plan["tasks"]) == expected
    _assert_lossless(manifest, plan)


@pytest.mark.parametrize("target, maximum", [(0, 4), (4, 0), (-1, 4), (True, 4), (4, 1.5)])
def test_invalid_task_limits_reject(target, maximum):
    with pytest.raises(ValueError, match="positive integers"):
        build_document_task_plan(_manifest(["原文"]), target_tasks=target, max_tasks=maximum)


@pytest.mark.parametrize("fault", ["duplicate_unit", "missing_text", "unknown_table", "unknown_parent", "parent_cycle", "unknown_header"])
def test_invalid_source_references_fail_explicitly(fault):
    manifest = _manifest(["原文", "正文"])
    if fault == "duplicate_unit":
        manifest["units"][1]["id"] = manifest["units"][0]["id"]
    elif fault == "missing_text":
        manifest["units"][0].pop("text")
    elif fault == "unknown_table":
        manifest["units"][0]["table_id"] = "missing"
    elif fault == "unknown_parent":
        manifest["tables"] = [{"id": "table", "parent_table_id": "missing"}]
    elif fault == "parent_cycle":
        manifest["tables"] = [{"id": "table", "parent_table_id": "table"}]
    else:
        manifest["tables"] = [{"id": "table", "header_evidence_ids": ["missing"]}]
    with pytest.raises(ValueError):
        build_document_task_plan(manifest)


def test_unlocated_table_and_unresolved_context_are_reported_without_invention():
    manifest = _manifest(["正文"])
    manifest["tables"] = [{"id": "empty-table", "parent_table_id": None, "header_evidence_ids": []}]
    manifest["units"][0]["context"]["section_heading_id"] = "absent-source"
    plan = build_document_task_plan(manifest)
    assert plan["coverage"]["unlocated_table_ids"] == ["empty-table"]
    assert plan["tasks"][0]["unresolved_context_references"] == [{"evidence_id": "absent-source", "role": "source_section_heading"}]
    assert "absent-source" not in plan["tasks"][0]["context_evidence_ids"]
    _assert_lossless(manifest, plan)
