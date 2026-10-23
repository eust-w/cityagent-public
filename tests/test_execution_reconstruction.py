"""Structural audit regressions; synthetic responses never prove real execution."""

import hashlib
import json
from copy import deepcopy
from dataclasses import asdict

import pytest
from docx import Document
from test_chunked_integration import CharTokenizer, FragmentBackend
from test_execution_audit import FINGERPRINT, make_record

from bridge_evidence.contest.candidate import validate_config
from bridge_evidence.contest.chunked_inference import infer_chunked
from bridge_evidence.contest.docx_writer import write_extraction_docx
from bridge_evidence.contest.execution_audit import validate_execution_record
from bridge_evidence.contest.long_document import DocumentSnapshot
from bridge_evidence.contest.model_preserving_fusion import preserve_model_report
from bridge_evidence.contest.run_state import file_digest, json_digest


@pytest.fixture
def reconstructed(tmp_path):
    source = tmp_path / "inputs" / "bridge.docx"
    source.parent.mkdir()
    doc = Document()
    doc.add_paragraph("测试桥")
    doc.save(source)
    snapshot = DocumentSnapshot.from_docx(source)
    raw = []
    backend = FragmentBackend()

    def generate(ids, maximum):
        response = backend(ids, maximum)
        raw.append(response["text"])
        return response

    result = infer_chunked(snapshot, CharTokenizer(), generate, config=validate_config(
        {"input_budget": 4000, "output_budget": 500}))
    output = tmp_path / "result" / "bridge.docx"
    output.parent.mkdir()
    final, fusion = preserve_model_report(result["data"], label_conventions=True)
    write_extraction_docx(final, str(output))
    record = make_record(source, output, engine="direct")
    record.update(schema_version=2, **{key: result[key] for key in (
        "snapshot", "chunk_plan", "fragments", "coverage", "provenance",
        "basic_observations", "basic_ambiguities", "conflicts", "contributing_call_ids",
        "snapshot_sha256", "plan_sha256", "model_calls", "seconds")})
    record.update(final_data=final, model_json_sha256=json_digest(result["data"]),
                  final_json_sha256=json_digest(final),
                  fusion={**fusion, "mode": "model_preserving", "result_sha256": json_digest(final)},
                  source_evidence={"schema_version": 1, "source_sha256": file_digest(source),
                                   "prepared_docx_path": str(source), "prepared_docx_sha256": file_digest(source),
                                   "conversion": None},
                  generation_evidence={"schema_version": 1, "responses": dict(zip(
                      [c["call_id"] for c in result["model_calls"]], raw))})
    record["workflow"]["stage_records"][1].update(
        model_call_ids=[c["call_id"] for c in result["model_calls"]],
        accepted_model_call_ids=[c["call_id"] for c in result["model_calls"]])
    return source, output, record


def check(value, *, strong=False):
    source, output, record = value
    return validate_execution_record(record, expected_document=record["document"],
        expected_run_fingerprint=FINGERPRINT, output_path=output, profile="legacy",
        require_strong=strong, source_path=source)


def test_rebuilds_valid_fragment_and_model_preserving_projection(reconstructed):
    assert check(reconstructed, strong=True) == []


def test_self_hashed_final_score_cannot_override_actual_generation(reconstructed):
    _, _, record = reconstructed
    record["final_data"]["基本信息"]["总体评分"] = "100"
    record["final_json_sha256"] = json_digest(record["final_data"])
    record["fusion"]["result_sha256"] = record["final_json_sha256"]
    assert any("projection" in error for error in check(reconstructed))


def test_fragment_and_raw_must_be_independently_reconstructed(reconstructed):
    _, _, record = reconstructed
    record["fragments"][0]["basic"]["总体评分"] = "100"
    assert any("fragment" in error for error in check(reconstructed))


def test_response_hash_is_bound_to_actual_call(reconstructed):
    _, _, record = reconstructed
    call_id = record["model_calls"][0]["call_id"]
    record["generation_evidence"]["responses"][call_id] += " "
    assert any("response" in error for error in check(reconstructed))


def test_delivered_docx_must_match_reconstructed_projection(reconstructed):
    _, output, record = reconstructed
    wrong = deepcopy(record["final_data"])
    wrong["基本信息"]["总体评分"] = "100"
    write_extraction_docx(wrong, str(output))
    record["output"]["sha256"] = file_digest(output)
    assert any("DOCX" in error for error in check(reconstructed))


def test_snapshot_cannot_be_replaced_by_unrelated_source(reconstructed):
    source, _, record = reconstructed
    unrelated = DocumentSnapshot.from_text("unrelated")
    record["snapshot"] = asdict(unrelated)
    record["snapshot_sha256"] = unrelated.digest
    assert any("source" in error for error in check(reconstructed, strong=True))


def test_legacy_v2_compatibility_does_not_grant_strong_admission(reconstructed):
    _, _, record = reconstructed
    record.pop("generation_evidence")
    record.pop("source_evidence")
    assert check(reconstructed) == []
    assert any("fresh" in error for error in check(reconstructed, strong=True))


def test_source_bytes_cannot_change_after_preparation(reconstructed):
    source, _, _ = reconstructed
    doc = Document()
    doc.add_paragraph("different actual input")
    doc.save(source)
    assert any("source" in error for error in check(reconstructed, strong=True))


def test_self_declared_conversion_does_not_substitute_for_actual_doc(reconstructed):
    source, output, record = reconstructed
    old = source.with_suffix(".doc")
    old.write_bytes(b"actual legacy input")
    record["document"].update(relative_path=old.name, sha256=file_digest(old))
    record["source_evidence"].update(source_sha256=file_digest(old), conversion={
        "source_sha256": file_digest(old), "output_sha256": file_digest(source),
        "converter": "mock", "conversion": "docx"})
    assert check((old, output, record), strong=True)


def test_actual_doc_conversion_can_have_different_file_hashes(reconstructed, monkeypatch):
    from bridge_evidence.contest import execution_audit

    source, output, record = reconstructed
    old = source.with_suffix(".doc")
    old.write_bytes(b"legacy input fixture")
    record["document"].update(relative_path=old.name, sha256=file_digest(old))
    record["source_evidence"].update(source_sha256=file_digest(old), conversion={
        "source_sha256": file_digest(old), "output_sha256": file_digest(source),
        "converter": "LibreOffice fixture", "conversion": "docx"})
    calls = []

    def controlled_converter(actual, directory, version):
        calls.append((actual, version))
        return source

    monkeypatch.setattr(execution_audit, "_reconvert_doc", controlled_converter)
    assert check((old, output, record), strong=True) == []
    assert calls == [(old, "LibreOffice fixture")]
    different = source.parent / "different.docx"
    doc = Document()
    doc.add_paragraph("converted evidence is different")
    doc.save(different)
    monkeypatch.setattr(execution_audit, "_reconvert_doc", lambda *args: different)
    assert any("fresh legacy conversion" in error for error in check((old, output, record), strong=True))


def test_fresh_whole_raw_response_can_be_reconstructed(reconstructed):
    from bridge_evidence.contest.text_extract import extract_docx_text

    source, output, record = reconstructed
    record["schema_version"] = 1
    record.pop("generation_evidence")
    model = deepcopy(record["final_data"])
    for row in model["病害明细"]:
        row["序号"] = int(row["序号"])
    # Whole records keep the historical projection without label conventions.
    final, fusion = preserve_model_report(model)
    write_extraction_docx(final, str(output))
    record.update(final_data=final, final_json_sha256=json_digest(final),
                  model_json_sha256=json_digest(model),
                  fusion={**fusion, "mode": "model_preserving", "result_sha256": json_digest(final)})
    record["output"]["sha256"] = file_digest(output)
    record["raw_output"] = json.dumps(model, ensure_ascii=False)
    record["model_calls"][-1]["output_sha256"] = hashlib.sha256(record["raw_output"].encode()).hexdigest()
    record["input_text_sha256"] = hashlib.sha256(extract_docx_text(str(source)).encode()).hexdigest()
    assert check(reconstructed, strong=True) == []
    record["input_text_sha256"] = hashlib.sha256(b"other report text").hexdigest()
    assert any("whole input text" in error for error in check(reconstructed, strong=True))


def repair_result():
    snapshot = DocumentSnapshot.from_text("2020年左幅总体评分80\n2020年左幅总体评分81")
    replies = []

    def generate(ids, maximum):
        if not replies:
            text = "".join(map(chr, ids))
            payload = json.JSONDecoder().raw_decode(text[text.index('{"owned":'):])[0]
            response = {"covered": [s["key"] for s in payload["owned"]],
                        "基本信息": {}, "basic_sources": {}, "维修建议": [], "病害明细": [],
                        "basic_observations": [{"field": "总体评分", "value": str(80 + i),
                            "source_id": s["source_id"], "source_key": s["key"], "quote": s["text"],
                            "scope": {"year": "2020", "span": "左幅"}}
                            for i, s in enumerate(payload["owned"])]}
        elif len(replies) == 1:
            response = {"tool": "read_evidence", "arguments": {"source_ids": ["p1"]}}
        else:
            response = {"resolved": {"总体评分": {"value": "81", "source_id": "p1"}}}
        replies.append(json.dumps(response, ensure_ascii=False))
        return {"text": replies[-1], "finish_reason": "stop", "output_tokens": 100}

    return infer_chunked(snapshot, CharTokenizer(), generate, config=validate_config({
        "workflow": "bounded-react", "input_budget": 5000, "output_budget": 500}))


def test_rebuilds_actual_readonly_repair_without_rewriting_observations():
    from bridge_evidence.contest.execution_audit import rebuild_chunked_result

    value = repair_result()
    rebuilt = rebuild_chunked_result(value)
    assert rebuilt["data"] == value["data"]
    assert [item["value"] for item in rebuilt["basic_observations"]] == ["80", "81"]
    assert rebuilt["repair_evidence"] == value["repair_evidence"]
    value["repair_evidence"][0]["value"] = "99"
    with pytest.raises(ValueError, match="repair evidence"):
        rebuild_chunked_result(value)


def test_legal_repair_requires_actual_selection_response_and_tool_ranges():
    from bridge_evidence.contest.execution_audit import rebuild_chunked_result

    value = repair_result()
    value["model_calls"][-1]["source_ranges"] = []
    with pytest.raises(ValueError, match="repair tool evidence"):
        rebuild_chunked_result(value)


def test_full_document_coverage_cannot_be_replaced_by_owned_range_coverage(reconstructed):
    _, _, record = reconstructed
    record["coverage"]["document_complete"] = False
    assert any("full document coverage" in error for error in check(reconstructed, strong=True))


def test_conversion_failure_returns_audit_error_instead_of_raising(reconstructed, monkeypatch):
    import subprocess

    from bridge_evidence.contest import execution_audit

    source, output, record = reconstructed
    old = source.with_suffix(".doc")
    old.write_bytes(b"legacy fixture")
    record["document"].update(relative_path=old.name, sha256=file_digest(old))
    record["source_evidence"].update(source_sha256=file_digest(old), conversion={
        "source_sha256": file_digest(old), "output_sha256": file_digest(source),
        "converter": "LibreOffice fixture", "conversion": "docx"})

    def failed(*args):
        raise subprocess.TimeoutExpired("soffice", 300)

    monkeypatch.setattr(execution_audit, "_reconvert_doc", failed)
    assert any("source:" in error for error in check((old, output, record), strong=True))


@pytest.mark.parametrize("part", ["source_evidence", "generation_evidence"])
def test_removing_new_binding_cannot_downgrade_formal_receipt(reconstructed, part):
    reconstructed[2].pop(part)
    assert check(reconstructed, strong=True)


def test_unneeded_repair_cannot_modify_an_unconflicted_field(reconstructed):
    record = reconstructed[2]
    record["repair_evidence"] = [{"field": "总体评分", "value": "100"}]
    assert any("no corresponding extraction conflict" in error for error in check(reconstructed))


def test_plan_cannot_change_visible_text_while_preserving_source_ranges(reconstructed):
    record = reconstructed[2]
    record["chunk_plan"]["tasks"][0]["owned"][0]["text"] = "other input source"
    record["plan_sha256"] = json_digest(record["chunk_plan"])
    assert any("task source text" in error for error in check(reconstructed))
