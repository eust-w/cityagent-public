#!/usr/bin/env python3
"""把官方训练集（报告 + 信息提取报告标签）构造成 SFT JSONL 教材。

输入：赛题一_训练集（按年份分目录的 .doc/.docx 报告）
      赛题一_训练集标签（按年份分目录的标签文档）
输出：out_dir/
      sft_train.jsonl   每行 {instruction, input, output, meta}
      stats.json        配对与字段统计
      anomalies.md      未配对/结构异常清单（人工复核用）

只读取官方训练集，绝不触碰测试集。原始文件不做任何修改，
.doc 转换缓存在 out_dir/cache 下。
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
import difflib
from functools import partial
import hashlib
import json
import re
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from bridge_evidence.contest.text_extract import extract_docx_text

YEARS = ["2012年", "2013年", "2014年", "2016年", "2017年", "2018年"]

INSTRUCTION = (
    "你是桥梁定期检测报告分析专家。阅读检测报告全文，提取关键信息，输出三部分JSON："
    "基本信息（桥梁名称、报告日期、总体及各部位技术状况评分与等级、上一次定检结果、"
    "病害发展趋势与具体说明、总体结论、主要风险点、建议统计）、"
    "维修建议（逐条含序号、建议类别、建议内容、病害部位）、"
    "病害明细（逐条含序号、病害部位、病害类型、病害描述、是否新增、上一次定检状态、发展程度）。"
    "所有信息必须来自报告原文，不得编造；封闭字段按规定取值。"
)

# 标签文件名后缀的各种变体
LABEL_SUFFIX = re.compile(
    r"[-—··\s]*(含有?|有|无)对比年度的?(信息提?提?取(报告)?|定检报告|提取报告)"
    r"(?:\d+|\s*\(\d+\))?$"
)
PAIRING_VERSION = "2-strict-unique"
LABEL_PARSER_VERSION = "2-xml-grid-strict"
REPORT_PARSER_VERSION = "2-contest-docx-libreoffice-txt"

ADVICE_FIELDS = ("序号", "建议类别", "建议内容", "病害部位")
DISEASE_FIELDS = ("序号", "病害部位", "病害类型", "病害描述", "是否新增", "上一次定检状态", "发展程度")
# 官方训练标签中已核对的同义表头；不靠列位置猜未知字段。
HEADER_ALIASES = {
    "建议内容(原文复刻)": "建议内容", "建议内容(复刻原文)": "建议内容",
    "位置": "病害部位", "病害位置": "病害部位", "病害种类": "病害类型",
    "病害情况": "病害描述", "病害描述(含照片编号)": "病害描述",
}


class LabelTableError(ValueError):
    """标签结构不足以确定字段映射，必须拒绝该记录。"""


def norm_name(path: Path, is_label: bool) -> str:
    # NFKC 将 Ⅱ/Ⅲ 等转为 ii/iii，而不是删除桥梁编号。
    # 只折叠空白，保留桩号、左右幅和有语义的标点；疑似别名留待复核。
    stem = unicodedata.normalize("NFKC", path.stem).casefold().strip()
    if is_label:
        stem = LABEL_SUFFIX.sub("", stem)
        # 部分标签由“原报告.docx-有对比年度的信息提取报告.docx”命名。
        stem = re.sub(r"\.docx?$", "", stem)
    return re.sub(r"\s+", "", stem)


def plan_pairs(reports: list[Path], labels: list[Path]) -> list[dict]:
    """同一年内只接收唯一的一对一精确配对，模糊相似仅用于人工复核。"""
    report_index: dict[str, list[Path]] = defaultdict(list)
    label_index: dict[str, list[Path]] = defaultdict(list)
    for path in sorted(reports):
        report_index[norm_name(path, False)].append(path)
    for path in sorted(labels):
        label_index[norm_name(path, True)].append(path)

    pairs = []
    for label in sorted(labels):
        key = norm_name(label, True)
        matches = report_index.get(key, [])
        entry = {"label": label.name, "report": None, "normalized_name": key,
                 "status": "unmatched", "reason": "no_exact_match", "candidates": []}
        if not key:
            entry["reason"] = "empty_normalized_name"
        elif len(matches) == 1 and len(label_index[key]) == 1:
            entry.update(report=matches[0].name, status="exact", reason="unique_exact")
        elif matches:
            entry.update(status="ambiguous", reason="normalized_name_collision",
                         candidates=[p.name for p in matches])
        else:
            # 不把最高相似度当作身份确认，也不因文件排序抢占某份报告。
            candidates = difflib.get_close_matches(key, sorted(report_index), n=3, cutoff=0.55)
            entry["candidates"] = [p.name for candidate in candidates for p in report_index[candidate]]
        pairs.append(entry)
    return pairs


def source_sha256(path: Path) -> str:
    with path.open("rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest()


def compare_kind(label_name: str) -> str:
    if "无对比年度" in label_name:
        return "无对比年度"
    if re.search(r"(?:含有?|有)对比年度", label_name):
        return "含有对比年度"
    return "未知"


def read_docx_tables(path: Path) -> list[list[list[str]]]:
    from docx import Document

    doc = Document(path)
    # 保留网格列，文本相同不代表同一单元格。标签字段映射另按 XML 身份处理。
    return [[[_cell_text(cell) for cell in row] for row in _table_grid(tb, f"表{index}")]
            for index, tb in enumerate(doc.tables, 1)]


def _cell_text(cell) -> str:
    return cell.text.strip().replace("\n", " ")


def _field_name(text: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))


def _table_grid(table, name: str) -> list[list]:
    """返回完整网格；python-docx 用同一 XML 单元格展开横向/纵向合并。"""
    from docx.oxml.ns import qn

    width = len(table.columns)
    if not width or not table.rows:
        raise LabelTableError(f"{name}: 空表")
    grid = []
    for index, row in enumerate(table.rows, 1):
        # rows.cells 不返回 gridBefore/gridAfter 省略的格子，不能直接 zip。
        omissions = row._tr.xpath("./w:trPr/w:gridBefore | ./w:trPr/w:gridAfter")
        if any(int(node.get(qn("w:val"), "0")) for node in omissions):
            raise LabelTableError(f"{name}第{index}行: 缺列（省略网格单元格）")
        span = sum(tc.grid_span for tc in row._tr.tc_lst)
        cells = list(row.cells)
        if span != width or len(cells) != width:
            raise LabelTableError(f"{name}第{index}行: 缺列或多列（应有{width}列，实际{span}列）")
        grid.append(cells)
    return grid


def _header_name(text: str) -> str:
    field = _field_name(text)
    if re.fullmatch(r"上一次定检状态\(\d{4}年\)", field):
        return "上一次定检状态"
    return HEADER_ALIASES.get(field, field)


def _label_rows(table, name: str, required: tuple[str, ...] = (), *, basic: bool = False) -> list[list[str]]:
    """按独立 XML 格的顺序映射字段；布局网格可变宽，逻辑列数必须一致。"""
    grid = _table_grid(table, name)
    spans = []
    for column, cell in enumerate(grid[0]):
        if spans and cell._tc is spans[-1][2]._tc:
            start, _, first = spans[-1]
            spans[-1] = (start, column + 1, first)
        else:
            spans.append((column, column + 1, cell))
    raw_headers = [_field_name(_cell_text(cell)) for _, _, cell in spans]
    headers = raw_headers if basic else [_header_name(field) for field in raw_headers]
    # 三份官方训练标签在该完整模板中把病害描述列命名为“具体位置”。
    # 已核对其单元格含现象、尺寸、照片信息；其他模板不作此推断。
    if name == "病害明细" and raw_headers == ["序号", "位置", "病害种类", "具体位置",
                                               "是否新增", "上一次定检状态", "发展程度"]:
        headers[3] = "病害描述"
    # 少量官方标签无“字段/内容/说明”表头，首行就是桥梁名称，不能丢掉它。
    has_header = not basic or headers[0] == "字段"
    if basic and not has_header and headers[0] != "桥梁名称":
        raise LabelTableError(f"{name}: 无法识别字段/内容表头或首行桥梁名称")
    if basic:
        if len(headers) not in (2, 3):
            raise LabelTableError(f"{name}: 应为字段/内容或字段/内容/说明，实际{len(headers)}个字段")
        if has_header and (headers[1] not in {"内容", "值", "示例", "提取内容"}
                           or len(headers) == 3 and headers[2] != "说明"):
            raise LabelTableError(f"{name}: 无法识别内容列或说明列")
        if not has_header:
            headers = ["字段", "内容"] + (["说明"] if len(headers) == 3 else [])
    if any(not field for field in headers):
        raise LabelTableError(f"{name}: 表头含空字段")
    if len(set(headers)) != len(headers):
        raise LabelTableError(f"{name}: 表头含重复字段")
    missing = [field for field in required if field not in headers]
    if missing:
        raise LabelTableError(f"{name}: 缺少字段: {', '.join(missing)}")

    rows = [headers]
    start_row = 1 if has_header else 0
    for index, cells in enumerate(grid[start_row:], start_row + 1):
        if not any(_cell_text(cell) for cell in cells):
            continue
        logical_cells = []
        for cell in cells:
            if not logical_cells or cell._tc is not logical_cells[-1]._tc:
                logical_cells.append(cell)
        # 官方标签存在行间列宽变化，gridSpan 可从末字段移到前一字段。
        # 只折叠同一 XML 格占用的网格位置，绝不按文本去重或截短 zip。
        if len(logical_cells) < len(headers):
            raise LabelTableError(f"{name}第{index}行: 跨字段合并或缺列（应有{len(headers)}个独立字段，实际{len(logical_cells)}个）")
        if len(logical_cells) > len(headers):
            raise LabelTableError(f"{name}第{index}行: 多列（应有{len(headers)}个独立字段，实际{len(logical_cells)}个）")
        if has_header and any(cell._tc is header._tc for cell in logical_cells for _, _, header in spans):
            raise LabelTableError(f"{name}第{index}行: 表头与数据纵向合并")
        rows.append([_cell_text(cell) for cell in logical_cells])
    return rows


def docx_text(path: Path) -> str:
    """复用正式推理提取器，保留独立同值格和段落/表格原顺序。"""
    return extract_docx_text(str(path))


def convert_doc(path: Path, to: str, cache: Path) -> Path | None:
    """用 soffice 把 .doc 转成 docx/txt，缓存在 cache 目录。"""
    cache.mkdir(parents=True, exist_ok=True)
    out = cache / (path.stem + "." + to)
    if out.exists() and out.stat().st_size > 0:
        return out
    try:
        filter_spec = "docx:MS Word 2007 XML" if to == "docx" else "txt:Text (encoded):UTF8"
        # LibreOffice 默认复用已启动的进程；并行转换需独立 profile，防止请求被吞掉。
        with tempfile.TemporaryDirectory(prefix=".soffice-", dir=cache) as profile:
            subprocess.run(
                ["soffice", f"-env:UserInstallation={Path(profile).resolve().as_uri()}",
                 "--headless", "--convert-to", filter_spec, "--outdir", str(cache), str(path)],
                check=True, capture_output=True, timeout=120,
            )
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
        return None
    return out if out.exists() and out.stat().st_size > 0 else None


def report_text(path: Path, cache: Path) -> str | None:
    if path.suffix.lower() == ".docx":
        try:
            return docx_text(path)
        except Exception:
            return None
    # 与 run.py 的 .doc 默认文本路径一致：统一 LibreOffice UTF8 TXT。
    # 不使用 macOS textutil，以免相同数据在不同平台生成不同训练输入。
    out = convert_doc(path, "txt", cache)
    try:
        # 正式推理直接按 UTF-8 读取；保留 BOM/首尾空白，不另做清洗。
        return out.read_text(encoding="utf-8") if out is not None else None
    except (OSError, UnicodeError):
        return None


def soffice_version() -> str:
    try:
        result = subprocess.run(["soffice", "--version"], check=True, capture_output=True,
                                text=True, timeout=15)
        return result.stdout.strip() or "unknown"
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
        return "unknown"


def label_answer(path: Path, cache: Path, *, diagnostics: list[str] | None = None) -> dict | None:
    """把标签文档三张表转成答案 JSON。"""
    real = path
    if path.suffix.lower() == ".doc":
        real = convert_doc(path, "docx", cache)
        if real is None:
            if diagnostics is not None:
                diagnostics.append("标签 .doc 转换失败")
            return None
    try:
        from docx import Document

        tables = Document(real).tables
        if len(tables) != 3:
            raise LabelTableError(f"应有基本信息、维修建议、病害明细三张表，实际{len(tables)}张")
        t1 = _label_rows(tables[0], "基本信息", basic=True)
        t2 = _label_rows(tables[1], "维修建议", ADVICE_FIELDS)
        t3 = _label_rows(tables[2], "病害明细", DISEASE_FIELDS)
        info = {}
        for row in t1[1:]:
            # 三列表格可能有仅“说明”列保留模板提示的空行，不是答案字段。
            if not row[0].strip() and not row[1].strip():
                continue
            key = _field_name(row[0])
            if not key:
                raise LabelTableError("基本信息: 数据含空字段名")
            if key in info:
                raise LabelTableError(f"基本信息: 重复字段: {key}")
            info[key] = row[1]
        if not info:
            raise LabelTableError("基本信息: 无数据")
        recs = [dict(zip(t2[0], row, strict=True)) for row in t2[1:]]
        diseases = [dict(zip(t3[0], row, strict=True)) for row in t3[1:]]
        return {"基本信息": info, "维修建议": recs, "病害明细": diseases}
    except Exception as exc:
        if diagnostics is not None:
            diagnostics.append(str(exc) if isinstance(exc, LabelTableError)
                               else f"标签文档读取失败（{type(exc).__name__}）")
        return None


def extract_training_pair(pair: dict, *, train_root: Path, label_root: Path, cache: Path) -> dict:
    """单对提取不修改共享统计；失败作为结果返回，不中断其他文件对。"""
    year = pair["year"]
    rp = train_root / year / pair["report"]
    lp = label_root / year / pair["label"]
    result = {"pair": pair, "report_path": rp, "label_path": lp, "error": None}
    try:
        report_hash, label_hash = source_sha256(rp), source_sha256(lp)
        # 文件内容相同但文件名不同的两对也隔离缓存，避免并发转换共用产物。
        pair_key = hashlib.sha256(json.dumps([year, rp.name, lp.name], ensure_ascii=False).encode("utf-8")).hexdigest()
        pair_cache = cache / year / report_hash / label_hash / pair_key
        text = report_text(rp, pair_cache / "report")
        label_errors: list[str] = []
        answer = label_answer(lp, pair_cache / "label", diagnostics=label_errors)
        result.update(report_hash=report_hash, label_hash=label_hash,
                      text=text, answer=answer, label_errors=label_errors)
    except Exception as exc:
        # 不把异常中的文档内容写入日志；文件名由主线程附加。
        result["error"] = type(exc).__name__
    return result


def main() -> None:
    ap = argparse.ArgumentParser(description="构建 SFT 训练集 JSONL")
    ap.add_argument("--train-root", type=Path, required=True)
    ap.add_argument("--label-root", type=Path, required=True)
    ap.add_argument("--out-dir", type=Path, required=True)
    ap.add_argument("--workers", type=int, choices=range(1, 5), default=1,
                    help="文档提取并发数（1–4），按原清单顺序写结果；默认1")
    ap.add_argument("--audit-only", action="store_true", help="只审计文件名配对，不读文档正文、不生成 SFT")
    args = ap.parse_args()

    for root in (args.train_root, args.label_root):
        if not root.is_dir():
            ap.error(f"输入目录不存在: {root}")
        if args.out_dir.resolve().is_relative_to(root.resolve()):
            ap.error("输出目录不能位于原始训练报告或标签目录内")
    if args.out_dir.exists() and any(args.out_dir.iterdir()):
        ap.error("输出目录必须为空，请使用新目录保留旧数据与审计记录")

    cache = args.out_dir / "cache"
    args.out_dir.mkdir(parents=True, exist_ok=True)
    anomalies: list[str] = []
    records = 0
    stats: dict = {"pairing_version": PAIRING_VERSION, "label_parser_version": LABEL_PARSER_VERSION,
                   "report_parser_version": REPORT_PARSER_VERSION, "conversion_tool_versions": {},
                   "label_parse_failures": 0, "extraction_failures": 0,
                   "workers": args.workers, "audit_only": args.audit_only,
                   "per_year": {}, "temporal": {}, "compare_kind": {}, "unmatched": 0,
                   "ambiguous": 0, "exact": 0}
    manifest: list[dict] = []

    for year in YEARS:
        train_dir = args.train_root / year
        label_dir = args.label_root / year
        if not train_dir.is_dir() or not label_dir.is_dir():
            anomalies.append(f"## {year}: 目录缺失")
            continue
        reports = sorted(p for p in train_dir.iterdir() if p.is_file() and p.suffix.lower() in {".doc", ".docx"})
        labels = sorted(p for p in label_dir.iterdir() if p.is_file() and p.suffix.lower() in {".doc", ".docx"})
        pairs = plan_pairs(reports, labels)
        year_stats = {"labels": len(labels), "reports": len(reports), "paired": 0,
                      "exact": 0, "ambiguous": 0, "unmatched": 0}
        for pair in pairs:
            pair["year"] = year
            manifest.append(pair)
            stats[pair["status"]] += 1
            year_stats[pair["status"]] += 1
            if pair["status"] != "exact":
                anomalies.append(f"- {year} {pair['status']} ({pair['reason']}): {pair['label']}；候选仅供复核: {pair['candidates']}")
        stats["per_year"][year] = year_stats

    (args.out_dir / "pairing_manifest.json").write_text(
        json.dumps({"pairing_version": PAIRING_VERSION, "pairs": manifest}, ensure_ascii=False, indent=2), encoding="utf-8")

    out_jsonl = args.out_dir / "sft_train.jsonl"
    if not args.audit_only:
        exact_pairs = [pair for pair in manifest if pair["status"] == "exact"]
        if any(Path(pair["report"]).suffix.lower() == ".doc" or Path(pair["label"]).suffix.lower() == ".doc"
               for pair in exact_pairs):
            stats["conversion_tool_versions"]["soffice"] = soffice_version()
            print(f"文档转换工具: {stats['conversion_tool_versions']['soffice']}", flush=True)
        processed = 0
        active_year = None
        extract = partial(extract_training_pair, train_root=args.train_root,
                          label_root=args.label_root, cache=cache)
        print(f"提取并发: {args.workers}；按清单顺序写出", flush=True)
        with out_jsonl.open("x", encoding="utf-8") as fout, ThreadPoolExecutor(max_workers=args.workers) as executor:
            # map 按输入顺序返回；只有主线程更新统计、异常清单和 JSONL。
            for extracted in executor.map(extract, exact_pairs):
                pair = extracted["pair"]
                year = pair["year"]
                if processed % 25 == 0 or year != active_year:
                    print(f"进度: 已检查{processed}/{stats['exact']}对，生成{records}条，标签失败{stats['label_parse_failures']}份；正在处理{year}", flush=True)
                processed += 1
                active_year = year
                lp, rp = extracted["label_path"], extracted["report_path"]
                if extracted["error"] is not None:
                    stats["extraction_failures"] += 1
                    anomalies.append(f"- {year} 文件对提取异常: {rp.name} / {lp.name}；{extracted['error']}")
                    continue
                report_hash, label_hash = extracted["report_hash"], extracted["label_hash"]
                text, answer, label_errors = extracted["text"], extracted["answer"], extracted["label_errors"]
                if answer is None or not answer["病害明细"]:
                    stats["label_parse_failures"] += 1
                    reason = "; ".join(label_errors) or "病害表为空"
                    anomalies.append(f"- {year} 标签解析失败: {lp.name}；{reason}")
                if not text or len(text) < 500:
                    anomalies.append(f"- {year} 报告文本提取失败/过短: {rp.name}")
                    continue
                if answer is None or not answer["病害明细"]:
                    continue

                kind = compare_kind(lp.name)
                stats["compare_kind"][kind] = stats["compare_kind"].get(kind, 0) + 1
                for row in answer["病害明细"]:
                    for col in ("是否新增", "上一次定检状态", "发展程度"):
                        v = str(row.get(col, "∅")).strip() or "∅"
                        bucket = stats["temporal"].setdefault(col, {})
                        bucket[v] = bucket.get(v, 0) + 1

                record = {
                    "instruction": INSTRUCTION,
                    "input": text,
                    "output": json.dumps(answer, ensure_ascii=False),
                    "meta": {
                        "year": year, "report": rp.name, "label": lp.name,
                        "compare_kind": kind, "pair_confidence": "exact",
                        "pairing_version": PAIRING_VERSION,
                        "label_parser_version": LABEL_PARSER_VERSION,
                        "report_parser_version": REPORT_PARSER_VERSION,
                        "conversion_tool_versions": (stats["conversion_tool_versions"]
                                                     if rp.suffix.lower() == ".doc" or lp.suffix.lower() == ".doc" else {}),
                        "report_sha256": report_hash, "label_sha256": label_hash,
                        "input_chars": len(text),
                        "disease_rows": len(answer["病害明细"]),
                    },
                }
                fout.write(json.dumps(record, ensure_ascii=False) + "\n")
                records += 1
                stats["per_year"][year]["paired"] += 1
    for year, counts in stats["per_year"].items():
        print(f"{year}: 标签{counts['labels']} 报告{counts['reports']} 唯一精确{counts['exact']} 歧义{counts['ambiguous']} 待复核{counts['unmatched']} 生成{counts['paired']}", flush=True)

    (args.out_dir / "stats.json").write_text(
        json.dumps({"records": records, **stats}, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out_dir / "anomalies.md").write_text(
        "# 配对与解析异常清单\n\n" + ("\n".join(anomalies) if anomalies else "无异常") + "\n",
        encoding="utf-8")
    if not args.audit_only:
        print(f"\n完成: {records} 条教材 -> {out_jsonl}")
    print(f"统计 -> {args.out_dir / 'stats.json'}  异常清单 -> {args.out_dir / 'anomalies.md'}")


if __name__ == "__main__":
    sys.exit(main())
