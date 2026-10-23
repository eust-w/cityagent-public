from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from bridge_evidence.adapters.word import FormatAdapter
from bridge_evidence.adapters.word_quality import (
    AxisResult,
    DeepAuditResult,
    LegacyDocPreflight,
)


class _TestQualityGate:
    cache_fingerprint = "test-environment"

    @staticmethod
    def preflight(_: Path) -> LegacyDocPreflight:
        return LegacyDocPreflight(
            valid_ole=True,
            has_word_stream=True,
            encrypted=False,
            obfuscated=False,
            has_macros=False,
            embedded_objects=0,
            stream_count=1,
        )

    @staticmethod
    def assess(*_args, **_kwargs) -> DeepAuditResult:
        return DeepAuditResult(
            structure=AxisResult(name="structure", score=1.0),
            visual=AxisResult(name="visual", score=1.0),
        )


def _adapter(tmp_path: Path) -> FormatAdapter:
    return FormatAdapter(tmp_path / "cache", quality_gate=_TestQualityGate())


def _write_docx(path: Path, text: str, *, with_table: bool = False) -> Path:
    document = Document()
    document.add_paragraph(text)
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "指标"
        table.cell(0, 1).text = "评分"
        table.cell(1, 0).text = "BCI"
        table.cell(1, 1).text = "88.82"
    document.save(path)
    return path


def _source_text() -> str:
    return (
        "桥梁定期检查报告。主梁存在裂缝，支座存在剪切变形，建议尽快维修。"
        "桥面系评分88.82，上部结构评分90.15，下部结构评分91.36，总体等级B级。"
    ) * 8


def test_native_docx_is_structurally_audited(tmp_path: Path) -> None:
    source = _write_docx(tmp_path / "native.docx", "桥梁检测结论", with_table=True)

    adapted = _adapter(tmp_path).adapt(source)

    assert adapted.docx_path == source.resolve()
    assert adapted.converted is False
    assert adapted.audit.status == "native"
    assert adapted.audit.accepted is True
    assert adapted.audit.tables == 1


def test_verified_conversion_uses_versioned_audited_cache(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy-word-placeholder")
    reference = _source_text()
    adapter = _adapter(tmp_path)
    calls = 0

    def convert(_: Path, temp_dir: Path) -> Path:
        nonlocal calls
        calls += 1
        return _write_docx(temp_dir / "legacy.docx", reference, with_table=True)

    monkeypatch.setattr(adapter, "_extract_legacy_text", lambda *_: (reference, "test-reader"))
    monkeypatch.setattr(adapter, "_convert_with_libreoffice", convert)
    monkeypatch.setattr(adapter, "_convert_with_wv", lambda *_: pytest.fail("fallback was not expected"))

    first = adapter.adapt(source)
    second = adapter.adapt(source)

    assert calls == 1
    assert first.docx_path == second.docx_path
    assert first.docx_path.name.startswith("v5-")
    assert first.audit.converter == "libreoffice"
    assert first.audit.status == "verified"
    assert first.audit.text_coverage == 1.0
    assert first.audit.numeric_coverage == 1.0
    assert first.audit.requires_review is False
    assert first.docx_path.with_suffix(".conversion.json").exists()


def test_silent_truncation_is_rejected_and_falls_back(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "truncated.doc"
    source.write_bytes(b"legacy-word-placeholder")
    reference = _source_text()
    adapter = _adapter(tmp_path)

    monkeypatch.setattr(adapter, "_extract_legacy_text", lambda *_: (reference, "test-reader"))
    monkeypatch.setattr(
        adapter,
        "_convert_with_libreoffice",
        lambda _path, temp_dir: _write_docx(temp_dir / "truncated.docx", "桥梁报告"),
    )
    monkeypatch.setattr(
        adapter,
        "_convert_with_wv",
        lambda _path, temp_dir: _write_docx(temp_dir / "wv-converted.docx", reference),
    )

    adapted = adapter.adapt(source)

    assert adapted.audit.converter == "wvware-pandoc"
    assert adapted.audit.accepted is True
    assert adapted.audit.requires_review is True
    assert "fallback_converter_may_change_layout_or_media" in adapted.audit.warnings


def test_all_content_losing_candidates_fail_closed(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "broken.doc"
    source.write_bytes(b"legacy-word-placeholder")
    reference = _source_text()
    adapter = _adapter(tmp_path)

    monkeypatch.setattr(adapter, "_extract_legacy_text", lambda *_: (reference, "test-reader"))
    monkeypatch.setattr(
        adapter,
        "_convert_with_libreoffice",
        lambda _path, temp_dir: _write_docx(temp_dir / "broken.docx", "仅剩封面"),
    )
    monkeypatch.setattr(
        adapter,
        "_convert_with_wv",
        lambda _path, temp_dir: _write_docx(temp_dir / "wv-converted.docx", "仍然只有封面"),
    )

    with pytest.raises(RuntimeError, match="failed fidelity gate"):
        adapter.adapt(source)


def test_wv_image_markup_is_not_counted_as_document_text() -> None:
    raw = (
        '<img alt="0x01 graphic" src="StrangeNoGraphicData"><br>'
        'the jpg file'
        '<img alt="0x01 graphic" src="StrangeNoGraphicData"><br>'
        'the png file'
    )

    cleaned = FormatAdapter._clean_reference_text(raw)

    assert FormatAdapter._normalize_text(cleaned) == "thejpgfilethepngfile"


def test_numeric_tokens_preserve_ooxml_table_cell_boundaries(tmp_path: Path) -> None:
    path = tmp_path / "numeric-table.docx"
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "15"
    table.cell(0, 1).text = "02"
    table.cell(0, 2).text = "31"
    document.save(path)

    stats = FormatAdapter._inspect_docx(path)

    assert FormatAdapter._number_tokens(stats.text) == {"15", "02", "31"}


def test_invalid_binary_doc_is_rejected_before_conversion(tmp_path: Path, monkeypatch) -> None:
    def unexpected_converter_lookup():
        pytest.fail("Invalid binary input must be rejected before locating LibreOffice")

    monkeypatch.setattr("bridge_evidence.adapters.word_quality._find_soffice", unexpected_converter_lookup)
    source = tmp_path / "not-word.doc"
    source.write_bytes(b"not-an-ole2-word-document")

    with pytest.raises(ValueError, match="OLE2"):
        FormatAdapter(tmp_path / "cache").adapt(source)
