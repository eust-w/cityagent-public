"""Pixel-preserving supported picture views and strict metadata boundaries."""

import base64
import hashlib
import io
import json
import zipfile
from copy import deepcopy

import pytest
from docx import Document
from lxml import etree
from PIL import Image

from bridge_evidence.contest.long_document import DocumentSnapshot
from bridge_evidence.contest.visual_evidence import (
    build_visual_request,
    execute_visual_jobs,
    plan_visual_jobs,
    rebuild_visual_snapshot,
)
from bridge_evidence.contest.visual_inventory import A
from bridge_evidence.contest.visual_views import VIEW_VERSION, original_bbox


def _document(tmp_path, edit, *, repeat=1):
    image = tmp_path / "pixels.png"
    raster = Image.new("RGB", (8, 6))
    raster.putdata([(x * 29, y * 37, (x + y) * 13) for y in range(6) for x in range(8)])
    raster.save(image)
    document = Document()
    document.add_paragraph("2024 桥梁图片证据")
    for _ in range(repeat):
        document.add_picture(str(image))
    path = tmp_path / "report.docx"
    document.save(path)
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = etree.fromstring(entries["word/document.xml"])
    pictures = root.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"
    )
    for index, picture in enumerate(pictures):
        edit(picture, index)
    entries["word/document.xml"] = etree.tostring(root)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return path, image.read_bytes()


def _plan(path, **config):
    native = DocumentSnapshot.from_docx(path)
    return plan_visual_jobs(
        path, native, "a" * 64, {"trial_id": "supported-views", **config}
    )


def _line(picture, attributes=None, children=("noFill",)):
    shape = picture.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/picture}spPr"
    )
    line = etree.SubElement(shape, f"{{{A}}}ln", attributes or {})
    for tag in children:
        etree.SubElement(line, f"{{{A}}}{tag}")
    return line


@pytest.mark.parametrize("attributes", [{}, {"w": "9525"}, {"w": "0"}])
def test_exact_no_fill_line_preserves_original_pixels_and_occurrences(
    tmp_path, attributes
):
    path, data = _document(
        tmp_path, lambda picture, _: _line(picture, attributes), repeat=2
    )
    plan = _plan(path)
    assert not plan["unresolved"]
    assert len(plan["jobs"]) == 1
    job = plan["jobs"][0]
    assert len(job["bindings"]) == 2
    request = build_visual_request(plan, job)
    assert base64.b64decode(request["image"]["base64"]) == data
    assert request["image"]["sha256"] == hashlib.sha256(data).hexdigest()


@pytest.mark.parametrize(
    "case",
    [
        "solid",
        "extra-child",
        "extra-attribute",
        "negative-width",
        "bad-width",
        "leaf-attribute",
        "leaf-child",
        "text",
        "foreign-namespace",
    ],
)
def test_only_exact_no_fill_line_is_accepted(tmp_path, case):
    def edit(picture, _):
        line = _line(picture)
        if case == "solid":
            line[0].tag = f"{{{A}}}solidFill"
        elif case == "extra-child":
            etree.SubElement(line, f"{{{A}}}prstDash", val="solid")
        elif case == "extra-attribute":
            line.set("cap", "flat")
        elif case == "negative-width":
            line.set("w", "-1")
        elif case == "bad-width":
            line.set("w", "nan")
        elif case == "leaf-attribute":
            line[0].set("val", "true")
        elif case == "leaf-child":
            etree.SubElement(line[0], f"{{{A}}}extLst")
        elif case == "text":
            line[0].text = "unknown"
        elif case == "foreign-namespace":
            line[0].tag = "{urn:unknown}noFill"

    path, _ = _document(tmp_path, edit)
    plan = _plan(path)
    assert not plan["jobs"]
    assert len(plan["unresolved"]) == 1


def _transform(picture, *, crop=None, rotation=0, flip_h=False, flip_v=False):
    fill = picture.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/picture}blipFill"
    )
    if crop is not None:
        fill.insert(1, etree.Element(f"{{{A}}}srcRect", **crop))
    xfrm = picture.find(f".//{{{A}}}xfrm")
    xfrm.set("rot", str(rotation))
    xfrm.set("flipH", "1" if flip_h else "0")
    xfrm.set("flipV", "1" if flip_v else "0")


@pytest.mark.parametrize("turns", [0, 1, 2, 3])
@pytest.mark.parametrize(
    "flips", [(False, False), (True, False), (False, True), (True, True)]
)
def test_positive_crop_quarter_turns_and_flips_preserve_exact_visible_pixels(
    tmp_path, turns, flips
):
    path, original = _document(
        tmp_path,
        lambda picture, _: _transform(
            picture,
            crop={"l": "25000", "r": "12500", "t": "50000"},
            rotation=turns * 5400000,
            flip_h=flips[0],
            flip_v=flips[1],
        ),
    )
    # Existing full-raster callers must never silently acquire new geometry.
    assert _plan(path)["unresolved"]
    plan = _plan(path, allow_display_transforms=True)
    assert not plan["unresolved"]
    assert plan["identity"]["view_validation_version"] == VIEW_VERSION
    job = plan["jobs"][0]
    assert job["view"]["source_pixel_box"] == [2, 3, 7, 6]
    assert (
        job["view"]["original_encoded_sha256"] == hashlib.sha256(original).hexdigest()
    )
    data = base64.b64decode(build_visual_request(plan, job)["image"]["base64"])
    actual = Image.open(io.BytesIO(data))
    source = Image.open(io.BytesIO(original)).convert("RGBA")
    # Independently derive each output pixel from the recorded edge inverse;
    # this also checks the inverse is correct after every flip/rotation ordering.
    a, b, c, d, e, f = job["view"]["view_to_original_pixel_edges"]
    for y in range(actual.height):
        for x in range(actual.width):
            sx = int(a * (x + 0.5) + b * (y + 0.5) + c)
            sy = int(d * (x + 0.5) + e * (y + 0.5) + f)
            assert actual.getpixel((x, y)) == source.getpixel((sx, sy))
    assert original_bbox(job["view"], [0, 0, 1, 1]) == [0.25, 0.5, 0.875, 1]
    assert (
        build_visual_request(plan, job)["image"]["sha256"]
        == hashlib.sha256(data).hexdigest()
    )


def test_clockwise_rotation_is_visually_clockwise(tmp_path):
    path, original = _document(
        tmp_path, lambda picture, _: _transform(picture, rotation=5400000)
    )
    plan = _plan(path, allow_display_transforms=True)
    image = Image.open(
        io.BytesIO(
            base64.b64decode(
                build_visual_request(plan, plan["jobs"][0])["image"]["base64"]
            )
        )
    )
    source = Image.open(io.BytesIO(original)).convert("RGBA")
    assert image.size == (6, 8)
    assert image.getpixel((0, 0)) == source.getpixel((0, 5))
    assert image.getpixel((5, 7)) == source.getpixel((7, 0))


def test_fractional_pixel_crop_retains_boundary_pixels_and_exact_crop_metadata(
    tmp_path,
):
    path, _ = _document(
        tmp_path,
        lambda picture, _: _transform(
            picture, crop={"l": "32004", "r": "21525", "t": "1060"}
        ),
    )
    job = _plan(path, allow_display_transforms=True)["jobs"][0]
    assert job["view"]["source_pixel_box"] == [2, 0, 7, 6]
    assert job["view"]["transform"]["crop_100000"] == {
        "l": 32004,
        "r": 21525,
        "t": 1060,
        "b": 0,
    }
    assert job["view"]["crop_rasterization"] == "all-intersecting-source-pixels-v1"


def test_empty_stretch_and_no_shape_fill_preserve_visible_raster(tmp_path):
    def edit(picture, _):
        _transform(picture, crop={"l": "25000"})
        stretch = picture.find(f".//{{{A}}}stretch")
        stretch.remove(stretch[0])
        shape = picture.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/picture}spPr"
        )
        shape.append(etree.Element(f"{{{A}}}noFill"))
        _line(picture)

    path, original = _document(tmp_path, edit)
    plan = _plan(path, allow_display_transforms=True)
    assert not plan["unresolved"]
    raster = Image.open(
        io.BytesIO(
            base64.b64decode(
                build_visual_request(plan, plan["jobs"][0])["image"]["base64"]
            )
        )
    )
    assert (
        raster.tobytes()
        == Image.open(io.BytesIO(original)).convert("RGBA").crop((2, 0, 8, 6)).tobytes()
    )


def test_same_blob_different_views_do_not_share_model_job(tmp_path):
    path, _ = _document(
        tmp_path,
        lambda picture, index: _transform(picture, rotation=(index % 2) * 5400000),
        repeat=3,
    )
    plan = _plan(path, allow_display_transforms=True)
    assert len(plan["jobs"]) == 2
    assert sorted(len(job["bindings"]) for job in plan["jobs"]) == [1, 2]
    assert len({job["blob"]["sha256"] for job in plan["jobs"]}) == 1
    assert len({job["view_sha256"] for job in plan["jobs"]}) == 2


@pytest.mark.parametrize(
    "case",
    [
        "negative-crop",
        "over-crop",
        "percent",
        "empty-visible",
        "non-quarter",
        "bad-flip",
        "duplicate-crop",
        "unknown-transform",
        "crop-child",
        "fill-offset",
        "foreign-crop",
        "picture-effect",
        "shape-extension",
        "shape-fill-attribute",
        "geometry-adjustment",
        "picture-style",
    ],
)
def test_unsupported_display_transforms_still_fail_closed(tmp_path, case):
    def edit(picture, _):
        _transform(picture, crop={"l": "25000"}, rotation=5400000)
        crop = picture.find(f".//{{{A}}}srcRect")
        transform = picture.find(f".//{{{A}}}xfrm")
        if case == "negative-crop":
            crop.set("l", "-1")
        elif case == "over-crop":
            crop.set("l", "100001")
        elif case == "percent":
            crop.set("l", "25%")
        elif case == "empty-visible":
            crop.set("r", "75000")
        elif case == "non-quarter":
            transform.set("rot", "60000")
        elif case == "bad-flip":
            transform.set("flipH", "yes")
        elif case == "duplicate-crop":
            crop.getparent().append(deepcopy(crop))
        elif case == "unknown-transform":
            transform.set("unknown", "1")
        elif case == "crop-child":
            etree.SubElement(crop, f"{{{A}}}extLst")
        elif case == "fill-offset":
            picture.find(f".//{{{A}}}fillRect").set("t", "1000")
        elif case == "foreign-crop":
            crop.tag = "{urn:unknown}srcRect"
        elif case == "picture-effect":
            etree.SubElement(transform.getparent(), f"{{{A}}}effectLst")
        elif case == "shape-extension":
            extension = etree.SubElement(transform.getparent(), f"{{{A}}}extLst")
            etree.SubElement(extension, f"{{{A}}}ext", uri="unknown")
        elif case == "shape-fill-attribute":
            etree.SubElement(transform.getparent(), f"{{{A}}}noFill", unknown="true")
        elif case == "geometry-adjustment":
            adjustments = etree.SubElement(
                picture.find(f".//{{{A}}}prstGeom"), f"{{{A}}}avLst"
            )
            etree.SubElement(
                adjustments,
                f"{{{A}}}gd",
                name="adjustment",
                fmla="val 50000",
            )
        elif case == "picture-style":
            etree.SubElement(
                picture,
                "{http://schemas.openxmlformats.org/drawingml/2006/picture}style",
            )

    path, _ = _document(tmp_path, edit)
    plan = _plan(path, allow_display_transforms=True)
    assert not plan["jobs"]
    assert len(plan["unresolved"]) == 1


def _response(request, maximum):
    from bridge_evidence.contest.rsi_runtime import REQUEST_CALL_ID

    call_id = REQUEST_CALL_ID.get()
    return {
        "text": json.dumps(
            {
                "schema_version": 2,
                "job_id": request["job_id"],
                "outcome": "complete",
                "bbox_space": "normalized_1000",
                "text_blocks": [
                    {"id": "t1", "bbox": [0, 0, 1000, 1000], "text": "test pixels"}
                ],
                "observations": [],
                "unresolved_regions": [],
            }
        ),
        "finish_reason": "stop",
        "call_id": call_id,
        "backend_request_id": "backend-" + call_id,
        "prompt_tokens": 100,
        "output_tokens": 100,
        "latency_seconds": 0.1,
    }


@pytest.mark.parametrize(
    "wrapper",
    ["json-fence", "plain-fence", "prose", "trailing", "duplicate", "nonfinite"],
)
def test_visual_pipeline_accepts_only_complete_json_wrapper_and_keeps_raw(
    tmp_path, wrapper
):
    path, _ = _document(tmp_path, lambda picture, _: None)
    native = DocumentSnapshot.from_docx(path)
    plan = _plan(path)

    def generate(request, maximum):
        receipt = _response(request, maximum)
        text = receipt["text"]
        if wrapper == "json-fence":
            text = "```json\n" + text + "\n```"
        elif wrapper == "plain-fence":
            text = "```\n" + text + "\n```"
        elif wrapper == "prose":
            text = "Here is the result:\n```json\n" + text + "\n```"
        elif wrapper == "trailing":
            text = "```json\n" + text[:-1] + ",}\n```"
        elif wrapper == "duplicate":
            text = "```json\n" + text[:-1] + ',"outcome":"complete"}\n```'
        elif wrapper == "nonfinite":
            text = (
                "```json\n"
                + text.replace("[0, 0, 1000, 1000]", "[0, 0, 1e400, 1000]")
                + "\n```"
            )
        receipt["text"] = text
        return receipt

    result = execute_visual_jobs(plan, generate, tmp_path / "cache")
    assert "```" in result["raw"][0]["response"]["text"]
    rebuilt = rebuild_visual_snapshot(
        native, plan, result["evidence"], result["calls"], result["raw"]
    )
    if wrapper in {"json-fence", "plain-fence"}:
        assert not result["errors"] and not rebuilt.unverified_content
    else:
        assert (
            result["errors"] and rebuilt.unverified_content == native.unverified_content
        )


def test_rebuild_binds_transformed_evidence_to_original_coordinates_and_rejects_mutation(
    tmp_path,
):
    from bridge_evidence.contest.run_state import json_digest

    path, original = _document(
        tmp_path,
        lambda picture, _: _transform(picture, crop={"l": "25000"}, rotation=5400000),
    )
    native = DocumentSnapshot.from_docx(path)
    plan = _plan(path, allow_display_transforms=True)
    result = execute_visual_jobs(plan, _response, tmp_path / "cache")
    rebuilt = rebuild_visual_snapshot(
        native, plan, result["evidence"], result["calls"], result["raw"]
    )
    assert not rebuilt.unverified_content
    region = next(r for r in rebuilt.regions if r.kind == "visual-transcription")
    assert (
        region.location["original_raster_sha256"]
        == hashlib.sha256(original).hexdigest()
    )
    assert region.location["original_raster_bbox"] == [0.25, 0, 1, 1]
    changed = deepcopy(plan)
    changed["jobs"][0]["view"]["view_to_original_pixel_edges"][2] += 1
    changed["plan_sha256"] = json_digest(
        {k: v for k, v in changed.items() if k != "plan_sha256"}
    )
    with pytest.raises(ValueError, match="transformed raster identity"):
        build_visual_request(changed, changed["jobs"][0])
    with pytest.raises(ValueError, match="reconstructed"):
        rebuild_visual_snapshot(
            native, changed, result["evidence"], result["calls"], result["raw"]
        )
