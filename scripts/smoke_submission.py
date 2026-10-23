#!/usr/bin/env python3
"""Run a real short report from staged submission code; fail closed without evidence.

This is an operator command, never invoked on normal inference or resume. Formal
runs reserve wall-clock time from the existing shared experiment ledger, require
an explicit idle window and do not load/restart the resident model service.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import sqlite3
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from bridge_evidence.contest.run_state import atomic_json, file_digest, json_digest

PROTOCOL = "staged-submission-real-smoke-v1"


def code_binding(code_dir):
    """Bind all packaged executable/configuration bytes, not timestamps."""
    root = Path(code_dir).resolve()
    paths = [p for p in root.rglob("*") if p.is_file() and "__pycache__" not in p.parts]
    if any(p.is_symlink() for p in root.rglob("*")):
        raise ValueError("staged code cannot contain symlinks")
    return {p.relative_to(root).as_posix(): file_digest(p) for p in sorted(paths)}


def build_command(
    *,
    code_dir,
    input_dir,
    output_dir,
    model_manifest,
    candidate_config,
    sdk_acceptance,
    python,
    backend="vllm-api",
    backend_url="http://127.0.0.1:8000/v1",
    rsi_bundle=None,
    rsi_sandbox_root=None,
    rsi_release=None,
    workflow="whole",
    runtime_options=None,
    sdk_parallel_acceptance=None,
):
    command = [
        str(python),
        str(Path(code_dir).resolve() / "run.py"),
        "--input-dir",
        str(Path(input_dir).resolve()),
        "--result-dir",
        str(Path(output_dir).resolve() / "result"),
        "--work-dir",
        str(Path(output_dir).resolve() / "work"),
        "--model-manifest",
        str(Path(model_manifest).resolve()),
        "--candidate-config",
        str(Path(candidate_config).resolve()),
        "--sdk-acceptance",
        str(Path(sdk_acceptance).resolve()),
        "--backend",
        backend,
        "--backend-url",
        backend_url,
        "--engine",
        "tongagents",
        "--workflow",
        workflow,
        "--rsi-mode",
        "frozen" if rsi_bundle else "off",
    ]
    for key, flag in (
        ("max_model_len", "max-model-len"),
        ("max_new_tokens", "max-new-tokens"),
        ("max_attempts", "max-attempts"),
        ("fusion", "fusion-policy"),
        ("report_workers", "report-workers"),
    ):
        if key in (runtime_options or {}):
            command += ["--" + flag, str(runtime_options[key])]
    if sdk_parallel_acceptance:
        command += [
            "--sdk-parallel-acceptance",
            str(Path(sdk_parallel_acceptance).resolve()),
        ]
    if rsi_bundle:
        if not rsi_sandbox_root or not rsi_release:
            raise ValueError(
                "frozen smoke requires sandbox and independently qualified release"
            )
        for flag, value in (
            ("rsi-bundle", rsi_bundle),
            ("rsi-sandbox-root", rsi_sandbox_root),
            ("rsi-release", rsi_release),
        ):
            command += ["--" + flag, str(Path(value).resolve())]
    return command


def validate_ledger_call(call, seconds, status, receipt, *, schema_version=2):
    errors = []
    if (
        type(seconds) not in (int, float)
        or not math.isfinite(seconds)
        or seconds <= 0
        or status not in ("success", "failed", "interrupted")
        or (call.get("status") == "success" and status != "success")
    ):
        errors.append("invalid/incomplete smoke ledger receipt")
    measured = call.get("seconds")
    if (
        type(measured) not in (int, float)
        or not math.isfinite(measured)
        or measured + 1e-6 < seconds
        or (
            schema_version in (2, 3)
            and not math.isclose(measured, seconds, rel_tol=1e-6, abs_tol=1e-6)
        )
    ):
        errors.append("smoke generation duration differs from actual ledger")
    if not receipt.get("prompt_sha256") or receipt["prompt_sha256"] != call.get(
        "prompt_sha256"
    ):
        errors.append("smoke generation receipt mismatch: prompt_sha256")
    if status == "success" or call.get("output_sha256") is not None:
        if not receipt.get("backend_request_id"):
            errors.append("actual backend request ID missing")
        for field in (
            "prompt_sha256",
            "output_sha256",
            "output_tokens",
            "finish_reason",
        ):
            if receipt.get(field) != call.get(field) or receipt.get(field) is None:
                errors.append("smoke generation receipt mismatch: " + field)
    if receipt.get("mock") or receipt.get("simulated"):
        errors.append("mock receipt cannot certify smoke")
    return errors


def validate_smoke_artifacts(
    code_dir, input_dir, output_dir, *, expected_identity=None
):
    """Recompute audits from raw files/SQLite; never trust a saved passed/verified bit."""
    from docx import Document

    from bridge_evidence.contest.execution_audit import audit_execution_records
    from bridge_evidence.contest.rsi_acceptance import backend_request_key

    code_dir, input_dir, output_dir = map(Path, (code_dir, input_dir, output_dir))
    result = {
        "passed": False,
        "protocol": PROTOCOL,
        "evidence_kind": "artifact-consistency-not-remote-attestation",
        "errors": [],
    }
    errors = result["errors"]
    try:
        work = output_dir / "work"
        manifest_path = work / "run_manifest.json"
        manifest = json.loads(manifest_path.read_text())
        identity = manifest["identity"]
        if json_digest(identity) != manifest["run_fingerprint"]:
            raise ValueError("smoke run_manifest identity mismatch")
        hashes = identity.get("code_files", {})
        if not hashes or "run.py" not in hashes:
            raise ValueError("smoke source identity missing")
        for name, digest in hashes.items():
            relative = Path(name)
            if relative.is_absolute() or ".." in relative.parts or "\\" in name:
                raise ValueError("unsafe smoke source path")
            path = code_dir / name
            if path.is_symlink() or file_digest(path) != digest:
                raise ValueError("staged source changed: " + name)
        options = identity.get("options", {})
        if options.get("engine") != "tongagents" or options.get("backend") not in (
            "vllm-api",
            "sglang-api",
        ):
            raise ValueError("smoke requires native SDK and resident API backend")
        if options.get("rsi_mode", "off") not in ("off", "frozen"):
            raise ValueError("experimental strategy cannot certify a submission")
        if expected_identity is not None:
            expected_options = expected_identity.get("options", {})
            for left, right, label in (
                (
                    identity.get("model_files"),
                    expected_identity.get("model_files"),
                    "base model",
                ),
                (
                    options.get("model_manifest", options.get("model")),
                    expected_options.get(
                        "model_manifest", expected_options.get("model")
                    ),
                    "model/adapter/tokenizer",
                ),
                (
                    options.get("candidate", options.get("config")),
                    expected_options.get("candidate", expected_options.get("config")),
                    "candidate",
                ),
                (
                    options.get("rsi_mode", "off"),
                    expected_options.get("rsi_mode", "off"),
                    "strategy mode",
                ),
                (
                    options.get("rsi_bundle"),
                    expected_options.get("rsi_bundle"),
                    "strategy bundle",
                ),
                (
                    options.get("sdk_acceptance_sha256"),
                    expected_options.get("sdk_acceptance_sha256"),
                    "SDK acceptance",
                ),
            ):
                if left != right or (
                    label
                    in (
                        "base model",
                        "model/adapter/tokenizer",
                        "candidate",
                        "SDK acceptance",
                    )
                    and not left
                ):
                    raise ValueError("smoke differs from packaged run: " + label)
            if identity["code_files"] != expected_identity.get("code_files"):
                raise ValueError("smoke source differs from packaged run")
            defaults = {
                "workflow": "whole",
                "backend": "vllm-api",
                "backend_url": "http://127.0.0.1:8000/v1",
                "max_model_len": 65536,
                "max_new_tokens": 32768,
                "max_attempts": 2,
                "fusion": "model-preserving",
                "report_workers": 1,
            }
            for key, default in defaults.items():
                expected = expected_options.get(key, default)
                if key == "workflow":
                    expected = expected_options.get(
                        "workflow",
                        expected_options.get("config", {}).get("workflow", default),
                    )
                if options.get(key, default) != expected:
                    raise ValueError(
                        "smoke runtime setting differs from packaged run: " + key
                    )
        execution = audit_execution_records(
            input_dir,
            output_dir / "result",
            work / "records",
            manifest["run_fingerprint"],
            profile="contest-sdk",
            require_strong=True,
        )
        errors.extend(execution["errors"])
        result["execution"] = execution
        records = [
            json.loads(p.read_text()) for p in sorted((work / "records").glob("*.json"))
        ]
        ids = [
            call["call_id"]
            for record in records
            for call in record.get("model_calls", [])
        ]
        if not ids or len(ids) != len(set(ids)):
            raise ValueError("missing or duplicated smoke model call IDs")
        ledger_path = work / "timing.sqlite"
        with sqlite3.connect(f"file:{ledger_path.resolve()}?mode=ro", uri=True) as db:
            rows = db.execute("SELECT id,seconds,status,receipt FROM calls").fetchall()
        ledger = {row[0]: row for row in rows}
        if set(ids) != set(ledger):
            raise ValueError("smoke receipts do not match actual call ledger")
        backend_ids = []
        for record in records:
            if not record.get("input_text_sha256"):
                errors.append("source text fingerprint missing")
            for call in record["model_calls"]:
                _, seconds, status, encoded = ledger[call["call_id"]]
                receipt = json.loads(encoded or "{}")
                errors.extend(
                    validate_ledger_call(
                        call,
                        seconds,
                        status,
                        receipt,
                        schema_version=record["schema_version"],
                    )
                )
                if status == "success" and receipt.get("backend_request_id"):
                    backend_ids.append(backend_request_key(receipt))
            output = output_dir / "result" / record["output"]["name"]
            if len(Document(output).tables) != 3:
                errors.append("smoke output must render exactly three tables")
        if len(backend_ids) != len(set(backend_ids)):
            errors.append("duplicate backend request receipt")
        seconds = json.loads((output_dir / "result/infer_time.json").read_text())[
            "infer_time"
        ]
        if (
            type(seconds) not in (int, float)
            or not math.isfinite(seconds)
            or seconds <= 0
        ):
            errors.append("positive official generation timing missing")
        # At least one real DOC conversion is part of the formal operator smoke.
        doc_sources = {
            file_digest(p)
            for p in input_dir.rglob("*")
            if p.is_file()
            and p.suffix.lower() == ".doc"
            and not p.name.startswith("~$")
        }
        converted = set()
        for metadata in work.rglob("conversion.json"):
            value = json.loads(metadata.read_text())
            if (
                value.get("conversion") == "docx"
                and value.get("source_sha256") in doc_sources
            ):
                files = [
                    p
                    for p in metadata.parent.glob("*.docx")
                    if file_digest(p) == value.get("output_sha256")
                ]
                if len(files) == 1 and value.get("converter"):
                    Document(files[0])
                    converted.add(value["source_sha256"])
        if not doc_sources or converted != doc_sources:
            errors.append("real DOC-to-DOCX conversion evidence missing/incomplete")
        result["binding"] = {
            "code": code_binding(code_dir),
            "run_manifest_sha256": file_digest(manifest_path),
            "model": options.get("model_manifest", options.get("model")),
            "inputs": {
                p.relative_to(input_dir).as_posix(): file_digest(p)
                for p in sorted(input_dir.rglob("*"))
                if p.is_file() and p.suffix.lower() in (".doc", ".docx")
            },
            "results": {
                p.name: file_digest(p)
                for p in sorted((output_dir / "result").glob("*"))
                if p.is_file()
            },
            "records": {
                p.name: file_digest(p)
                for p in sorted((work / "records").glob("*.json"))
            },
            "timing_sha256": file_digest(ledger_path),
            "call_ids": ids,
        }
    except (OSError, ValueError, TypeError, KeyError, sqlite3.Error) as error:
        errors.append(str(error))
    result["passed"] = not errors
    return result


def run_staged_process(command, *, cwd, env, stdout, timeout):
    """Own a fresh process group so a timeout also stops our converter children."""
    import signal

    with subprocess.Popen(
        command,
        cwd=cwd,
        env=env,
        stdout=stdout,
        stderr=subprocess.STDOUT,
        start_new_session=True,
    ) as process:
        try:
            process.wait(timeout=timeout)
            if process.returncode:
                raise subprocess.CalledProcessError(process.returncode, command)
        except BaseException:
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
            process.wait()
            raise


def run_smoke(
    *,
    code_dir,
    input_dir,
    output_dir,
    model_manifest,
    candidate_config,
    sdk_acceptance,
    python=sys.executable,
    backend="vllm-api",
    backend_url="http://127.0.0.1:8000/v1",
    timeout=1200,
    budget_ledger=None,
    idle_window=False,
    formal=True,
    expected_identity=None,
    rsi_bundle=None,
    rsi_sandbox_root=None,
    rsi_release=None,
    sdk_parallel_acceptance=None,
):
    """Execute new model requests from exactly the staged cwd; no cache reuse."""
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=False)
    result = {
        "protocol": PROTOCOL,
        "passed": False,
        "formal_accepted": False,
        "errors": [],
    }
    started, budget, job = time.perf_counter(), None, None
    try:
        if type(timeout) is not int or not 1 <= timeout <= 7200:
            raise ValueError("bounded integer smoke timeout required")
        if formal:
            if (
                not budget_ledger
                or not Path(budget_ledger).is_file()
                or not idle_window
            ):
                raise ValueError(
                    "formal smoke requires existing shared budget and explicit idle window"
                )
            if any(
                not p.resolve().is_relative_to(Path("/workspace"))
                for p in (Path(code_dir), output_dir, Path(budget_ledger))
            ):
                raise ValueError(
                    "formal smoke source/artifacts/shared budget must be under /workspace"
                )
            from bridge_evidence.contest.meta_optimization import CloudBudget

            budget = CloudBudget(budget_ledger)
            job = budget.start("staged-package-smoke", timeout)
        inputs = [
            p
            for p in Path(input_dir).rglob("*")
            if p.is_file()
            and p.suffix.lower() in (".doc", ".docx")
            and not p.name.startswith("~$")
        ]
        if not 1 <= len(inputs) <= 2 or not any(
            p.suffix.lower() == ".doc" for p in inputs
        ):
            raise ValueError(
                "select one or two approved short reports including a real .doc"
            )
        before = code_binding(code_dir)
        config = json.loads(Path(candidate_config).read_text())
        if config.get("workflow") not in ("whole", "chunked", "bounded-react"):
            raise ValueError("formal smoke needs supported recorded workflow")
        for path in (model_manifest, sdk_acceptance):
            if not Path(path).is_file():
                raise ValueError("required environment file missing: " + str(path))
        env = {
            k: v for k, v in os.environ.items() if k not in ("PYTHONPATH", "PYTHONHOME")
        }
        env["PYTHONNOUSERSITE"] = "1"
        env["PYTHONDONTWRITEBYTECODE"] = "1"
        # Real imports run in the same interpreter/cwd as staged inference.
        probe = "import importlib,json,sys; from pathlib import Path; sys.path.insert(0,str(Path.cwd()/'src')); names=('docx','transformers','tongagents','bridge_evidence.contest.inference_core'); mods={name:importlib.import_module(name).__file__ for name in names}; assert Path(mods[names[-1]]).resolve().is_relative_to(Path.cwd().resolve()); print(json.dumps(mods))"
        imported = subprocess.run(
            [str(python), "-c", probe],
            cwd=Path(code_dir).resolve(),
            env=env,
            capture_output=True,
            text=True,
            timeout=min(timeout, 60),
            check=True,
        )
        result["imports"] = json.loads(imported.stdout)
        command = build_command(
            code_dir=code_dir,
            input_dir=input_dir,
            output_dir=output_dir,
            model_manifest=model_manifest,
            candidate_config=candidate_config,
            sdk_acceptance=sdk_acceptance,
            python=python,
            backend=backend,
            backend_url=backend_url,
            rsi_bundle=rsi_bundle,
            rsi_sandbox_root=rsi_sandbox_root,
            rsi_release=rsi_release,
            workflow=config["workflow"],
            runtime_options=(expected_identity or {}).get("options", {}),
            sdk_parallel_acceptance=sdk_parallel_acceptance,
        )
        remaining = timeout - (time.perf_counter() - started)
        if remaining <= 0:
            raise TimeoutError("smoke preparation exhausted its wall-clock budget")
        with (output_dir / "runtime.log").open("w") as log:
            run_staged_process(
                command,
                cwd=Path(code_dir).resolve(),
                env=env,
                stdout=log,
                timeout=remaining,
            )
        result.update(
            validate_smoke_artifacts(
                Path(code_dir),
                Path(input_dir),
                output_dir,
                expected_identity=expected_identity,
            )
        )
        if code_binding(code_dir) != before:
            result["errors"].append("staged code changed during smoke")
        result["passed"] = not result["errors"]
        result["formal_accepted"] = bool(formal and result["passed"])
    except Exception as error:
        result["errors"].append(type(error).__name__ + ": " + str(error))
        result["passed"] = False
    finally:
        result["wall_seconds"] = time.perf_counter() - started
        result["code_dir"] = str(Path(code_dir).resolve())
        result["budget_job"] = job
        if job:
            budget.finish(job, status="complete" if result["passed"] else "failed")
        atomic_json(output_dir / "smoke-evidence.json", result)
    return result


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    for name in (
        "code-dir",
        "input-dir",
        "output-dir",
        "model-manifest",
        "candidate-config",
        "sdk-acceptance",
        "budget-ledger",
    ):
        ap.add_argument("--" + name, type=Path, required=True)
    ap.add_argument("--python", default=sys.executable)
    ap.add_argument("--backend", choices=("vllm-api", "sglang-api"), default="vllm-api")
    ap.add_argument("--backend-url", default="http://127.0.0.1:8000/v1")
    ap.add_argument("--timeout", type=int, default=1200)
    ap.add_argument("--idle-window", action="store_true")
    for name in (
        "rsi-bundle",
        "rsi-sandbox-root",
        "rsi-release",
        "sdk-parallel-acceptance",
    ):
        ap.add_argument("--" + name, type=Path)
    result = run_smoke(**vars(ap.parse_args()))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["formal_accepted"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
