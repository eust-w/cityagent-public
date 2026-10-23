"""Render synthetic history values; no model, historical dataset or GPU needed."""
from copy import deepcopy

from docx import Document
import pytest

from bridge_evidence.contest.docx_writer import write_extraction_docx


def render(tmp_path, basic, *, defects=None):
    data = {"基本信息": {"桥梁名称": "合成历史判定桥", **basic},
            "维修建议": [], "病害明细": [] if defects is None else defects}
    before = deepcopy(data)
    output = tmp_path / "report.docx"
    write_extraction_docx(data, str(output))
    assert data == before
    document = Document(output)
    subtitle = document.paragraphs[1].text
    fields = {row.cells[0].text: row.cells[1].text for row in document.tables[0].rows[1:]}
    return document, subtitle, fields


@pytest.mark.parametrize("basic", [
    {"上一次总体评分": "无"},
    {"上一次总体评分": "无", "上一次定检结果": "无"},
    {"上一次总体评分": "无", "上一次总体等级": "无", "上一次定检结果": "无"},
    {"上一次总体评分": None, "上一次定检结果": "无"},
    {"上一次总体评分": " \n无\t", "上一次定检结果": "  无  "},
])
def test_explicit_absent_history_does_not_render_as_a_comparison(tmp_path, basic):
    _, subtitle, fields = render(tmp_path, basic)
    assert subtitle == "（无对比年度）"
    for key, value in basic.items():
        if value not in (None, ""):
            assert fields[key] == str(value)  # Classification does not rewrite evidence.


@pytest.mark.parametrize("basic", [
    {},
    {"上一次总体评分": None},
    {"上一次总体评分": "", "上一次定检结果": ""},
    {"上一次总体评分": "  ", "上一次定检结果": "\n"},
    {"上一次总体评分": "未知", "上一次总体等级": "未知", "上一次定检结果": "未知"},
    {"上一次总体评分": "未提供"},
    {"上一次总体等级": "无法判断"},
    {"上一次定检结果": "不详"},
    {"上一次总体评分": "无", "上一次定检结果": "无法确定"},
    {"上一次总体评分": " Unknown ", "上一次总体等级": "无"},
])
def test_missing_unknown_and_conflicting_absence_remain_unconfirmed(tmp_path, basic):
    _, subtitle, fields = render(tmp_path, basic)
    assert subtitle == "（对比年度信息待确认）"
    for key, value in basic.items():
        if value not in (None, ""):
            assert fields[key] == str(value)


@pytest.mark.parametrize("zero", [0, 0.0, "0", "0.00"])
def test_zero_previous_score_is_valid_history_and_stays_in_the_table(tmp_path, zero):
    _, subtitle, fields = render(tmp_path, {"上一次总体评分": zero, "上一次定检结果": "无"})
    assert subtitle == "（有对比年度）"
    assert fields["上一次总体评分"] == str(zero)
    assert fields["上一次定检结果"] == "无"


@pytest.mark.parametrize("known", [{"上一次总体评分": 0}, {"上一次总体等级": "二类"},
                                   {"上一次定检结果": "上次发现 P1 裂缝"}])
def test_positive_history_evidence_takes_precedence_over_unknown_fields(tmp_path, known):
    basic = {"上一次总体评分": "未知", "上一次总体等级": "未提供", "上一次定检结果": "无法判断", **known}
    _, subtitle, fields = render(tmp_path, basic)
    assert subtitle == "（有对比年度）"
    assert all(fields[key] == str(value) for key, value in basic.items())


def test_result_text_beginning_with_none_is_not_the_absent_history_sentinel(tmp_path):
    _, subtitle, fields = render(tmp_path, {"上一次总体评分": "无", "上一次定检结果": "无明显病害"})
    assert subtitle == "（有对比年度）"
    assert fields["上一次定检结果"] == "无明显病害"


def test_explicit_previous_result_counts_without_a_previous_score(tmp_path):
    _, subtitle, fields = render(tmp_path, {"上一次定检结果": "上次检测发现 P1 裂缝"})
    assert subtitle == "（有对比年度）"
    assert fields["上一次定检结果"] == "上次检测发现 P1 裂缝"
    assert "上一次总体评分" not in fields


@pytest.mark.parametrize("grade", ["二类", "C级"])
def test_explicit_previous_grade_counts_with_no_score_or_result(tmp_path, grade):
    _, subtitle, fields = render(tmp_path, {"上一次总体评分": "无", "上一次总体等级": grade})
    assert subtitle == "（有对比年度）"
    assert fields["上一次总体评分"] == "无"
    assert fields["上一次总体等级"] == grade
    assert "上一次定检结果" not in fields


def test_unknown_or_unobservable_history_is_preserved_without_inventing_status(tmp_path):
    defects = [{"序号": "1", "病害描述": "墩底剥落", "是否新增": "无法判断",
                "上一次定检状态": "上次水位高无法观测", "发展程度": "无法比较"}]
    document, subtitle, fields = render(tmp_path, {
        "上一次总体评分": "80.00", "上一次定检结果": "局部状态未知",
        "病害发展趋势与具体说明": "上次无法观测，不能据此判断是否新增。",
    }, defects=defects)
    assert subtitle == "（有对比年度）"  # Known prior score, despite uncertain local history.
    assert fields["上一次定检结果"] == "局部状态未知"
    assert fields["病害发展趋势与具体说明"] == "上次无法观测，不能据此判断是否新增。"
    row = dict(zip([cell.text for cell in document.tables[-1].rows[0].cells],
                   [cell.text for cell in document.tables[-1].rows[1].cells]))
    for key, value in defects[0].items():
        assert row[key] == value
    assert "上一次定检未发现" not in "\n".join(row.values())
