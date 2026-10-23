from __future__ import annotations

# Hand-constructed fixtures use synthetic bridge and organization names.

import pytest

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw

from bridge_evidence.agents.vision import VisionEvidenceReviewer
from bridge_evidence.config import Settings
from bridge_evidence.evidence import extract_visual_candidates
from bridge_evidence.parsers import DocxStructureParser


def _write_text_image(path: Path) -> None:
    image = Image.new("RGB", (900, 420), "white")
    draw = ImageDraw.Draw(image)
    for y in range(40, 381, 50):
        draw.line((30, y, 870, y), fill="black", width=2)
    for x in range(30, 871, 210):
        draw.line((x, 40, x, 380), fill="black", width=2)
    draw.text((55, 60), "IAIC TRACK 01  BCI=87.35  A GRADE", fill="black")
    image.save(path, format="PNG", compress_level=0)


def test_extract_visual_candidates_finds_text_like_docx_image(tmp_path: Path) -> None:
    image_path = tmp_path / "rating-table.png"
    _write_text_image(image_path)
    document = Document()
    document.add_picture(str(image_path))
    docx_path = tmp_path / "report.docx"
    document.save(docx_path)

    candidates = extract_visual_candidates(docx_path)

    assert len(candidates) == 1
    assert candidates[0].media_path.startswith("word/media/")
    assert candidates[0].width == 900
    assert candidates[0].height == 420
    assert candidates[0].text_likelihood > 0.5
    assert candidates[0].data_url.startswith("data:image/png;base64,")


def test_visual_candidate_maps_reused_media_to_all_image_occurrences(tmp_path: Path) -> None:
    image_path = tmp_path / "reused-rating-table.png"
    _write_text_image(image_path)
    document = Document()
    document.add_picture(str(image_path))
    document.add_picture(str(image_path))
    docx_path = tmp_path / "reused.docx"
    document.save(docx_path)
    parsed = DocxStructureParser().parse(docx_path)

    candidates = extract_visual_candidates(docx_path, parsed=parsed)

    assert len(candidates) == 1
    assert candidates[0].occurrence_count == 2
    assert candidates[0].image_indices == (0, 1)
    assert candidates[0].stories == ("body",)


class _VisionClient:
    def __init__(self) -> None:
        self.calls: list[tuple[str, str]] = []

    def complete_vision_json(self, *, model: str, task: str, **_: object):
        self.calls.append((task, model))
        if task == "visual-evidence-primary":
            return {
                "visible_text": "总体 BCI=87.35，技术状况等级 A级；桥面系 90.00。",
                "facts": [
                    {"field": "scores.overall.score", "value": "87.35", "evidence_quote": "BCI=87.35"},
                    {"field": "scores.deck.score", "value": "90.00", "evidence_quote": "桥面系 90.00"},
                ],
            }
        return {
            "visible_text": "总体 BCI=87.35，技术状况等级 A级；桥面系 89.00。",
            "facts": [
                {"field": "scores.overall.score", "value": "87.35", "evidence_quote": "BCI=87.35"},
                {"field": "scores.deck.score", "value": "89.00", "evidence_quote": "桥面系 89.00"},
            ],
        }


def test_visual_reviewer_keeps_only_cross_model_grounded_facts() -> None:
    client = _VisionClient()
    settings = Settings(llm_mode="off", vision_model="minimax-m3", vision_critic_model="kimi-k3")
    reviewer = VisionEvidenceReviewer(client, settings, "只提取图中原文")

    review = reviewer.review("data:image/png;base64,AA==")

    assert client.calls == [
        ("visual-evidence-primary", "minimax-m3"),
        ("visual-evidence-critic", "kimi-k3"),
    ]
    assert [(fact.field, fact.value) for fact in review.verified_facts] == [
        ("scores.overall.score", "87.35")
    ]


class _AliasedVisionClient:
    def complete_vision_json(self, *, task: str, **_: object):
        if task == "visual-evidence-primary":
            return {
                "visible_text": "总体 BCI=87.35，等级 A GRADE。",
                "facts": [
                    {"field": "BCI", "value": "BCI=87.350", "evidence_quote": "BCI=87.35"},
                    {"field": "等级", "value": "A GRADE", "evidence_quote": "A GRADE"},
                    {"field": "模型猜测", "value": "无", "evidence_quote": "无"},
                ],
            }
        return {
            "visible_text": "总体 BCI=87.35，等级 A级。",
            "facts": [
                {"field": "scores.overall.score", "value": "87.35", "evidence_quote": "BCI=87.35"},
                {"field": "scores.overall.grade", "value": "A级", "evidence_quote": "A级"},
            ],
        }


def test_visual_reviewer_normalizes_only_allowlisted_schema_aliases() -> None:
    reviewer = VisionEvidenceReviewer(
        _AliasedVisionClient(),
        Settings(llm_mode="off", vision_model="minimax-m3", vision_critic_model="kimi-k3"),
        "只提取图中原文",
    )

    review = reviewer.review("data:image/png;base64,AA==")

    assert [(fact.field, fact.value) for fact in review.verified_facts] == [
        ("scores.overall.score", "BCI=87.350"),
        ("scores.overall.grade", "A GRADE"),
    ]


def test_local_ocr_fact_extraction() -> None:
    from bridge_evidence.evidence.visuals import LocalOcrEngine

    text = """
    工程名称：合成跨河大桥
    检测报告日期：2019年09月15日
    总体评定：BCI=88.50，等级评定为 B级
    桥面系：BCIm=85.20
    上部结构：BCIs=90.10
    下部结构：BCIx=89.00
    """
    facts = LocalOcrEngine.extract_facts(text)
    fact_map = {f["field"]: f["value"] for f in facts}

    assert fact_map["bridge.name"] == "合成跨河大桥"
    assert fact_map["bridge.report_date"] == "2019年09月15日"
    assert fact_map["scores.overall.score"] == "88.50"
    assert fact_map["scores.overall.grade"] == "B级"
    assert fact_map["scores.deck.score"] == "85.20"
    assert fact_map["scores.superstructure.score"] == "90.10"
    assert fact_map["scores.substructure.score"] == "89.00"


def test_visual_reviewer_offline_ocr_fallback() -> None:
    reviewer = VisionEvidenceReviewer(
        None,
        Settings(llm_mode="off", vision_model=""),
        "只提取图中原文",
    )
    # Empty data should return without exception
    review = reviewer.review("data:image/png;base64,")
    assert isinstance(review.verified_facts, list)
from bridge_evidence.evidence.visuals import OcrSpan, reconstruct_ocr_table


def test_reconstruct_ocr_disease_table_from_normalized_spans() -> None:
    spans = [
        OcrSpan("病害位置", 0.99, (0.05, 0.10, 0.22, 0.16)),
        OcrSpan("病害类型", 0.98, (0.31, 0.10, 0.47, 0.16)),
        OcrSpan("病害描述", 0.97, (0.56, 0.10, 0.78, 0.16)),
        OcrSpan("左幅1跨1#梁底", 0.96, (0.05, 0.24, 0.25, 0.30)),
        OcrSpan("裂缝", 0.95, (0.32, 0.24, 0.42, 0.30)),
        OcrSpan("纵向裂缝0.5m", 0.94, (0.56, 0.24, 0.82, 0.30)),
        OcrSpan("右幅2跨桥面", 0.93, (0.05, 0.39, 0.25, 0.45)),
        OcrSpan("破损", 0.92, (0.32, 0.39, 0.42, 0.45)),
        OcrSpan("铺装局部破损", 0.91, (0.56, 0.39, 0.82, 0.45)),
    ]

    table = reconstruct_ocr_table(spans)

    assert table is not None
    assert table.kind == "disease_table"
    assert table.rows == [
        ["病害位置", "病害类型", "病害描述"],
        ["左幅1跨1#梁底", "裂缝", "纵向裂缝0.5m"],
        ["右幅2跨桥面", "破损", "铺装局部破损"],
    ]
    assert table.bbox == pytest.approx((0.05, 0.10, 0.82, 0.45))
    assert table.confidence == pytest.approx(0.95)


def test_reconstruct_ocr_table_rejects_non_table_text() -> None:
    spans = [
        OcrSpan("桥梁定期检查报告", 0.99, (0.10, 0.10, 0.80, 0.18)),
        OcrSpan("总体技术状况等级为B级", 0.98, (0.10, 0.30, 0.80, 0.38)),
    ]

    assert reconstruct_ocr_table(spans) is None
