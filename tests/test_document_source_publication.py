"""Synthetic publication contracts only; no real report/model/SDK/network."""
from copy import deepcopy
import json
from pathlib import Path
import sys

from docx import Document
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from bridge_evidence.contest import document_cli as cli
from bridge_evidence.contest.docx_writer import write_extraction_docx
from bridge_evidence.contest.document_delivery_validation import validate_delivery_docx


@pytest.fixture
def env(tmp_path):
    class Env:
        def __init__(self):
            self.directory = tmp_path / 'work'
            self.directory.mkdir()
            self.output = tmp_path / 'synthetic.docx'
            self.fingerprint = cli.json_digest('synthetic-run')
            self.source_hash = cli.json_digest('synthetic-source')
            self.data = {'基本信息': {'桥梁名称': '合成桥', '总体评分': 95}, '病害明细': [], '维修建议': []}
            self.manifest = {'document_sha256': self.source_hash, 'source_kind': 'docx', 'units': [], 'images': [],
                'unverified_content': [{'source_path': f'word/document.xml/p[{i}]/object[1]', 'kind': 'object',
                    'verification': 'unverified', 'reason': 'SYNTHETIC unparsed native equation'} for i in (1, 2)]}
            scope_hash, catalog_hash = cli.json_digest('synthetic-scope'), cli.json_digest('synthetic-catalog')
            rows = [dict(object_id='so_'+cli.json_digest(f'object-{i}')[:32], representation_state='unread',
                disposition='not_relied_on' if i == 0 else 'outside_task', review_status='accepted',
                object_gate_passed=True, proposal_id='sp_'+cli.json_digest(f'proposal-{i}')[:32],
                model_call_id='synthetic-independent-'+str(i),
                source_limitations=[f'合成限制 {i}：原公式运算符未恢复，未使用它复算评分。']) for i in range(2)]
            coverage = {'scope_sha256': scope_hash, 'objects': rows, 'total_objects': 2, 'object_gate_passed': True,
                'accepted_count': 2, 'blocking_count': 0, 'unverified_representation_count': 2,
                'all_source_semantics_verified': False}
            processing = {'task': {'name': 'three-table-faithful-extraction'}, 'source_catalog_sha256': catalog_hash,
                'scope_sha256': scope_hash, 'object_gate_passed': True, 'all_source_semantics_verified': False,
                'unverified_content': deepcopy(self.manifest['unverified_content']),
                'objects': [{**deepcopy(row), 'locator': self.manifest['unverified_content'][i]['source_path'], 'kind': 'ole'}
                            for i, row in enumerate(rows)], 'limitation': 'Synthetic scope only, no actual model was called.'}
            self.source = {'passed': True, 'blockers': [], 'catalog_ready': True,
                'source_catalog_sha256': catalog_hash, 'scope_sha256': scope_hash,
                'semantic_input_sha256': cli.json_digest('synthetic-semantic-input'), 'inventory_count': 2,
                'coverage': coverage, 'processing_scope': processing}
            self.source['attestation_sha256'] = cli.json_digest({'catalog_sha256': catalog_hash, 'scope_sha256': scope_hash,
                'objects': rows, 'passed': True, 'blockers': []})
            self.receipt = {'status': 'completed', 'workspace_version': 9, 'draft_sha256': cli.json_digest(self.data),
                'final_data': deepcopy(self.data), 'source_objects': deepcopy(self.source),
                'delivery': {'ready': True, 'workspace_version': 9, 'draft_sha256': cli.json_digest(self.data),
                             'source_objects': deepcopy(self.source)},
                'accepted_finish': {'workspace_version': 9, 'draft_sha256': cli.json_digest(self.data),
                    'source_object_attestation_sha256': self.source['attestation_sha256'],
                    'source_object_semantic_input_sha256': self.source['semantic_input_sha256'],
                    'source_object_processing_scope_sha256': cli.json_digest(processing)}}
            self.bind_manifest()

        def bind_manifest(self):
            cli._bind_source(self.directory, self.manifest, None, self.fingerprint, self.source_hash)

        def result(self):
            return {'data': self.data, 'receipt': self.receipt}

        def publish(self):
            return cli._publish(self.result(), self.directory, self.output, self.fingerprint, self.source_hash)

        def completed(self):
            return cli._completed('synthetic', self.directory, self.output, self.fingerprint, self.source_hash)

        def read(self, name):
            return json.loads((self.directory / name).read_text())

        def rehash_execution(self, receipt):
            # Update untrusted outer hashes deliberately. Content/scope bindings
            # must still catch corruption, not merely rely on file checksums.
            cli.atomic_json(self.directory / 'execution.json', receipt)
            marker = self.read('publication.json')
            marker['execution_sha256'] = cli.file_digest(self.directory / 'execution.json')
            cli.atomic_json(self.directory / 'publication.json', marker)

    return Env()


def test_publish_and_resume_bind_full_scope_and_visible_per_occurrence_note(env):
    digest = env.publish()
    execution = env.read('execution.json')
    scope = env.read('processing_scope.json')
    marker = env.read('publication.json')
    assert scope == env.source['processing_scope']
    assert execution['processing_scope_sha256'] == marker['processing_scope_sha256'] == cli.json_digest(scope)
    assert execution['processing_note_sha256'] == marker['processing_note_sha256']
    assert execution['source_object_attestation_sha256'] == marker['source_object_attestation_sha256'] == env.source['attestation_sha256']
    assert execution['accepted_finish']['source_object_processing_scope_sha256'] == cli.json_digest(scope)
    doc = Document(env.output)
    assert len(doc.tables) == 3
    assert doc.paragraphs[-2].text == '处理范围说明'
    note = doc.paragraphs[-1].text
    assert note == cli.processing_scope_note(scope)
    assert '未识别' in note and '本次三表未依赖此对象' in note and '不属于本次三表任务范围' in note
    for row in scope['objects']:
        assert row['object_id'] in note and row['locator'] in note
        assert all(value in note for value in row['source_limitations'])
    assert execution['delivery_validation'] == validate_delivery_docx(env.data, env.output, processing_note=note)
    assert env.completed()['execution_sha256'] == digest


@pytest.mark.parametrize('field', ['source_objects', 'accepted_finish'])
def test_native_inventory_missing_receipt_field_blocks_before_writer(env, monkeypatch, field):
    env.receipt.pop(field)
    monkeypatch.setattr(cli, 'write_extraction_docx', lambda *_a, **_k: pytest.fail('Writer must not run'))
    with pytest.raises(ValueError): env.publish()
    assert not env.output.exists()


@pytest.mark.parametrize('field', ['source_object_attestation_sha256', 'source_object_semantic_input_sha256',
                                  'source_object_processing_scope_sha256', 'draft_sha256', 'workspace_version'])
def test_finish_binding_change_blocks_publication(env, field):
    env.receipt['accepted_finish'][field] = 'wrong'
    with pytest.raises(ValueError, match='same state'): env.publish()
    assert not env.output.exists()


@pytest.mark.parametrize('field', ['source_catalog_sha256', 'scope_sha256', 'semantic_input_sha256', 'attestation_sha256'])
def test_coverage_hash_missing_or_delivery_disagrees_blocks(env, field):
    env.receipt['source_objects'].pop(field)
    with pytest.raises(ValueError): env.publish()
    assert not env.output.exists()


def test_scope_locator_swap_is_rejected_even_while_inventory_set_still_matches(env):
    for source in (env.receipt['source_objects'], env.receipt['delivery']['source_objects']):
        rows = source['processing_scope']['objects']
        rows[0]['locator'], rows[1]['locator'] = rows[1]['locator'], rows[0]['locator']
    with pytest.raises(ValueError, match='same state'): env.publish()


@pytest.mark.parametrize('change', ['partial_inventory', 'partial_rows', 'unaccepted', 'forged_attestation', 'locator_missing', 'no_limits'])
def test_malformed_or_incomplete_scope_never_publishes(env, change):
    source = deepcopy(env.source)
    if change == 'partial_inventory': source['processing_scope']['unverified_content'].pop()
    elif change == 'partial_rows': source['processing_scope']['objects'].pop()
    elif change == 'unaccepted': source['processing_scope']['objects'][0]['review_status'] = 'unresolved'
    elif change == 'forged_attestation': source['attestation_sha256'] = cli.json_digest('wrong')
    elif change == 'locator_missing': source['processing_scope']['objects'][0].pop('locator')
    elif change == 'no_limits':
        source['processing_scope']['objects'][0]['source_limitations'] = []
        source['coverage']['objects'][0]['source_limitations'] = []
        source['attestation_sha256'] = cli.json_digest({'catalog_sha256': source['source_catalog_sha256'],
            'scope_sha256': source['scope_sha256'], 'objects': source['coverage']['objects'], 'passed': True, 'blockers': []})
    env.receipt['source_objects'] = source
    env.receipt['delivery']['source_objects'] = deepcopy(source)
    env.receipt['accepted_finish']['source_object_attestation_sha256'] = source['attestation_sha256']
    env.receipt['accepted_finish']['source_object_processing_scope_sha256'] = cli.json_digest(source['processing_scope'])
    with pytest.raises(ValueError): env.publish()
    assert not env.output.exists()


def test_formatter_never_states_complete_semantic_verification(env):
    scope = deepcopy(env.source['processing_scope'])
    scope['all_source_semantics_verified'] = True
    with pytest.raises(ValueError, match='all source semantics'): cli.processing_scope_note(scope)


def test_silent_writer_note_drop_cannot_publish(env, monkeypatch):
    def writer(data, path, *, processing_note):
        write_extraction_docx(data, path)  # Simulate silent omission while tables are perfect.
    monkeypatch.setattr(cli, 'write_extraction_docx', writer)
    with pytest.raises(ValueError, match='block count'): env.publish()
    assert not env.output.exists() and not list(env.output.parent.glob('.pending-*.docx'))


def test_writer_same_length_note_rewrite_cannot_publish(env, monkeypatch):
    def writer(data, path, *, processing_note):
        write_extraction_docx(data, path, processing_note=processing_note.replace('未识别', '已识别'))
    monkeypatch.setattr(cli, 'write_extraction_docx', writer)
    with pytest.raises(ValueError, match='paragraph'): env.publish()
    assert not env.output.exists()


def test_scope_file_mutation_at_last_publication_guard_prevents_link(env):
    def guard(stage):
        if stage == 'before_atomic_publication':
            path = env.directory / 'processing_scope.json'
            value = env.read(path.name)
            value['objects'][0]['source_limitations'] = ['changed after DOCX readback']
            cli.atomic_json(path, value)
    with pytest.raises(ValueError, match='Saved processing scope'):
        cli._publish(env.result(), env.directory, env.output, env.fingerprint, env.source_hash, deadline_guard=guard)
    assert not env.output.exists()


@pytest.mark.parametrize('artifact', ['scope_missing', 'scope_corrupt', 'scope_symlink', 'marker_scope', 'marker_attestation', 'receipt_scope'])
def test_resume_rejects_missing_changed_scope_or_publication_binding(env, artifact):
    env.publish()
    if artifact == 'scope_missing': (env.directory / 'processing_scope.json').unlink()
    elif artifact == 'scope_corrupt':
        scope = env.read('processing_scope.json'); scope['limitation'] = 'changed'
        cli.atomic_json(env.directory / 'processing_scope.json', scope)
    elif artifact == 'scope_symlink':
        path = env.directory / 'processing_scope.json'; target = env.directory / 'copied_scope.json'
        target.write_bytes(path.read_bytes()); path.unlink(); path.symlink_to(target)
    elif artifact in {'marker_scope', 'marker_attestation'}:
        marker = env.read('publication.json')
        marker.pop('processing_scope_sha256' if artifact == 'marker_scope' else 'source_object_attestation_sha256')
        cli.atomic_json(env.directory / 'publication.json', marker)
    else:
        receipt = env.read('execution.json'); receipt.pop('processing_scope_sha256'); env.rehash_execution(receipt)
    with pytest.raises(ValueError): env.completed()


def test_resume_word_note_rewrite_rejected_even_when_all_file_hashes_updated(env):
    env.publish()
    doc = Document(env.output)
    doc.paragraphs[-1].text = doc.paragraphs[-1].text.replace('未识别', '已识别')
    doc.save(env.output)
    receipt = env.read('execution.json')
    new_hash = cli.file_digest(env.output)
    receipt['output']['sha256'] = new_hash
    receipt['delivery_validation']['docx_sha256'] = new_hash
    env.rehash_execution(receipt)
    with pytest.raises(ValueError, match='paragraph'): env.completed()


def test_resume_prepared_marker_reconciles_only_after_scope_and_note_recheck(env):
    env.publish()
    marker = env.read('publication.json'); marker['status'] = 'prepared'
    cli.atomic_json(env.directory / 'publication.json', marker)
    assert env.completed()['status'] == 'completed'
    assert env.read('publication.json')['status'] == 'committed'


def test_no_native_objects_preserve_old_two_argument_writer_and_no_note_fields(tmp_path, monkeypatch):
    directory = tmp_path / 'work'; directory.mkdir()
    manifest = {'units': [], 'images': [], 'unverified_content': []}
    cli._bind_source(directory, manifest, None, 'fp', 'src')
    data = {'基本信息': {'桥梁名称': '纯文本合成桥'}, '维修建议': [], '病害明细': []}
    receipt = {'status': 'completed', 'delivery': {'ready': True}, 'draft_sha256': cli.json_digest(data), 'final_data': data}
    calls = []
    def old_writer(data, path):
        calls.append(path); write_extraction_docx(data, path)
    monkeypatch.setattr(cli, 'write_extraction_docx', old_writer)
    output = tmp_path / 'legacy.docx'
    cli._publish({'data': data, 'receipt': receipt}, directory, output, 'fp', 'src')
    assert len(calls) == 1 and not (directory / 'processing_scope.json').exists()
    execution = json.loads((directory / 'execution.json').read_text())
    assert not any(key.startswith('processing_') or key == 'source_object_attestation_sha256' for key in execution)
    assert '处理范围说明' not in [p.text for p in Document(output).paragraphs]
    assert cli._completed('synthetic', directory, output, 'fp', 'src')['reused']


def test_no_native_objects_reject_unexpected_scope_sidecar(tmp_path):
    directory = tmp_path / 'work'; directory.mkdir()
    cli._bind_source(directory, {'unverified_content': []}, None, 'fp', 'src')
    cli.atomic_json(directory / 'processing_scope.json', {'unexpected': 'stale scope'})
    data = {'基本信息': {'桥梁名称': '合成桥'}, '维修建议': [], '病害明细': []}
    result = {'data': data, 'receipt': {'status': 'completed', 'delivery': {'ready': True},
              'draft_sha256': cli.json_digest(data), 'final_data': data}}
    with pytest.raises(ValueError, match='Unexpected processing scope'):
        cli._publish(result, directory, tmp_path / 'no.docx', 'fp', 'src')


def test_scope_recheck_remains_inside_total_publication_deadline(env):
    stages = []
    def guard(stage):
        stages.append(stage)
        if stage == 'after_processing_scope_recheck':
            raise RuntimeError('synthetic total deadline exhausted after scope readback')
    with pytest.raises(RuntimeError, match='deadline exhausted'):
        cli._publish(env.result(), env.directory, env.output, env.fingerprint, env.source_hash, deadline_guard=guard)
    assert stages[-2:] == ['before_atomic_publication', 'after_processing_scope_recheck']
    assert not env.output.exists()
