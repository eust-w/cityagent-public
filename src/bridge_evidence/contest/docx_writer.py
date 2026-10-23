"""把模型输出的三部分 JSON 渲染成"信息提取报告" docx（对齐官方标签格式）。

版式参照训练集标签：
  标题《桥名·信息提取报告》+（有/无对比年度）
  1、简要信息      -> 表0（字段/内容/说明）
  2、详细信息      -> 详细结论段落 + 建议明细表 + 病害列表表
依赖 python-docx（赛机镜像已装；若无则 pip 安装即可）。
"""
from __future__ import annotations

from docx import Document
from docx.shared import Pt

# 基本信息表的字段顺序与说明列，对齐标签样式
BASIC_FIELDS = [
    ("桥梁名称", "桥梁全称"),
    ("报告日期", "报告出具日期"),
    ("总体评分", "总体技术状况评分"),
    ("总体等级", "总体技术状况等级"),
    ("上部结构评分", "上部结构技术状况评分"),
    ("上部结构等级", "上部结构技术状况等级"),
    ("下部结构评分", "下部结构技术状况评分"),
    ("下部结构等级", "下部结构技术状况等级"),
    ("桥面系评分", "桥面系技术状况评分"),
    ("桥面系等级", "桥面系技术状况等级"),
    ("上一次总体评分", "上一次定检总体评分"),
    ("上一次总体等级", "上一次定检总体等级"),
    ("上一次定检结果", "上一次定检结果"),
    ("病害发展趋势与具体说明", "病害发展趋势与具体说明"),
    ("总体结论", "总体结论"),
    ("主要风险点", "主要风险点"),
    ("建议", "建议分类统计"),
    ("建议统计", "按建议明细条目统计"),
]

ADVICE_FIELDS = ["序号", "建议类别", "建议内容", "病害部位"]
DEFECT_FIELDS = ["序号", "病害部位", "病害类型", "病害描述", "是否新增", "上一次定检状态", "发展程度"]


def _get(d: dict, *names: str) -> str:
    for n in names:
        v = d.get(n)
        if v not in (None, ""):
            return str(v)
    return ""


def numbered_delivery_rows(rows: list[dict]) -> list[dict]:
    """Number displayed rows in final order without changing saved source data.

    These are local table numbers, not fact IDs or numbers from the report.
    Each table starts at one, including after rows are merged or removed.
    """
    return [{**row, "序号": str(index)}
            for index, row in enumerate((row for row in rows if isinstance(row, dict)), 1)]


def write_extraction_docx(data: dict, out_path: str, *, processing_note: str | None = None,
                          submission_format: bool = False, renumber_rows: bool = False) -> None:
    """Render saved values verbatim unless display numbering is explicitly requested.

    Whole/chunked inference and its reconstruction audit retain source sequence
    values. Document-agent exports opt into their separate display-number contract.
    """
    if submission_format and processing_note is not None:
        raise ValueError("Submission documents keep processing notes in sidecar records only")
    if processing_note is not None and (not isinstance(processing_note, str) or not processing_note.strip()):
        raise ValueError("Processing note must be a nonempty string")
    basic = data.get("基本信息") or {}
    advice = data.get("维修建议") or []
    defects = data.get("病害明细") or []
    if renumber_rows:
        advice = numbered_delivery_rows(advice)
        defects = numbered_delivery_rows(defects)

    name = _get(basic, "桥梁名称") or "未知桥梁"
    prev_res = _get(basic, "上一次定检结果")
    prev_score = _get(basic, "上一次总体评分")
    prev_grade = _get(basic, "上一次总体等级")
    # Missing/unknown history is not proof of absence or presence. Classify
    # only exact sentinels; keep every original field and numeric zero intact.
    history_values = [value.strip().casefold() for value in (prev_score, prev_grade, prev_res)]
    unknown_history = {"未知", "未提供", "无法判断", "无法确定", "不详", "unknown",
                       "不确定", "未说明", "未明确", "未提及", "无法比较"}
    if any(value not in {"", "无", *unknown_history} for value in history_values):
        subtitle = "（有对比年度）"
    elif "无" in history_values and not any(value in unknown_history for value in history_values):
        subtitle = "（无对比年度）"
    else:
        subtitle = "（对比年度信息待确认）"

    doc = Document()
    doc.add_heading(f"{name}·信息提取报告", level=0)
    if not submission_format:
        doc.add_paragraph(subtitle)

    # 1、简要信息
    doc.add_heading("1、简要信息", level=1)
    if not submission_format:
        doc.add_paragraph("从报告中提取桥梁定检的概要结论与关键指标，输出内容包括：")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["字段", "内容", "说明"]):
        table.rows[0].cells[i].text = h
    for field, gloss in BASIC_FIELDS:
        val = _get(basic, field)
        if not val:
            continue
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = field, val, gloss

    # 2、详细信息
    doc.add_heading("2、详细信息", level=1)
    doc.add_heading("（1）详细结论", level=2)
    for para_field in ["总体结论", "病害发展趋势与具体说明", "主要风险点"]:
        val = _get(basic, para_field)
        if val:
            doc.add_paragraph(val)

    doc.add_heading("（2）建议明细", level=2)
    table = doc.add_table(rows=1, cols=len(ADVICE_FIELDS))
    table.style = "Table Grid"
    for i, h in enumerate(ADVICE_FIELDS):
        table.rows[0].cells[i].text = h
    for row in advice:
        if not isinstance(row, dict):
            continue
        cells = table.add_row().cells
        for i, f in enumerate(ADVICE_FIELDS):
            cells[i].text = _get(row, f)

    doc.add_heading("（3）病害列表", level=2)
    table = doc.add_table(rows=1, cols=len(DEFECT_FIELDS))
    table.style = "Table Grid"
    for i, h in enumerate(DEFECT_FIELDS):
        table.rows[0].cells[i].text = h
    for row in defects:
        if not isinstance(row, dict):
            continue
        cells = table.add_row().cells
        for i, f in enumerate(DEFECT_FIELDS):
            cells[i].text = _get(row, f)

    if processing_note is not None:
        doc.add_heading("处理范围说明", level=1)
        doc.add_paragraph(processing_note)

    for p in doc.paragraphs:
        for run in p.runs:
            run.font.size = Pt(10.5)

    doc.save(out_path)


def write_numbered_extraction_docx(data: dict, out_path: str, *, processing_note: str | None = None,
                                  submission_format: bool = False) -> None:
    """Use the document-agent display-number contract without changing saved JSON."""
    write_extraction_docx(data, out_path, processing_note=processing_note,
                          submission_format=submission_format, renumber_rows=True)
