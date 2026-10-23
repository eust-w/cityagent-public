"""Synthetic contract tests only: no server, processor, model or real report."""
import base64
from copy import deepcopy
import hashlib
import io
import json

from docx import Document
from PIL import Image
import pytest

from bridge_evidence.contest.document_tools import DocumentTools
from bridge_evidence.contest.document_vision_backend import (
    MODEL, OwnModelVisionBackend, VisionObservationError, canonical_bytes, json_digest,
)
from bridge_evidence.contest.evidence_manifest import build_manifest


def png(size=(8, 6), *, orientation=None):
    image = Image.new("RGB", size)
    for x in range(size[0]):
        for y in range(size[1]):
            image.putpixel((x, y), (x * 17 % 256, y * 29 % 256, (x + y) * 11 % 256))
    out = io.BytesIO()
    options = {}
    if orientation is not None:
        exif = Image.Exif()
        exif[274] = orientation
        options["exif"] = exif
    image.save(out, format="PNG", **options)
    return out.getvalue()


def identity():
    return {"model": MODEL, "max_model_len": 65536,
            "model_identity_sha256": "1" * 64, "service_identity_sha256": "2" * 64,
            "processor_identity_sha256": "3" * 64}


def arguments(blob=None):
    blob = png() if blob is None else blob
    return blob, {"region": {}, "question": "图片能辨认出什么？",
                  "context": [{"evidence_id": "E000001", "start": 0, "end": 3,
                               "text": "图注甲", "source_path": "word/document.xml#p[1]"}],
                  "image": {"id": "I00001", "source_path": "word/document.xml#p[2]/drawing",
                            "target": "media/image1.png", "relationship_id": "rId7",
                            "image_sha256": hashlib.sha256(blob).hexdigest()}}


class Clock:
    value = 10.0

    def __call__(self):
        return self.value


class FakeGeneration:
    """Fabricated test records, explicitly not real-model acceptance evidence."""
    def __init__(self):
        self.calls = []
        self.mutate = None
        self.content = {"visible_content": ["合成色块"], "readable_text": [],
                        "uncertainties": ["没有尺度，不能判断尺寸。"]}

    def __call__(self, payload, *, provenance, deadline):
        self.calls.append({"payload": deepcopy(payload), "provenance": deepcopy(provenance), "deadline": deadline})
        request_sha = json_digest(payload)
        usage = {"prompt_tokens": 662, "completion_tokens": 50, "total_tokens": 712}
        response = {"id": "synthetic-response", "model": MODEL, "usage": usage,
                    "choices": [{"finish_reason": "stop", "message": {
                        "role": "assistant", "content": json.dumps(self.content, ensure_ascii=False)}}]}
        call_id, request_id = f"synthetic-call-{len(self.calls)}", f"synthetic-request-{len(self.calls)}"
        receipt = {"call_id": call_id, "request_id": request_id, "status": "success", "http_status": 200,
                   "kind": "generation", "model": MODEL, "endpoint": "/v1/chat/completions",
                   "input_sha256": request_sha, "output_sha256": hashlib.sha256(json.dumps(response).encode()).hexdigest(),
                   "usage": deepcopy(usage), "prompt_tokens_checked": 662, "finish_reason": "stop"}
        audit = {"call_id": call_id, "request_id": request_id, "status": "success", "generation_attempted": True,
                 "http_status": 200, "request_kind": "generation", "model": MODEL,
                 "request_sha256": request_sha, "response_sha256": json_digest(response),
                 "usage": deepcopy(usage), "prompt_tokens_checked": 662, "finish_reason": "stop",
                 "response_callback_completed": True}
        admission = {"status": "processor_counted", "method": "installed_vllm_native_processor",
                     "request_sha256": request_sha,
                     "image_payload_sha256": provenance["representation"]["payload_image_sha256"],
                     "processor_identity_sha256": "3" * 64, "service_identity_sha256": "2" * 64,
                     "expanded_prompt_tokens": 662, "max_model_len": 65536,
                     "reserved_output_tokens": payload["max_tokens"], "final_prompt_token_ids_sha256": "4" * 64}
        value = {"response": response, "receipt": receipt, "audit": audit, "admission": admission}
        if self.mutate:
            self.mutate(value, payload, provenance)
        return value


def backend(callback=None, clock=None, **kwargs):
    return OwnModelVisionBackend(callback or FakeGeneration(), identity=identity(), document_sha256="a" * 64,
                                 deadline=100, clock=clock or Clock(), **kwargs)


def sync_response_records(value):
    response = value["response"]
    value["audit"]["response_sha256"] = json_digest(response)
    value["receipt"]["output_sha256"] = hashlib.sha256(json.dumps(response).encode()).hexdigest()
    for record in (value["audit"], value["receipt"]):
        record["usage"] = deepcopy(response["usage"])
        record["finish_reason"] = response["choices"][0]["finish_reason"]


def test_original_png_preserved_with_full_occurrence_context_and_question_binding():
    blob, args = arguments()
    model = backend()
    request, provenance = model.prepare_request(blob, **args)
    assert base64.b64decode(request["messages"][0]["content"][1]["image_url"]["url"].split(",", 1)[1]) == blob
    assert request["max_tokens"] == 512 and len(request["messages"]) == 1
    assert provenance["request_sha256"] == json_digest(request)
    assert provenance["source"]["zip_member"] == "word/media/image1.png"
    assert provenance["representation"]["raw_bytes_unchanged"] is True
    assert provenance["context"] == args["context"]
    for change in (lambda a: a["image"].update(id="I00002"),
                   lambda a: a.update(question="还有可读文字吗？"),
                   lambda a: a["context"][0].update(text="图注乙"),
                   lambda a: a["image"].update(source_path="word/document.xml#p[9]/drawing")):
        changed = deepcopy(args)
        change(changed)
        other_request, other_provenance = model.prepare_request(blob, **changed)
        assert json_digest(other_request) != provenance["request_sha256"]
        assert other_provenance != provenance


def test_response_preview_retains_complete_real_provenance_without_any_generation():
    fake = FakeGeneration()
    model = backend(fake)
    blob, args = arguments()
    args["region"] = {"x": 1, "y": 1, "width": 3, "height": 2}
    _, expected_provenance = model.prepare_request(blob, **args)
    preview = model.preview_response(blob, **args)
    assert fake.calls == []
    assert preview["response"]["provenance"] == expected_provenance
    assert preview["response"]["provenance"]["context"] == args["context"]
    assert preview["response"]["provenance"]["representation"]["crop"] == args["region"]
    assert preview["output_reserve_tokens"] == 1024
    assert preview["preview_only"] is True and preview["reserve_is_worst_case_guarantee"] is False
    actual = model(blob, **args)
    assert set(preview["response"]) == set(actual)
    assert set(preview["response"]["generation"]) == set(actual["generation"])
    assert preview["response"]["provenance"] == actual["provenance"]
    assert preview["response"]["limitations"] == actual["limitations"]
    assert len(fake.calls) == 1  # Only the explicit actual call generates.


@pytest.mark.parametrize("output_tokens", [128, 512])
def test_preview_reserve_tracks_explicit_output_allowance(output_tokens):
    fake = FakeGeneration()
    blob, args = arguments()
    preview = backend(fake, output_tokens=output_tokens).preview_response(blob, **args)
    assert preview["output_reserve_tokens"] == output_tokens * 2
    assert preview["response"]["generation"]["usage"]["completion_tokens"] == output_tokens
    assert fake.calls == []


def test_preview_is_not_accepted_as_generation_proof_and_empty_observations_still_fail():
    fake = FakeGeneration()
    model = backend(fake)
    blob, args = arguments()
    request, provenance = model.prepare_request(blob, **args)
    preview = model.preview_response(blob, **args)
    assert preview["response"]["kind"] == "image_observation_preview"
    assert preview["response"]["verification"] == "not_generated"
    assert not preview["response"]["observation_id"].startswith("io_")
    for value in (preview, preview["response"]):
        with pytest.raises(VisionObservationError, match="CALLBACK_CONTRACT"):
            model.validate_generation(request, provenance, value)
    assert not fake.calls
    fake.content = deepcopy(preview["response"]["observation"])
    with pytest.raises(VisionObservationError, match="RESPONSE_FORMAT"):
        model(blob, **args)
    assert len(fake.calls) == 1


def test_preview_does_not_mutate_inputs_or_share_state_with_future_real_provenance():
    fake = FakeGeneration()
    model = backend(fake)
    blob, args = arguments()
    original = deepcopy(args)
    first = model.preview_response(blob, **args)
    first["response"]["provenance"]["context"][0]["text"] = "tampered preview"
    assert args == original
    second = model.preview_response(blob, **args)
    assert second["response"]["provenance"]["context"] == original["context"]
    assert fake.calls == []


def test_complete_preview_gate_can_reject_long_context_before_generation_without_truncation():
    # Synthetic counter and caller only; the real tool must use its exact model
    # tokenizer and add bind_image_observation to its complete outer envelope.
    count_tokens = lambda value: (len(json.dumps(value, ensure_ascii=False, sort_keys=True)) + 1) // 2
    fake = FakeGeneration()
    model = backend(fake)
    blob, args = arguments()
    args["context"] = []
    short = model.preview_response(blob, **args)
    def complete_outer(preview, context_refs):
        return {"ok": True, "status": "observed", "image_id": args["image"]["id"],
                "question": args["question"], "region": args["region"], "context_refs": context_refs,
                "judgment": preview["response"], "observation_id": "io_" + "a" * 64,
                "observation_sha256": "b" * 64, "cache": {"base_hit": False, "judgment_hit": False}}
    assert count_tokens(complete_outer(short, [])) + short["output_reserve_tokens"] < 4096
    text = "合成报告上下文；" * 400
    args["context"] = [{"evidence_id": "E000001", "start": 0, "end": len(text),
                        "text": text, "source_path": "word/document.xml#p[1]"}]
    assert count_tokens({"question": args["question"], "context": args["context"]}) < 4096
    preview = model.preview_response(blob, **args)
    outer = complete_outer(preview, ["E000001"])
    allowed = count_tokens(outer) + preview["output_reserve_tokens"] <= 4096
    if allowed:
        model(blob, **args)
    assert not allowed and fake.calls == []
    assert outer["judgment"]["provenance"]["context"][0]["text"] == text


@pytest.mark.parametrize("which", ["image", "context", "deadline"])
def test_preview_uses_same_input_validation_without_generation(which):
    fake, clock = FakeGeneration(), Clock()
    model = backend(fake, clock)
    blob, args = arguments()
    if which == "image":
        args["image"]["image_sha256"] = "0" * 64
    elif which == "context":
        args["context"][0]["end"] = 999
    else:
        clock.value = 101
    with pytest.raises(VisionObservationError):
        model.preview_response(blob, **args)
    assert not fake.calls


def test_crop_pixels_and_transformation_are_exact_without_silent_resize():
    blob, args = arguments()
    args["region"] = {"x": 2, "y": 1, "width": 3, "height": 4}
    request, provenance = backend().prepare_request(blob, **args)
    transmitted = base64.b64decode(request["messages"][0]["content"][1]["image_url"]["url"].split(",", 1)[1])
    actual = Image.open(io.BytesIO(transmitted))
    expected = Image.open(io.BytesIO(blob)).crop((2, 1, 5, 5)).convert("RGBA")
    assert actual.size == (3, 4) and actual.tobytes() == expected.tobytes()
    rep = provenance["representation"]
    assert rep["crop"] == args["region"] and rep["resize"] is None
    assert rep["original_blob_sha256"] == hashlib.sha256(blob).hexdigest()
    assert rep["payload_image_sha256"] == hashlib.sha256(transmitted).hexdigest()
    assert rep["raw_bytes_unchanged"] is False and rep["color_conversion"] == "RGBA"
    assert rep["transform_sha256"] == json_digest({k: v for k, v in rep.items() if k != "transform_sha256"})


@pytest.mark.parametrize("region", [None, [], {"x": True, "y": 0, "width": 1, "height": 1},
    {"x": 0, "y": 0, "width": 0, "height": 1}, {"x": -1, "y": 0, "width": 1, "height": 1},
    {"x": 7, "y": 0, "width": 2, "height": 1}, {"x": 0, "y": 0, "width": 2}])
def test_bad_regions_never_generate(region):
    fake = FakeGeneration()
    blob, args = arguments()
    args["region"] = region
    with pytest.raises(VisionObservationError, match="VISION_REGION"):
        backend(fake)(blob, **args)
    assert not fake.calls


@pytest.mark.parametrize("kind", ["truncated", "animated", "exif", "pixels"])
def test_unsupported_images_never_generate(kind):
    fake = FakeGeneration()
    options = {}
    if kind == "jpeg":
        out = io.BytesIO()
        Image.new("RGB", (2, 2)).save(out, format="JPEG")
        data = out.getvalue()
    elif kind == "truncated":
        data = png()[:35]
    elif kind == "animated":
        out = io.BytesIO()
        Image.new("RGB", (2, 2), "red").save(out, format="PNG", save_all=True,
                                               append_images=[Image.new("RGB", (2, 2), "blue")])
        data = out.getvalue()
    elif kind == "exif":
        data = png(orientation=6)
    else:
        data = png()
        options["max_pixels"] = 47
    blob, args = arguments(data)
    with pytest.raises(VisionObservationError, match="VISION_IMAGE"):
        backend(fake, **options)(blob, **args)
    assert not fake.calls


@pytest.mark.parametrize("amount", [0, 513, True, 1.5])
def test_output_budget_must_be_explicit_at_most_512(amount):
    with pytest.raises(VisionObservationError, match="CONFIGURATION"):
        backend(output_tokens=amount)


@pytest.mark.parametrize("mutation", [lambda i: i.update(model="external-vision"),
    lambda i: i.update(max_model_len=262144), lambda i: i.pop("processor_identity_sha256"),
    lambda i: i.update(service_identity_sha256="unknown")])
def test_configuration_cannot_fall_back_to_unpinned_model_or_processor(mutation):
    value = identity()
    mutation(value)
    with pytest.raises(VisionObservationError, match="IDENTITY"):
        OwnModelVisionBackend(FakeGeneration(), identity=value, document_sha256="a" * 64, deadline=100)
    with pytest.raises(VisionObservationError, match="IDENTITY"):
        OwnModelVisionBackend(None, identity=identity(), document_sha256="a" * 64, deadline=100)


@pytest.mark.parametrize("mutation", [lambda a: a["image"].update(image_sha256="b" * 64),
    lambda a: a["image"].update(target="https://external.example/image.png"),
    lambda a: a["image"].update(target="../../outside.png"),
    lambda a: a["image"].update(region={"x": 0, "y": 0, "width": 1, "height": 1}),
    lambda a: a["context"][0].update(end=99), lambda a: a["context"][0].update(start=True),
    lambda a: a["context"][0].update(text=float("nan"))])
def test_mismatched_source_and_context_fail_before_callback(mutation):
    fake = FakeGeneration()
    blob, args = arguments()
    mutation(args)
    with pytest.raises(VisionObservationError):
        backend(fake)(blob, **args)
    assert not fake.calls


def test_entire_json_request_byte_limit_includes_base64_and_context():
    blob, args = arguments()
    request, _ = backend().prepare_request(blob, **args)
    size = len(canonical_bytes(request))
    # The configured byte limit is part of backend identity, not prompt content.
    exact = backend(max_request_bytes=size)
    exact.prepare_request(blob, **args)
    fake = FakeGeneration()
    with pytest.raises(VisionObservationError, match="VISION_REQUEST_BYTES"):
        backend(fake, max_request_bytes=size - 1)(blob, **args)
    assert not fake.calls


def test_success_remains_uncertain_observation_and_no_cache_hides_calls():
    fake = FakeGeneration()
    fake.content = {"visible_content": [], "readable_text": [], "uncertainties": ["无法判读。"]}
    blob, args = arguments()
    model = backend(fake)
    first, second = model(blob, **args), model(blob, **args)
    assert len(fake.calls) == 2
    assert first["verification"] == "unverified" and first["semantic_review_required"] is True
    assert first["observation"] == fake.content and "facts" not in first
    assert first["generation"]["call_id"] == "synthetic-call-1"
    assert second["generation"]["call_id"] == "synthetic-call-2"
    assert first["observation_id"] != second["observation_id"]


@pytest.mark.parametrize("mutation", [
    lambda v: v["admission"].update(status="estimated"),
    lambda v: v["admission"].update(method="tokenize"),
    lambda v: v["admission"].update(image_payload_sha256="b" * 64),
    lambda v: v["admission"].update(processor_identity_sha256="b" * 64),
    lambda v: v["admission"].update(service_identity_sha256="b" * 64),
    lambda v: v["admission"].update(request_sha256="b" * 64),
    lambda v: v["admission"].update(expanded_prompt_tokens=65025),
    lambda v: v["admission"].update(expanded_prompt_tokens=True),
    lambda v: v["admission"].pop("final_prompt_token_ids_sha256"),
    lambda v: v["admission"].update(reserved_output_tokens=4096),
])
def test_missing_wrong_or_over_context_admission_cannot_be_claimed_success(mutation):
    fake = FakeGeneration()
    fake.mutate = lambda v, p, s: mutation(v)
    blob, args = arguments()
    with pytest.raises(VisionObservationError, match="ADMISSION"):
        backend(fake)(blob, **args)
    assert len(fake.calls) == 1  # The owner retains attempts even if validation fails.


@pytest.mark.parametrize("mutation", [
    lambda v: v["response"]["usage"].update(prompt_tokens=65),
    lambda v: v["response"]["usage"].update(completion_tokens=513),
    lambda v: v["response"]["usage"].update(total_tokens=0),
    lambda v: v["response"]["usage"].update(completion_tokens=True),
])
def test_actual_usage_must_match_processor_not_placeholder_count(mutation):
    fake = FakeGeneration()
    fake.mutate = lambda v, p, s: (mutation(v), sync_response_records(v))
    blob, args = arguments()
    with pytest.raises(VisionObservationError, match="USAGE"):
        backend(fake)(blob, **args)


def test_exact_expanded_context_boundary_reserves_all_512_output_tokens():
    fake = FakeGeneration()
    def boundary(value, payload, provenance):
        value["admission"]["expanded_prompt_tokens"] = 65024
        value["response"]["usage"].update(prompt_tokens=65024, total_tokens=65074)
        sync_response_records(value)
        for record in (value["audit"], value["receipt"]):
            record["prompt_tokens_checked"] = 65024
    fake.mutate = boundary
    blob, args = arguments()
    result = backend(fake)(blob, **args)
    assert result["generation"]["usage"]["prompt_tokens"] == 65024


@pytest.mark.parametrize("section", ["admission", "receipt", "audit"])
def test_no_synthetic_default_is_created_when_callback_proof_is_missing(section):
    fake = FakeGeneration()
    fake.mutate = lambda value, payload, provenance: value.pop(section)
    blob, args = arguments()
    with pytest.raises(VisionObservationError, match="CALLBACK_CONTRACT"):
        backend(fake)(blob, **args)
    assert len(fake.calls) == 1


@pytest.mark.parametrize("mutation", [
    lambda v: v["audit"].update(generation_attempted=False),
    lambda v: v["audit"].update(status="rejected"),
    lambda v: v["audit"].update(response_callback_completed=False),
    lambda v: v["audit"].update(response_sha256="b" * 64),
    lambda v: v["receipt"].update(call_id="another-call"),
    lambda v: v["receipt"].update(request_id="another-request"),
    lambda v: v["receipt"].update(kind="tokenize"),
    lambda v: v["receipt"].update(input_sha256="b" * 64),
    lambda v: v["receipt"].update(status="failed"),
    lambda v: v["receipt"].pop("output_sha256"),
])
def test_success_requires_matching_real_call_and_persisted_response_contract(mutation):
    fake = FakeGeneration()
    fake.mutate = lambda v, p, s: mutation(v)
    blob, args = arguments()
    with pytest.raises(VisionObservationError, match="RECEIPT"):
        backend(fake)(blob, **args)


@pytest.mark.parametrize("value", ["truncated", "duplicate", "extra", "empty", "tool", "other_model"])
def test_incomplete_or_non_observation_responses_rejected_without_semantic_keyword_rules(value):
    fake = FakeGeneration()
    def mutate(result, payload, provenance):
        response = result["response"]
        message = response["choices"][0]["message"]
        if value == "truncated":
            response["choices"][0]["finish_reason"] = "length"
        elif value == "duplicate":
            message["content"] = '{"visible_content":[],"visible_content":["x"],"readable_text":[],"uncertainties":[]}'
        elif value == "extra":
            message["content"] = '{"visible_content":[],"readable_text":[],"uncertainties":["x"],"recommendation":"repair"}'
        elif value == "empty":
            message["content"] = '{"visible_content":[],"readable_text":[],"uncertainties":[]}'
        elif value == "tool":
            message["tool_calls"] = [{"id": "x"}]
        else:
            response["model"] = "external-model"
        sync_response_records(result)
    fake.mutate = mutate
    blob, args = arguments()
    with pytest.raises(VisionObservationError) as caught:
        backend(fake)(blob, **args)
    if value == "truncated":
        assert caught.value.code == "VISION_RESPONSE_INCOMPLETE"
    assert len(fake.calls) == 1


def test_late_success_is_rejected_but_attempt_survives_and_no_automatic_retry():
    fake, clock = FakeGeneration(), Clock()
    fake.mutate = lambda v, p, s: setattr(clock, "value", 101)
    blob, args = arguments()
    with pytest.raises(VisionObservationError, match="DEADLINE"):
        backend(fake, clock)(blob, **args)
    assert len(fake.calls) == 1 and fake.calls[0]["deadline"] == 100
    with pytest.raises(VisionObservationError, match="DEADLINE"):
        backend(fake, clock)(blob, **args)
    assert len(fake.calls) == 1


def test_callback_failure_is_bounded_and_does_not_expose_private_error():
    calls = []
    def fail(payload, **kwargs):
        calls.append(payload)
        raise RuntimeError("PRIVATE_BASE64_OR_REPORT")
    blob, args = arguments()
    with pytest.raises(VisionObservationError) as captured:
        backend(fail)(blob, **args)
    error = captured.value.as_tool_error()
    assert error["error"]["code"] == "VISION_CALLBACK_FAILED"
    assert "PRIVATE" not in str(captured.value) and "PRIVATE" not in json.dumps(error)
    assert error["verification"] == "unverified" and len(calls) == 1


def test_callback_cannot_mutate_bound_request_or_provenance():
    fake = FakeGeneration()
    def mutate(v, payload, provenance):
        provenance["source"]["image_id"] = "I99999"
        payload["messages"][0]["content"][0]["text"] = "different request"
        v["audit"]["request_sha256"] = json_digest(payload)
        v["receipt"]["input_sha256"] = json_digest(payload)
        v["admission"]["request_sha256"] = json_digest(payload)
    fake.mutate = mutate
    blob, args = arguments()
    original = deepcopy(args)
    with pytest.raises(VisionObservationError, match="ADMISSION"):
        backend(fake)(blob, **args)
    assert args == original


def test_existing_document_tools_callable_contract_and_error_no_read(tmp_path):
    original = tmp_path / "synthetic.docx"
    image = tmp_path / "synthetic.png"
    image.write_bytes(png())
    doc = Document()
    doc.add_paragraph("合成图片的上下文")
    doc.add_picture(str(image))
    doc.save(original)
    manifest = build_manifest(original)
    fake = FakeGeneration()
    own = OwnModelVisionBackend(fake, identity=identity(), document_sha256=manifest["document_sha256"],
                                deadline=100, clock=Clock())
    reads = []
    tools = DocumentTools(manifest, count_tokens=len, source_path=original,
                          max_tool_tokens=15000, vision_backend=own, on_read=lambda *v: reads.append(v))
    image_id = manifest["images"][0]["id"]
    bad = tools.read_image(image_id, "观察颜色", '{"x":7,"y":0,"width":2,"height":1}', [])
    assert bad["ok"] is False and not reads and not fake.calls
    good = tools.read_image(image_id, "观察颜色", "{}", [])
    assert good["ok"] and good["verification"] == "unverified"
    assert good["judgment"]["kind"] == "image_observation"
    assert reads == [(image_id, 0, 1)] and len(fake.calls) == 1
    assert tools.image_observations[0]["judgment"]["generation"]["call_id"] == "synthetic-call-1"
    # This test does not claim the unchanged workspace now versions observations.
