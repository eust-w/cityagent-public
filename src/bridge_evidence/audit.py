from __future__ import annotations

import json
import re
import tarfile
import zipfile
from collections import Counter
from pathlib import Path, PurePosixPath
from typing import Any, Iterable

from docx import Document


REQUIRED_ACTIONS = (
    "review_summary",
    "review_details",
    "normalize_diseases",
    "write_decision",
    "finalize",
)
REQUIRED_TASKS = (
    "react-plan",
    "summary-evidence-embedding",
    "summary-evidence-rerank",
    "detail-evidence-embedding",
    "detail-evidence-rerank",
    "summary-review",
    "detail-review",
    "decision",
)
ALLOWED_TOP_LEVEL = {"code", "design", "result"}
FORBIDDEN_PATH_PARTS = {
    ".git",
    ".pytest_cache",
    ".mypy_cache",
    ".ruff_cache",
    "__pycache__",
    ".venv",
    "venv",
    "cache",
    "state",
    "logs",
}
TEXT_SUFFIXES = {
    ".cfg",
    ".ini",
    ".json",
    ".jsonl",
    ".md",
    ".py",
    ".sh",
    ".toml",
    ".txt",
    ".yaml",
    ".yml",
}
SECRET_PATTERNS = (
    re.compile(r"(?i)(?:api[_-]?key|authorization|bearer|token)\s*[:=]\s*['\"]?(?!\.\.\.|<[^>]+>|\$\{)[A-Za-z0-9_-]{16,}"),
    re.compile(r"\bsk-[A-Za-z0-9_-]{16,}\b"),
)


def _word_files(directory: Path) -> list[Path]:
    return sorted(path for path in directory.iterdir() if path.is_file() and path.suffix.lower() in {".doc", ".docx"})


def _expected_result_names(input_dir: Path) -> set[str]:
    return {f"{path.stem}.docx" for path in _word_files(input_dir)}


def _docx_signature(path: Path) -> dict[str, Any]:
    if not zipfile.is_zipfile(path):
        raise ValueError("不是有效的 DOCX/ZIP 文件")
    document = Document(path)
    tables: list[list[list[str]]] = []
    for table in document.tables:
        tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
    return {
        "tables": tables,
        "table_count": len(tables),
        "row_counts": [len(table) for table in tables],
        "paragraph_count": len(document.paragraphs),
    }


def _state_invocations(payload: dict[str, Any]) -> Counter[str]:
    metrics = payload.get("metrics") or {}
    counts: Counter[str] = Counter(metrics.get("task_calls") or {})
    counts.update(metrics.get("task_cache_hits") or {})
    return counts


def _successful_actions(payload: dict[str, Any]) -> Counter[str]:
    actions: Counter[str] = Counter()
    agent_run = payload.get("agent_run") or {}
    for step in agent_run.get("trajectory") or []:
        command = step.get("command") or {}
        result = step.get("result") or {}
        if result.get("success") is True:
            actions[str(command.get("action") or "")] += 1
    return actions


def online_audit_requirements(
    workflow_profile: str = "full",
    planner_mode: str = "model",
    semantic_fallback: bool = False,
) -> dict[str, set[str]]:
    """Return the network evidence required by an online workflow profile."""
    profile = str(workflow_profile or "full").strip().lower()
    planner = str(planner_mode or "model").strip().lower()
    if profile not in {"full", "summary", "details"}:
        raise ValueError(f"unsupported online workflow profile: {workflow_profile}")
    if planner not in {"model", "deterministic"}:
        raise ValueError(f"unsupported planner mode: {planner_mode}")

    actions: set[str] = {"finalize"}
    tasks: set[str] = set()
    if profile in {"full", "summary"}:
        actions.add("review_summary")
        tasks.update(
            {"summary-evidence-embedding", "summary-evidence-rerank", "summary-review"}
        )
    if profile in {"full", "details"}:
        actions.add("review_details")
        tasks.add("detail-review")
        if not semantic_fallback:
            tasks.update({"detail-evidence-embedding", "detail-evidence-rerank"})
    if profile == "full":
        actions.update({"normalize_diseases", "write_decision"})
        tasks.add("decision")
    if planner == "model":
        tasks.add("react-plan")

    return {
        "planner_models": {"deterministic"} if planner == "deterministic" else {"qwen", "deepseek"},
        "actions": actions,
        "tasks": tasks,
    }


def _state_has_semantic_fallback(agent_run: dict[str, Any]) -> bool:
    for step in agent_run.get("trajectory") or []:
        command = step.get("command") or {}
        result = step.get("result") or {}
        content = result.get("content") or {}
        if (
            command.get("action") == "review_details"
            and result.get("success") is True
            and isinstance(content, dict)
            and content.get("semantic_fallback")
        ):
            return True
    return False


def audit_states(
    state_dir: Path,
    expected_sources: set[str],
    *,
    require_online: bool,
    workflow_profile: str = "full",
    planner_mode: str = "model",
) -> tuple[dict[str, Any], list[str]]:
    errors: list[str] = []
    files = sorted(path for path in state_dir.glob("*.json") if path.is_file())
    source_names: set[str] = set()
    planners: Counter[str] = Counter()
    task_calls: Counter[str] = Counter()
    task_cache_hits: Counter[str] = Counter()
    model_calls: Counter[str] = Counter()
    model_cache_hits: Counter[str] = Counter()
    task_model_routes: Counter[str] = Counter()
    action_calls: Counter[str] = Counter()
    audit_passed = 0
    parse_failures: list[str] = []

    for path in files:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            parse_failures.append(path.name)
            errors.append(f"状态文件无法解析：{path.name}")
            continue
        source_name = Path(str(payload.get("source_path") or "")).name
        if source_name:
            source_names.add(source_name)
        agent_run = payload.get("agent_run") or {}
        planners[str(agent_run.get("planner") or "missing")] += 1
        metrics = payload.get("metrics") or {}
        task_calls.update(metrics.get("task_calls") or {})
        task_cache_hits.update(metrics.get("task_cache_hits") or {})
        model_calls.update(metrics.get("model_calls") or {})
        model_cache_hits.update(metrics.get("model_cache_hits") or {})
        task_model_routes.update(
            f"{task} -> {model}"
            for task, model in (metrics.get("task_models") or {}).items()
        )
        successful = _successful_actions(payload)
        action_calls.update(successful)

        if require_online:
            requirements = online_audit_requirements(
                workflow_profile=workflow_profile,
                planner_mode=planner_mode,
                semantic_fallback=_state_has_semantic_fallback(agent_run),
            )
            if agent_run.get("planner") not in requirements["planner_models"]:
                errors.append(
                    f"{path.name}：planner 与审计模式不匹配 "
                    f"(expected={sorted(requirements['planner_models'])}, actual={agent_run.get('planner')})"
                )
            if metrics.get("online_audit_passed") is not True:
                errors.append(f"{path.name}：online_audit_passed 不为 true")
            else:
                audit_passed += 1
            for task in sorted(requirements["tasks"]):
                if (metrics.get("task_calls") or {}).get(task, 0) <= 0:
                    errors.append(f"{path.name}：缺少真实网络模型任务 {task}")
                routed_model = (metrics.get("task_models") or {}).get(task)
                if not routed_model:
                    errors.append(f"{path.name}：模型任务 {task} 没有路由记录")
                elif (metrics.get("model_calls") or {}).get(routed_model, 0) <= 0:
                    errors.append(f"{path.name}：路由模型 {routed_model} 的真实网络调用为 0")
            for action in sorted(requirements["actions"]):
                if successful.get(action, 0) <= 0:
                    errors.append(f"{path.name}：缺少成功动作 {action}")

    if len(files) != len(expected_sources):
        errors.append(f"状态文件数量不匹配：expected={len(expected_sources)}, actual={len(files)}")
    missing_sources = sorted(expected_sources - source_names)
    extra_sources = sorted(source_names - expected_sources)
    if missing_sources:
        errors.append(f"状态缺少输入来源：{len(missing_sources)} 个")
    if extra_sources:
        errors.append(f"状态包含未知来源：{len(extra_sources)} 个")

    return {
        "files": len(files),
        "parse_failures": parse_failures,
        "planners": dict(planners),
        "task_calls": dict(task_calls),
        "task_cache_hits": dict(task_cache_hits),
        "model_calls": dict(model_calls),
        "model_cache_hits": dict(model_cache_hits),
        "task_model_routes": dict(task_model_routes),
        "successful_actions": dict(action_calls),
        "online_audit_passed_files": audit_passed,
        "missing_sources": missing_sources,
        "extra_sources": extra_sources,
    }, errors


def audit_results(result_dir: Path, expected_names: set[str]) -> tuple[dict[str, Any], list[str]]:
    errors: list[str] = []
    files = sorted(path for path in result_dir.iterdir() if path.is_file())
    # 官方新格式：result/ = 与测试集同名的 docx + infer_time.json（必需伴侣文件）
    infer_time_path = result_dir / "infer_time.json"
    infer_time_ok = False
    if infer_time_path.exists():
        try:
            payload = json.loads(infer_time_path.read_text(encoding="utf-8"))
            infer_time_ok = isinstance(payload.get("infer_time"), (int, float)) and payload["infer_time"] > 0
        except Exception:
            pass
    if not infer_time_ok:
        errors.append("infer_time.json 缺失或格式非法（要求 {\"infer_time\": 正数秒数}）")

    docx_files = [p for p in files if p.name != "infer_time.json"]
    actual_names = {path.name for path in docx_files}
    missing = sorted(expected_names - actual_names)
    extra = sorted(actual_names - expected_names)
    if missing:
        errors.append(f"result 缺少文件：{len(missing)} 个")
    if extra:
        errors.append(f"result 多出文件：{len(extra)} 个")

    invalid: list[dict[str, str]] = []
    table_counts: Counter[str] = Counter()
    for path in docx_files:
        if path.suffix.lower() != ".docx":
            invalid.append({"file": path.name, "error": "扩展名不是 .docx"})
            continue
        if path.stem.endswith("-信息提取报告"):
            invalid.append({"file": path.name, "error": "错误继承标签后缀"})
            continue
        try:
            signature = _docx_signature(path)
            table_counts[str(signature["table_count"])] += 1
            if signature["table_count"] != 3:
                invalid.append({"file": path.name, "error": f"表格数量为 {signature['table_count']}，要求 3"})
        except Exception as exc:
            invalid.append({"file": path.name, "error": str(exc)})
    if invalid:
        errors.append(f"result 中有 {len(invalid)} 个结构无效文件")

    return {
        "files": len(files),
        "expected": len(expected_names),
        "missing": missing,
        "extra": extra,
        "invalid": invalid,
        "table_counts": dict(table_counts),
    }, errors


def _safe_archive_name(name: str) -> PurePosixPath | None:
    normalized = name[2:] if name.startswith("./") else name
    if not normalized or normalized == ".":
        return None
    path = PurePosixPath(normalized)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError(f"不安全的 tar 成员路径：{name}")
    return path


def audit_archive(archive: Path) -> tuple[dict[str, Any], list[str]]:
    errors: list[str] = []
    members: list[str] = []
    top_levels: set[str] = set()
    forbidden: list[str] = []
    unsafe: list[str] = []
    try:
        archive_members: list[tuple[str, bool]] = []
        if zipfile.is_zipfile(archive):
            with zipfile.ZipFile(archive) as handle:
                archive_members = [
                    (
                        member.filename,
                        ((member.external_attr >> 16) & 0o170000) == 0o120000,
                    )
                    for member in handle.infolist()
                ]
        else:
            with tarfile.open(archive, "r:gz") as handle:
                archive_members = [
                    (member.name, member.issym() or member.islnk())
                    for member in handle.getmembers()
                ]

        for member_name, is_link in archive_members:
            try:
                path = _safe_archive_name(member_name)
            except ValueError:
                unsafe.append(member_name)
                continue
            if path is None:
                continue
            members.append(path.as_posix())
            top_levels.add(path.parts[0])
            lowered_parts = {part.lower() for part in path.parts}
            if lowered_parts & FORBIDDEN_PATH_PARTS or "信息提取报告" in path.name and path.parts[0] == "result":
                forbidden.append(path.as_posix())
            if is_link:
                forbidden.append(path.as_posix())
    except Exception as exc:
        errors.append(f"压缩包无法读取：{exc}")
        return {"members": 0, "top_levels": [], "forbidden": [], "unsafe": []}, errors

    missing_top = sorted(ALLOWED_TOP_LEVEL - top_levels)
    extra_top = sorted(top_levels - ALLOWED_TOP_LEVEL)
    if missing_top:
        errors.append("压缩包缺少顶层目录：" + ", ".join(missing_top))
    if extra_top:
        errors.append("压缩包存在额外顶层目录：" + ", ".join(extra_top))
    if forbidden:
        errors.append(f"压缩包包含 {len(set(forbidden))} 个禁止成员")
    if unsafe:
        errors.append(f"压缩包包含 {len(unsafe)} 个不安全路径")
    return {
        "members": len(members),
        "top_levels": sorted(top_levels),
        "missing_top_levels": missing_top,
        "extra_top_levels": extra_top,
        "forbidden": sorted(set(forbidden)),
        "unsafe": unsafe,
    }, errors


def _text_payloads(paths: Iterable[Path]) -> Iterable[tuple[str, str]]:
    for root in paths:
        if root.is_file():
            candidates = [root]
        elif root.is_dir():
            candidates = [path for path in root.rglob("*") if path.is_file()]
        else:
            continue
        for path in candidates:
            if path.suffix.lower() in TEXT_SUFFIXES:
                try:
                    yield str(path), path.read_text(encoding="utf-8")
                except UnicodeDecodeError:
                    continue
            elif path.suffix.lower() == ".docx" and zipfile.is_zipfile(path):
                try:
                    with zipfile.ZipFile(path) as archive:
                        xml = "\n".join(
                            archive.read(name).decode("utf-8", errors="ignore")
                            for name in archive.namelist()
                            if name.endswith(".xml")
                        )
                    yield str(path), xml
                except Exception:
                    continue


def audit_secrets(paths: Iterable[Path]) -> tuple[dict[str, Any], list[str]]:
    findings: list[str] = []
    scanned = 0
    for name, text in _text_payloads(paths):
        scanned += 1
        if any(pattern.search(text) for pattern in SECRET_PATTERNS):
            findings.append(name)
    errors = [f"发现疑似密钥材料：{len(findings)} 个文件"] if findings else []
    return {"scanned_files": scanned, "finding_files": sorted(findings)}, errors


def compare_result_directories(current: Path, reference: Path) -> dict[str, Any]:
    current_files = {path.name: path for path in current.glob("*.docx")}
    reference_files = {path.name: path for path in reference.glob("*.docx")}
    shared = sorted(current_files.keys() & reference_files.keys())
    changes: list[dict[str, Any]] = []
    unchanged = 0
    for name in shared:
        current_signature = _docx_signature(current_files[name])
        reference_signature = _docx_signature(reference_files[name])
        if current_signature == reference_signature:
            unchanged += 1
            continue
        current_tables = current_signature["tables"]
        reference_tables = reference_signature["tables"]
        changed_cells = 0
        for table_index in range(max(len(current_tables), len(reference_tables))):
            left = current_tables[table_index] if table_index < len(current_tables) else []
            right = reference_tables[table_index] if table_index < len(reference_tables) else []
            for row_index in range(max(len(left), len(right))):
                left_row = left[row_index] if row_index < len(left) else []
                right_row = right[row_index] if row_index < len(right) else []
                changed_cells += sum(
                    (left_row[index] if index < len(left_row) else None)
                    != (right_row[index] if index < len(right_row) else None)
                    for index in range(max(len(left_row), len(right_row)))
                )
        changes.append({
            "file": name,
            "changed_cells": changed_cells,
            "current_row_counts": current_signature["row_counts"],
            "reference_row_counts": reference_signature["row_counts"],
        })
    return {
        "shared_files": len(shared),
        "unchanged_files": unchanged,
        "changed_files": len(changes),
        "missing_from_current": sorted(reference_files.keys() - current_files.keys()),
        "new_in_current": sorted(current_files.keys() - reference_files.keys()),
        "changes": changes,
    }


def run_submission_audit(
    *,
    input_dir: Path,
    result_dir: Path,
    state_dir: Path,
    require_online: bool,
    online_workflow_profile: str = "full",
    planner_mode: str = "model",
    package_root: Path | None = None,
    archive: Path | None = None,
    reference_result_dir: Path | None = None,
) -> dict[str, Any]:
    expected_sources = {path.name for path in _word_files(input_dir)}
    expected_results = _expected_result_names(input_dir)
    errors: list[str] = []

    if state_dir.exists():
        states, state_errors = audit_states(
            state_dir,
            expected_sources,
            require_online=require_online,
            workflow_profile=online_workflow_profile,
            planner_mode=planner_mode,
        )
    else:
        # 新 vllm 微调模型流水线（run.py）不产出逐份状态文件，跳过该项
        states, state_errors = {"skipped": "vllm pipeline, no state files"}, []
    results, result_errors = audit_results(result_dir, expected_results)
    errors.extend(state_errors)
    errors.extend(result_errors)
    report: dict[str, Any] = {
        "passed": False,
        "require_online": require_online,
        "online_workflow_profile": online_workflow_profile,
        "planner_mode": planner_mode,
        "expected_inputs": len(expected_sources),
        "states": states,
        "results": results,
    }

    secret_targets = [state_dir, result_dir]
    if package_root:
        actual_top = {path.name for path in package_root.iterdir()}
        missing_top = sorted(ALLOWED_TOP_LEVEL - actual_top)
        extra_top = sorted(actual_top - ALLOWED_TOP_LEVEL)
        package_section = {"top_levels": sorted(actual_top), "missing": missing_top, "extra": extra_top}
        report["package_root"] = package_section
        if missing_top:
            errors.append("打包目录缺少顶层目录：" + ", ".join(missing_top))
        if extra_top:
            errors.append("打包目录存在额外顶层目录：" + ", ".join(extra_top))
        secret_targets.append(package_root)
    if archive:
        archive_section, archive_errors = audit_archive(archive)
        report["archive"] = archive_section
        errors.extend(archive_errors)
    secrets, secret_errors = audit_secrets(secret_targets)
    report["secrets"] = secrets
    errors.extend(secret_errors)
    if reference_result_dir:
        report["reference_diff"] = compare_result_directories(result_dir, reference_result_dir)

    report["errors"] = errors
    report["passed"] = not errors
    return report


def render_audit_markdown(report: dict[str, Any]) -> str:
    status = "通过" if report.get("passed") else "失败"
    states = report.get("states") or {}
    results = report.get("results") or {}
    lines = [
        "# IAIC 提交与 Agent 运行审计",
        "",
        f"- 结论：{status}",
        f"- 预期输入：{report.get('expected_inputs', 0)}",
        f"- 状态文件：{states.get('files', 0)}",
        f"- 结果文件：{results.get('files', 0)}",
        f"- DOCX 表格计数：{json.dumps(results.get('table_counts', {}), ensure_ascii=False)}",
        f"- Planner：{json.dumps(states.get('planners', {}), ensure_ascii=False)}",
        f"- 模型任务（网络）：{json.dumps(states.get('task_calls', {}), ensure_ascii=False)}",
        f"- 模型任务（缓存）：{json.dumps(states.get('task_cache_hits', {}), ensure_ascii=False)}",
        f"- 模型调用（网络）：{json.dumps(states.get('model_calls', {}), ensure_ascii=False)}",
        f"- 模型调用（缓存）：{json.dumps(states.get('model_cache_hits', {}), ensure_ascii=False)}",
        f"- 任务模型路由：{json.dumps(states.get('task_model_routes', {}), ensure_ascii=False)}",
        "",
        "## 错误",
        "",
    ]
    errors = report.get("errors") or []
    lines.extend(f"- {error}" for error in errors)
    if not errors:
        lines.append("- 无")
    if report.get("reference_diff"):
        diff = report["reference_diff"]
        lines.extend([
            "",
            "## 相对参考版本",
            "",
            f"- 共同文件：{diff.get('shared_files', 0)}",
            f"- 内容变化文件：{diff.get('changed_files', 0)}",
            f"- 内容未变化文件：{diff.get('unchanged_files', 0)}",
        ])
    return "\n".join(lines) + "\n"
