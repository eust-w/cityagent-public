"""No model: document-bound material navigation and exact delivery receipts."""
from copy import deepcopy
import json

from docx import Document
import pytest

from bridge_evidence.contest.document_material_index import DocumentMaterialIndex
from bridge_evidence.contest.document_source_delivery import SourceDeliveryLedger
from bridge_evidence.contest.document_tools import DocumentTools
from bridge_evidence.contest.evidence_manifest import build_manifest
from bridge_evidence.contest.tong_react_adapter import wrap_tool_response
from bridge_evidence.contest import tong_react_adapter as adapter
from test_document_material_index import make_material_case


def _build(tmp_path, *, ledger=False, budget=12000, mutate=None):
    case = make_material_case(tmp_path)
    if mutate:
        mutate(case)
    index = DocumentMaterialIndex(case['manifest'], case['materials'], source_path=case['source_path'])
    delivery = SourceDeliveryLedger(case['manifest'], tmp_path/'delivery', run_id='run', agent_id='child') if ledger else None
    tools = DocumentTools(case['manifest'], count_tokens=len, max_tool_tokens=budget,
        material_index=index, delivery_ledger=delivery, native_response_budget=True)
    return case, index, tools, delivery


def _confirm(ledger, page, request):
    payload = {'messages': [{'role': 'tool', 'tool_call_id': 'native', 'content': json.dumps(
        wrap_tool_response('read_source', page), ensure_ascii=False)}]}
    ledger.observe_request(payload, {'request_id': request})
    ledger.confirm_request(request, evidence={'input_delivery_confirmed': True, 'request_id': request})


def _read_all(tools, name, args, ledger=None):
    pages = []
    for i in range(300):
        page = tools.automatic_handlers()[name](**args)
        assert page['ok'], page
        assert len(json.dumps(wrap_tool_response(name, page), ensure_ascii=False,
            sort_keys=True, separators=(',', ':'))) <= tools.max_tool_tokens
        pages.append(page)
        if ledger:
            _confirm(ledger, page, f'request-{len(ledger.state()["requests"])}')
        if page['continuation'] is None:
            return pages
        assert 'cursor' not in page['continuation']['arguments']
        assert 'start' not in page['continuation']['arguments']
    pytest.fail('Material read did not terminate')


def test_cell_packet_and_old_e_return_exact_text_and_compact_context(tmp_path):
    case, index, tools, _ = _build(tmp_path)
    ref = case['refs']['nested_table']+'/r1/c1'
    page = tools.automatic_handlers()['read_source'](ref)
    assert page['ok'] and any(s['text']=='94.57' for s in page['spans'])
    assert any(s['role']=='header_candidate' for s in page['context'])
    assert any(s['role']=='parent_cell' for s in page['context'])
    metadata = json.dumps(page['material_context'], ensure_ascii=False)
    assert 'xml' not in metadata and 'source_provenance' not in metadata
    assert '94.57' not in metadata and '父表说明' not in metadata
    eid = index.resolve(ref)['spans'][0]['evidence_id']
    assert tools.automatic_handlers()['read_source'](eid)['ok']
    packet = _read_all(tools, 'read_source', {'evidence_id': '#/packets/P0001'})
    assert packet[0]['material_context']['unresolved']
    assert any(s['text']=='桥梁甲，2018年检测。' for p in packet for s in p['spans'])


def test_unresolved_reference_preserves_safe_fallback_and_does_not_advance(tmp_path):
    case, index, tools, _ = _build(tmp_path)
    before = tools.navigation_state()
    result = tools.automatic_handlers()['read_source'](case['refs']['ambiguous'])
    assert not result['ok'] and result['status']=='unresolved_reference'
    assert len(result['fallback']['evidence_ids'])==2
    assert result['unresolved'][0]['code']=='ambiguous_xml_text_mapping'
    assert all('text' not in span for issue in result['unresolved'] for span in issue.get('candidate_spans', []))
    assert tools.navigation_state()==before
    pages = _read_all(tools, 'read_section', {'section_id':'#/sections/S0001', 'max_units':2})
    assert pages[0]['material_context']['unresolved']


def test_local_qualified_aliases_share_progress_foreign_sha_rejected_and_restart_explicit(tmp_path):
    case, index, tools, _ = _build(tmp_path)
    read = tools.automatic_handlers()['read_source']
    first = read(case['refs']['unique'])
    assert first['ok']
    assert read(index.qualify_ref(case['refs']['unique']))['status']=='already_delivered'
    before = tools.navigation_state()
    assert not read('mineru:'+'0'*64+case['refs']['unique'])['ok']
    assert tools.navigation_state()==before
    assert read(case['refs']['unique'], restart=True)['spans']==first['spans']


def test_table_alias_has_one_canonical_automatic_position(tmp_path):
    case, index, tools, _ = _build(tmp_path)
    first = tools.automatic_handlers()['read_table_rows'](case['refs']['table'], max_rows=1)
    second = tools.automatic_handlers()['read_table_rows']('T00001', max_rows=1)
    assert first['requested_rows']==[0,1]
    assert second['requested_rows']==[1,2]
    assert '轻微锈蚀' in second['spans'][0]['text']


def test_pending_material_pages_replay_and_restore_before_confirming(tmp_path):
    case, index, tools, ledger = _build(tmp_path, ledger=True)
    ref='#/sections/S0001'
    first=tools.automatic_handlers()['read_section'](ref,max_units=1)
    again=tools.automatic_handlers()['read_section'](ref,max_units=1)
    assert again['replayed'] and again['spans']==first['spans']
    assert not ledger.confirmed_ranges()
    restored=DocumentTools(case['manifest'],count_tokens=len,max_tool_tokens=12000,
        material_index=index,delivery_ledger=ledger,native_response_budget=True)
    restored.automatic_handlers(state=tools.navigation_state())
    replay=restored.automatic_handlers()['read_section'](ref,max_units=1)
    assert replay['source_delivery_receipt']==first['source_delivery_receipt']
    _confirm(ledger,replay,'accepted')
    next_page=restored.automatic_handlers()['read_section'](ref,max_units=1)
    assert next_page['source_delivery_receipt']!=first['source_delivery_receipt']
    changed=deepcopy(case['materials']); changed['navigation_note']='new material revision'
    other=DocumentMaterialIndex(case['manifest'],changed,source_path=case['source_path'])
    with pytest.raises(ValueError,match='material index'):
        DocumentTools(case['manifest'],count_tokens=len,material_index=other).automatic_handlers(state=tools.navigation_state())


def _long(case):
    doc=Document(case['source_path'])
    doc.tables[0].cell(0,0).text='长表头甲乙'*1500
    doc.tables[0].cell(1,1).text='病害长记录甲乙'*1800
    doc.save(case['source_path'])
    case['manifest']=build_manifest(case['source_path'])
    case['materials']['docx_sha256']=case['manifest']['document_sha256']
    case['materials']['packets'][0]['docx_sha256']=case['manifest']['document_sha256']
    case['materials']['sections'][0]['title']='长章节导航标题'*1500


def test_long_metadata_header_and_cell_paginate_losslessly_with_real_receipt(tmp_path):
    case,index,tools,ledger=_build(tmp_path,ledger=True,budget=2100,mutate=_long)
    pages=_read_all(tools,'read_section',{'section_id':'#/sections/S0001','max_units':2},ledger)
    assert len(pages)>10 and any(p.get('context_part') for p in pages)
    assert all('source_delivery_receipt' in p for p in pages)
    expected=index.section('#/sections/S0001')
    confirmed=ledger.confirmed_ranges()
    for span in expected['spans']+expected['context']:
        assert any(a<=span['start']<=span['end']<=b for a,b in confirmed[span['evidence_id']])
    assert not any('xml' in json.dumps(p.get('material_context',{})) for p in pages)


def test_material_search_keeps_scope_and_original_excerpt(tmp_path):
    case,index,tools,_=_build(tmp_path)
    result=tools.automatic_handlers()['search_in_document']('锈蚀',section_id='#/sections/S0001')
    assert result['ok'] and result['matches']
    for hit in result['matches']:
        unit=tools._index[hit['evidence_id']]
        assert hit['text']==unit['text'][hit['start']:hit['end']]
        assert all(ref.startswith('mineru:'+index.document_sha256) for ref in hit['material_refs'])
    assert not tools.automatic_handlers()['search_in_document']('重复说明',section_id='#/sections/S0001')['matches']


def test_partial_mapped_paragraph_search_does_not_include_unmapped_tail(tmp_path):
    def partial(case):
        segment=next(s for s in case['materials']['packets'][0]['segments'] if s['ref']=='#/texts/11')
        segment['text']='桥梁甲';segment['offset']=[0,3]
    case,index,tools,_=_build(tmp_path,mutate=partial)
    search=tools.automatic_handlers()['search_in_document']
    assert not search('2018',section_id='#/sections/S0001')['matches']
    assert search('2018',scope='document')['matches']
    matched=search('桥梁',section_id='#/sections/S0001')
    assert matched['matches'][0]['text']=='桥梁甲'


def test_native_material_arguments_work_without_cursor_or_source_hash_fields(tmp_path):
    case,index,tools,_=_build(tmp_path)
    session={'handlers':tools.automatic_handlers(),'tool_calls':[],'max_calls':100,
        'tool_budget_exceeded':False,'count_tokens':len,'max_tool_tokens':12000}
    for name,arguments in [
        ('read_source',{'evidence_id':'#/packets/P0001'}),
        ('read_section',{'section_id':index.qualify_ref('#/sections/S0001'),'max_units':1}),
        ('read_table_rows',{'table_id':case['refs']['table'],'max_rows':1}),
        ('search_in_document',{'query':['锈蚀'],'section_id':'#/sections/S0001'}),
    ]:
        result=adapter._invoke_active(session,name,arguments)
        assert result['ok'],result
        assert tools.receipts[-1]['result_tokens']==len(adapter._json_bytes(result).decode())


def test_material_page_failure_does_not_publish_or_advance(tmp_path,monkeypatch):
    case,index,tools,ledger=_build(tmp_path,ledger=True)
    before=tools.navigation_state()
    original=ledger.prepare_page
    def fail(*args,**kwargs):
        raise OSError('simulated disk failure before page publication')
    monkeypatch.setattr(ledger,'prepare_page',fail)
    read=tools.automatic_handlers()['read_source']
    result=read(case['refs']['unique'])
    assert not result['ok'] and tools.navigation_state()==before
    assert not ledger.confirmed_ranges()
    monkeypatch.setattr(ledger,'prepare_page',original)
    result=read(case['refs']['unique'])
    assert result['ok'] and result['spans'][0]['start']==0


def test_material_alias_corrupt_canonical_text_is_rejected_without_advancing(tmp_path,monkeypatch):
    case,index,tools,ledger=_build(tmp_path,ledger=True)
    original=index.resolve
    def corrupt(ref,**kwargs):
        result=original(ref,**kwargs)
        result['spans'][0]['text']='模型补写'
        return result
    monkeypatch.setattr(index,'resolve',corrupt)
    before=tools.navigation_state()
    result=tools.automatic_handlers()['read_source'](case['refs']['unique'])
    assert not result['ok'] and 'differs from the canonical' in result['error']['message']
    assert tools.navigation_state()==before and not ledger.confirmed_ranges()
