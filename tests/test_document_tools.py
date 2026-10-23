import hashlib
import json

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
import pytest

from bridge_evidence.contest.document_tools import DocumentTools, make_tesseract_backend
from bridge_evidence.contest.evidence_manifest import build_manifest, manifest_from_text
from bridge_evidence.contest import tong_react_adapter as adapter


def _tokens(result):
    return len(json.dumps(result, ensure_ascii=False, sort_keys=True, separators=(",", ":")))


def _native_session(tools, counter, budget):
    # Both production dispatch layers run; no SDK import/model request occurs.
    return {"handlers": tools.automatic_handlers(), "tool_calls": [], "max_calls": 1000,
            "tool_budget_exceeded": False, "count_tokens": counter, "max_tool_tokens": budget}


@pytest.mark.parametrize("counter,budget", [(len, 4096), (lambda text: len(text.encode("utf-8")), 750)])
def test_native_source_pages_fit_actual_dispatch_envelope_and_reconstruct_without_loss(counter, budget):
    source = '甲😀带"引号"与\\路径' * 600
    reads = []
    tools = DocumentTools(manifest_from_text(source), count_tokens=counter, max_tool_tokens=budget,
                          native_response_budget=True, on_read=lambda *span: reads.append(span))
    session = _native_session(tools, counter, budget)
    call = {"tool": "read_source", "arguments": {"evidence_id": "E000001"}}
    received = []
    while call:
        result = adapter._invoke_active(session, call["tool"], call["arguments"])
        assert result["ok"], result
        payload = result["result"]
        actual = counter(adapter._json_bytes(result).decode("utf-8"))
        assert actual <= budget and tools.receipts[-1]["result_tokens"] == actual
        assert tools.receipts[-1]["result_token_scope"] == "native_response"
        received.extend(payload["spans"])
        call = payload["continuation"]
    assert "".join(span["text"] for span in received) == source
    assert reads == [(span["evidence_id"], span["start"], span["end"]) for span in received]
    assert reads[0][1] == 0 and reads[-1][2] == len(source)
    assert all(left[2] == right[1] for left, right in zip(reads, reads[1:]))
    assert len(received) > 1 and all(row["status"] == "success" for row in session["tool_calls"])
    assert tools._response_tool is None


def test_native_section_continuation_counts_the_current_tool_name():
    def counter(text):
        # Synthetic tokenizer makes fixed reserves/per-tool-name guesses fail.
        obj = json.loads(text)
        return len(text) + {"read_section": 73, "read_source": 11}.get(obj.get("tool"), 0)

    source = "长段原文" * 1000
    reads = []
    tools = DocumentTools(manifest_from_text(source + "\n末段"), count_tokens=counter,
                          max_tool_tokens=900, native_response_budget=True,
                          on_read=lambda *span: reads.append(span))
    session = _native_session(tools, counter, 900)
    result = adapter._invoke_active(session, "read_section", {"section_id": "root", "max_units": 2})
    assert result["ok"], result
    payload = result["result"]
    assert tools.receipts[-1]["result_tokens"] == counter(adapter._json_bytes(result).decode())
    # Long source pages retain their private section resume position; the
    # model keeps calling the same high-level section tool without a cursor.
    assert "resume_section" not in payload
    text = "".join(span["text"] for span in payload["spans"])
    call = payload["continuation"]
    while call:
        assert call["tool"] == "read_section" and "cursor" not in call["arguments"]
        result = adapter._invoke_active(session, call["tool"], call["arguments"])
        assert result["ok"], result
        assert tools.receipts[-1]["result_tokens"] == counter(adapter._json_bytes(result).decode()) <= 900
        text += "".join(span["text"] for span in result["result"]["spans"])
        call = result["result"]["continuation"]
    assert text == source + "\n末段"  # The exact long unit and section tail both arrive.
    assert reads[-1] == ("E000002", 0, 2)


def test_native_over_budget_nonpageable_read_does_not_grant_coverage():
    reads = []
    tools = DocumentTools(manifest_from_text("甲" * 2000), count_tokens=len, max_tool_tokens=700,
                          native_response_budget=True, on_read=lambda *span: reads.append(span))
    result = adapter._invoke_active(_native_session(tools, len, 700), "get_neighbors", {"evidence_id": "E000001"})
    assert not result["ok"] and result["result"]["status"] == "budget_exceeded"
    assert not reads and tools.receipts[-1]["ranges"] == []
    assert len(adapter._json_bytes(result).decode()) <= 700
    assert tools._response_tool is None


def test_standalone_budget_default_remains_payload_only():
    source = "甲" * 10000
    legacy = DocumentTools(manifest_from_text(source), count_tokens=len, max_tool_tokens=4096)
    native = DocumentTools(manifest_from_text(source), count_tokens=len, max_tool_tokens=4096,
                           native_response_budget=True)
    bare = legacy.read_source("E000001", 0, len(source))
    wrapped = native.read_source("E000001", 0, len(source))
    assert _tokens(bare) == 4096
    assert _tokens(adapter.wrap_tool_response("read_source", bare)) > 4096
    assert _tokens(adapter.wrap_tool_response("read_source", wrapped)) == 4096
    assert bare["spans"][0]["end"] > wrapped["spans"][0]["end"]
    assert legacy.receipts[-1]["result_token_scope"] == "payload"


@pytest.mark.parametrize("value", [None, 0, 1, "true"])
def test_native_budget_flag_requires_a_boolean(value):
    with pytest.raises(ValueError, match="native_response_budget"):
        DocumentTools(manifest_from_text("甲"), count_tokens=len, native_response_budget=value)


def _image_document(tmp_path):
    path = tmp_path / "images.docx"
    picture = tmp_path / "image.png"
    Image.new("RGB", (10, 10), "blue").save(picture)
    doc = Document()
    doc.add_paragraph("左幅 2018 年")
    doc.add_picture(str(picture))
    doc.add_paragraph("右幅 2020 年")
    doc.add_picture(str(picture))
    doc.save(path)
    return path, build_manifest(path)


def test_source_pagination_exactly_preserves_characters_and_receipt_ranges():
    source = "桥梁😀裂缝\r\n" * 300
    manifest = manifest_from_text(source.replace("\r\n", "长行"))
    seen = []
    tools = DocumentTools(manifest, count_tokens=len, max_tool_tokens=750,
                          on_read=lambda *span: seen.append(span))
    call = {"tool": "read_source", "arguments": {"evidence_id": "E000001", "start": 0, "end": len(manifest["units"][0]["text"])}}
    chunks = []
    while call:
        result = tools.handlers()[call["tool"]](**call["arguments"])
        assert result["ok"], result
        assert _tokens(result) <= 750
        chunks.extend(span["text"] for span in result["spans"])
        call = result["continuation"]
    assert "".join(chunks) == manifest["units"][0]["text"]
    cursor = 0
    for evidence_id, start, end in seen:
        assert evidence_id == "E000001" and start == cursor and end > start
        cursor = end
    assert cursor == len(manifest["units"][0]["text"])
    receipt = tools.receipts[-1]
    assert receipt["arguments"]["end"] == cursor
    assert receipt["ranges"][0]["end"] == cursor
    assert receipt["seconds"] >= 0
    assert len(receipt["result_sha256"]) == 64


def test_section_without_outline_and_oversize_single_unit_has_recovery():
    manifest = manifest_from_text("很长的正文" * 300 + "\n最后一段")
    tools = DocumentTools(manifest, count_tokens=len, max_tool_tokens=1200)
    page = tools.read_section("root", 0, 3)
    assert page["ok"] and page["continuation"]["tool"] == "read_source"
    assert page["resume_section"]["arguments"]["cursor"] == 1
    reconstructed = page["spans"][0]["text"]
    continuation = page["continuation"]
    while continuation:
        page2 = tools.handlers()[continuation["tool"]](**continuation["arguments"])
        assert page2["ok"]
        reconstructed += page2["spans"][0]["text"]
        continuation = page2["continuation"]
    assert reconstructed == manifest["units"][0]["text"]
    tail = tools.handlers()[page["resume_section"]["tool"]](**page["resume_section"]["arguments"])
    assert tail["spans"][0]["text"] == "最后一段"
    assert tail["continuation"] is None


def test_sections_include_subheadings_but_stop_at_next_peer(tmp_path):
    path = tmp_path / "headings.docx"
    doc = Document()
    doc.add_paragraph("前言")
    doc.add_heading("章节一", 1)
    doc.add_paragraph("正文甲")
    doc.add_heading("子章节", 2)
    doc.add_paragraph("正文乙")
    doc.add_heading("章节二", 1)
    doc.save(path)
    manifest = build_manifest(path)
    tools = DocumentTools(manifest, count_tokens=len, max_tool_tokens=5000)
    ids = []
    call = {"tool": "read_section", "arguments": {"section_id": "E000002", "cursor": 0, "max_units": 2}}
    while call:
        result = tools.handlers()[call["tool"]](**call["arguments"])
        ids.extend(span["evidence_id"] for span in result["spans"])
        call = result["continuation"]
    assert ids == ["E000002", "E000003", "E000004", "E000005"]
    assert tools.read_section("root", 6, 1)["spans"] == []
    assert tools.read_section("root", -1, 1)["ok"] is False


def test_table_context_carries_headers_units_notes_and_merged_instance(tmp_path):
    path = tmp_path / "table.docx"
    doc = Document()
    doc.add_paragraph("表 3 检测记录，单位：mm")
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "位置"
    table.cell(0, 1).text = "宽度（mm）"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(1, 0).merge(table.cell(2, 0)).text = "左幅"
    table.cell(1, 1).text = "0.1"
    table.cell(2, 1).text = "0.2"
    table.cell(3, 0).text = "右幅"
    table.cell(3, 1).text = "0.3"
    doc.add_paragraph("注：宽度为实测最大值。")
    doc.add_paragraph("备注：本次检测为 2018 年。")
    doc.save(path)
    manifest = build_manifest(path)
    seen = []
    tools = DocumentTools(manifest, count_tokens=len, max_tool_tokens=10000,
                          on_read=lambda *span: seen.append(span))
    result = tools.read_table_rows("T00001", 2, 3)
    assert result["ok"]
    context = "\n".join(span["text"] for span in result["context"])
    assert "宽度（mm）" in context and "单位：mm" in context
    assert "实测最大值" in context and "2018 年" in context
    merged = result["spans"][0]["cell_context"][0]
    assert merged["merged_context_text"] == "左幅"
    assert "word/document.xml#" in merged["merged_from_source"]
    assert seen == [(result["spans"][0]["evidence_id"], 0, len(" | 0.2"))]
    assert all(item["grants_coverage"] is False for item in result["context"])


def test_table_rows_paginate_at_physical_row_boundary(tmp_path):
    path = tmp_path / "table-pagination.docx"
    doc = Document()
    table = doc.add_table(rows=8, cols=1)
    for index, row in enumerate(table.rows):
        row.cells[0].text = "字段" if index == 0 else f"第 {index} 行 " + "很长的病害描述" * 20
    doc.save(path)
    tools = DocumentTools(build_manifest(path), count_tokens=len, max_tool_tokens=2800)
    seen = []
    call = {"tool": "read_table_rows", "arguments": {"table_id": "T00001", "start": 1, "end": 8}}
    while call:
        result = tools.handlers()[call["tool"]](**call["arguments"])
        assert result["ok"], result
        assert _tokens(result) <= 2800
        seen.extend(span["row_index"] for span in result["spans"])
        call = result["continuation"]
    assert seen == list(range(1, 8))


def test_oversize_header_does_not_silently_disappear(tmp_path):
    path = tmp_path / "large-header.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "表头" * 1000
    table.cell(1, 0).text = "正文"
    doc.save(path)
    seen = []
    tools = DocumentTools(build_manifest(path), count_tokens=len, max_tool_tokens=1100,
                          on_read=lambda *span: seen.append(span))
    result = tools.read_source("E000002", 0, 2)
    assert result["ok"] is False and result["status"] == "budget_exceeded"
    assert result["error"]["recoverable"] is True and not seen
    assert _tokens(result) <= 1100
    pages = _follow_recovery(tools, result["continuation"])
    texts = _reconstructed_originals(pages)
    assert texts == {"E000001": "表头" * 1000, "E000002": "正文"}
    metadata = _reconstructed_context(pages)
    assert metadata["tables"][0]["header_evidence_ids"] == ["E000001"]
    assert metadata["primary"][0]["evidence_id"] == "E000002"
    assert metadata["primary"][0]["cell_context"][0]["column_index"] == 0
    assert all(not identifier.startswith("@") for identifier, _, _ in seen)


def _follow_recovery(tools, call, *, recreate=None):
    pages = []
    for _ in range(500):
        if call is None:
            return pages
        if recreate is not None:
            tools = recreate()
        page = tools.handlers()[call["tool"]](**call["arguments"])
        assert page["ok"], page
        assert _tokens(page) <= tools.max_tool_tokens
        pages.append(page)
        call = page["continuation"]
    pytest.fail("Recovery must advance rather than repeat the same over-budget read")


def _reconstructed_originals(pages):
    pieces = {}
    for page in pages:
        for span in page["spans"]:
            pieces.setdefault(span["evidence_id"], []).append(span)
    result = {}
    for identifier, spans in pieces.items():
        cursor = spans[0]["start"]
        for span in spans:
            assert span["start"] == cursor
            assert span["end"] - span["start"] == len(span["text"])
            cursor = span["end"]
        result[identifier] = "".join(span["text"] for span in spans)
    return result


def _reconstructed_context(pages):
    parts = [page["context_part"] for page in pages if "context_part" in page]
    cursor = 0
    for part in parts:
        assert part["start"] == cursor
        cursor = part["end"]
    assert cursor == parts[-1]["total_chars"]
    return json.loads("".join(part["text"] for part in parts))


def test_large_headers_notes_and_merged_values_recover_with_all_original_provenance(tmp_path):
    from bridge_evidence.contest.document_workspace import DocumentWorkspace

    path = tmp_path / "large-context.docx"
    doc = Document()
    doc.add_paragraph("单位：mm；" + "本表列间对应说明。" * 200)
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "位置及对象" * 200
    table.cell(0, 1).text = "宽度（mm）；" + "完整列名限定" * 180
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(1, 0).merge(table.cell(2, 0)).text = "左侧构件身份" * 200
    table.cell(1, 1).text = "0.1"
    table.cell(2, 1).text = "0.2"
    doc.add_paragraph("注：数值为最大宽度；" + "必须保留的长注释。" * 180)
    doc.add_paragraph("备注：2018 年；" + "历史对比单位说明。" * 200)
    doc.save(path)
    manifest = build_manifest(path)
    workspace = DocumentWorkspace(manifest, tmp_path / "workspace-context", {"test": "synthetic"})
    tools = DocumentTools(manifest, count_tokens=len, max_tool_tokens=1000, on_read=workspace.record_read)
    body = next(unit for unit in manifest["units"] if unit.get("row_index") == 2)
    failed = tools.read_source(body["id"], 0, len(body["text"]))
    assert not failed["ok"] and workspace.snapshot()["reads"] == {}
    pages = _follow_recovery(tools, failed["continuation"])
    originals = _reconstructed_originals(pages)
    assert originals == {unit["id"]: unit["text"] for unit in manifest["units"]}
    metadata = _reconstructed_context(pages)
    primary = metadata["primary"][0]
    assert primary["table_id"] == body["table_id"] and primary["row_index"] == 2
    assert primary["cell_context"][0]["merged_context_text"] == "左侧构件身份" * 200
    assert primary["cell_context"][0]["merged_from_source"] == body["cells"][0]["merged_from_source"]
    assert primary["cell_context"][1]["column_index"] == 1
    assert {item["evidence_id"] for item in metadata["context"]} == set(originals) - {body["id"]}
    assert all(page["spans"] == [] and page["grants_coverage"] is False
               for page in pages if "context_part" in page)
    for unit in manifest["units"]:
        assert workspace.snapshot()["reads"][unit["id"]] == [[0, len(unit["text"])]]
    assert not workspace.delivery_status()["ready"]  # Reading is not semantic acceptance.


def test_recovery_continuations_survive_tool_reconstruction_and_keep_requested_subrange(tmp_path):
    path = tmp_path / "recovery-restore.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "完整表头😀与单位" * 200
    table.cell(1, 0).text = "前缀-要读取的病害-后缀"
    doc.save(path)
    manifest = build_manifest(path)
    seen = []

    def recreate():
        return DocumentTools(manifest, count_tokens=len, max_tool_tokens=900,
                             on_read=lambda *args: seen.append(args))

    tools = recreate()
    failed = tools.read_source("E000002", 3, 10)
    assert not failed["ok"]
    pages = _follow_recovery(tools, failed["continuation"], recreate=recreate)
    originals = _reconstructed_originals(pages)
    assert originals["E000001"] == manifest["units"][0]["text"]
    assert originals["E000002"] == manifest["units"][1]["text"][3:10]
    body_ranges = [(start, end) for identifier, start, end in seen if identifier == "E000002"]
    assert body_ranges[0][0] == 3 and body_ranges[-1][1] == 10
    metadata = _reconstructed_context(pages)
    assert metadata["primary"][0]["start"] == 3 and metadata["primary"][0]["end"] == 10


@pytest.mark.parametrize("budget", [180, 350])
def test_recovery_envelope_too_large_requires_configuration_change_without_retry_loop(tmp_path, budget):
    path = tmp_path / "small-budget.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "完整表头" * 400
    table.cell(1, 0).text = "病害"
    doc.save(path)
    tools = DocumentTools(build_manifest(path), count_tokens=len, max_tool_tokens=budget)
    result = tools.read_source("E000002", 0, 2)
    for _ in range(200):
        assert _tokens(result) <= budget
        call = result.get("continuation")
        if call is None:
            break
        result = tools.handlers()[call["tool"]](**call["arguments"])
    else:
        pytest.fail("An impossible protocol envelope must not cause an endless retry")
    assert not result["ok"] and result["status"] == "budget_exceeded"
    assert result["error"]["recoverable"] is False
    assert "budget" in result["error"]["message"].lower()


def test_large_table_row_recovery_continues_every_requested_physical_row(tmp_path):
    path = tmp_path / "table-recovery.docx"
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text = "对象限定" * 300
    table.cell(0, 1).text = "长度（m）"
    for row in range(1, 4):
        table.cell(row, 0).text = f"第{row}项对象"
        table.cell(row, 1).text = str(row)
    doc.save(path)
    manifest = build_manifest(path)
    tools = DocumentTools(manifest, count_tokens=len, max_tool_tokens=950)
    failed = tools.read_table_rows("T00001", 1, 4)
    assert not failed["ok"] and failed["continuation"]["tool"] == "read_section"
    pages = _follow_recovery(tools, failed["continuation"])
    assert _reconstructed_originals(pages) == {unit["id"]: unit["text"] for unit in manifest["units"]}
    metadata = _reconstructed_context(pages)
    assert [span["row_index"] for span in metadata["primary"]] == [1, 2, 3]
    assert all(span["table_id"] == "T00001" for span in metadata["primary"])


def test_section_context_recovery_retains_outer_resume_without_skipping_text(tmp_path):
    path = tmp_path / "section-recovery.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "单位和表头" * 250
    table.cell(1, 0).text = "正文病害"
    doc.add_paragraph("章节尾部")
    doc.save(path)
    manifest = build_manifest(path)
    tools = DocumentTools(manifest, count_tokens=len, max_tool_tokens=900)
    failed = tools.read_section("root", 1, 1)
    assert not failed["ok"]
    assert failed["resume_section"] == {"tool": "read_section", "arguments": {
        "section_id": "root", "cursor": 2, "max_units": 1}}
    pages = _follow_recovery(tools, failed["continuation"])
    assert _reconstructed_originals(pages)["E000002"] == "正文病害"
    tail = tools.handlers()[failed["resume_section"]["tool"]](**failed["resume_section"]["arguments"])
    assert tail["ok"] and tail["spans"][0]["text"] == "章节尾部"


def test_search_returns_exact_occurrences_without_claiming_full_coverage():
    original = "正文" * 500 + "这里有 CrAcK 病害" + "后文" * 500
    seen = []
    tools = DocumentTools(manifest_from_text(original), count_tokens=len, max_tool_tokens=1200,
                          on_read=lambda *span: seen.append(span))
    result = tools.search_in_document("crack")
    hit = result["matches"][0]
    assert original[hit["start"]:hit["end"]] == hit["text"]
    assert original[hit["match_start"]:hit["match_end"]] == "CrAcK"
    assert result["grants_coverage"] is False and seen == []
    assert tools.receipts[-1]["ranges"] == []
    assert tools.search_in_document(" ")["ok"] is False


def test_search_reports_omitted_matches_and_outline_preserves_root_navigation():
    tools = DocumentTools(manifest_from_text("病害\n" * 100), count_tokens=len, max_tool_tokens=650)
    result = tools.search_in_document("病害")
    assert result["total_matches"] == 100 and result["omitted_matches"] > 0
    assert len(result["matches"]) + result["omitted_matches"] == 100
    outline = tools.read_outline()
    assert outline["ok"]
    # The public outline attribute remains complete for local navigation.
    assert tools.outline["unit_count"] == 100
    assert tools.outline["sections"][0]["id"] == "root"


def test_full_text_has_stable_ids_and_never_implies_reading():
    seen = []
    tools = DocumentTools(manifest_from_text("甲\r\n乙"), count_tokens=len,
                          on_read=lambda *span: seen.append(span))
    assert "[E000001 chars 0:3; source=" in tools.full_text()
    assert "甲\r\n" in tools.full_text() and "E000002" in tools.full_text()
    assert not seen and not tools.receipts
    assert set(tools.handlers()) == {"read_outline", "read_source", "read_section", "read_table_rows",
                                     "search_in_document", "get_neighbors", "read_image"}


@pytest.mark.parametrize("start,end", [(-1, 1), (0, 100), (2, 1), (True, 1)])
def test_bad_character_ranges_are_recoverable_and_recorded(start, end):
    tools = DocumentTools(manifest_from_text("甲乙"), count_tokens=len)
    result = tools.read_source("E000001", start, end)
    assert result["ok"] is False
    assert len(tools.receipts) == 1 and tools.receipts[0]["ranges"] == []


def test_unavailable_image_backend_does_not_imply_understanding(tmp_path):
    path, manifest = _image_document(tmp_path)
    seen = []
    tools = DocumentTools(manifest, source_path=path, count_tokens=len,
                          on_read=lambda *span: seen.append(span))
    result = tools.read_image("I00001", "裂缝在哪里？", "{}", ["E000001"])
    assert result["ok"] is False and result["status"] == "unavailable"
    assert result["verification"] == "unverified" and not seen
    assert tools.outline["unverified_image_count"] == 2
    assert tools.image_observations == []


def test_image_base_cache_deduplicates_pixels_but_judgment_binds_context_and_instance(tmp_path):
    path, manifest = _image_document(tmp_path)

    class Backend:
        cache_key = "synthetic-v1"

        def __init__(self):
            self.recognitions, self.judgments = [], []

        def recognize(self, blob, *, region):
            self.recognitions.append((hashlib.sha256(blob).hexdigest(), region))
            return {"text": "合成图片"}

        def interpret(self, observation, *, question, context, image):
            self.judgments.append((question, context, image["id"]))
            return {"statement": "待复核", "image": image["id"]}

    backend = Backend()
    read_receipts = []
    tools = DocumentTools(manifest, source_path=path, count_tokens=len, vision_backend=backend,
                          on_read=lambda *args: read_receipts.append(args))
    first = tools.read_image("I00001", "描述", "{}", ["E000001"])
    repeat = tools.read_image("I00001", "描述", "{}", ["E000001"])
    other = tools.read_image("I00002", "描述", "{}", ["E000003"])
    changed = tools.read_image("I00001", "另一个问题", "{}", ["E000001"])
    assert all(result["ok"] for result in (first, repeat, other, changed))
    assert len(backend.recognitions) == 1 and len(backend.judgments) == 3
    assert repeat["cache"] == {"base_hit": True, "judgment_hit": True}
    assert other["cache"] == {"base_hit": True, "judgment_hit": False}
    assert first["source_path"] != other["source_path"]
    assert first["verification"] == other["verification"] == "unverified"
    assert tools.receipts[-1]["ranges"] == []
    assert tools.receipts[-1]["image_ids"] == ["I00001"]
    assert read_receipts == [("I00001", 0, 1), ("I00001", 0, 1), ("I00002", 0, 1), ("I00001", 0, 1)]
    assert tools.image_observations == [first, repeat, other, changed]
    assert first["question"] == "描述" and first["backend"]["identity"] == "synthetic-v1"
    assert first["backend"]["identity_source"] == "cache_key"
    assert first["backend"]["version"] is None  # Never invent a model/backend version.
    first["judgment"]["statement"] = "调用方修改"
    assert tools.image_observations[0]["judgment"]["statement"] == "待复核"
    assert tools.read_image("I00001", "描述", '{"x":0,"y":0,"width":5,"height":5}', ["E000001"])["ok"]
    assert len(backend.recognitions) == 2
    backend.version = "v2"
    changed_backend = tools.read_image("I00001", "描述", "{}", ["E000001"])
    assert changed_backend["ok"] and changed_backend["backend"]["version"] == "v2"
    assert changed_backend["cache"] == {"base_hit": False, "judgment_hit": False}
    assert len(backend.recognitions) == 3


@pytest.mark.parametrize("target", ["../../../etc/passwd", "https://example.test/image.png", "/etc/passwd", "..\\secret"])
def test_unsafe_image_targets_never_reach_backend(tmp_path, target):
    path, manifest = _image_document(tmp_path)
    manifest["images"][0]["target"] = target
    calls = []
    tools = DocumentTools(manifest, source_path=path, count_tokens=len,
                          vision_backend=lambda *args, **kwargs: calls.append(True))
    result = tools.read_image("I00001", "描述", "{}", [])
    assert result["ok"] is False and not calls


def test_changed_docx_is_rejected_even_after_an_image_cache_hit(tmp_path):
    path, manifest = _image_document(tmp_path)
    tools = DocumentTools(manifest, source_path=path, count_tokens=len,
                          vision_backend=lambda *args, **kwargs: {"observation": "合成"})
    assert tools.read_image("I00001", "描述", "{}", [])["ok"]
    path.write_bytes(path.read_bytes() + b"changed")
    result = tools.read_image("I00001", "描述", "{}", [])
    assert result["ok"] is False and "hash" in result["error"]["message"]


def test_explicit_ocr_factory_does_not_fake_missing_executable_or_languages(monkeypatch):
    import bridge_evidence.contest.document_tools as module

    monkeypatch.setattr(module.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="not installed"):
        make_tesseract_backend()
    monkeypatch.setattr(module.shutil, "which", lambda name: "/fake/tesseract")
    monkeypatch.setattr(module.subprocess, "run", lambda *args, **kwargs: type("Result", (), {"stdout": "eng\n", "stderr": ""})())
    with pytest.raises(RuntimeError, match="chi_sim"):
        make_tesseract_backend()


def test_failed_vision_response_is_not_cached_as_a_success(tmp_path):
    path, manifest = _image_document(tmp_path)
    calls = []

    def unavailable(*args, **kwargs):
        calls.append(True)
        return {"ok": False, "error": "backend failed"}

    tools = DocumentTools(manifest, source_path=path, count_tokens=len, vision_backend=unavailable)
    for _ in range(2):
        result = tools.read_image("I00001", "描述", "{}", [])
        assert result["ok"] is False
        assert tools.receipts[-1]["status"] == "error"
    assert len(calls) == 2
    assert tools.image_observations == []


def test_over_budget_image_observation_does_not_grant_read_or_store_success(tmp_path):
    path, manifest = _image_document(tmp_path)
    reads = []
    tools = DocumentTools(manifest, source_path=path, count_tokens=len, max_tool_tokens=700,
                          on_read=lambda *args: reads.append(args),
                          vision_backend=lambda *args, **kwargs: {"text": "图片内容" * 1000})
    result = tools.read_image("I00001", "描述", "{}", [])
    assert result["ok"] is False and result["status"] == "budget_exceeded"
    assert not reads and not tools.image_observations
    assert tools.receipts[-1]["image_ids"] == []


def test_native_image_borderline_envelope_rejects_without_coverage_or_observation(tmp_path):
    path, manifest = _image_document(tmp_path)

    def backend(*args, **kwargs):
        return {"text": "合成图片观察，不代表语义正确" * 10}

    bare_tools = DocumentTools(manifest, source_path=path, count_tokens=len, vision_backend=backend)
    bare = bare_tools.read_image("I00001", "描述", "{}", [])
    budget = _tokens(bare)
    assert bare["ok"] and _tokens(adapter.wrap_tool_response("read_image", bare)) > budget
    reads = []
    tools = DocumentTools(manifest, source_path=path, count_tokens=len, vision_backend=backend,
                          max_tool_tokens=budget, native_response_budget=True,
                          on_read=lambda *span: reads.append(span))
    arguments = {"image_id": "I00001", "question": "描述", "region_json": "{}", "context_refs": []}
    result = adapter._invoke_active(_native_session(tools, len, budget), "read_image", arguments)
    assert not result["ok"] and result["result"]["status"] == "budget_exceeded"
    assert not reads and not tools.image_observations and not tools.receipts[-1]["image_ids"]
    assert _tokens(result) <= budget

    # Grant coverage only when the complete observation really survives both layers.
    bigger_budget = _tokens(adapter.wrap_tool_response("read_image", bare))
    accepted_reads = []
    accepted = DocumentTools(manifest, source_path=path, count_tokens=len, vision_backend=backend,
                            max_tool_tokens=bigger_budget, native_response_budget=True,
                            on_read=lambda *span: accepted_reads.append(span))
    delivered = adapter._invoke_active(_native_session(accepted, len, bigger_budget), "read_image", arguments)
    assert delivered["ok"] and delivered["result"] == bare
    assert accepted_reads == [("I00001", 0, 1)]
    assert accepted.image_observations == [bare]
    assert accepted.receipts[-1]["result_tokens"] == _tokens(delivered) == bigger_budget


def test_successful_image_read_satisfies_workspace_read_receipt_only(tmp_path):
    from bridge_evidence.contest.document_workspace import DocumentWorkspace

    path, manifest = _image_document(tmp_path)
    workspace = DocumentWorkspace(manifest, tmp_path / "workspace", {"test": "synthetic"})
    tools = DocumentTools(manifest, source_path=path, count_tokens=len,
                          on_read=workspace.record_read,
                          vision_backend=lambda *args, **kwargs: {"text": "识别结果待复核"})
    result = tools.read_image("I00001", "描述", "{}", [])
    assert result["ok"] and result["verification"] == "unverified"
    snapshot = workspace.snapshot()
    assert snapshot["image_reads"] == ["I00001"]
    assert "I00001" in snapshot["reading_progress"]["review_required_evidence_ids"]
    assert not workspace.delivery_status()["ready"]


def test_smallest_accepted_budget_still_bounds_error_responses():
    for budget in range(100, 600):
        try:
            tools = DocumentTools(manifest_from_text("甲"), count_tokens=len, max_tool_tokens=budget)
        except ValueError:
            continue
        result = tools.read_source("X" * 1000, 0, 1)
        assert result["ok"] is False and _tokens(result) <= budget
        result = tools.read_source("E000001", 0, 1)
        assert _tokens(result) <= budget
        break
    else:
        pytest.fail("No valid tool budget")


def test_ocr_only_observation_stays_semantically_unverified(tmp_path):
    path, manifest = _image_document(tmp_path)

    class OCR:
        def recognize(self, blob, *, region):
            return {"text": "识别字符", "verification": "unverified"}

    tools = DocumentTools(manifest, source_path=path, count_tokens=len, vision_backend=OCR())
    result = tools.read_image("I00001", "能判断裂缝原因吗", "{}", [])
    assert result["ok"] and result["verification"] == "unverified"
    assert result["judgment"]["status"] == "unavailable"
    assert result["semantic_review_required"]


def test_read_neighbors_does_not_cross_body_and_header_story(tmp_path):
    path = tmp_path / "header.docx"
    doc = Document()
    doc.add_paragraph("正文")
    doc.sections[0].header.paragraphs[0].text = "页眉"
    doc.save(path)
    tools = DocumentTools(build_manifest(path), count_tokens=len)
    result = tools.get_neighbors("E000001")
    assert [span["text"] for span in result["spans"]] == ["正文"]
