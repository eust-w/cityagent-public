"""Receipt-driven continuation and lossless evidence; no model fact verdicts."""
from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from bridge_evidence.contest.document_merge_progress import (
    build_handoff_evidence, child_fact_id, MergeWorkspaceTools, next_merge_action,
    inspection_progress, semantic_state_fingerprint, MergeProgressStalled, ConsecutiveProgressGuard, next_blocker_action)
from bridge_evidence.contest.document_multiagent import MergeWorkflow, import_candidates
from bridge_evidence.contest.document_entries import EntryStore, EntryStoreError
from bridge_evidence.contest.document_workspace import DocumentWorkspace
from bridge_evidence.contest.evidence_manifest import manifest_from_text


def current_file():
    state = {'version': 7, 'draft_sha256': 'draft', 'questions': {},
             'draft': {'基本信息': {}, '病害明细': [], '维修建议': []}}
    artifact = {'workspace_version': 7, 'draft_sha256': 'draft', 'file_sha256': 'file',
                'structure_valid': True, 'content_validated': True}
    pages = [{**artifact, 'view': view, 'cursor': 0, 'end': 20, 'total_chars': 20}
             for view in ('basic', 'defects')]
    return state, artifact, pages


def test_only_missing_advice_inspection_is_requested_not_whole_task():
    state, artifact, pages = current_file()
    action = next_merge_action(state, {'ready': True}, [], artifact, pages)
    assert action['stage'] == 'inspect_current_file'
    assert action['next_action'] == {'tool': 'inspect_delivery', 'arguments': {'view': 'advice', 'limit': 8}}
    assert action['file_inspection']['pending_reads'] == [{'view': 'advice', 'cursor': 0, 'limit': 8}]
    assert action['file_inspection']['completed_views'] == ['basic', 'defects']
    assert action['semantic_approval'] is False


def test_partial_pages_continue_at_first_gap_and_finish_needs_model_judgment():
    state, artifact, pages = current_file()
    pages += [{**artifact, 'view': 'advice', 'cursor': a, 'end': b, 'total_chars': 30}
              for a, b in [(20, 30), (0, 10)]]
    action = next_merge_action(state, {'ready': True}, [], artifact, pages)
    assert action['next_action'] == {'tool': 'inspect_delivery', 'arguments': {'view': 'advice', 'limit': 8}}
    assert action['file_inspection']['pending_reads'] == [{'view': 'advice', 'cursor': 10, 'limit': 8}]
    pages.append({**artifact, 'view': 'advice', 'cursor': 10, 'end': 20, 'total_chars': 30})
    action = next_merge_action(state, {'ready': True}, [], artifact, pages)
    assert action['next_action'] == {'tool': 'finish', 'program_binds_inspected_file': True}
    assert action['file_inspection']['pending_reads'] == []
    assert action['self_check_submitted'] is False and action['self_check_accepted'] is False


@pytest.mark.parametrize('field,value', [('workspace_version', 6), ('draft_sha256', 'old'),
    ('file_sha256', 'old'), ('requested_failure_marker', True), ('content_validated', False)])
def test_stale_or_failed_export_never_counts_as_current_inspection(field, value):
    state, artifact, pages = current_file()
    artifact[field] = value
    progress = inspection_progress(state, artifact, pages)
    assert not progress['completed_views']
    assert len(progress['pending_reads']) == 3


def test_invalid_page_does_not_falsely_complete_empty_view():
    state, artifact, pages = current_file()
    pages.append({**artifact, 'view': 'advice', 'cursor': -1, 'end': 0, 'total_chars': 0})
    assert next_merge_action(state, {'ready': True}, [], artifact, pages)['stage'] == 'inspect_current_file'
    pages.append({**artifact, 'view': 'advice', 'cursor': 0, 'end': 0, 'total_chars': 0})
    assert next_merge_action(state, {'ready': True}, [], artifact, pages)['next_action']['tool'] == 'finish'


def test_pending_question_navigation_skips_already_resolved_questions():
    state, artifact, pages = current_file()
    state['questions'] = {'a': {'valid': True, 'status': 'resolved'}, 'b': {'valid': True, 'status': 'open'}}
    action = next_merge_action(state, {'ready': False}, [], artifact, pages)
    assert action['next_action']['arguments'] == {'view': 'unresolved_questions', 'limit': 4}
    assert action['next_action']['target_question_id'] == 'b'
    assert action['unresolved_question_ids'] == ['b']


def test_noop_versions_and_notes_cannot_mask_stall_but_real_progress_resets_it(tmp_path):
    state, artifact, pages = current_file()
    merge = object.__new__(MergeWorkflow)
    merge.directory = tmp_path
    merge.workspace = SimpleNamespace(snapshot=lambda: deepcopy(state), delivery_status=lambda: {'ready': True})
    merge.child_summary = []
    merge._check_budget = lambda: None
    merge._finish_is_current = lambda: False
    merge.artifact_delivery, merge.delivery_inspections = artifact, pages
    merge.model_self_check = None
    merge._progress_guard = ConsecutiveProgressGuard()
    merge._progress_pages, merge._progress_events = set(), []
    merge._active_agent_call_id = None
    assert merge._before_native_model() is None
    for i in range(1, 7):
        state['working_notes'] = {'note': str(i)}
        merge._active_agent_call_id = str(i)
        result = merge._before_native_model()
        assert (result is not None) == (i == 6)
    # Multiple calls at the same generation boundary cannot increment the guard.
    assert merge._before_native_model() is None
    state['draft']['基本信息']['桥梁名称'] = '模型实际保存的值'
    merge._active_agent_call_id = '7'
    assert merge._before_native_model() is None
    assert merge._progress_guard.state['unchanged'] == 0
    for i in range(8, 19):
        merge._active_agent_call_id = str(i)
        merge._before_native_model()
    merge._active_agent_call_id = '19'
    with pytest.raises(MergeProgressStalled, match='Twelve model rounds'):
        merge._before_native_model()
    receipt = json.loads((tmp_path / 'merge-stalled.json').read_text())
    assert receipt['actual_delivery_complete'] is False
    assert receipt['pending_action']['next_action']['tool'] == 'inspect_delivery'
    assert state['draft']['基本信息']['桥梁名称'] == '模型实际保存的值'


def test_valid_page_progress_never_hits_stall_and_batch_save_wins():
    guard = ConsecutiveProgressGuard()
    guard.observe(None, 'initial')
    for i in range(20):
        # Fingerprint is taken after the complete tool batch, including saves.
        assert guard.observe(str(i), 'new-page-or-save-' + str(i)) is None
    assert guard.state['unchanged'] == 0


def progress_harness(tmp_path):
    state, artifact, pages = current_file()
    state['reads'] = {}
    delivery = {'ready': False, 'blockers': ['unread_source_ranges'],
        'reading_progress': {'unread_ranges': [{'evidence_id': 'E473', 'start': 0, 'end': 100}]}}
    merge = object.__new__(MergeWorkflow)
    merge.directory = tmp_path
    merge.workspace = SimpleNamespace(snapshot=lambda: deepcopy(state), delivery_status=lambda: deepcopy(delivery))
    merge.child_summary, merge.context_state = [], None
    merge._check_budget, merge._finish_is_current = lambda: None, lambda: False
    merge.artifact_delivery, merge.delivery_inspections = artifact, pages
    merge.model_self_check = None
    merge._progress_guard = ConsecutiveProgressGuard()
    merge._progress_pages, merge._progress_events = set(), []
    merge.gateway_receipts, merge._active_agent_call_id = [], None
    assert merge._before_native_model() is None

    def boundary():
        merge.gateway_receipts.append({'generation_attempted': True,
            'request_id': 'request-' + str(len(merge.gateway_receipts) + 1)})
        return merge._before_native_model()
    return merge, state, boundary


def source_page(start=0, *, replayed=False):
    return {'ok': True, 'tool': 'read_source', 'result': {'ok': True, 'replayed': replayed,
        'source_delivery_receipt': 'synthetic-page-' + str(start),
        'spans': [{'evidence_id': 'E473', 'start': start, 'end': start + 1, 'text': '源'}]}}


def test_new_pending_source_at_eleven_stalls_reaches_next_model_without_faking_read(tmp_path):
    merge, state, boundary = progress_harness(tmp_path)
    for _ in range(11):
        boundary()
    assert merge._progress_guard.state['unchanged'] == 11
    merge._observe_native_tool('read_source', {'evidence_id': 'E473'}, source_page())
    # Liveness only: actual source coverage is still awaiting gateway input proof.
    assert state['reads'] == {}
    assert boundary() is None
    assert merge._progress_guard.state['unchanged'] == 0


@pytest.mark.parametrize('replayed', [False, True])
def test_same_source_page_replay_cannot_reset_stall(tmp_path, replayed):
    merge, state, boundary = progress_harness(tmp_path)
    merge._observe_native_tool('read_source', {'evidence_id': 'E473'}, source_page())
    boundary()
    for _ in range(11):
        merge._observe_native_tool('read_source', {'evidence_id': 'E473'}, source_page(replayed=replayed))
        boundary()
    assert len(merge._progress_pages) == 1 and state['reads'] == {}
    with pytest.raises(MergeProgressStalled):
        boundary()


@pytest.mark.parametrize('kind', ['source', 'workspace_fragment', 'workspace_items'])
def test_more_than_twelve_distinct_pages_count_as_exploration_even_when_text_repeats(tmp_path, kind):
    merge, state, boundary = progress_harness(tmp_path)
    for index in range(20):
        if kind == 'source':
            merge._observe_native_tool('read_source', {'evidence_id': 'E473'}, source_page(index))
        else:
            page = ({'mode': 'canonical_json_fragment', 'item_index': 0, 'item_start': index,
                'item_end': index + 1, 'fragment': '同'} if kind == 'workspace_fragment' else
                {'items': [{'value': '同'}], 'start': index, 'end': index + 1})
            merge._observe_native_tool('read_workspace_page', {'view': 'facts'},
                {'ok': True, 'tool': 'read_workspace_page', 'result': {'ok': True, **page}})
        assert boundary() is None
    assert merge._progress_guard.state['unchanged'] == 0
    assert len(merge._progress_pages) == 20


@pytest.mark.parametrize('replayed', [False, True])
def test_main_first_read_of_twenty_child_covered_pages_is_progress_but_replay_still_stops(tmp_path, replayed):
    merge, state, boundary = progress_harness(tmp_path)
    # This union records input delivered to children, not the merger's personal
    # reading. The main Agent may need a long reread to resolve a conflict.
    state['reads'] = {'E473': [[0, 100]]}
    state['questions'] = {'year-conflict': {'id': 'year-conflict', 'valid': True,
        'status': 'conflict', 'question': '子Agent对两次检测年份的关联不一致'}}
    merge.workspace.delivery_status = lambda: {'ready': False,
        'blockers': ['unresolved_question:year-conflict']}
    before_reads = deepcopy(state['reads'])
    for index in range(20):
        merge._observe_native_tool('read_source', {'evidence_id': 'E473'}, source_page(index))
        assert boundary() is None
    assert merge._progress_guard.state['unchanged'] == 0
    assert len(merge._progress_pages) == 20 and state['reads'] == before_reads
    for _ in range(11):
        merge._observe_native_tool('read_source', {'evidence_id': 'E473'}, source_page(19, replayed=replayed))
        boundary()
    assert len(merge._progress_pages) == 20
    with pytest.raises(MergeProgressStalled):
        boundary()


def test_only_audit_record_versions_on_same_workspace_page_do_not_count_as_progress(tmp_path):
    merge, state, boundary = progress_harness(tmp_path)
    for version in range(12):
        merge._observe_native_tool('read_workspace_page', {'view': 'facts'},
            {'ok': True, 'tool': 'read_workspace_page', 'result': {'ok': True, 'version': version,
                'start': 0, 'end': 1, 'items': [{'id': 'fact', 'value': '同', 'record_version': version}]}})
        boundary()
    assert len(merge._progress_pages) == 1
    assert merge._progress_guard.state['unchanged'] == 11
    with pytest.raises(MergeProgressStalled):
        boundary()


def test_failed_then_successful_save_in_one_native_tool_batch_resets_only_at_next_model(tmp_path):
    merge, state, boundary = progress_harness(tmp_path)
    for _ in range(11):
        boundary()
    merge._observe_native_tool('update_workspace', {'update_json': '{}'},
        {'ok': False, 'tool': 'update_workspace', 'error': 'synthetic argument failure'})
    # A later sibling in the same model output saves useful content. There is
    # no intervening model request, so failure may not stop this tool batch.
    assert merge._progress_guard.state['unchanged'] == 11
    state['draft']['基本信息']['桥梁名称'] = '本批次成功保存的桥名'
    merge._observe_native_tool('update_workspace', {'update_json': 'synthetic valid update'},
        {'ok': True, 'tool': 'update_workspace', 'result': {'ok': True}})
    assert merge._progress_guard.state['unchanged'] == 11
    assert boundary() is None
    assert merge._progress_guard.state['unchanged'] == 0


def test_unread_blocker_gives_exact_source_and_preserves_valid_unknown():
    state, artifact, pages = current_file()
    state['questions'] = {'unknown': {'valid': True, 'status': 'unknown', 'reason': '原文未交代'}}
    delivery = {'ready': False, 'blockers': ['unread_source_ranges'],
                'reading_progress': {'unread_ranges': [{'evidence_id': 'E473', 'start': 0, 'end': 9}]}}
    action = next_merge_action(state, delivery, [], artifact, pages)
    assert action['stage'] == 'complete_source_delivery'
    assert action['next_action'] == {'tool': 'read_source', 'arguments': {'evidence_id': 'E473'}}
    assert action['missing_ranges_preview'] == delivery['reading_progress']['unread_ranges']
    assert next_blocker_action(state, delivery)['next_action'] == action['next_action']


def test_stall_failure_run_exports_real_saved_word_and_concrete_blocker(tmp_path):
    from docx import Document
    from bridge_evidence.contest.document_workflow import DocumentSettings, DocumentWorkflowFailure
    from test_document_workflow import FakeModel
    client = FakeModel()
    client.max_model_len = 262144
    client.dynamic_output_budget = True
    manifest = manifest_from_text('合成桥已保存，建议待核。')
    merge = MergeWorkflow(manifest, client, {'model': client.model}, tmp_path / 'run', child_summary=[],
        settings=DocumentSettings(max_model_len=262144, output_tokens=262144, dynamic_output_budget=True,
            single_document_agent=True, require_semantic_review=False))
    unit = manifest['units'][0]
    merge.workspace.record_read(unit['id'], 0, len(unit['text']))
    merge.workspace.apply_update({'base_version': merge.workspace.snapshot()['version'],
        'facts': [{'id': 'name', 'value': '合成桥', 'source_refs': [{'evidence_id': unit['id'], 'start': 0, 'end': len(unit['text'])}], 'status': 'explicit'}],
        'draft_patch': [{'op': 'set_basic', 'field': '桥梁名称', 'value': '合成桥', 'fact_ids': ['name']}]})
    def stopped_entry():
        merge._active_agent_call_id = None
        merge._before_native_model()
        for i in range(12):
            merge._active_agent_call_id = str(i)
            merge._before_native_model()
        pytest.fail('Stall must stop before model dispatch')
    merge._entry = stopped_entry
    with pytest.raises(DocumentWorkflowFailure, match='Twelve model rounds'):
        merge.run()
    receipt = merge.artifact_delivery
    assert receipt['complete'] is False and receipt['phase'] == 'incomplete'
    document = Document(receipt['docx_path'])
    text = '\n'.join(p.text for p in document.paragraphs) + '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert '合成桥' in text
    assert 'Twelve model rounds' in json.loads(Path(receipt['status_path']).read_text())['reason']
    assert len(client.calls) == 0


def test_handoff_contains_exact_original_ranges_headers_context_without_mutating_workspace(tmp_path):
    manifest = manifest_from_text('年份 | 构件 | 本次与前次\n2018 | 左幅支座 | 已有锈蚀，本次未新增')
    header, unit = manifest['units']
    unit.update(table_id='T1', row_index=1, cells=[{'column_index': 1, 'text': '左幅支座'}])
    manifest['tables'] = [{'id': 'T1', 'header_evidence_ids': [header['id']]}]
    store = EntryStore(manifest, 'child', tmp_path / 'child')
    store.record_read(unit['id'], 0, len(unit['text']))
    entry = {'id': 'bearing', 'section': '病害明细', 'row': {'病害描述': '锈蚀'}, 'status': 'explicit',
             'source_refs': [{'evidence_id': unit['id'], 'start': 7, 'end': len(unit['text'])}],
             'context': {'year': '2018', 'side': '左幅', 'period': '本次与前次', 'location': None}}
    store.save_entries([entry])
    result = {'task_id': 'child', 'entries': store.snapshot(), 'handoff': {'unresolved': ['年份角色待主模型核对']}}
    ledger = build_handoff_evidence(manifest, [result])
    evidence = ledger[child_fact_id('child', 'bearing')]
    assert evidence['source_excerpts'][0]['text'] == unit['text'][7:]
    assert evidence['table_headers'][0]['text'] == header['text']
    assert evidence['model_context'] == entry['context']
    workspace = DocumentWorkspace(manifest, tmp_path / 'main', {'model': 'synthetic'})
    import_candidates(workspace, [result])
    original = workspace.snapshot()
    tools = MergeWorkspaceTools(workspace, count_tokens=len, max_tool_tokens=20000, handoff_evidence=ledger)
    page = tools.read_page('facts', 0, 1)
    assert page['items'][0]['handoff_evidence'] == evidence
    assert workspace.snapshot() == original
    changed = deepcopy(original)
    changed['facts'][child_fact_id('child', 'bearing')]['value'] = {'病害描述': '模型后来修订'}
    tools.workspace = SimpleNamespace(snapshot=lambda: changed, delivery_status=workspace.delivery_status)
    assert 'handoff_evidence' not in tools.read_page('facts', 0, 1)['items'][0]


@pytest.mark.parametrize('context', [{'year': 2018}, {'guessed': 'no'}, ['2018'], {'year': True}])
def test_context_rejects_unrepresentable_metadata_without_saving(tmp_path, context):
    manifest = manifest_from_text('合成桥')
    store = EntryStore(manifest, 'part', tmp_path)
    unit = manifest['units'][0]
    store.record_read(unit['id'], 0, len(unit['text']))
    entry = {'id': 'bridge', 'section': '基本信息', 'field': '桥梁名称', 'value': '合成桥',
             'status': 'explicit', 'source_refs': [{'evidence_id': unit['id']}], 'context': context}
    before = store.path.read_bytes()
    with pytest.raises(EntryStoreError, match='context|Context'):
        store.save_entries([entry])
    assert store.path.read_bytes() == before
