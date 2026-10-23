from __future__ import annotations

import hashlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from bridge_evidence.contest.evidence_manifest import (
    build_manifest,
    get_neighbors,
    make_chunks,
    manifest_from_text,
    outline,
    read_evidence,
    read_table_rows,
    render_spans,
    search_evidence,
)


def test_indexes_interleaved_body_and_nested_table_in_source_order(tmp_path):
    path = tmp_path / "nested.docx"
    document = Document()
    document.add_heading("2018 年左幅检查", 1)
    document.add_paragraph("正文开始\t保留空白  内容")
    outer = document.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "嵌套表之前"
    nested = cell.add_table(rows=2, cols=2)
    nested.cell(0, 0).text = "位置"
    nested.cell(0, 1).text = "病害"
    nested.cell(1, 0).text = "左幅 1 号梁"
    nested.cell(1, 1).text = "裂缝"
    cell.paragraphs[-1].text = "嵌套表之后"
    document.add_paragraph("末尾建议")
    document.save(path)

    manifest = build_manifest(path)
    assert [unit["text"] for unit in manifest["units"]] == [
        "2018 年左幅检查", "正文开始\t保留空白  内容", "嵌套表之前",
        "位置 | 病害", "左幅 1 号梁 | 裂缝", "嵌套表之后", "末尾建议",
    ]
    assert manifest["document_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    outer_info, nested_info = manifest["tables"]
    assert nested_info["parent_table_id"] == outer_info["id"]
    assert outer_info["header_kind"] == "none"
    assert nested_info["header_kind"] == "first_row_candidate"
    outer_units = read_table_rows(manifest, outer_info["id"], 0, 1)
    assert [unit["segment_index"] for unit in outer_units] == [0, 1]
    assert all(unit["context"]["section_heading"] == "2018 年左幅检查" for unit in outer_units)
    assert all("word/document.xml#" in unit["source_path"] for unit in manifest["units"])


def test_physical_cells_do_not_duplicate_horizontal_or_vertical_merges(tmp_path):
    path = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(0, 2).text = "备注"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "左幅"
    table.cell(1, 1).text = "第一条病害"
    table.cell(2, 1).text = "第二条病害"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.save(path)

    manifest = build_manifest(path)
    units = manifest["units"]
    assert units[0]["text"] == "合并表头 | 备注"
    assert units[0]["cells"][0]["column_span"] == 2
    assert units[0]["cells"][1]["column_index"] == 2
    assert units[2]["cells"][0]["text"] == ""
    assert units[2]["cells"][0]["merged_context_text"] == "左幅"
    assert units[2]["cells"][0]["merged_from_source"] == units[1]["cells"][0]["source_path"]
    assert "左幅" not in units[2]["text"]
    assert manifest["tables"][0]["header_kind"] == "explicit"
    chunks = make_chunks(manifest, len, 450)
    last_chunk = next(chunk for chunk in chunks if units[2]["id"] in chunk["evidence_ids"])
    assert "合并表头" in last_chunk["text"]
    assert "[Merged-cell context] 左幅" in last_chunk["text"]


def test_header_footer_and_repeated_images_are_accounted_not_claimed_verified(tmp_path):
    path = tmp_path / "images.docx"
    image_path = tmp_path / "photo.png"
    Image.new("RGB", (12, 12), "green").save(image_path)
    document = Document()
    document.add_paragraph("正文")
    document.sections[0].header.paragraphs[0].text = "桥名页眉"
    document.sections[0].footer.paragraphs[0].text = "报告编号页脚"
    document.add_picture(str(image_path))
    document.add_picture(str(image_path))
    document.save(path)

    manifest = build_manifest(path)
    assert {(unit["context"]["story"], unit["text"]) for unit in manifest["units"]} >= {
        ("body", "正文"), ("header", "桥名页眉"), ("footer", "报告编号页脚"),
    }
    assert len(manifest["images"]) == 2
    assert {image["verification"] for image in manifest["images"]} == {"unverified"}
    assert len({image["target"] for image in manifest["images"]}) == 1
    assert outline(manifest)["unverified_image_count"] == 2


def test_training_text_retains_every_character_and_source_binding():
    original = " 桥名\r\n\r\n第一行\t数值  0\n最后一行 "
    manifest = manifest_from_text(original, source_sha256="a" * 64)
    assert "".join(unit["text"] for unit in manifest["units"]) == original
    assert manifest["document_sha256"] == "a" * 64
    assert manifest["text_sha256"] == hashlib.sha256(original.encode()).hexdigest()
    for unit in manifest["units"]:
        assert original[unit["source_start"]:unit["source_end"]] == unit["text"]
    assert manifest["source_kind"] == "training_text"


@pytest.mark.parametrize("text", ["", "短文\n", "病害跨页😀" * 300 + "\n尾部必须存在", "a\n\n" * 60])
def test_chunk_spans_cover_each_character_exactly_once_within_budget(text):
    manifest = manifest_from_text(text)
    chunks = make_chunks(manifest, len, 768)
    assert chunks
    assert all(chunk["token_count"] <= 768 and len(chunk["text"]) <= 768 for chunk in chunks)
    intervals = {unit["id"]: [] for unit in manifest["units"]}
    for chunk in chunks:
        rendered = render_spans(manifest, chunk["spans"])
        assert rendered["text"] == chunk["text"]
        assert rendered["context_evidence_ids"] == chunk["context_evidence_ids"]
        for span in chunk["spans"]:
            intervals[span["coverage_id"]].append((span["start"], span["end"]))
    for unit in manifest["units"]:
        position = 0
        recovered = []
        for start, end in intervals[unit["id"]]:
            assert start == position
            recovered.append(unit["text"][start:end])
            position = end
        assert position == len(unit["text"])
        assert "".join(recovered) == unit["text"]


def test_oversize_table_header_fails_instead_of_dropping_it(tmp_path):
    path = tmp_path / "header.docx"
    document = Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "很长的表头" * 100
    table.cell(1, 0).text = "病害记录"
    document.save(path)
    manifest = build_manifest(path)
    with pytest.raises(ValueError, match="header/context"):
        make_chunks(manifest, len, 100)
    assert manifest["units"][1]["text"] == "病害记录"


def test_reads_are_explicit_and_neighbor_queries_include_target():
    manifest = manifest_from_text("第一段\n第二段病害\n第三段\n第四段病害")
    units = manifest["units"]
    assert read_evidence(manifest, units[0]["id"]) == units[:1]
    assert read_evidence(manifest, [units[2]["id"], units[0]["id"]]) == [units[2], units[0]]
    assert get_neighbors(manifest, units[1]["id"]) == units[:3]
    assert [unit["id"] for unit in search_evidence(manifest, "病害", limit=1)] == [units[1]["id"]]
    with pytest.raises(KeyError, match="Unknown evidence"):
        read_evidence(manifest, "made-up")
    with pytest.raises(KeyError, match="Unknown evidence"):
        get_neighbors(manifest, "made-up")
    with pytest.raises(ValueError, match="blank"):
        search_evidence(manifest, " ")
    with pytest.raises(KeyError, match="Unknown table"):
        read_table_rows(manifest, "made-up", 0, 1)


def test_row_reader_uses_half_open_physical_row_range(tmp_path):
    path = tmp_path / "rows.docx"
    document = Document()
    table = document.add_table(rows=3, cols=1)
    for index, row in enumerate(table.rows):
        row.cells[0].text = str(index)
    document.save(path)
    manifest = build_manifest(path)
    table_id = manifest["tables"][0]["id"]
    assert [unit["text"] for unit in read_table_rows(manifest, table_id, 1, 3)] == ["1", "2"]
    with pytest.raises(ValueError, match="Row range"):
        read_table_rows(manifest, table_id, 0, 4)


def test_textbox_table_preserves_surrounding_text_order(tmp_path):
    path = tmp_path / "textbox.docx"
    document = Document()
    paragraph = document.add_paragraph("段前")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "文本框里的表格"
    box = OxmlElement("w:txbxContent")
    box.append(table._tbl)  # Moving the XML element removes the top-level copy.
    paragraph._p.append(box)
    paragraph.add_run("段后")
    document.save(path)
    manifest = build_manifest(path)
    assert [unit["text"] for unit in manifest["units"]] == ["段前", "文本框里的表格", "段后"]


def test_unparsed_embedded_objects_prevent_silent_complete_coverage_claim(tmp_path):
    path = tmp_path / "embedded.docx"
    document = Document()
    paragraph = document.add_paragraph("正文")
    paragraph._p.append(OxmlElement("w:object"))
    document.save(path)
    manifest = build_manifest(path)
    assert manifest["unverified_content"][0]["kind"] == "object"
    assert outline(manifest)["unverified_content_count"] == 1
    assert any("not complete document coverage" in item for item in manifest["limitations"])


def test_public_span_renderer_keeps_bridge_year_heading_and_table_context(tmp_path):
    path = tmp_path / "context.docx"
    document = Document()
    document.add_paragraph("甲桥")
    document.add_paragraph("2018 年检测报告")
    document.add_paragraph("报告正文开场")
    document.add_heading("左幅上部结构", 1)
    document.add_paragraph("说明段")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "位置"
    table.cell(0, 1).text = "病害"
    table.cell(1, 0).text = "1 号梁"
    table.cell(1, 1).text = "跨中裂缝"
    document.save(path)
    manifest = build_manifest(path)
    row = manifest["units"][-1]
    heading = next(unit for unit in manifest["units"] if unit["text"] == "左幅上部结构")
    header = manifest["units"][-2]
    spans = [{"coverage_id": row["id"], "start": 0, "end": len(row["text"])}]
    rendered = render_spans(manifest, spans)
    assert all(text in rendered["text"] for text in ("甲桥", "2018 年检测报告", "左幅上部结构", "位置 | 病害", "1 号梁 | 跨中裂缝"))
    assert {heading["id"], header["id"]} <= set(rendered["context_evidence_ids"])
    assert row["id"] not in rendered["context_evidence_ids"]
    assert spans == [{"coverage_id": row["id"], "start": 0, "end": len(row["text"])}]


def test_document_opening_context_is_bounded_and_never_primary_coverage():
    manifest = manifest_from_text("甲桥 2018 左幅 " + "长正文" * 1000 + "\n末行病害")
    unit = manifest["units"][-1]
    spans = [{"coverage_id": unit["id"], "start": 0, "end": len(unit["text"])}]
    rendered = render_spans(manifest, spans)
    assert "甲桥 2018 左幅" in rendered["text"]
    assert "chars 0:160; not primary coverage" in rendered["text"]
    assert len(rendered["text"]) < 400
    assert rendered["context_evidence_ids"] == [manifest["units"][0]["id"]]


@pytest.mark.parametrize("start,end", [(-1, 1), (0, 99), (2, 1), (False, 1)])
def test_span_renderer_rejects_invalid_character_intervals(start, end):
    manifest = manifest_from_text("原文")
    with pytest.raises(ValueError, match="Invalid character interval"):
        render_spans(manifest, [{"coverage_id": "E000001", "start": start, "end": end}])
