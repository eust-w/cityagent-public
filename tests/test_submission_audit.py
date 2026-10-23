from __future__ import annotations

import json
import tarfile
from pathlib import Path

from docx import Document

from bridge_evidence.audit import audit_archive, run_submission_audit


def _write_result(path: Path) -> None:
    document = Document()
    for _ in range(3):
        document.add_table(rows=1, cols=1).cell(0, 0).text = "ok"
    document.save(path)
    (path.parent / "infer_time.json").write_text(
        json.dumps({"infer_time": 1}), encoding="utf-8"
    )


def _write_online_state(path: Path, source: Path) -> None:
    actions = ["review_summary", "review_details", "normalize_diseases", "write_decision", "finalize"]
    tasks = {
        "react-plan": 5,
        "summary-evidence-embedding": 1,
        "summary-evidence-rerank": 1,
        "detail-evidence-embedding": 1,
        "detail-evidence-rerank": 1,
        "summary-review": 1,
        "detail-review": 1,
        "decision": 1,
    }
    payload = {
        "source_path": str(source),
        "agent_run": {
            "planner": "qwen",
            "trajectory": [
                {"command": {"action": action}, "result": {"success": True}}
                for action in actions
            ],
        },
        "metrics": {
            "task_calls": tasks,
            "task_cache_hits": {},
            "task_models": {
                "react-plan": "deepseek/deepseek-v4-pro",
                "summary-evidence-embedding": "Qwen3-VL-Embedding-8B",
                "summary-evidence-rerank": "Qwen3-VL-Reranker-8B",
                "detail-evidence-embedding": "Qwen3-VL-Embedding-8B",
                "detail-evidence-rerank": "Qwen3-VL-Reranker-8B",
                "summary-review": "deepseek/deepseek-v4-pro",
                "detail-review": "deepseek/deepseek-v4-pro",
                "decision": "deepseek/deepseek-v4-pro",
            },
            "model_calls": {
                "deepseek/deepseek-v4-pro": 8,
                "Qwen3-VL-Embedding-8B": 2,
                "Qwen3-VL-Reranker-8B": 2,
            },
            "online_audit_passed": True,
        },
    }
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def test_submission_audit_accepts_doc_input_with_docx_result(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    result_dir = tmp_path / "result"
    state_dir = tmp_path / "state"
    for directory in (input_dir, result_dir, state_dir):
        directory.mkdir()
    source = input_dir / "旧报告.doc"
    source.write_bytes(b"legacy-placeholder")
    _write_result(result_dir / "旧报告.docx")
    _write_online_state(state_dir / "旧报告.json", source)

    report = run_submission_audit(
        input_dir=input_dir,
        result_dir=result_dir,
        state_dir=state_dir,
        require_online=True,
    )

    assert report["passed"], report["errors"]
    assert report["results"]["table_counts"] == {"3": 1}


def test_submission_audit_rejects_missing_model_task(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    result_dir = tmp_path / "result"
    state_dir = tmp_path / "state"
    for directory in (input_dir, result_dir, state_dir):
        directory.mkdir()
    source = input_dir / "report.docx"
    source.write_bytes(b"placeholder")
    _write_result(result_dir / "report.docx")
    _write_online_state(state_dir / "report.json", source)
    payload = json.loads((state_dir / "report.json").read_text(encoding="utf-8"))
    del payload["metrics"]["task_calls"]["summary-evidence-rerank"]
    (state_dir / "report.json").write_text(json.dumps(payload), encoding="utf-8")

    report = run_submission_audit(
        input_dir=input_dir,
        result_dir=result_dir,
        state_dir=state_dir,
        require_online=True,
    )

    assert not report["passed"]
    assert any("summary-evidence-rerank" in error for error in report["errors"])


def test_archive_audit_rejects_extra_top_level(tmp_path: Path) -> None:
    staging = tmp_path / "staging"
    for name in ("code", "design", "result", "state"):
        (staging / name).mkdir(parents=True)
        (staging / name / "keep.txt").write_text("ok", encoding="utf-8")
    archive = tmp_path / "submission.tar.gz"
    with tarfile.open(archive, "w:gz") as handle:
        for child in staging.iterdir():
            handle.add(child, arcname=child.name)

    section, errors = audit_archive(archive)

    assert "state" in section["extra_top_levels"]
    assert errors
def test_online_audit_requirements_support_details_deterministic_profile():
    from bridge_evidence.audit import online_audit_requirements

    requirements = online_audit_requirements(
        workflow_profile="details",
        planner_mode="deterministic",
        semantic_fallback=False,
    )

    assert requirements["planner_models"] == {"deterministic"}
    assert requirements["actions"] == {"review_details", "finalize"}
    assert requirements["tasks"] == {
        "detail-evidence-embedding",
        "detail-evidence-rerank",
        "detail-review",
    }


def test_online_audit_requirements_allow_recorded_semantic_fallback():
    from bridge_evidence.audit import online_audit_requirements

    requirements = online_audit_requirements(
        workflow_profile="details",
        planner_mode="deterministic",
        semantic_fallback=True,
    )

    assert requirements["tasks"] == {"detail-review"}
    assert "review_details" in requirements["actions"]


def test_audit_archive_supports_zip(tmp_path):
    import zipfile

    from bridge_evidence.audit import audit_archive

    archive_path = tmp_path / "submission.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("code/main.py", "print('ok')\n")
        archive.writestr("design/README.md", "design\n")
        archive.writestr("result/sample.docx", b"placeholder")

    report, errors = audit_archive(archive_path)

    assert errors == []
    assert report["top_levels"] == ["code", "design", "result"]
