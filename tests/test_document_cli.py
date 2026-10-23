"""Synthetic boundary tests only: no real reports, SDK, server or training."""
from __future__ import annotations

import argparse
from concurrent.futures import Future
from dataclasses import make_dataclass
import json
import os
from pathlib import Path
import sys

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from bridge_evidence.contest import document_cli as cli

Settings = make_dataclass("Settings", [(name, type(value), value) for name, value in cli._DEFAULTS.items()])


@pytest.fixture
def boundary(tmp_path, monkeypatch):
    parser = argparse.ArgumentParser()
    cli.divide_cli.add_arguments(parser)
    cli.add_arguments(parser)
    args = parser.parse_args([])
    for key, value in dict(input_dir=None, result_dir=tmp_path / "result", work_dir=tmp_path / "work",
                           text_dir=None, engine="tongagents", fusion_policy="model-preserving",
                           resume=False, limit=None, check=False).items():
        setattr(args, key, value)
    args.served_adapter = tmp_path / "adapter"
    args.served_adapter.mkdir()
    (args.served_adapter / "adapter_model.safetensors").write_bytes(b"synthetic weights")
    (args.served_adapter / "adapter_config.json").write_text('{"r": 16, "lora_alpha": 32}')
    args.served_model = "synthetic-model"
    args.training_evidence_json = tmp_path / "training.json"
    args.training_evidence_json.write_text(json.dumps({"evidence_id": "TRAIN-1", "training_line": 1,
                                                      "text": "桥梁名称：合成桥"}, ensure_ascii=False))
    code_root = tmp_path / "code"
    contest = code_root / "src/bridge_evidence/contest"
    contest.mkdir(parents=True)
    (code_root / "run.py").write_text("# synthetic entry")
    (contest / "document_prompt.py").write_text("# synthetic prompt v1")
    calls = {"workflows": 0, "clients": 0, "writes": 0, "pool": [], "resumes": []}
    catalog = [{"id": args.served_model, "parent": "synthetic-base", "root": str(args.served_adapter)},
               {"id": "synthetic-base", "parent": None, "root": "/synthetic/base-model"}]

    class Client:
        base_url, api_key = "http://127.0.0.1:8000", "synthetic-secret"

        def __init__(self, client_args):
            calls["clients"] += 1
            calls["client_args"] = client_args
            self.model_catalog = catalog
            # Vary calls every invocation to detect accidental identity binding
            # to timestamps or request IDs on resume.
            self.calls = [{"kind": "models", "call_id": str(calls["clients"])}]

        def models(self):
            return [item["id"] for item in self.model_catalog]

    class Workflow:
        def __init__(self, manifest, client, config, directory, *, settings, source_path=None, run_identity=None,
                     vision_backend=None):
            calls["workflows"] += 1
            calls["config"], calls["settings"] = config, settings
            calls["source_path"], calls["directory"] = source_path, directory
            calls["run_identity"] = run_identity
            calls["vision_backend"] = vision_backend
            calls["cwd"] = Path.cwd()
            calls["sdk_log"] = os.environ.get("LOG_TRACE_DB_PATH")
            self.directory = directory

        def run(self):
            data = {"基本信息": {"桥梁名称": "合成桥"}, "病害明细": [], "维修建议": []}
            return {"data": data,
                    "receipt": {"status": "completed", "synthetic": True,
                                "delivery": {"ready": True}, "draft_sha256": cli.json_digest(data),
                                "final_data": data,
                                "model_calls": [{"call_id": "synthetic-generation"}]}}

    class Pool:
        def __init__(self, **kwargs):
            calls["pool"].append(kwargs)

        def __enter__(self):
            return self

        def __exit__(self, *exc):
            pass

        def submit(self, function, *job):
            future = Future()
            try:
                future.set_result(function(*job))
            except BaseException as exc:
                future.set_exception(exc)
            return future

    def writer(data, path):
        from bridge_evidence.contest.docx_writer import write_numbered_extraction_docx as write_extraction_docx
        calls["writes"] += 1
        write_extraction_docx(data, path)

    monkeypatch.setattr(cli.divide_cli, "_make_client", Client)
    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, Workflow))
    monkeypatch.setattr(cli, "ProcessPoolExecutor", Pool)
    monkeypatch.setattr(cli, "write_extraction_docx", writer)
    monkeypatch.setattr(cli.importlib.metadata, "version", lambda name: "synthetic-1.0")
    return args, code_root, calls, Workflow, catalog


def run(boundary):
    args, code_root, *_ = boundary
    return cli.run_document_agent(args, code_root=code_root, convert_doc=Path)


def read(path):
    return json.loads(path.read_text())


def use_training_batch(boundary, count=3):
    args = boundary[0]
    directory = args.training_evidence_json.parent / "batch-inputs"
    directory.mkdir()
    entries, originals = [], []
    for index in range(count):
        original = directory / f"input-{index}.json"
        cli.atomic_json(original, {"evidence_id": f"BATCH-{index}", "training_line": index + 1,
                                   "text": f"合成第{index}篇报告"})
        originals.append(original)
        entries.append({"path": original.name, "sha256": cli.file_digest(original)})
    args.training_evidence_manifest = directory / "manifest.json"
    cli.atomic_json(args.training_evidence_manifest, {"schema_version": 1,
        "scope": "official_training_text_only", "entries": entries})
    args.training_evidence_json = None
    return originals


@pytest.mark.parametrize("workers", [1, 2, 8])
def test_explicit_training_batch_identity_workers_and_per_worker_original(boundary, monkeypatch, workers):
    args, _, calls, Workflow, _ = boundary
    originals = use_training_batch(boundary)
    args.document_workers = workers
    parsed = []
    original_parser = cli.divide_cli._training_manifest
    def tracked(path):
        parsed.append((Path.cwd(), path))
        return original_parser(path)
    monkeypatch.setattr(cli.divide_cli, "_training_manifest", tracked)
    result = run(boundary)
    identity = read(args.work_dir / "run_manifest.json")["identity"]
    assert result["reports"] == result["completed_reports"] == calls["workflows"] == 3
    assert identity["document_workers"] == workers
    assert identity["effective_document_workers"] == calls["pool"][0]["max_workers"] == min(workers, 3)
    assert identity["training_evidence_batch"]["manifest_sha256"] == cli.file_digest(args.training_evidence_manifest)
    assert identity["sources"] == {str(p): cli.file_digest(p) for p in originals}
    assert [path for _, path in parsed[:3]] == originals
    assert parsed[3:] == [(args.work_dir / "documents" / f"BATCH-{i}", p) for i, p in enumerate(originals)]
    for i, original in enumerate(originals):
        manifest = read(args.work_dir / "documents" / f"BATCH-{i}" / "source_manifest.json")
        assert manifest["source_json_sha256"] == cli.file_digest(original)
        assert manifest["training_source"]["evidence_id"] == f"BATCH-{i}"
    assert result["input_scope"] == "official_training_text_only"
    assert any("image coverage" in item for item in result["input_limitations"])


@pytest.mark.parametrize("other", ["input_dir", "training_evidence_json", "limit"])
def test_batch_is_mutually_exclusive_and_never_silently_limited(boundary, other):
    args, _, calls, _, _ = boundary
    originals = use_training_batch(boundary)
    setattr(args, other, 1 if other == "limit" else originals[0].parent if other == "input_dir" else originals[0])
    with pytest.raises(ValueError, match="exactly one|cannot truncate"):
        run(boundary)
    assert calls["clients"] == calls["workflows"] == 0


@pytest.mark.parametrize("changed", ["manifest", "source", "source_and_declared_hash", "workers"])
def test_batch_resume_rejects_changed_manifest_sources_or_concurrency(boundary, changed):
    args, _, calls, _, _ = boundary
    originals = use_training_batch(boundary)
    run(boundary)
    args.resume = True
    if changed == "workers":
        args.document_workers = 2
    elif changed == "manifest":
        args.training_evidence_manifest.write_bytes(args.training_evidence_manifest.read_bytes() + b"\n")
    else:
        originals[0].write_bytes(originals[0].read_bytes() + b"\n")
        if changed == "source_and_declared_hash":
            value = read(args.training_evidence_manifest)
            value["entries"][0]["sha256"] = cli.file_digest(originals[0])
            cli.atomic_json(args.training_evidence_manifest, value)
    with pytest.raises(ValueError, match="Run identity changed|SHA256 mismatch"):
        run(boundary)
    assert calls["workflows"] == 3


def test_batch_resume_reuses_each_bound_result_without_new_workflows(boundary):
    args, _, calls, _, _ = boundary
    use_training_batch(boundary)
    run(boundary)
    args.resume = True
    result = run(boundary)
    assert result["reused_reports"] == 3
    assert result["newly_completed_reports"] == 0
    assert calls["workflows"] == 3


def test_batch_failure_keeps_all_inputs_in_throughput_reconciliation(boundary, monkeypatch):
    args, _, calls, Workflow, _ = boundary
    use_training_batch(boundary)
    args.document_workers = 2
    class OneFailure(Workflow):
        def run(self):
            if self.directory.name == "BATCH-1":
                raise RuntimeError("synthetic one-report failure")
            return super().run()
    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, OneFailure))
    with pytest.raises(RuntimeError, match="1 document-agent report"):
        run(boundary)
    summary = read(args.work_dir / "summary.json")
    assert summary["reports"] == calls["workflows"] == 3
    assert summary["completed_reports"] == 2
    assert summary["failed_reports"] == 1
    assert summary["unfinished_reports"] == 0
    assert summary["reports_per_hour"] == pytest.approx(2 * 3600 / summary["seconds"])
    assert [item["name"] for item in summary["documents"]] == [f"BATCH-{i}" for i in range(3)]


@pytest.mark.parametrize("target", ["manifest", "own_source"])
def test_batch_change_during_workflow_blocks_publication_and_final_acceptance(boundary, monkeypatch, target):
    args, _, calls, Workflow, _ = boundary
    originals = use_training_batch(boundary)
    class Mutation(Workflow):
        def run(self):
            if calls["workflows"] == 1:
                path = args.training_evidence_manifest if target == "manifest" else originals[0]
                path.write_bytes(path.read_bytes() + b"\n")
            return super().run()
    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, Mutation))
    with pytest.raises(ValueError, match="changed during"):
        run(boundary)
    summary = read(args.work_dir / "summary.json")
    assert summary["reports"] == summary["failed_reports"] == 3
    assert summary["completed_reports"] == 0
    assert not (args.result_dir / "BATCH-0.docx").exists()


def test_batch_rechecks_identity_after_completed_resume_readback(boundary, monkeypatch):
    args, _, calls, _, _ = boundary
    use_training_batch(boundary)
    run(boundary)
    args.resume = True
    completed = cli._completed
    def changed(*values):
        result = completed(*values)
        args.training_evidence_manifest.write_bytes(args.training_evidence_manifest.read_bytes() + b"\n")
        return result
    monkeypatch.setattr(cli, "_completed", changed)
    with pytest.raises(ValueError, match="manifest changed"):
        run(boundary)
    summary = read(args.work_dir / "summary.json")
    assert summary["completed_reports"] == summary["reused_reports"] == 0
    assert summary["failed_reports"] == 3
    assert calls["workflows"] == 3


def test_success_binds_provenance_and_isolates_sdk(boundary):
    args, code_root, calls, _, _ = boundary
    before = Path.cwd(), os.environ.get("LOG_TRACE_DB_PATH")
    args.document_output_tokens = 1234
    result = run(boundary)
    directory = args.work_dir / "documents/TRAIN-1"
    receipt = read(directory / "execution.json")
    identity = read(args.work_dir / "run_manifest.json")
    assert result["completed_reports"] == result["newly_completed_reports"] == 1
    assert result["reports_per_hour"] == pytest.approx(3600 / result["seconds"])
    assert result["failed_reports"] == result["unfinished_reports"] == 0
    assert receipt["run_fingerprint"] == identity["run_fingerprint"]
    assert receipt["output"]["sha256"] == cli.file_digest(args.result_dir / "TRAIN-1.docx")
    assert receipt["report_sha256"] == cli.json_digest(read(directory / "report.json"))
    assert receipt["model_calls"] == [{"call_id": "synthetic-generation"}]
    assert calls["run_identity"] == {"run_fingerprint": identity["run_fingerprint"],
                                      "source_sha256": cli.file_digest(args.training_evidence_json)}
    assert receipt["source_binding_sha256"] == cli.file_digest(directory / "source_binding.json")
    assert receipt["delivery_validation"]["report_sha256"] == receipt["report_sha256"]
    assert receipt["delivery_validation"]["docx_sha256"] == receipt["output"]["sha256"]
    assert receipt["delivery_validation"]["table_data_rows"] == [1, 0, 0]
    assert "synthetic-secret" not in json.dumps(identity)
    assert calls["client_args"].chunk_output_tokens == calls["config"]["max_tokens"] == 1234
    assert calls["settings"].max_model_len == 65536
    assert calls["cwd"] == directory
    assert calls["sdk_log"] == str(directory / "sdk-logs/traces.db")
    assert (Path.cwd(), os.environ.get("LOG_TRACE_DB_PATH")) == before
    assert calls["pool"][0]["mp_context"].get_start_method() == "spawn"
    assert calls["pool"][0]["max_tasks_per_child"] == 1
    assert not result["submission_ready"]


def test_resume_reuses_verified_results_without_faking_calls(boundary):
    args, _, calls, _, _ = boundary
    run(boundary)
    receipt_before = (args.work_dir / "documents/TRAIN-1/execution.json").read_bytes()
    args.resume = True
    result = run(boundary)
    assert calls["workflows"] == calls["writes"] == 1
    assert result["completed_reports"] == result["reused_reports"] == 1
    assert result["newly_completed_reports"] == result["reports_per_hour"] == 0
    assert result["documents"][0]["model_calls_this_attempt"] == 0
    assert (args.work_dir / "documents/TRAIN-1/execution.json").read_bytes() == receipt_before
    assert len(list((args.work_dir / "attempts").glob("*.summary.json"))) == 2


def test_review_iterations_flag_is_independent_and_bound_to_resume(boundary):
    args, _, calls, _, _ = boundary
    assert args.document_review_iterations == 32
    parser = argparse.ArgumentParser()
    cli.add_arguments(parser)
    args.document_review_iterations = parser.parse_args(["--document-review-iterations", "40"]).document_review_iterations
    run(boundary)
    settings = read(args.work_dir / "run_manifest.json")["identity"]["settings"]
    assert settings["review_iterations"] == calls["settings"].review_iterations == 40
    assert settings["session_iterations"] == calls["settings"].session_iterations == 8
    assert settings["max_model_len"] == calls["settings"].max_model_len == 65536
    args.resume = True
    args.document_review_iterations = 41
    with pytest.raises(ValueError, match="Run identity changed"):
        run(boundary)
    assert calls["workflows"] == 1


@pytest.mark.parametrize("change", ["source", "prompt", "weights", "adapter_config", "settings", "endpoint"])
def test_resume_rejects_changed_bound_identity(boundary, change):
    args, code_root, calls, _, _ = boundary
    run(boundary)
    args.resume = True
    paths = {"source": args.training_evidence_json,
             "prompt": code_root / "src/bridge_evidence/contest/document_prompt.py",
             "weights": args.served_adapter / "adapter_model.safetensors",
             "adapter_config": args.served_adapter / "adapter_config.json"}
    if change in paths:
        path = paths[change]
        path.write_bytes(path.read_bytes() + b"\n")
    elif change == "settings":
        args.document_tool_tokens += 1
    else:
        # The endpoint is taken from the actual client, not copied from flags.
        cli.divide_cli._make_client.base_url = "http://127.0.0.1:8888"
    with pytest.raises(ValueError, match="Run identity changed"):
        run(boundary)
    assert calls["workflows"] == 1


@pytest.mark.parametrize("artifact", ["output", "report", "receipt", "missing_output"])
def test_resume_hash_checks_outputs_data_and_execution(boundary, artifact):
    args, _, calls, _, _ = boundary
    run(boundary)
    directory = args.work_dir / "documents/TRAIN-1"
    if artifact == "output":
        (args.result_dir / "TRAIN-1.docx").write_bytes(b"tampered")
    elif artifact == "report":
        (directory / "report.json").write_text('{"changed": true}')
    elif artifact == "receipt":
        record = read(directory / "execution.json")
        record["model_calls"] = []
        (directory / "execution.json").write_text(json.dumps(record))
    else:
        (args.result_dir / "TRAIN-1.docx").unlink()
    args.resume = True
    with pytest.raises(RuntimeError, match="1 document-agent report"):
        run(boundary)
    assert calls["workflows"] == calls["writes"] == 1
    assert read(args.work_dir / "summary.json")["failed_reports"] == 1


def test_incomplete_workspace_is_passed_back_to_workflow(boundary, monkeypatch):
    args, _, calls, Workflow, _ = boundary

    class Failure(RuntimeError):
        receipt = {"status": "failed", "model_calls": [{"call_id": "first-attempt"}]}

    class ResumeWorkflow(Workflow):
        def run(self):
            checkpoint = self.directory / "checkpoint.json"
            if not checkpoint.exists():
                checkpoint.write_text('{"unfinished": true}')
                raise Failure("synthetic partial execution")
            calls["resumes"].append(read(checkpoint))
            return super().run()

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, ResumeWorkflow))
    with pytest.raises(RuntimeError):
        run(boundary)
    failures = list((args.work_dir / "documents/TRAIN-1/attempts").glob("*.failure.json"))
    assert read(failures[0])["workflow_receipt"]["model_calls"] == [{"call_id": "first-attempt"}]
    assert not list(args.result_dir.glob("*.docx"))
    args.resume = True
    result = run(boundary)
    assert result["newly_completed_reports"] == 1
    assert calls["resumes"] == [{"unfinished": True}]


def test_batch_continues_after_one_report_fails(boundary, monkeypatch):
    args, _, calls, Workflow, _ = boundary
    args.training_evidence_json = None
    args.input_dir = args.work_dir.parent / "inputs"
    args.input_dir.mkdir()
    for name in ("bad", "good", "last"):
        (args.input_dir / (name + ".docx")).write_bytes(b"synthetic fixture")
    args.document_workers = 2
    monkeypatch.setattr(cli, "build_manifest", lambda path: {"name": path.stem})

    class MixedWorkflow(Workflow):
        def run(self):
            if self.directory.name == "bad":
                raise ValueError("synthetic bad report")
            return super().run()

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, MixedWorkflow))
    with pytest.raises(RuntimeError, match="1 document-agent report"):
        run(boundary)
    summary = read(args.work_dir / "summary.json")
    assert summary["completed_reports"] == 2 and summary["failed_reports"] == 1
    assert summary["unfinished_reports"] == 0 and calls["workflows"] == 3
    assert {p.name for p in args.result_dir.glob("*.docx")} == {"good.docx", "last.docx"}
    assert calls["pool"][0]["max_workers"] == 2
    assert all((args.work_dir / "documents" / d["name"] / "document_summary.json").exists()
               for d in summary["documents"])


def test_export_failure_does_not_publish_or_leave_pending(boundary, monkeypatch):
    args, _, _, _, _ = boundary

    def broken_writer(data, path):
        Path(path).write_bytes(b"partial")
        raise OSError("synthetic export failure")

    monkeypatch.setattr(cli, "write_extraction_docx", broken_writer)
    with pytest.raises(RuntimeError):
        run(boundary)
    assert not list(args.result_dir.glob("*.docx"))
    assert not list(args.result_dir.glob(".pending-*"))
    assert read(args.work_dir / "summary.json")["failed_reports"] == 1


def test_source_mutation_during_workflow_blocks_export(boundary, monkeypatch):
    args, _, calls, Workflow, _ = boundary

    class Mutating(Workflow):
        def run(self):
            args.training_evidence_json.write_text("changed")
            return super().run()

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, Mutating))
    with pytest.raises(RuntimeError):
        run(boundary)
    assert calls["writes"] == 0
    assert not list(args.result_dir.glob("*.docx"))


def test_worker_pool_failure_reconciles_every_input(boundary, monkeypatch):
    args, _, _, _, _ = boundary

    class BrokenPool:
        def __init__(self, **kwargs):
            raise OSError("synthetic process allocation failure")

    monkeypatch.setattr(cli, "ProcessPoolExecutor", BrokenPool)
    with pytest.raises(OSError):
        run(boundary)
    result = read(args.work_dir / "summary.json")
    assert result["unfinished_reports"] == 1 and result["completed_reports"] == 0
    assert result["documents"][0]["status"] == "unfinished"
    assert (args.work_dir / "documents/TRAIN-1/document_summary.json").exists()


@pytest.mark.parametrize("setting,value", [("engine", "direct"), ("fusion_policy", "legacy"),
    ("document_workers", 0), ("document_output_tokens", 0), ("document_report_timeout", float("nan")),
    ("model_timeout", float("inf")), ("document_max_model_len", 4096), ("document_review_iterations", 0)])
def test_invalid_configuration_never_calls_model(boundary, setting, value):
    args, _, calls, _, _ = boundary
    setattr(args, setting, value)
    with pytest.raises(ValueError):
        run(boundary)
    assert calls["clients"] == calls["workflows"] == 0


@pytest.mark.parametrize("mode", ["none", "both", "labels", "new_resume"])
def test_explicit_input_and_training_identity_required(boundary, mode):
    args, _, calls, _, _ = boundary
    if mode == "none":
        args.training_evidence_json = None
    elif mode == "both":
        args.input_dir = args.training_evidence_json.parent
    elif mode == "labels":
        record = read(args.training_evidence_json)
        record["output"] = "forbidden label"
        args.training_evidence_json.write_text(json.dumps(record))
    else:
        args.resume = True
    with pytest.raises(ValueError):
        run(boundary)
    assert calls["clients"] == calls["workflows"] == 0


def test_check_performs_identity_validation_without_artifacts(boundary, capsys):
    args, _, calls, _, _ = boundary
    args.check = True
    run(boundary)
    assert calls["clients"] == 1 and calls["workflows"] == 0
    assert not args.work_dir.exists() and not args.result_dir.exists()
    assert "real SDK/model execution is not yet verified" in capsys.readouterr().out


def test_dynamic_lora_path_mismatch_rejected(boundary):
    _, _, calls, _, catalog = boundary
    catalog[0]["root"] = "/unrelated/adapter"
    with pytest.raises(ValueError, match="does not bind"):
        run(boundary)
    assert calls["workflows"] == 0


def test_dynamic_lora_accepts_same_id_base_and_adapter_with_traceable_parent(boundary):
    args, _, _, _, catalog = boundary
    catalog[0]["parent"] = args.served_model
    catalog[1]["id"] = args.served_model
    run(boundary)
    identity = read(args.work_dir / "run_manifest.json")["identity"]
    deployment = identity["model_identity"]["deployment"]
    assert identity["model_catalog"] == catalog
    assert deployment["base_model_id"] == args.served_model
    assert deployment["base_model_root"] == "/synthetic/base-model"
    assert deployment["adapter_path"] == str(args.served_adapter)
    assert deployment["runtime_weights_checksum_verified"] is False
    assert deployment["base_weights_full_checksum_verified"] is False


@pytest.mark.parametrize("fault", ["base_only", "duplicate_adapter", "other_adapter_path", "missing_parent",
                                  "duplicate_base", "empty_parent", "parent_is_another_adapter", "base_missing_root"])
def test_dynamic_lora_rejects_ambiguous_or_untraceable_catalog(boundary, fault):
    args, _, calls, _, catalog = boundary
    if fault == "base_only":
        catalog[:] = [{"id": args.served_model, "parent": None, "root": str(args.served_adapter)}]
    elif fault in {"duplicate_adapter", "other_adapter_path"}:
        catalog.append({**catalog[0], "root": str(args.served_adapter) if fault == "duplicate_adapter" else "/other/adapter"})
    elif fault == "missing_parent":
        catalog.pop()
    elif fault == "duplicate_base":
        catalog.append({**catalog[1], "root": "/other/base-model"})
    elif fault == "empty_parent":
        catalog[0]["parent"] = ""
    elif fault == "parent_is_another_adapter":
        catalog[1]["parent"] = "another-model"
    else:
        catalog[1].pop("root")
    with pytest.raises(ValueError, match="catalog"):
        run(boundary)
    assert calls["workflows"] == 0
    assert not args.work_dir.exists()


def test_resume_binds_full_dynamic_catalog_including_unrelated_entries(boundary):
    args, _, calls, _, catalog = boundary
    run(boundary)
    catalog.append({"id": "other-model", "parent": None, "root": "/other/base"})
    args.resume = True
    with pytest.raises(ValueError, match="Run identity changed"):
        run(boundary)
    assert calls["workflows"] == 1


def test_merged_identity_and_tokenizer_validation_are_reused(boundary, monkeypatch):
    args, _, calls, _, _ = boundary
    args.served_merged_model = args.work_dir.parent / "merged"
    seen = []
    files = ("tokenizer.json", "tokenizer_config.json", "chat_template.jinja")

    def merged(passed_args, client, digest):
        seen.append((passed_args.served_merged_model, digest))
        return {"binding_source": "model_catalog.alias+model_info.model_path", "mode": "w8a8_int8",
                "verified_artifacts": {name: {"sha256": "synthetic-" + name} for name in files}}

    monkeypatch.setattr(cli.divide_cli, "_merged_identity", merged)
    monkeypatch.setattr(cli.divide_cli, "_MergedChatTokenizer", lambda path, hashes: seen.append((path, hashes)))
    run(boundary)
    assert len(seen) == 2
    assert set(calls["client_args"].chat_tokenizer_hashes) == set(files)
    identity = read(args.work_dir / "run_manifest.json")["identity"]
    assert identity["model_identity"]["deployment"]["mode"] == "w8a8_int8"


def test_resume_preserves_previous_failed_attempt_wall_time(boundary, monkeypatch):
    args, _, _, Workflow, _ = boundary
    failures = [True]

    class Retry(Workflow):
        def run(self):
            if failures.pop() if failures else False:
                raise RuntimeError("synthetic first attempt")
            return super().run()

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, Retry))
    with pytest.raises(RuntimeError):
        run(boundary)
    previous = read(args.work_dir / "summary.json")
    args.resume = True
    result = run(boundary)
    assert result["cumulative_seconds"] == pytest.approx(previous["seconds"] + result["seconds"])
    assert result["cumulative_reports_per_hour"] == pytest.approx(3600 / result["cumulative_seconds"])


def test_failed_export_preserves_successful_model_receipt(boundary, monkeypatch):
    args, _, _, _, _ = boundary
    monkeypatch.setattr(cli, "write_extraction_docx", lambda *a: (_ for _ in ()).throw(OSError("export failed")))
    with pytest.raises(RuntimeError):
        run(boundary)
    failure = read(next((args.work_dir / "documents/TRAIN-1/attempts").glob("*.failure.json")))
    assert failure["workflow_receipt"]["model_calls"] == [{"call_id": "synthetic-generation"}]


def test_persistent_completion_marker_survives_repeated_hash_failures(boundary):
    args, _, calls, _, _ = boundary
    run(boundary)
    (args.result_dir / "TRAIN-1.docx").unlink()
    args.resume = True
    for _ in range(2):
        with pytest.raises(RuntimeError):
            run(boundary)
    assert calls["workflows"] == 1


def test_resume_handles_interrupted_after_link_before_commit(boundary):
    args, _, calls, _, _ = boundary
    run(boundary)
    directory = args.work_dir / "documents/TRAIN-1"
    marker = read(directory / "publication.json")
    marker["status"] = "prepared"
    cli.atomic_json(directory / "publication.json", marker)
    (directory / "document_summary.json").unlink()
    args.resume = True
    result = run(boundary)
    assert result["reused_reports"] == 1 and calls["workflows"] == 1
    assert read(directory / "publication.json")["status"] == "committed"


def _synthetic_spawned_worker(name, root):
    """Importable spawn target; all substitutions occur inside the fresh child."""
    from types import SimpleNamespace
    root = Path(root)

    class ChildClient:
        base_url, api_key = "http://127.0.0.1:8000", "synthetic-only"
        calls = []

    class ChildWorkflow:
        def __init__(self, manifest, client, config, directory, **kwargs):
            self.directory = directory

        def run(self):
            Path("sdk-relative-marker.txt").write_text(str(os.getpid()))
            data = {"基本信息": {"桥梁名称": "合成桥"}, "病害明细": [], "维修建议": []}
            return {"data": data,
                    "receipt": {"status": "completed", "synthetic": True,
                                "delivery": {"ready": True}, "draft_sha256": cli.json_digest(data),
                                "final_data": data,
                                "sdk_db": os.environ["LOG_TRACE_DB_PATH"], "cwd": str(Path.cwd())}}

    cli.divide_cli._make_client = lambda args: ChildClient()
    cli._workflow_types = lambda: (None, ChildWorkflow)
    source = root / "input.json"
    args = SimpleNamespace(work_dir=root / "work", result_dir=root / "result", resume=False,
                           served_model="synthetic-only", model_timeout=1)
    settings = SimpleNamespace(**cli._DEFAULTS)
    return cli._run_document(name, {}, None, source, args, settings, "synthetic-fingerprint",
                             cli.file_digest(source), Path)


def test_real_spawn_uses_fresh_process_and_private_sdk_paths(tmp_path):
    from concurrent.futures import ProcessPoolExecutor
    import multiprocessing
    (tmp_path / "input.json").write_text("synthetic source only")
    (tmp_path / "result").mkdir()
    with ProcessPoolExecutor(max_workers=1, mp_context=multiprocessing.get_context("spawn"),
                             max_tasks_per_child=1) as pool:
        outcomes = [pool.submit(_synthetic_spawned_worker, name, str(tmp_path)) for name in ("A", "B")]
        records = [future.result(timeout=30) for future in outcomes]
    assert all(record["status"] == "completed" for record in records)
    assert len({record["worker_pid"] for record in records}) == 2
    for name in ("A", "B"):
        directory = tmp_path / "work/documents" / name
        receipt = read(directory / "execution.json")
        assert receipt["cwd"] == str(directory)
        assert receipt["sdk_db"] == str(directory / "sdk-logs/traces.db")
        assert (directory / "sdk-relative-marker.txt").is_file()


def test_resume_does_not_invent_timing_for_killed_attempt(boundary):
    args, _, _, _, _ = boundary
    run(boundary)
    fingerprint = read(args.work_dir / "run_manifest.json")["run_fingerprint"]
    cli.atomic_json(args.work_dir / "attempts/killed.preflight.json", {"run_fingerprint": fingerprint})
    args.resume = True
    result = run(boundary)
    assert result["cumulative_timing_complete"] is False
    assert result["attempts_with_unknown_wall_time"] == ["killed"]
    assert result["cumulative_reports_per_hour"] is None
    assert read(args.result_dir / "infer_time.json")["timing_complete"] is False


def test_code_mutation_during_workflow_is_not_published(boundary, monkeypatch):
    _, code_root, calls, Workflow, _ = boundary

    class Mutating(Workflow):
        def run(self):
            (code_root / "src/bridge_evidence/contest/document_prompt.py").write_text("changed during run")
            return super().run()

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, Mutating))
    with pytest.raises(RuntimeError):
        run(boundary)
    assert calls["writes"] == 0


@pytest.mark.parametrize("fault", ["unreviewed", "changed_draft", "changed_digest"])
def test_publication_requires_exact_model_accepted_draft(boundary, monkeypatch, fault):
    args, _, calls, Workflow, _ = boundary

    class InvalidDelivery(Workflow):
        def run(self):
            result = super().run()
            if fault == "unreviewed":
                result["receipt"]["delivery"]["ready"] = False
            elif fault == "changed_draft":
                result["data"] = {**result["data"], "基本信息": {"桥梁名称": "另一座桥"}}
            else:
                result["receipt"]["draft_sha256"] = "unreviewed-hash"
            return result

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, InvalidDelivery))
    with pytest.raises(RuntimeError):
        run(boundary)
    assert calls["writes"] == 0
    assert not list(args.result_dir.glob("*.docx"))
    assert "accepted draft" in read(args.work_dir / "summary.json")["documents"][0]["error"]


@pytest.mark.parametrize("artifact", ["source_manifest", "source_binding"])
def test_resume_requires_unchanged_parsed_source_receipt(boundary, artifact):
    args, _, calls, _, _ = boundary
    run(boundary)
    path = args.work_dir / "documents/TRAIN-1" / (artifact + ".json")
    value = read(path)
    value["manifest_sha256" if artifact == "source_binding" else "units"] = "tampered"
    cli.atomic_json(path, value)
    args.resume = True
    with pytest.raises(RuntimeError):
        run(boundary)
    assert calls["workflows"] == 1


@pytest.mark.parametrize("mutation", ["during_workflow", "between_attempts"])
def test_converted_docx_bytes_are_immutable(boundary, monkeypatch, mutation):
    args, code_root, calls, Workflow, _ = boundary
    args.training_evidence_json = None
    args.input_dir = args.work_dir.parent / "originals"
    args.input_dir.mkdir()
    (args.input_dir / "synthetic.doc").write_bytes(b"synthetic Word 97 source")
    monkeypatch.setattr(cli.shutil, "which", lambda command: "/synthetic/soffice")
    monkeypatch.setattr(cli, "build_manifest", lambda path: {"document_sha256": cli.file_digest(path), "units": []})
    converted = []

    def convert(source, directory):
        directory.mkdir(parents=True, exist_ok=True)
        target = directory / "synthetic.docx"
        if not target.exists():
            target.write_bytes(b"synthetic converted bytes")
        converted.append(target)
        return target

    class ConversionWorkflow(Workflow):
        def run(self):
            if mutation == "during_workflow":
                calls["source_path"].write_bytes(b"changed after parsing")
            elif not args.resume:
                raise RuntimeError("synthetic interruption before completion")
            return super().run()

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, ConversionWorkflow))
    with pytest.raises(RuntimeError):
        cli.run_document_agent(args, code_root=code_root, convert_doc=convert)
    if mutation == "between_attempts":
        converted[0].write_bytes(b"another conversion of same original")
        args.resume = True
        with pytest.raises(RuntimeError):
            cli.run_document_agent(args, code_root=code_root, convert_doc=convert)
        assert calls["workflows"] == 1
    assert calls["writes"] == 0
    assert not list(args.result_dir.glob("*.docx"))


def test_real_docx_writer_exports_three_tables_without_completing_unknown_history(boundary, monkeypatch):
    from docx import Document
    from bridge_evidence.contest.docx_writer import write_numbered_extraction_docx as write_extraction_docx
    args, _, _, Workflow, _ = boundary
    data = {"基本信息": {"桥梁名称": "合成桥", "报告日期": "2026年", "上一次定检结果": "未知"},
            "维修建议": [{"建议类别": "保养", "建议内容": "清理排水孔", "病害部位": "左幅第2孔"}],
            "病害明细": [{"病害部位": "左幅第2孔", "病害类型": "堵塞", "病害描述": "排水孔杂物堵塞",
                          "是否新增": "未知", "上一次定检状态": "", "发展程度": ""}]}

    class ReportWorkflow(Workflow):
        def run(self):
            result = super().run()
            result["data"] = data
            result["receipt"].update(final_data=data, draft_sha256=cli.json_digest(data))
            return result

    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, ReportWorkflow))
    monkeypatch.setattr(cli, "write_extraction_docx", write_extraction_docx)
    run(boundary)
    document = Document(args.result_dir / "TRAIN-1.docx")
    assert len(document.tables) == 3
    basic = {row.cells[0].text: row.cells[1].text for row in document.tables[0].rows[1:]}
    assert basic == data["基本信息"]
    assert [c.text for c in document.tables[1].rows[1].cells] == ["1", "保养", "清理排水孔", "左幅第2孔"]
    assert [c.text for c in document.tables[2].rows[1].cells] == ["1", "左幅第2孔", "堵塞", "排水孔杂物堵塞", "未知", "", ""]
    assert "序号" not in data["维修建议"][0] and "序号" not in data["病害明细"][0]
    assert "（对比年度信息待确认）" in [p.text for p in document.paragraphs]


def test_actual_workflow_constructor_receipt_and_docx_boundary_with_synthetic_review(boundary, monkeypatch):
    """The execution double supplies a review; actual state/receipt/export APIs run."""
    from docx import Document
    from bridge_evidence.contest.document_workflow import DocumentSettings, DocumentWorkflow
    from bridge_evidence.contest.docx_writer import write_numbered_extraction_docx as write_extraction_docx
    args, _, calls, _, _ = boundary
    # This receipt/export double has no HTTP transport; select its legacy path.
    args.document_reliable_generation = False
    args.document_max_model_calls = 27
    Client = cli.divide_cli._make_client
    monkeypatch.setattr(Client, "model", args.served_model, raising=False)
    monkeypatch.setattr(Client, "max_model_len", 65536, raising=False)
    monkeypatch.setattr(Client, "timeout", 180.0, raising=False)
    monkeypatch.setattr(Client, "count_tokens", staticmethod(lambda text: max(1, len(text) // 4)), raising=False)
    data = {"基本信息": {"桥梁名称": "合成桥"}, "病害明细": [], "维修建议": []}

    class SyntheticExecution(DocumentWorkflow):
        def run(self):
            assert self.settings.max_model_calls == 27
            unit = self.manifest["units"][0]
            for item in self.manifest["units"]:
                self.workspace.record_read(item["id"], 0, len(item["text"]))
            state = self.workspace.apply_update({
                "base_version": self.workspace.snapshot()["version"],
                "facts": [{"id": "F1", "value": "合成桥", "status": "explicit", "depends_on": [],
                           "source_refs": [{"evidence_id": unit["id"], "start": 0, "end": len(unit["text"])}]}],
                "draft": data, "draft_fact_ids": ["F1"], "draft_bindings": {"基本信息.桥梁名称": ["F1"]}})
            # Explicitly synthetic receipt. No agent, generation or HTTP ran.
            self.client.calls.append({"kind": "generation", "call_id": "synthetic-review", "synthetic": True})
            self.workspace.record_review({"workspace_version": state["version"],
                "draft_sha256": state["draft_sha256"], "passed": True,
                "checked_evidence_ids": [u["id"] for u in self.manifest["units"]],
                "unsupported_claims": [], "missing_items": [], "conflicts": [], "unresolved_items": [],
                "model_call_ids": ["synthetic-review"]})
            return {"data": data, "receipt": self._receipt("completed", "synthetic")}

    monkeypatch.setattr(cli, "_workflow_types", lambda: (DocumentSettings, SyntheticExecution))
    monkeypatch.setattr(cli, "write_extraction_docx", write_extraction_docx)
    result = run(boundary)
    directory = args.work_dir / "documents/TRAIN-1"
    receipt = read(directory / "execution.json")
    identity = read(args.work_dir / "run_manifest.json")
    assert receipt["status"] == "completed" and receipt["delivery"]["ready"]
    assert receipt["identity"]["run_identity"] == {"run_fingerprint": identity["run_fingerprint"],
                                                   "source_sha256": cli.file_digest(args.training_evidence_json)}
    assert receipt["draft_sha256"] == receipt["report_sha256"] == cli.json_digest(data)
    assert len(Document(args.result_dir / "TRAIN-1.docx").tables) == 3
    assert result["completed_reports"] == 1
    args.resume = True
    resumed = run(boundary)
    assert resumed["reused_reports"] == 1
    assert read(directory / "execution.json")["model_calls"] == receipt["model_calls"]


@pytest.fixture
def ocr_runtime(boundary, monkeypatch):
    from types import SimpleNamespace
    from bridge_evidence.contest import document_tools
    args, _, calls, _, _ = boundary
    args.document_ocr = "tesseract"
    directory = args.work_dir.parent / "synthetic-tessdata"
    directory.mkdir()
    executable = directory.parent / "synthetic-tesseract"
    executable.write_bytes(b"synthetic tesseract executable; never run")
    for language in ("chi_sim", "eng"):
        (directory / (language + ".traineddata")).write_bytes(("synthetic " + language).encode())
    calls["ocr_probes"], calls["ocr_reads"] = [], []

    def recognize(blob, *, region):
        calls["ocr_reads"].append({"prefix": os.environ.get("TESSDATA_PREFIX"), "blob": blob, "region": region})
        return {"kind": "ocr", "text": "合成文字", "verification": "unverified"}

    def factory(**kwargs):
        return SimpleNamespace(executable=str(executable), recognize=recognize, **kwargs)

    def command(argv, **kwargs):
        assert argv[0] == str(executable)
        assert kwargs["capture_output"] and kwargs["check"]
        calls["ocr_probes"].append(argv)
        if argv[1:] == ["--version"]:
            return SimpleNamespace(stdout="tesseract synthetic-5\nleptonica synthetic", stderr="")
        assert argv[1:] == ["--list-langs"]
        return SimpleNamespace(stdout=f'List of available languages in "{directory}" (2):\nchi_sim\neng\n', stderr="")

    monkeypatch.setattr(document_tools, "make_tesseract_backend", factory)
    monkeypatch.setattr(cli.subprocess, "run", command)
    return executable, directory


def test_default_ocr_does_not_probe_or_create_backend(boundary, monkeypatch):
    from bridge_evidence.contest import document_tools
    args, _, calls, _, _ = boundary
    monkeypatch.setattr(document_tools, "make_tesseract_backend", lambda **kwargs: pytest.fail("OCR is disabled"))
    run(boundary)
    assert calls["vision_backend"] is None
    assert read(args.work_dir / "run_manifest.json")["identity"]["vision_backend"]["status"] == "disabled"


def test_explicit_ocr_binds_real_artifacts_and_reaches_workflow(boundary, ocr_runtime, monkeypatch):
    args, _, calls, _, _ = boundary
    executable, directory = ocr_runtime
    monkeypatch.setenv("TESSDATA_PREFIX", "previous-prefix")
    run(boundary)
    identity = read(args.work_dir / "run_manifest.json")["identity"]["vision_backend"]
    assert identity["status"] == "available" and identity["device"] == "cpu"
    assert identity["executable"] == {"path": str(executable), "sha256": cli.file_digest(executable)}
    assert identity["traineddata"]["chi_sim"]["sha256"] == cli.file_digest(directory / "chi_sim.traineddata")
    assert identity["version"] == "tesseract synthetic-5\nleptonica synthetic"
    assert identity["pillow_version"] == "synthetic-1.0"
    assert identity["external_vision_api"] is False and identity["shared_libraries_checksum_verified"] is False
    backend = calls["vision_backend"]
    assert backend.identity == identity and cli.json_digest(identity) in backend.cache_key
    assert not hasattr(backend, "interpret")
    assert backend.recognize(b"synthetic pixels", region={})["verification"] == "unverified"
    assert calls["ocr_reads"] == [{"prefix": str(directory), "blob": b"synthetic pixels", "region": {}}]
    assert os.environ["TESSDATA_PREFIX"] == "previous-prefix"


def test_ocr_check_reports_availability_and_limits_without_recognition(boundary, ocr_runtime, capsys):
    args, _, calls, _, _ = boundary
    args.check = True
    run(boundary)
    output = capsys.readouterr().out
    assert "CPU OCR available: tesseract" in output
    assert "does not understand diagrams" in output
    assert calls["ocr_reads"] == [] and calls["workflows"] == 0
    assert not args.work_dir.exists()


def test_ocr_check_reports_unavailable_and_does_not_silently_disable(boundary, monkeypatch, capsys):
    from bridge_evidence.contest import document_tools
    args, _, calls, _, _ = boundary
    args.document_ocr, args.check = "tesseract", True
    def unavailable(**kwargs):
        raise RuntimeError("Tesseract language data unavailable: chi_sim")
    monkeypatch.setattr(document_tools, "make_tesseract_backend", unavailable)
    with pytest.raises(RuntimeError, match="chi_sim"):
        run(boundary)
    output = capsys.readouterr().out
    assert "CPU OCR unavailable" in output and "does not understand diagrams" in output
    assert calls["clients"] == calls["workflows"] == 0
    assert not args.work_dir.exists()


@pytest.mark.parametrize("change", ["executable", "traineddata", "language", "timeout", "disabled", "pillow"])
def test_resume_rejects_changed_ocr_artifacts_or_configuration(boundary, ocr_runtime, monkeypatch, change):
    args, _, calls, _, _ = boundary
    executable, directory = ocr_runtime
    run(boundary)
    args.resume = True
    if change in {"executable", "traineddata"}:
        path = executable if change == "executable" else directory / "chi_sim.traineddata"
        path.write_bytes(path.read_bytes() + b"changed")
    elif change == "language":
        args.document_ocr_language = "eng"
    elif change == "timeout":
        args.document_ocr_timeout += 1
    elif change == "disabled":
        args.document_ocr = "none"
    else:
        monkeypatch.setattr(cli.importlib.metadata, "version", lambda name: "changed" if name == "Pillow" else "synthetic-1.0")
    with pytest.raises(ValueError, match="Run identity changed"):
        run(boundary)
    assert calls["workflows"] == 1


def test_ocr_mutation_during_workflow_blocks_publication(boundary, ocr_runtime, monkeypatch):
    args, _, calls, Workflow, _ = boundary
    _, directory = ocr_runtime
    class MutatingOCR(Workflow):
        def run(self):
            (directory / "eng.traineddata").write_bytes(b"changed during workflow")
            return super().run()
    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, MutatingOCR))
    with pytest.raises(RuntimeError):
        run(boundary)
    assert calls["writes"] == 0
    assert not list(args.result_dir.glob("*.docx"))


def test_bound_ocr_refuses_changed_binary_before_recognition(boundary, ocr_runtime):
    args, _, calls, _, _ = boundary
    executable, _ = ocr_runtime
    backend, _ = cli._make_ocr_backend(args)
    executable.write_bytes(b"replaced executable")
    with pytest.raises(ValueError, match="changed"):
        backend.recognize(b"synthetic pixels", region={})
    assert calls["ocr_reads"] == []


def test_bound_ocr_tracks_language_pack_symlink_target(boundary, ocr_runtime):
    args, _, calls, _, _ = boundary
    _, directory = ocr_runtime
    lookup = directory / "chi_sim.traineddata"
    first, second = directory.parent / "language-one", directory.parent / "language-two"
    first.write_bytes(lookup.read_bytes())
    second.write_bytes(first.read_bytes())
    lookup.unlink()
    lookup.symlink_to(first)
    backend, identity = cli._make_ocr_backend(args)
    assert identity["traineddata"]["chi_sim"]["resolved_path"] == str(first)
    lookup.unlink()
    lookup.symlink_to(second)
    with pytest.raises(ValueError, match="changed"):
        backend.recognize(b"synthetic pixels", region={})
    assert calls["ocr_reads"] == []


def test_same_size_docx_content_rewrite_blocks_publication(boundary, monkeypatch):
    from copy import deepcopy
    from bridge_evidence.contest.docx_writer import write_numbered_extraction_docx as write_extraction_docx
    args, _, _, _, _ = boundary
    def changed_writer(data, path):
        changed = deepcopy(data)
        changed["基本信息"]["桥梁名称"] = "另一座桥"
        write_extraction_docx(changed, path)
    monkeypatch.setattr(cli, "write_extraction_docx", changed_writer)
    with pytest.raises(RuntimeError):
        run(boundary)
    assert not list(args.result_dir.glob("*.docx"))
    assert not list(args.result_dir.glob(".pending-*"))
    summary = read(args.work_dir / "summary.json")
    assert summary["failed_reports"] == 1
    assert "differs from reviewed delivery" in summary["documents"][0]["error"]


def test_resume_reads_word_values_again_even_when_file_hashes_were_rebound(boundary):
    from docx import Document
    args, _, calls, _, _ = boundary
    run(boundary)
    directory = args.work_dir / "documents/TRAIN-1"
    output = args.result_dir / "TRAIN-1.docx"
    document = Document(output)
    document.tables[0].rows[1].cells[1].text = "另一座桥"
    document.save(output)
    receipt = read(directory / "execution.json")
    receipt["output"]["sha256"] = cli.file_digest(output)
    receipt["delivery_validation"]["docx_sha256"] = cli.file_digest(output)
    cli.atomic_json(directory / "execution.json", receipt)
    digest = cli.file_digest(directory / "execution.json")
    for path in (directory / "publication.json", directory / "document_summary.json"):
        previous = read(path)
        previous["execution_sha256"] = digest
        cli.atomic_json(path, previous)
    args.resume = True
    with pytest.raises(RuntimeError):
        run(boundary)
    assert calls["workflows"] == 1
    assert "differs from reviewed delivery" in read(args.work_dir / "summary.json")["documents"][0]["error"]


def test_export_rejects_reviewed_fields_that_the_writer_cannot_represent(boundary, monkeypatch):
    args, _, calls, Workflow, _ = boundary
    class UnsupportedField(Workflow):
        def run(self):
            result = super().run()
            result["data"]["基本信息"]["未支持字段"] = "会被静默丢弃"
            result["receipt"]["draft_sha256"] = cli.json_digest(result["data"])
            return result
    monkeypatch.setattr(cli, "_workflow_types", lambda: (Settings, UnsupportedField))
    with pytest.raises(RuntimeError):
        run(boundary)
    assert calls["writes"] == 0
    assert "unsupported fields" in read(args.work_dir / "summary.json")["documents"][0]["error"]
