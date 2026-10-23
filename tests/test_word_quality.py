from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from PIL import Image

from bridge_evidence.adapters.word_quality import (
    DocumentFingerprint,
    compare_page_images,
    compare_structure,
    environment_fingerprint,
    fingerprint_docx,
    fingerprint_wv_html,
    parse_fib_base,
)


def test_fib_preflight_detects_encryption_and_obfuscation() -> None:
    fib = bytearray(32)
    fib[0:2] = (0xA5EC).to_bytes(2, "little")
    fib[2:4] = (0x00D9).to_bytes(2, "little")
    fib[10:12] = ((1 << 8) | (1 << 15)).to_bytes(2, "little")

    result = parse_fib_base(bytes(fib))

    assert result.valid is True
    assert result.encrypted is True
    assert result.obfuscated is True


def test_wv_html_fingerprint_counts_source_structure_without_markup() -> None:
    fingerprint = fingerprint_wv_html(
        """
        <html><body>
          <table>
            <tr><th>构件</th><th>评分</th></tr>
            <tr><td>主梁</td><td>88.50</td></tr>
          </table>
          <img src="one.png"><img src="two.png">
        </body></html>
        """
    )

    assert fingerprint.tables == 1
    assert fingerprint.rows == 2
    assert fingerprint.cells == 4
    assert fingerprint.image_occurrences == 2
    assert fingerprint.text_chars == len("构件评分主梁88.50")


def test_docx_fingerprint_counts_stories_drawings_and_external_links(tmp_path: Path) -> None:
    image_path = tmp_path / "sample.png"
    Image.new("RGB", (20, 20), "red").save(image_path)
    path = tmp_path / "fingerprint.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "页眉评分表"
    document.sections[0].footer.paragraphs[0].text = "第1页"
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "构件"
    table.cell(0, 1).text = "评分"
    document.add_picture(str(image_path))
    document.add_picture(str(image_path))
    document.part.relate_to(
        "https://example.invalid/evidence",
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    document.save(path)

    fingerprint = fingerprint_docx(path)

    assert fingerprint.tables == 1
    assert fingerprint.rows == 2
    assert fingerprint.cells == 4
    assert fingerprint.image_occurrences == 2
    assert fingerprint.unique_images == 1
    assert fingerprint.story_parts == ("body", "footer", "header")
    assert fingerprint.external_relationships == 1


def test_structure_gate_rejects_missing_tables_or_images() -> None:
    source = DocumentFingerprint(
        provider="wvHtml",
        text_chars=100,
        tables=2,
        rows=10,
        cells=40,
        image_occurrences=3,
    )
    target = DocumentFingerprint(
        provider="ooxml",
        text_chars=100,
        tables=0,
        rows=0,
        cells=0,
        image_occurrences=0,
    )

    result = compare_structure(source, target)

    assert result.accepted is False
    assert result.requires_review is True
    assert "source_tables_missing_after_conversion" in result.warnings
    assert "source_images_missing_after_conversion" in result.warnings


def test_visual_gate_accepts_identical_pages_and_rejects_changed_or_missing_pages() -> None:
    white = Image.new("RGB", (120, 160), "white")
    black = Image.new("RGB", (120, 160), "black")

    identical = compare_page_images([white], [white.copy()])
    changed = compare_page_images([white], [black])
    missing = compare_page_images([white, white.copy()], [white.copy()])

    assert identical.accepted is True
    assert identical.score == 1.0
    assert changed.accepted is False
    assert changed.score is not None and changed.score < 0.5
    assert missing.accepted is False
    assert "visual_page_count_mismatch" in missing.warnings


def test_visual_gate_aligns_small_page_reflow_in_long_documents() -> None:
    source = [Image.new("RGB", (120, 160), (240 - index, 240, 240)) for index in range(40)]
    target = [page.copy() for page in source]
    target.insert(20, Image.new("RGB", (120, 160), "black"))

    result = compare_page_images(source, target)

    assert result.accepted is True
    assert result.requires_review is True
    assert "visual_page_count_mismatch" in result.warnings
    assert "visual_page_alignment_applied" in result.warnings
    assert len(result.details["unmatched_target_pages"]) == 1
    assert result.details["unmatched_source_pages"] == []


def test_environment_fingerprint_is_stable_and_sensitive() -> None:
    first = environment_fingerprint({
        "soffice": "LibreOffice 26.2",
        "pandoc": "pandoc 3.7",
        "platform": "darwin-arm64",
        "fonts": "font-set-a",
    })
    reordered = environment_fingerprint({
        "fonts": "font-set-a",
        "platform": "darwin-arm64",
        "pandoc": "pandoc 3.7",
        "soffice": "LibreOffice 26.2",
    })
    changed = environment_fingerprint({
        "soffice": "LibreOffice 26.3",
        "pandoc": "pandoc 3.7",
        "platform": "darwin-arm64",
        "fonts": "font-set-a",
    })

    assert first == reordered
    assert first != changed
