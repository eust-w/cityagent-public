"""Real DOCX write/read checks on synthetic advice summaries; no model or GPU."""
from copy import deepcopy

from docx import Document
import pytest

from bridge_evidence.contest.docx_writer import write_extraction_docx, write_numbered_extraction_docx


def render(tmp_path, basic, *, numbered=False):
    # Source numbers may be absent or out of order. Display numbering must not
    # sort/deduplicate content or mutate the saved rows.
    advice = [
        {"序号": "8", "建议类别": "立即处置", "建议内容": "更换伸缩缝（保留限定）", "病害部位": "P2"},
        {"建议类别": "尽快维修", "建议内容": "2#台加固", "病害部位": "2#台"},
    ]
    advice.append(deepcopy(advice[0]))
    defects = [
        {"序号": "9", "病害部位": "P2", "病害类型": "裂缝", "病害描述": "原有裂缝",
         "是否新增": "无法判断", "上一次定检状态": "无", "发展程度": "无法比较"},
        {"序号": None, "病害部位": "P1", "病害类型": "剥落", "病害描述": "局部剥落",
         "是否新增": "否", "上一次定检状态": "已有", "发展程度": "无"},
    ]
    defects.append(deepcopy(defects[0]))
    data = {"基本信息": {"桥梁名称": "合成建议字段桥", **basic},
            "维修建议": advice, "病害明细": defects}
    before = deepcopy(data)
    path = tmp_path / "advice-fields.docx"
    writer = write_numbered_extraction_docx if numbered else write_extraction_docx
    writer(data, str(path))
    assert data == before
    document = Document(path)
    assert len(document.tables) == 3
    rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows[1:]]
    return document, rows, before


@pytest.mark.parametrize("original", [
    "1条立即处置（伸缩缝断裂更换）、4条尽快维修、2条预防性养护建议",
    "1条立即处置（2#台加固）、4条尽快维修、1条预防性养护建议",
])
def test_original_advice_keeps_maintenance_qualifiers_without_generating_derived_stats(tmp_path, original):
    _, rows, _ = render(tmp_path, {"建议": original})
    assert [row for row in rows if row[0] == "建议"] == [["建议", original, "建议分类统计"]]
    assert not any(row[0] == "建议统计" for row in rows)


def test_derived_statistics_keep_their_own_field_without_fabricating_original_advice(tmp_path):
    derived = "共提出维护建议 3 项：立即处置 2 项，尽快维修 1 项。"
    _, rows, _ = render(tmp_path, {"建议统计": derived})
    assert [row for row in rows if row[0] == "建议统计"] == [["建议统计", derived, "按建议明细条目统计"]]
    assert not any(row[0] == "建议" for row in rows)


def test_conflicting_original_and_derived_summaries_remain_separate_and_verbatim(tmp_path):
    original = "1条立即处置（2#台加固）、4条尽快维修"
    derived = "共提出维护建议 3 项：立即处置 2 项，尽快维修 1 项。"
    _, rows, _ = render(tmp_path, {"建议": original, "建议统计": derived})
    assert [row for row in rows if row[0] in {"建议", "建议统计"}] == [
        ["建议", original, "建议分类统计"],
        ["建议统计", derived, "按建议明细条目统计"],
    ]


@pytest.mark.parametrize("derived", [
    "1条立即处置、2条尽快维修",
    "共提出维护建议 3 项：立即处置 1 项，尽快维修 2 项。",
])
def test_equal_values_or_counts_do_not_deduplicate_two_distinct_source_fields(tmp_path, derived):
    original = "1条立即处置、2条尽快维修"
    _, rows, _ = render(tmp_path, {"建议": original, "建议统计": derived})
    assert [row[0] for row in rows if row[0] in {"建议", "建议统计"}] == ["建议", "建议统计"]
    assert {row[0]: row[1] for row in rows}["建议"] == original
    assert {row[0]: row[1] for row in rows}["建议统计"] == derived


@pytest.mark.parametrize("basic", [{}, {"建议": "", "建议统计": ""}, {"建议": None, "建议统计": None}])
def test_missing_or_empty_summaries_do_not_generate_either_field_from_detail_rows(tmp_path, basic):
    _, rows, _ = render(tmp_path, basic)
    assert not any(row[0] in {"建议", "建议统计"} for row in rows)


@pytest.mark.parametrize("numbered", [False, True])
def test_numbering_is_explicit_and_preserves_content_order_and_saved_rows(tmp_path, numbered):
    document, _, before = render(tmp_path, {
        "建议": "1条立即处置（伸缩缝断裂更换）",
        "建议统计": "共提出维护建议 3 项。",
    }, numbered=numbered)
    for table, section, fields in (
        (document.tables[1], "维修建议", ["序号", "建议类别", "建议内容", "病害部位"]),
        (document.tables[2], "病害明细", ["序号", "病害部位", "病害类型", "病害描述", "是否新增", "上一次定检状态", "发展程度"]),
    ):
        assert [cell.text for cell in table.rows[0].cells] == fields
        actual = [[cell.text for cell in row.cells] for row in table.rows[1:]]
        expected_numbers = (["1", "2", "3"] if numbered else
                            [str(row.get("序号") or "") for row in before[section]])
        assert [row[0] for row in actual] == expected_numbers
        assert [row[1:] for row in actual] == [
            [row[field] for field in fields[1:]] for row in before[section]
        ]
        assert len(actual) == 3 and actual[0][1:] == actual[2][1:]
