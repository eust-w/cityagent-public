"""Create a fictional report and exercise the public CLI without model access."""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import subprocess
import sys

from docx import Document


def create_report(path: Path) -> None:
    document = Document()
    document.add_heading("云样桥定期检测报告", 0)
    document.add_paragraph("本报告完全虚构，仅用于软件演示，不对应任何真实桥梁。")
    document.add_paragraph("2026年10月22日")
    table = document.add_table(rows=0, cols=2)
    for values in (
        ("工程名称", "云样桥"),
        ("检测日期", "2026年10月22日"),
        ("检测结论", "桥梁总体BCI=92.0，技术状况等级为A级。"),
    ):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    document.add_paragraph("本桥无往年检测报告，故无上一次检测报告。")
    document.add_heading("技术状况评定", 1)
    table = document.add_table(rows=0, cols=3)
    for values in (
        ("部位", "评分", "等级"),
        ("桥面系", "91.0", "A级"),
        ("上部结构", "92.0", "A级"),
        ("下部结构", "93.0", "A级"),
        ("全桥", "92.0", "A级"),
    ):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    document.add_heading("主梁检测结果表", 1)
    table = document.add_table(rows=0, cols=4)
    for values in (
        ("序号", "位置", "病害种类", "病害描述"),
        ("1", "1跨1号主梁", "裂缝", "1跨1号主梁出现一条裂缝，长度0.2m，宽度0.1mm。"),
    ):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    document.add_heading("维修建议", 1)
    document.add_paragraph("建议对主梁裂缝进行封闭处理，并加强日常巡查。")
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True, help="New directory; existing directories are never overwritten")
    args = parser.parse_args()
    root = args.output_dir.expanduser().resolve()
    try:
        root.mkdir(parents=True, exist_ok=False)
    except FileExistsError:
        parser.error(f"Output directory already exists: {root}")
    source = root / "synthetic-report.docx"
    create_report(source)
    config = root / "offline.json"
    config.write_text(json.dumps({
        "llm_mode": "off",
        "work_dir": str(root / "work"),
        "cache_dir": str(root / "cache"),
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    # Ignore user gateway/external document-expert overrides for this fixed,
    # offline demonstration. The regular CLI still supports those overrides.
    environment = {key: value for key, value in os.environ.items()
                   if not key.startswith("BRIDGE_EVIDENCE_")}
    subprocess.run([
        sys.executable, "-m", "bridge_evidence.cli", "--config", str(config),
        "single", "--input", str(source), "--output-dir", str(root / "result"),
        "--json-dir", str(root / "json"), "--llm-mode", "off",
    ], check=True, env=environment)
    print(f"\nFictional source: {source}")
    print(f"Word result: {root / 'result' / source.name}")
    print(f"Evidence JSON: {root / 'json' / (source.stem + '.json')}")
    print("Synthetic smoke test only; this is not a model quality benchmark.")


if __name__ == "__main__":
    main()
