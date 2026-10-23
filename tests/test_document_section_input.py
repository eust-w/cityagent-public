from __future__ import annotations

import copy
import json

import pytest
from docx import Document

from bridge_evidence.contest.document_model_source import ModelSourceView
from bridge_evidence.contest.document_section_input import SectionInputBudgetError, compact_section_prompt
from bridge_evidence.contest.document_task_plan import build_document_task_plan
from bridge_evidence.contest.evidence_manifest import build_manifest, manifest_from_text


def _data(prompt):
    return json.loads(prompt.split("\n", 1)[1])


def _decoded_cells(data, row):
    return [dict(zip(data["cell_fields"], values, strict=True)) for values in row["cells"]]


def _task(manifest, *, primary=None, context=None):
    return {"id": "synthetic-document-task", "primary_evidence_ids": primary if primary is not None else [u["id"] for u in manifest["units"]],
            "context_evidence_ids": context or [], "headings": [], "table_ids": []}


def _row_metadata(manifest, table_id, *, parent=None):
    return {"id": table_id, "parent_table_id": parent, "row_count": 3,
            "header_evidence_ids": [manifest["units"][0]["id"]], "header_kind": "explicit",
            "source_path": f"original.xml#/table[{table_id}]", "context": {"story": "body"}}


def test_initial_representation_keeps_every_source_character_and_offset_without_mutation():
    original = "  原文😀\r\n\r\n保留\t制表符  与空格\n最后一段 "
    manifest = manifest_from_text(original)
    task = _task(manifest)
    before = copy.deepcopy((manifest, task))
    prompt, exposed = compact_section_prompt(manifest, task)
    data = _data(prompt)
    assert "".join(row["text"] for row in data["source_units"]).encode() == original.encode()
    assert (manifest, task) == before
    mapping = ModelSourceView(manifest).source_map()
    for unit, row, span in zip(manifest["units"], data["source_units"], exposed):
        assert row["evidence_id"] == unit["id"]
        assert row["text"].encode() == unit["text"].encode()
        assert row["source_start"] == unit["source_start"]
        assert row["source_end"] == unit["source_end"]
        assert span == {"evidence_id": unit["id"], "start": 0, "end": len(unit["text"])}
        assert mapping[row["source_id"]] == unit["source_path"]
    assert len(prompt) <= 60000


def test_long_repeated_paths_neighbor_metadata_and_task_ranges_are_not_embedded():
    manifest = manifest_from_text("甲\n乙\n丙")
    long_path = "original.xml#/" + "nested-wrapper/" * 500
    for i, unit in enumerate(manifest["units"]):
        unit["source_path"] = long_path + str(i)
        unit["context"].update(section_heading="冗长重复标题" * 500,
                               preceding_evidence_ids=[u["id"] for u in manifest["units"]])
    task = _task(manifest)
    task.update(primary_ranges=[{"source_path": long_path, "repeated": "冗余" * 3000}] * 30,
                table_context=[{"unused": "冗余" * 3000}], unrelated_key="冗余" * 3000)
    prompt, exposed = compact_section_prompt(manifest, task, 5000)
    data = _data(prompt)
    assert len(exposed) == len(manifest["units"])
    assert long_path not in prompt
    assert "primary_ranges" not in prompt and "preceding_evidence_ids" not in prompt
    assert "unrelated_key" not in prompt and "table_context" not in prompt
    assert all("context" not in row for row in data["source_units"])


def test_context_ids_have_priority_are_embedded_once_and_do_not_repeat_responsibility():
    manifest = manifest_from_text("P" * 100 + "\n" + "单位:mm\n" + "注:原文限定")
    ids = [u["id"] for u in manifest["units"]]
    task = _task(manifest, primary=[ids[0]], context=[ids[1], ids[2], ids[1]])
    full, _ = compact_section_prompt(manifest, task)
    data = _data(full)
    assert [row["evidence_id"] for row in data["source_units"]] == [ids[1], ids[2], ids[0]]
    assert data["task"]["context_evidence_ids"] == ids[1:]
    first_row_size = len(json.dumps(data["source_units"][0], ensure_ascii=False, separators=(",", ":")))
    empty = copy.deepcopy(data)
    empty["source_units"] = []
    cap = len(full) - len(json.dumps(data, ensure_ascii=False, separators=(",", ":"))) + len(json.dumps(empty, ensure_ascii=False, separators=(",", ":"))) + first_row_size
    limited, exposed = compact_section_prompt(manifest, task, cap)
    assert len(limited) == cap
    assert exposed == [{"evidence_id": ids[1], "start": 0, "end": len(manifest["units"][1]["text"])}]
    assert _data(limited)["task"]["primary_evidence_ids"] == [ids[0]]


def test_exact_whole_prompt_budget_includes_prefix_notes_and_navigation():
    manifest = manifest_from_text("原文😀")
    task = _task(manifest)
    full, exposed = compact_section_prompt(manifest, task)
    exact, exact_exposed = compact_section_prompt(manifest, task, len(full))
    assert exact == full and exact_exposed == exposed
    short, short_exposed = compact_section_prompt(manifest, task, len(full) - 1)
    assert len(short) < len(full) - 1
    assert short_exposed == []
    assert _data(short)["source_units"] == []
    assert _data(short)["task"]["primary_evidence_ids"] == task["primary_evidence_ids"]


def test_navigation_cannot_fit_causes_clear_failure_without_silent_id_truncation():
    manifest = manifest_from_text("原文\n" * 100)
    task = _task(manifest)
    before = copy.deepcopy(task)
    with pytest.raises(SectionInputBudgetError) as caught:
        compact_section_prompt(manifest, task, 100)
    assert caught.value.limit == 100
    assert caught.value.metadata_characters > 100
    assert "no evidence IDs were dropped" in str(caught.value)
    assert task == before


def test_giant_unit_is_not_truncated_and_later_whole_units_can_still_fit():
    manifest = manifest_from_text("很长😀" * 20000 + "\n短文\n\n")
    prompt, exposed = compact_section_prompt(manifest, _task(manifest), 4000)
    assert len(prompt) <= 4000
    rows = _data(prompt)["source_units"]
    assert [row["text"] for row in rows] == [u["text"] for u in manifest["units"][1:]]
    assert all(span["evidence_id"] != manifest["units"][0]["id"] for span in exposed)
    assert _data(prompt)["task"]["primary_evidence_ids"] == [u["id"] for u in manifest["units"]]


def test_merge_metadata_is_distinct_from_actual_blank_cell_and_normal_text_is_not_duplicated():
    manifest = manifest_from_text("继承原值 | 数据\n | 下一行\n | 实际空白")
    table = _row_metadata(manifest, "T-original")
    manifest["tables"] = [table]
    for i, unit in enumerate(manifest["units"]):
        unit.update(kind="table_row", table_id="T-original", row_index=i, segment_index=0, is_header=i == 0)
        unit["context"]["parent_table_id"] = None
        unit["cells"] = [{"text": "继承原值" if i == 0 else "", "column_index": 0, "column_span": 1,
                          "vertical_merge": "restart" if i == 0 else "continue" if i == 1 else None,
                          "merged_from_source": "original.xml#/cell[0]" if i == 1 else None,
                          "merged_context_text": "继承原值" if i == 1 else "",
                          "source_path": f"original.xml#/cell[{i}]"},
                         {"text": "数据", "column_index": 1, "column_span": 2, "vertical_merge": None,
                          "merged_from_source": None, "merged_context_text": "", "source_path": f"original.xml#/cell[{i}b]"}]
    task = _task(manifest)
    task["table_ids"] = ["T-original"]
    prompt, _ = compact_section_prompt(manifest, task)
    data = _data(prompt)
    rows = data["source_units"]
    first, merged, blank = [_decoded_cells(data, row)[0] for row in rows]
    assert "text" not in first and "text" not in merged and "text" not in blank
    assert merged["vertical_merge"] == "continue"
    assert merged["merged_context_text"] == "继承原值"
    assert merged["merged_from_id"] == first["source_id"]
    assert blank["vertical_merge"] is None and blank["merged_from_id"] is None
    assert blank["merged_context_text"] == ""
    assert _decoded_cells(data, rows[1])[1]["column_span"] == 2
    mapping = ModelSourceView(manifest).source_map()
    assert mapping[merged["merged_from_id"]] == manifest["units"][1]["cells"][0]["merged_from_source"]
    assert data["merged_origin_evidence_ids"][merged["merged_from_id"]] == [manifest["units"][0]["id"]]
    assert rows[1]["row_index"] == 1 and rows[1]["segment_index"] == 0 and rows[1]["parent_table_id"] is None


def test_nested_tables_row_segments_and_parent_addresses_are_preserved(tmp_path):
    path = tmp_path / "nested.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "原文前段"
    child = cell.add_table(rows=2, cols=1)
    child.cell(0, 0).text = "表头"
    child.cell(1, 0).text = "记录"
    cell.paragraphs[-1].text = "原文后段"
    document.save(path)
    manifest = build_manifest(path)
    task = _task(manifest)
    prompt, exposed = compact_section_prompt(manifest, task)
    data = _data(prompt)
    table_by_id = {table["id"]: table for table in data["task"]["tables"]}
    outer_info, nested_info = manifest["tables"]
    assert table_by_id[nested_info["id"]]["parent_table_id"] == outer_info["id"]
    for unit, row in zip(manifest["units"], data["source_units"]):
        assert row["text"].encode() == unit["text"].encode()
        assert row["table_id"] == unit["table_id"]
        assert row["row_index"] == unit["row_index"]
        assert row["segment_index"] == unit["segment_index"]
        assert row["parent_table_id"] == unit["context"]["parent_table_id"]
        for original, cell in zip(unit["cells"], _decoded_cells(data, row), strict=True):
            start, end = cell["text_binding"]
            assert row["text"][start:end].encode() == original["text"].encode()
    assert len(exposed) == len(manifest["units"])


def test_blank_paragraphs_between_table_units_and_notes_preserve_context_once():
    manifest = manifest_from_text("封面\n前言\n导航\n" + "表格记录" * 200 + "\n\n单位:mm\n\n注:仅左侧\n\n2 后续章节\n后续正文")
    manifest["tables"] = [_row_metadata(manifest, "T-original")]
    manifest["tables"][0]["header_evidence_ids"] = [manifest["units"][3]["id"]]
    manifest["units"][3].update(kind="table_row", table_id="T-original", row_index=0, segment_index=0)
    plan = build_document_task_plan(manifest)
    task = next(task for task in plan["tasks"] if manifest["units"][3]["id"] in task["primary_evidence_ids"])
    prompt, exposed = compact_section_prompt(manifest, task)
    rows = _data(prompt)["source_units"]
    row_by_id = {row["evidence_id"]: row for row in rows}
    assert len(row_by_id) == len(rows) == len(exposed)
    for unit in manifest["units"][4:9]:
        assert row_by_id[unit["id"]]["text"] == unit["text"]
        assert row_by_id[unit["id"]]["end"] == len(unit["text"])
    assert manifest["units"][9]["id"] not in row_by_id


def test_heading_navigation_is_by_original_id_and_title_is_not_duplicated():
    manifest = manifest_from_text("1 原始章节标题\n正文")
    task = _task(manifest)
    task["headings"] = [{"evidence_id": manifest["units"][0]["id"], "title": "1 原始章节标题",
                          "level": 1, "basis": "regex_candidate", "story": "training_text", "unit_index": 0}]
    prompt, _ = compact_section_prompt(manifest, task)
    heading = _data(prompt)["task"]["headings"][0]
    assert heading == {"evidence_id": manifest["units"][0]["id"], "level": 1,
                       "basis": "regex_candidate", "story": "training_text"}
    assert prompt.count("1 原始章节标题") == 1


def test_empty_source_unit_is_explicitly_embedded_with_zero_range():
    manifest = manifest_from_text("")
    prompt, exposed = compact_section_prompt(manifest, _task(manifest))
    assert _data(prompt)["source_units"][0]["text"] == ""
    assert exposed == [{"evidence_id": manifest["units"][0]["id"], "start": 0, "end": 0}]


def test_cell_delimiter_occurrences_keep_distinct_exact_column_boundaries():
    outputs = []
    for texts in (["A | B", "C"], ["A", "B | C"]):
        manifest = manifest_from_text("A | B | C")
        manifest["units"][0]["cells"] = [{"text": text, "column_index": i, "column_span": 1,
                                           "source_path": f"source.xml#/cell[{i}]"} for i, text in enumerate(texts)]
        prompt, _ = compact_section_prompt(manifest, _task(manifest))
        data = _data(prompt)
        row = data["source_units"][0]
        decoded = _decoded_cells(data, row)
        assert [row["text"][c["text_binding"][0]:c["text_binding"][1]] for c in decoded] == list(texts)
        outputs.append([c["text_binding"] for c in decoded])
    assert outputs == [[[0, 5], [8, 9]], [[0, 1], [4, 9]]]


@pytest.mark.parametrize("texts", [["  前\n后\t", "", " 尾 "], ["", "", ""], ["😀 | 字", "\n"]])
def test_cell_intervals_preserve_whitespace_newlines_unicode_and_empty_cells(texts):
    manifest = manifest_from_text("")
    manifest["units"][0]["text"] = " | ".join(texts)
    manifest["units"][0]["cells"] = [{"text": text, "column_index": i, "column_span": 1} for i, text in enumerate(texts)]
    prompt, _ = compact_section_prompt(manifest, _task(manifest))
    data = _data(prompt)
    row = data["source_units"][0]
    decoded = _decoded_cells(data, row)
    assert all(len(values) == len(data["cell_fields"]) for values in row["cells"])
    assert [row["text"][c["text_binding"][0]:c["text_binding"][1]].encode() for c in decoded] == [t.encode() for t in texts]


def test_cell_text_falls_back_losslessly_when_join_cannot_be_verified():
    manifest = manifest_from_text("这是原行文字，无法与单元格拼接校验")
    manifest["units"][0]["cells"] = [{"text": "原单元格 | 含分隔符\n", "column_index": 0, "column_span": 1}]
    prompt, _ = compact_section_prompt(manifest, _task(manifest))
    data = _data(prompt)
    cell = _decoded_cells(data, data["source_units"][0])[0]
    assert cell["text_binding"] == {"text": "原单元格 | 含分隔符\n"}
    assert data["source_units"][0]["text"] == manifest["units"][0]["text"]


def test_absent_cell_fields_remain_distinct_from_explicit_null_and_unindexed_origin():
    manifest = manifest_from_text(" | ")
    manifest["units"][0]["cells"] = [{"text": "", "column_index": 0, "vertical_merge": None},
                                       {"text": "", "column_index": 1, "vertical_merge": "continue",
                                        "merged_from_source": "unemitted.xml#/original-cell",
                                        "merged_context_text": "原继承值"}]
    prompt, _ = compact_section_prompt(manifest, _task(manifest))
    data = _data(prompt)
    first, second = _decoded_cells(data, data["source_units"][0])
    assert first["vertical_merge"] is None
    assert first["column_span"] == {"absent_source_field": True}
    assert first["merged_from_id"] == {"absent_source_field": True}
    assert second["merged_context_text"] == "原继承值"
    assert data["merged_origin_evidence_ids"][second["merged_from_id"]] is None
    assert ModelSourceView(manifest).resolve(second["merged_from_id"]) == "unemitted.xml#/original-cell"


def test_unresolved_context_and_roles_retain_original_navigation_without_long_ranges():
    manifest = manifest_from_text("原文\n单位:mm")
    ids = [unit["id"] for unit in manifest["units"]]
    task = _task(manifest, primary=[ids[0]], context=[ids[1]])
    task["context_references"] = [{"evidence_id": ids[1], "roles": ["table_boundary_context", "source_preceding_context"],
                                    "source_path": "verylongpath/" * 100, "start": 0, "end": 5}]
    task["unresolved_context_references"] = [{"evidence_id": "unindexed-heading", "role": "source_section_heading"}]
    prompt, _ = compact_section_prompt(manifest, task)
    navigation = _data(prompt)["task"]
    assert navigation["context_roles"] == {ids[1]: ["table_boundary_context", "source_preceding_context"]}
    assert navigation["unresolved_context_references"] == task["unresolved_context_references"]
    assert "verylongpath" not in prompt


def test_merged_origin_maps_all_segments_of_the_same_physical_source_cell():
    manifest = manifest_from_text("嵌套表前\n嵌套表后\n")
    for unit in manifest["units"]:
        unit["cells"] = [{"text": unit["text"], "column_index": 0, "column_span": 1,
                          "source_path": "source.xml#/same-physical-cell"}]
    last = copy.deepcopy(manifest["units"][-1])
    last.update(id="original-continued-row", text="", source_path="source.xml#/following-row")
    last["cells"] = [{"text": "", "column_index": 0, "column_span": 1, "vertical_merge": "continue",
                       "source_path": "source.xml#/continued-cell",
                       "merged_from_source": "source.xml#/same-physical-cell",
                       "merged_context_text": "嵌套表前\n嵌套表后\n"}]
    manifest["units"].append(last)
    prompt, _ = compact_section_prompt(manifest, _task(manifest))
    data = _data(prompt)
    cell = _decoded_cells(data, data["source_units"][-1])[0]
    assert data["merged_origin_evidence_ids"][cell["merged_from_id"]] == [u["id"] for u in manifest["units"][:-1]]
    assert cell["merged_context_text"] == "嵌套表前\n嵌套表后\n"


@pytest.mark.parametrize("cap", [0, -1, True, 5.5, None])
def test_invalid_budget_rejected(cap):
    manifest = manifest_from_text("原文")
    with pytest.raises(ValueError, match="positive integer"):
        compact_section_prompt(manifest, _task(manifest), cap)


@pytest.mark.parametrize("fault", ["unknown_primary", "duplicate_primary", "unknown_context", "unknown_heading", "unknown_table", "table_cycle"])
def test_invalid_source_task_references_rejected(fault):
    manifest = manifest_from_text("原文")
    task = _task(manifest)
    if fault == "unknown_primary":
        task["primary_evidence_ids"] = ["missing"]
    elif fault == "duplicate_primary":
        task["primary_evidence_ids"] *= 2
    elif fault == "unknown_context":
        task["context_evidence_ids"] = ["missing"]
    elif fault == "unknown_heading":
        task["headings"] = [{"evidence_id": "missing"}]
    elif fault == "unknown_table":
        task["table_ids"] = ["missing"]
    else:
        manifest["tables"] = [{"id": "T-original", "parent_table_id": "T-original"}]
        task["table_ids"] = ["T-original"]
    with pytest.raises(ValueError):
        compact_section_prompt(manifest, task)
