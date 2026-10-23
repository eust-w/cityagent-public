"""Local real DOCX/CPU planning with explicit fake model fixtures, not acceptance."""

import hashlib
import importlib.util
import json
import os
import signal
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image
from test_contest_visual_evidence import response

from bridge_evidence.contest.meta_optimization import CloudBudget
from bridge_evidence.contest.run_state import file_digest

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def setup(tmp_path, monkeypatch):
    name = "probe_report_visual_evidence"
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / (name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    workspace, dataset = tmp_path / "workspace", tmp_path / "dataset"
    workspace.mkdir()
    dataset.mkdir()
    monkeypatch.setattr(module, "WORKSPACE", workspace)
    monkeypatch.setattr(module, "APPROVED_SOURCE_ROOT", dataset)
    monkeypatch.setattr(module, "EXPECTED_DENOMINATOR", 3)
    budget_path = workspace / "budget.sqlite"
    monkeypatch.setattr(module, "ORIGINAL_BUDGET", budget_path)
    budget = CloudBudget(budget_path)
    old = budget.start("previous-cloud-work", 100)
    budget.finish(old, 17.25)
    base, adapter = workspace / "Qwen3.6-27B", workspace / "adapter"
    base.mkdir()
    adapter.mkdir()
    (base / "config.json").write_text('{}')
    (adapter / "adapter_config.json").write_text(json.dumps({"base_model_name_or_path": str(base)}))
    (adapter / "adapter_model.safetensors").write_bytes(b"fixture-model")
    adapter_sha = hashlib.sha256(b"fixture-model").hexdigest()
    monkeypatch.setattr(module, "FIXED_ADAPTER_SHA256", adapter_sha)
    model = workspace / "model.json"
    model.write_text(json.dumps({
        "base": str(base), "tokenizer": str(base), "adapter": str(adapter),
        "adapter_sha256": adapter_sha, "served_model": module.FIXED_ALIAS,
    }))
    samples = []
    for identifier, count in (("z-many", 3), ("a-few", 2), ("b-few", 2)):
        document = Document()
        document.add_paragraph("Fixture bridge source only")
        for i in range(count):
            image = dataset / f"{identifier}-{i}.png"
            Image.new("RGB", (20 + i, 16), (i * 50, 64, 128)).save(image)
            document.add_paragraph().add_run().add_picture(str(image))
        path = dataset / f"{identifier}.docx"
        document.save(path)
        samples.append({"id": identifier, "report": str(path), "report_sha256": file_digest(path),
                        "gold": str(dataset / "DO-NOT-READ-MISSING-GOLD.json"), "gold_sha256": "f" * 64})
    manifest = workspace / "manifest.json"
    manifest.write_text(json.dumps({"schema_version": 2, "purpose": "development-regression",
                                    "metric_denominator": 3, "samples": samples}))
    proof = workspace / "proof.json"
    proof.write_text('{}')
    args = SimpleNamespace(
        manifest=manifest, manifest_sha256=file_digest(manifest), model_manifest=model,
        output_dir=workspace / "output", budget_ledger=budget_path,
        visual_capability=proof, prepared_identity=None, prepared_identity_sha256=None,
        planning_only=False, report_id=None, max_jobs=1, limit_seconds=30, max_tokens=2048,
        request_timeout=180, workers=2, api_url="http://127.0.0.1:8000/v1",
    )
    calls, proof_checks = [], []

    class FixtureBackend:
        def __init__(self, url, alias, tokenizer, **kwargs):
            assert alias == json.loads(args.model_manifest.read_text())["served_model"] and tokenizer is None

        def probe(self):
            return {"fixture_only": True}

        def __call__(self, request, maximum):
            calls.append(request)
            assert "gold" not in json.dumps(request)
            return response(request, maximum)

    monkeypatch.setattr(module, "LoopbackBackend", FixtureBackend)
    monkeypatch.setattr(module, "validate_vision_capability",
                        lambda path, identity: proof_checks.append(identity) or {"fixture_only": True})
    return module, args, budget, calls, proof_checks


def settled(budget):
    with budget.connect() as db:
        rows = db.execute("SELECT phase,seconds,status FROM jobs ORDER BY start").fetchall()
    assert rows[0] == ("previous-cloud-work", 17.25, "complete")
    assert len(rows) == 2 and rows[-1][1] > 0 and rows[-1][2] != "running"


def test_full_source_plan_selects_by_jobs_then_id_and_never_claims_report_complete(setup):
    module, args, budget, calls, proof_checks = setup
    result = module.run_probe(args)
    assert result.get("error") is None, result.get("error")
    assert result["status"] == "complete" and result["planning_complete"] is True
    assert result["source_hashes_verified"] == 3 and result["planning"]["supported_unique_jobs"] == 7
    assert result["selected_report_id"] == "a-few"
    assert result["visual"]["attempted_jobs"] == len(calls) == result["actual_model_calls"] == 1
    assert result["visual"]["planned_jobs"] == 2 and result["visual"]["unexecuted_jobs"] == 1
    assert proof_checks
    assert result["document_complete"] is False and result["full_report_evaluation"] is False
    assert result["formal_release_evidence"] is False and result["rsi_enabled"] is False
    assert result["labels_read"] is False and result["semantic_quality_verified"] is False
    assert len(list((args.output_dir / "plans").glob("*/plan.json"))) == 3
    raw = json.loads((args.output_dir / "visual-evidence.json").read_text())
    assert len(raw["plan"]["jobs"]) == 2  # Full plan never clipped to fit the sample.
    assert len(raw["raw"]) == 1 and raw["calls"][0]["call_id"]
    assert not list(args.output_dir.rglob("*.png"))  # Images remain in original approved source.
    settled(budget)


def test_planning_only_never_checks_proof_or_contacts_backend(setup):
    module, args, budget, calls, proof_checks = setup
    args.planning_only, args.visual_capability = True, None
    result = module.run_probe(args)
    assert result["planning_complete"] is True and result["status"] == "complete"
    assert result["actual_model_calls"] == 0 and not calls and not proof_checks
    assert not (args.output_dir / "timing.sqlite").exists()
    settled(budget)


def test_explicit_report_id_changes_only_pilot_not_complete_planning_denominator(setup):
    module, args, budget, calls, proof_checks = setup
    args.report_id = "z-many"
    result = module.run_probe(args)
    assert result["selected_report_id"] == "z-many"
    assert result["planning"]["planned_reports"] == 3
    assert result["visual"]["planned_jobs"] == 3 and result["visual"]["unexecuted_jobs"] == 2
    settled(budget)


@pytest.mark.parametrize("change", ["source", "manifest", "denominator"])
def test_frozen_input_failure_starts_no_planning_or_model(setup, change):
    module, args, budget, calls, proof_checks = setup
    manifest = json.loads(args.manifest.read_text())
    if change == "source":
        Path(manifest["samples"][0]["report"]).write_bytes(b"changed")
    elif change == "manifest":
        args.manifest.write_text('{}')
    else:
        manifest["metric_denominator"] = 2
        args.manifest.write_text(json.dumps(manifest))
        args.manifest_sha256 = file_digest(args.manifest)
    result = module.run_probe(args)
    assert result["status"] == "failed" and result["error_type"] == "ValueError"
    assert not calls and not (args.output_dir / "plans").exists()
    settled(budget)


def test_active_budget_is_never_recovered_or_overlapped(setup):
    module, args, budget, calls, proof_checks = setup
    active = budget.start("existing-experiment", 100)
    with pytest.raises(RuntimeError, match="already running"):
        module.run_probe(args)
    assert not args.output_dir.exists() and not calls
    with budget.connect() as db:
        assert db.execute("SELECT seconds,status FROM jobs WHERE id=?", (active,)).fetchone() == (None, "running")


@pytest.mark.parametrize("key,value", [
    ("api_url", "http://example.invalid/v1"), ("api_url", "http://localhost/v1"),
    ("max_jobs", 0), ("max_jobs", 65), ("workers", 4), ("limit_seconds", 0),
    ("request_timeout", float("nan")), ("visual_capability", None),
])
def test_invalid_request_rejected_before_budget_lease(setup, key, value):
    module, args, budget, calls, proof_checks = setup
    setattr(args, key, value)
    with pytest.raises(ValueError):
        module.run_probe(args)
    with budget.connect() as db:
        assert db.execute("SELECT count(*) FROM jobs").fetchone()[0] == 1
    assert not calls


def test_different_ledger_cannot_reset_original_budget(setup):
    module, args, budget, calls, proof_checks = setup
    args.budget_ledger = args.budget_ledger.parent / "new-budget.sqlite"
    with pytest.raises(ValueError, match="original cumulative"):
        module.run_probe(args)
    assert not args.budget_ledger.exists()


def test_failed_cpu_plan_prevents_model_pilot_and_keeps_all_report_statuses(setup, monkeypatch):
    module, args, budget, calls, proof_checks = setup
    original = module.plan_all

    def fail_one(*values, **kwargs):
        rows = original(*values, **kwargs)
        rows[0]["status"] = "failed"
        kwargs["progress"](rows)
        return rows

    monkeypatch.setattr(module, "plan_all", fail_one)
    result = module.run_probe(args)
    assert result["planning_complete"] is False and result["status"] == "failed"
    assert len(result["reports"]) == 3 and not calls
    settled(budget)


@pytest.mark.skipif(os.name != "posix", reason="real child termination and wall timer")
def test_cpu_planner_deadline_terminates_only_owned_children(tmp_path, monkeypatch):
    spec = importlib.util.spec_from_file_location("report_probe_termination", ROOT / "scripts/probe_report_visual_evidence.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    original, children = subprocess.Popen, []

    def slow_child(command, **kwargs):
        process = original([sys.executable, "-c", "import time; time.sleep(30)"], **kwargs)
        children.append(process)
        return process

    monkeypatch.setattr(module.subprocess, "Popen", slow_child)
    with pytest.raises(TimeoutError):
        with module.deadline(0.2):
            module.plan_all([{"id": "a"}, {"id": "b"}], "a" * 64, {}, tmp_path / "plans")
    assert len(children) == 2
    assert all(child.poll() == -signal.SIGTERM for child in children)
    for child in children:
        with pytest.raises(ProcessLookupError):
            os.kill(child.pid, 0)


def test_absolute_global_cap_stops_later_images_even_if_first_receipt_is_success(setup, monkeypatch):
    module, args, budget, calls, proof_checks = setup
    args.max_jobs = 8
    original_time, original_backend = module.time, module.LoopbackBackend

    class Clock:
        now = 100.0
        sleep = staticmethod(original_time.sleep)

        def monotonic(self):
            return self.now

    clock = Clock()
    monkeypatch.setattr(module, "time", clock)
    observed = []

    class ExpiresAfterFirstResponse(original_backend):
        def probe(self):
            observed.append(module.REPORT_DEADLINE.get())
            return super().probe()

        def __call__(self, request, maximum):
            response = super().__call__(request, maximum)
            clock.now += args.limit_seconds + 1
            return response

    monkeypatch.setattr(module, "LoopbackBackend", ExpiresAfterFirstResponse)
    result = module.run_probe(args)
    assert observed == [100 + args.limit_seconds]  # Metadata GET inherits the global cap too.
    assert len(calls) == result["actual_model_calls"] == 1
    assert result["visual"]["parsed_jobs"] == 1 and result["visual"]["unexecuted_jobs"] == 1
    assert result["status"] == "failed" and result["error_type"] == "TimeoutError"
    assert module.REPORT_DEADLINE.get() is None
    settled(budget)


def test_models_probe_exhausting_global_cap_prevents_all_image_requests(setup, monkeypatch):
    module, args, budget, calls, proof_checks = setup
    original_time, original_backend = module.time, module.LoopbackBackend

    class Clock:
        now = 100.0
        sleep = staticmethod(original_time.sleep)

        def monotonic(self):
            return self.now

    clock = Clock()
    monkeypatch.setattr(module, "time", clock)

    class SlowModelsProbe(original_backend):
        def probe(self):
            assert module.REPORT_DEADLINE.get() == 100 + args.limit_seconds
            clock.now += args.limit_seconds + 1
            return super().probe()

    monkeypatch.setattr(module, "LoopbackBackend", SlowModelsProbe)
    result = module.run_probe(args)
    assert result["status"] == "failed" and not calls
    assert result["actual_model_calls"] == 0 and result["error_type"] == "TimeoutError"
    settled(budget)


@pytest.mark.skipif(os.name != "posix", reason="actual SIGALRM during an image request")
def test_actual_one_shot_alarm_caught_by_visual_core_cannot_start_next_request(setup, monkeypatch):
    module, args, budget, calls, proof_checks = setup
    args.max_jobs, args.limit_seconds = 8, 3
    original_backend = module.LoopbackBackend

    class BlockFirstRequest(original_backend):
        def __call__(self, request, maximum):
            calls.append(request)
            if len(calls) == 1:
                module.time.sleep(5)
            return response(request, maximum)

    monkeypatch.setattr(module, "LoopbackBackend", BlockFirstRequest)
    result = module.run_probe(args)
    assert len(calls) == result["actual_model_calls"] == 1
    assert result["status"] == "failed" and result["error_type"] == "TimeoutError"
    assert result["timing"]["failures"] == 1 and result["wall_seconds"] < 4.5
    assert result["visual"]["unexecuted_jobs"] == 1
    assert module.REPORT_DEADLINE.get() is None
    settled(budget)


def _generic_development_inputs(args, *, count=2):
    """Copy frozen source fixtures, leaving missing gold paths deliberately unreadable."""
    original = json.loads(args.manifest.read_text())
    source_root = args.manifest.parent / "development-sources"
    source_root.mkdir()
    samples = []
    for i in range(count):
        row = dict(original["samples"][i % len(original["samples"])] )
        source = Path(row["report"])
        copied = source_root / f"report-{i}.docx"
        copied.write_bytes(source.read_bytes())
        samples.append({**row, "id": f"report-{i}", "report": str(copied)})
    args.manifest.write_text(json.dumps({**original, "samples": samples, "metric_denominator": count}))
    args.manifest_sha256 = file_digest(args.manifest)
    args.model_policy = "manifest-locked"
    args.source_root = source_root
    args.expected_denominator = count
    model = json.loads(args.model_manifest.read_text())
    model.pop("adapter")
    model.pop("adapter_sha256")
    model.update(weight_kind="base", served_model="cityagent-qwen35-9b-base")
    base = Path(model["base"])
    (base / "model.safetensors").write_bytes(b"synthetic-fixture-base")
    (base / "tokenizer_config.json").write_text('{}')
    (base / "tokenizer.json").write_text('{}')
    args.model_manifest.write_text(json.dumps(model))
    return samples


def test_manifest_locked_base_plans_full_declared_set_with_same_model_pilot(setup):
    module, args, budget, calls, proof_checks = setup
    _generic_development_inputs(args)
    args.visual_display_transforms = True
    result = module.run_probe(args)
    assert result.get("error") is None, result.get("error")
    assert result["status"] == "complete" and result["planning_complete"] is True
    assert result["expected_denominator"] == result["source_hashes_verified"] == 2
    assert result["planning"]["planned_reports"] == 2
    assert result["model_policy"] == "manifest-locked"
    assert result["source_root"] == str(args.source_root.resolve())
    assert result["model_identity"]["manifest"]["weight_kind"] == "base"
    assert result["visual_config"]["allow_display_transforms"] is True
    assert result["actual_model_calls"] == len(calls) == args.max_jobs
    assert result["document_complete"] is False and result["full_report_evaluation"] is False
    assert result["labels_read"] is False
    for path in (args.output_dir / "plans").glob("*/plan.json"):
        plan = json.loads(path.read_text())
        assert plan["identity"]["config"]["allow_display_transforms"] is True
    assert proof_checks
    settled(budget)


def test_manifest_locked_all_81_source_hashes_checked_without_reading_gold(setup):
    module, args, budget, calls, proof_checks = setup
    _generic_development_inputs(args, count=81)
    module.validate_inputs(args)
    cases = module.source_cases(args)
    assert len(cases) == 81
    assert all(set(case) == {"id", "report", "report_sha256"} for case in cases)
    assert not calls and not proof_checks


@pytest.mark.parametrize("field,value", [
    ("expected_denominator", None), ("expected_denominator", 0),
    ("expected_denominator", True), ("source_root", None),
    ("model_policy", "automatic"), ("visual_display_transforms", "true"),
])
def test_generic_policy_requires_explicit_valid_scope_before_lease(setup, field, value):
    module, args, budget, calls, proof_checks = setup
    _generic_development_inputs(args)
    setattr(args, field, value)
    with pytest.raises(ValueError):
        module.run_probe(args)
    with budget.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 1
    assert not calls


@pytest.mark.parametrize("scope", ["workspace-root", "outside", "symlink-outside"])
def test_generic_source_root_cannot_expand_to_unapproved_files(setup, tmp_path, scope):
    module, args, budget, calls, proof_checks = setup
    _generic_development_inputs(args)
    outside = tmp_path / "unapproved"
    outside.mkdir()
    if scope == "workspace-root":
        args.source_root = module.WORKSPACE
    elif scope == "outside":
        args.source_root = outside
    else:
        alias = module.WORKSPACE / "escaped-source-root"
        alias.symlink_to(outside, target_is_directory=True)
        args.source_root = alias
    with pytest.raises(ValueError, match="source root"):
        module.run_probe(args)
    assert not calls


@pytest.mark.parametrize("change", ["denominator", "source-outside", "source-symlink", "source-bytes", "purpose"])
def test_generic_frozen_manifest_rejects_changed_or_escaped_sources(setup, tmp_path, change):
    module, args, budget, calls, proof_checks = setup
    rows = _generic_development_inputs(args)
    value = json.loads(args.manifest.read_text())
    if change == "denominator":
        args.expected_denominator += 1
    elif change == "purpose":
        value["purpose"] = "official-test"
    elif change == "source-bytes":
        Path(rows[0]["report"]).write_bytes(b"changed")
    else:
        outside = tmp_path / "unapproved.docx"
        outside.write_bytes(Path(rows[0]["report"]).read_bytes())
        if change == "source-outside":
            value["samples"][0]["report"] = str(outside)
        else:
            linked = args.source_root / "escape.docx"
            linked.symlink_to(outside)
            value["samples"][0]["report"] = str(linked)
    args.manifest.write_text(json.dumps(value))
    args.manifest_sha256 = file_digest(args.manifest)
    result = module.run_probe(args)
    assert result["status"] == "failed" and result["error_type"] == "ValueError"
    assert not calls and not (args.output_dir / "plans").exists()
    settled(budget)


@pytest.mark.parametrize("change", ["denominator", "source-root", "model"])
def test_legacy_default_does_not_silently_accept_new_baseline_or_source_scope(setup, change):
    module, args, budget, calls, proof_checks = setup
    if change == "denominator":
        args.expected_denominator = 2
    elif change == "source-root":
        args.source_root = module.WORKSPACE / "development-sources"
        args.source_root.mkdir()
    else:
        model = json.loads(args.model_manifest.read_text())
        model["served_model"] = "cityagent-qwen35-9b-base"
        args.model_manifest.write_text(json.dumps(model))
    with pytest.raises(ValueError):
        module.run_probe(args)
    assert not calls
