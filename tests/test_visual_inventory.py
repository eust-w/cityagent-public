from __future__ import annotations

import importlib.util
import json
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from lxml import etree
from PIL import Image

from bridge_evidence.contest.run_state import file_digest
from bridge_evidence.contest.visual_inventory import (
    MC,
    A,
    InventoryLimits,
    O,
    R,
    V,
    W,
    inventory_docx,
)


def _document(tmp_path):
    image = tmp_path / "photo.png"
    Image.new("RGB", (12, 12), "green").save(image)
    document = Document()
    document.add_picture(str(image))
    document.add_picture(str(image))
    document.sections[0].header.paragraphs[0].add_run().add_picture(str(image))
    document.sections[0].footer.paragraphs[0].add_run().add_picture(str(image))
    outer = document.add_table(rows=1, cols=1)
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).paragraphs[0].add_run().add_picture(str(image))
    path = tmp_path / "report.docx"
    document.save(path)
    return path


def _mutate(path, transform=None, extra=None):
    with zipfile.ZipFile(path) as archive:
        values = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    if transform:
        transform(values)
    values.update(extra or {})
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in values.items():
            archive.writestr(name, value)


def test_preserves_repeated_occurrences_nested_tables_headers_and_footers(tmp_path):
    path = _document(tmp_path)
    value = inventory_docx(path, expected_sha256=file_digest(path))
    assert value == inventory_docx(path)
    assert value["counts"]["images"] == 5
    assert value["counts"]["unique_blobs"] == 1
    occurrences = value["occurrences"]
    assert len({o["id"] for o in occurrences}) == 5
    assert {o["story"] for o in occurrences} == {"body", "header", "footer"}
    assert any("w:tbl/w:tr/w:tc/w:tbl" in o["source_path"] for o in occurrences)
    assert all(o["verification"] == "unverified" for o in occurrences)
    assert not value["document_complete"] and not value["interpretation_verified"]
    assert value["task_cost_inputs"]["model_calls"] == 0
    assert all(
        r["blob"]["size_bytes"] > 0 for o in occurrences for r in o["relationships"]
    )


def test_same_image_distinct_crop_has_distinct_display_hash(tmp_path):
    path = _document(tmp_path)

    def crop(values):
        root = etree.fromstring(values["word/document.xml"])
        fills = root.findall(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/picture}blipFill"
        )
        for index, fill in enumerate(fills):
            node = etree.Element(f"{{{A}}}srcRect", l=str(index * 1000))
            fill.append(node)
        values["word/document.xml"] = etree.tostring(root)

    _mutate(path, crop)
    rows = [r for r in inventory_docx(path)["occurrences"] if r["story"] == "body"]
    assert len({r["relationships"][0]["blob"]["sha256"] for r in rows}) == 1
    assert len({r["display_xml_sha256"] for r in rows}) == 3
    assert [
        next(
            m["attributes"]["l"]
            for m in r["transform_metadata"]
            if m["tag"] == f"{{{A}}}srcRect"
        )
        for r in rows
    ] == ["0", "1000", "2000"]


def test_external_and_missing_relationships_are_unresolved_without_fetch(tmp_path):
    path = _document(tmp_path)

    def mutate(values):
        root = etree.fromstring(values["word/document.xml"])
        images = root.findall(f".//{{{A}}}blip")
        images[0].set(f"{{{R}}}embed", "missing")
        rid = images[1].get(f"{{{R}}}embed")
        rels = etree.fromstring(values["word/_rels/document.xml.rels"])
        relation = next(n for n in rels if n.get("Id") == rid)
        relation.set("Target", "https://example.invalid/private.png")
        relation.set("TargetMode", "External")
        values["word/document.xml"] = etree.tostring(root)
        values["word/_rels/document.xml.rels"] = etree.tostring(rels)

    _mutate(path, mutate)
    rows = [r for r in inventory_docx(path)["occurrences"] if r["story"] == "body"]
    assert rows[0]["relationships"][0]["status"] == "missing_relationship"
    assert rows[1]["relationships"][0]["status"] == "external_not_fetched"
    assert rows[1]["relationships"][0]["blob"] is None


def test_special_objects_ole_preview_and_alternate_content_keep_linkage(tmp_path):
    path = _document(tmp_path)

    def mutate(values):
        root = etree.fromstring(values["word/document.xml"])
        body = root.find(f"{{{W}}}body")
        paragraph = etree.SubElement(body, f"{{{W}}}p")
        run = etree.SubElement(paragraph, f"{{{W}}}r")
        obj = etree.SubElement(run, f"{{{W}}}object")
        shape = etree.SubElement(obj, f"{{{V}}}shape")
        image = etree.SubElement(shape, f"{{{V}}}imagedata")
        rid = root.find(f".//{{{A}}}blip").get(f"{{{R}}}embed")
        image.set(f"{{{R}}}id", rid)
        ole = etree.SubElement(obj, f"{{{O}}}OLEObject")
        ole.set(f"{{{R}}}id", "ole-rel")
        etree.SubElement(
            run, f"{{{W}}}sym", {f"{{{W}}}char": "F0B7", f"{{{W}}}font": "Symbol"}
        )
        etree.SubElement(
            run, "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
        )
        alternate = etree.SubElement(run, f"{{{MC}}}AlternateContent")
        for branch in ("Choice", "Fallback"):
            node = etree.SubElement(alternate, f"{{{MC}}}{branch}")
            etree.SubElement(node, f"{{{A}}}blip", {f"{{{R}}}embed": rid})
        rels = etree.fromstring(values["word/_rels/document.xml.rels"])
        etree.SubElement(
            rels,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="ole-rel",
            Target="embeddings/object.bin",
            Type="oleObject",
        )
        values["word/document.xml"] = etree.tostring(root)
        values["word/_rels/document.xml.rels"] = etree.tostring(rels)
        values["word/embeddings/object.bin"] = b"opaque OLE not executed"

    _mutate(path, mutate)
    value = inventory_docx(path)
    kinds = value["counts"]["by_kind"]
    assert kinds["object"] == kinds["OLEObject"] == kinds["sym"] == kinds["oMath"] == 1
    rows = value["occurrences"]
    ole = next(r for r in rows if r["kind"] == "OLEObject")
    parent_id = ole["parent_objects"][0]["id"]
    assert any(
        r["kind"] == "image" and parent_id in [p["id"] for p in r["parent_objects"]]
        for r in rows
    )
    assert sum(bool(r["alternate_content"]) for r in rows) == 2
    assert value["counts"]["unresolved_types"] == kinds


@pytest.mark.parametrize(
    "name", ["../escape", "/absolute", "word/../../escape", "word\\escape", "C:/drive"]
)
def test_rejects_unsafe_zip_names(tmp_path, name):
    path = _document(tmp_path)
    _mutate(path, extra={name: b"ignored"})
    with pytest.raises(ValueError, match="unsafe"):
        inventory_docx(path)


def test_hash_mismatch_and_legacy_doc_are_failures(tmp_path):
    path = _document(tmp_path)
    with pytest.raises(ValueError, match="SHA256"):
        inventory_docx(path, expected_sha256="0" * 64)
    with pytest.raises(ValueError, match="DOCX"):
        inventory_docx(tmp_path / "source.doc")


def test_metadata_bomb_and_dtd_are_rejected(tmp_path):
    path = _document(tmp_path)
    with pytest.raises(ValueError, match="metadata limit"):
        inventory_docx(path, limits=replace(InventoryLimits(), max_metadata_bytes=20))
    _mutate(
        path,
        extra={
            "word/document.xml": b'<!DOCTYPE x [<!ENTITY test SYSTEM "file:///etc/passwd">]><x>&test;</x>'
        },
    )
    with pytest.raises(ValueError, match="DTD"):
        inventory_docx(path)


def test_manifest_never_opens_gold_and_keeps_failed_source_in_denominator(tmp_path):
    script = (
        Path(__file__).resolve().parents[1] / "scripts/inventory_visual_evidence.py"
    )
    spec = importlib.util.spec_from_file_location("inventory_visual_script", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    path = _document(tmp_path)
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "purpose": "development-regression",
                "metric_denominator": 2,
                "samples": [
                    {
                        "id": "good",
                        "report": str(path),
                        "report_sha256": file_digest(path),
                        "gold": "/does-not-exist-never-read",
                    },
                    {
                        "id": "missing",
                        "report": str(tmp_path / "missing.docx"),
                        "report_sha256": "a" * 64,
                        "gold": "/does-not-exist-never-read",
                    },
                ],
            }
        )
    )
    result = module.inventory_manifest(
        manifest, tmp_path / "inventory", file_digest(manifest)
    )
    assert result["metric_denominator"] == result["processed"] == 2
    assert result["failed"] == 1 and not result["inventory_complete"]
    assert result["pending_report_ids"] == []
    assert [r["id"] for r in result["reports"]] == ["good", "missing"]
    assert result["reports"][1]["status"] == "failed"
    assert result["model_calls"] == 0 and not result["document_complete"]


def test_missing_target_and_unsupported_image_type_remain_visible(tmp_path):
    path = _document(tmp_path)

    def mutate(values):
        root = etree.fromstring(values["word/document.xml"])
        image = root.find(f".//{{{A}}}blip")
        old_rid = image.get(f"{{{R}}}embed")
        image.set(f"{{{R}}}embed", "unknown-format")
        rels = etree.fromstring(values["word/_rels/document.xml.rels"])
        next(n for n in rels if n.get("Id") == old_rid).set(
            "Target", "media/missing.png"
        )
        etree.SubElement(
            rels,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            Id="unknown-format",
            Target="media/image.unsupported",
            Type="image",
        )
        values["word/document.xml"] = etree.tostring(root)
        values["word/_rels/document.xml.rels"] = etree.tostring(rels)
        values["word/media/image.unsupported"] = b"not decoded"

    _mutate(path, mutate)
    result = inventory_docx(path)
    rows = [o for o in result["occurrences"] if o["story"] == "body"]
    assert rows[0]["relationships"][0]["blob"]["extension"] == ".unsupported"
    assert rows[1]["relationships"][0]["status"] == "missing_target"
    assert result["counts"]["unresolved_types"]["image"] == 5


def test_duplicate_members_symlinks_ratio_and_occurrence_limits_fail_closed(tmp_path):
    path = _document(tmp_path)
    with pytest.raises(ValueError, match="compression ratio"):
        inventory_docx(path, limits=replace(InventoryLimits(), max_compression_ratio=1))
    with pytest.raises(ValueError, match="occurrence limit"):
        inventory_docx(path, limits=replace(InventoryLimits(), max_occurrences=1))
    with zipfile.ZipFile(path, "a") as archive:
        info = zipfile.ZipInfo("word/media/link")
        info.create_system = 3
        info.external_attr = 0o120777 << 16
        archive.writestr(info, b"/etc/passwd")
    with pytest.raises(ValueError, match="special file"):
        inventory_docx(path)
    path = _document(tmp_path)
    with pytest.warns(UserWarning, match="Duplicate"):
        with zipfile.ZipFile(path, "a") as archive:
            archive.writestr("word/document.xml", b"duplicate")
    with pytest.raises(ValueError, match="duplicate"):
        inventory_docx(path)


@pytest.mark.parametrize(
    "discovery", ["content_type_and_relationship", "relationship", "content_type"]
)
def test_custom_story_part_names_are_never_silently_omitted(tmp_path, discovery):
    path = _document(tmp_path)

    def rename(values):
        for old, new in (("header1", "customHeader"), ("footer1", "customFooter")):
            values[f"word/{new}.xml"] = values.pop(f"word/{old}.xml")
            values[f"word/_rels/{new}.xml.rels"] = values.pop(
                f"word/_rels/{old}.xml.rels"
            )
            values["[Content_Types].xml"] = values["[Content_Types].xml"].replace(
                f"/{old}.xml".encode(), f"/{new}.xml".encode()
            )
            values["word/_rels/document.xml.rels"] = values[
                "word/_rels/document.xml.rels"
            ].replace(f'"{old}.xml"'.encode(), f'"{new}.xml"'.encode())
        if discovery == "relationship":
            types = etree.fromstring(values["[Content_Types].xml"])
            for entry in types:
                if entry.get("PartName") in {
                    "/word/customHeader.xml",
                    "/word/customFooter.xml",
                }:
                    entry.set("ContentType", "application/xml")
            values["[Content_Types].xml"] = etree.tostring(types)
        elif discovery == "content_type":
            rels = etree.fromstring(values["word/_rels/document.xml.rels"])
            for entry in list(rels):
                if entry.get("Target") in {"customHeader.xml", "customFooter.xml"}:
                    rels.remove(entry)
            values["word/_rels/document.xml.rels"] = etree.tostring(rels)

    _mutate(path, rename)
    if discovery == "content_type_and_relationship":
        doc = Document(path)
        assert len(list(doc.sections[0].header._element.iter(f"{{{A}}}blip"))) == 1
    value = inventory_docx(path)
    assert value["counts"]["images"] == 5
    assert {row["part_name"] for row in value["occurrences"]} == {
        "word/document.xml",
        "word/customHeader.xml",
        "word/customFooter.xml",
    }
    assert {row["story"] for row in value["occurrences"]} == {
        "body",
        "header",
        "footer",
    }


def test_missing_story_target_is_not_a_complete_inventory(tmp_path):
    path = _document(tmp_path)

    def remove(values):
        values.pop("word/header1.xml")

    _mutate(path, remove)
    with pytest.raises(ValueError, match="story part missing"):
        inventory_docx(path)
