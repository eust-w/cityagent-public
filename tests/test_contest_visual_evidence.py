from __future__ import annotations

import base64
import hashlib
import json
import subprocess
import sys
import zipfile
from copy import deepcopy

import pytest
from docx import Document
from lxml import etree
from PIL import Image

from bridge_evidence.contest.long_document import DocumentSnapshot
from bridge_evidence.contest.run_state import json_digest
from bridge_evidence.contest.visual_evidence import (
    build_visual_request,
    execute_visual_jobs,
    plan_visual_jobs,
    rebuild_visual_snapshot,
)
from bridge_evidence.contest.visual_inventory import A, W


def document(tmp_path, *, fmt="PNG", repeat=2):
    photo = tmp_path / ("image.png" if fmt == "PNG" else "image.jpg")
    Image.new("RGB", (20, 16), "green").save(photo, format=fmt)
    doc = Document()
    doc.add_paragraph("2024 桥梁报告")
    for _ in range(repeat):
        doc.add_paragraph().add_run().add_picture(str(photo))
    path = tmp_path / "report.docx"
    doc.save(path)
    return path


def mutate(path, callback):
    with zipfile.ZipFile(path) as archive:
        values = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    callback(values)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for key, value in values.items():
            archive.writestr(key, value)


def plan_for(path, **config):
    native = DocumentSnapshot.from_docx(path)
    return native, plan_visual_jobs(
        path, native, "a" * 64, {"trial_id": "trial-1", **config}
    )


def response(request, maximum, *, outcome="complete", text=True):
    from bridge_evidence.contest.rsi_runtime import REQUEST_CALL_ID

    call_id = REQUEST_CALL_ID.get() or "call-1"
    result = {
        "schema_version": 1,
        "job_id": request["job_id"],
        "outcome": outcome,
        "text_blocks": [{"id": "t1", "bbox": [0, 0, 1, 1], "text": "墩身裂缝"}]
        if text
        else [],
        "observations": []
        if text
        else [
            {
                "id": "o1",
                "bbox": [0, 0, 1, 1],
                "description": "绿色表面照片，未见可读文字",
            }
        ],
        "unresolved_regions": []
        if outcome == "complete"
        else [{"id": "u1", "bbox": [0, 0, 0.5, 0.5], "reason": "模糊"}],
    }
    return {
        "text": json.dumps(result, ensure_ascii=False),
        "finish_reason": "stop",
        "output_tokens": 100,
        "prompt_tokens": 150,
        "backend_request_id": "backend-" + call_id,
        "call_id": call_id,
        "latency_seconds": 0.2,
    }


def rebuild(native, plan, result):
    return rebuild_visual_snapshot(
        native, plan, result["evidence"], result["calls"], result["raw"]
    )


@pytest.mark.parametrize("fmt", ["PNG", "JPEG"])
def test_full_raster_shared_call_preserves_each_occurrence_and_native_anchor(
    tmp_path, fmt
):
    path = document(tmp_path, fmt=fmt)
    native, plan = plan_for(path)
    assert plan == plan_visual_jobs(path, native, "a" * 64, {"trial_id": "trial-1"})
    assert len(plan["jobs"]) == 1
    assert len(plan["jobs"][0]["occurrence_ids"]) == 2
    request = build_visual_request(plan, plan["jobs"][0])
    assert "base64" not in json.dumps(plan)
    assert set(request) == {"kind", "job_id", "image", "instruction"}
    assert (
        hashlib.sha256(base64.b64decode(request["image"]["base64"])).hexdigest()
        == request["image"]["sha256"]
    )
    budgets = []
    result = execute_visual_jobs(plan, response, tmp_path / "cache", budgets.append)
    assert len(budgets) == len(result["calls"]) == 1
    assert result["calls"][0]["prompt_sha256"] == json_digest(request)
    assert result["calls"][0]["source_ranges"] == plan["jobs"][0]["occurrence_ids"]
    effective = rebuild(native, plan, result)
    derived = [r for r in effective.regions if r.kind == "visual-transcription"]
    assert len(derived) == len({r.source_id for r in derived}) == 2
    assert not effective.unverified_content and native.unverified_content
    for region in derived:
        assert region.location["evidence_domain"] == "visual-transcription"
        assert (
            effective.regions[region.order - 1].source_id
            == region.location["native_anchor"]
        )
    assert [
        r.source_id for r in effective.regions if not r.kind.startswith("visual-")
    ] == [r.source_id for r in native.regions]


def test_photo_observation_never_becomes_literal_text_evidence(tmp_path):
    native, plan = plan_for(document(tmp_path))
    result = execute_visual_jobs(
        plan, lambda p, n: response(p, n, text=False), tmp_path / "cache"
    )
    effective = rebuild(native, plan, result)
    assert {r.kind for r in effective.regions if r.kind.startswith("visual-")} == {
        "visual-observation"
    }
    assert "visual-evidence" in effective.parser


@pytest.mark.parametrize(
    "mutation",
    ["empty", "bbox", "duplicate", "extra", "job", "truncated", "usage", "json"],
)
def test_invalid_or_empty_model_output_is_durable_but_cannot_resolve(
    tmp_path, mutation
):
    native, plan = plan_for(document(tmp_path))

    def generate(request, maximum):
        value = response(request, maximum)
        data = json.loads(value["text"])
        if mutation == "empty":
            data["text_blocks"] = []
        if mutation == "bbox":
            data["text_blocks"][0]["bbox"] = [0, 0, 2, 1]
        if mutation == "duplicate":
            data["text_blocks"] *= 2
        if mutation == "extra":
            data["verified"] = True
        if mutation == "job":
            data["job_id"] = "wrong"
        value["text"] = json.dumps(data)
        if mutation == "truncated":
            value["finish_reason"] = "length"
        if mutation == "usage":
            value["prompt_tokens"] = 0
        if mutation == "json":
            value["text"] = "not JSON"
        return value

    cache = tmp_path / "cache"
    result = execute_visual_jobs(plan, generate, cache)
    assert result["errors"] and result["raw"] and result["calls"]
    assert list(cache.glob("raw/*.json"))
    assert rebuild(native, plan, result).unverified_content == native.unverified_content


def test_partial_result_keeps_occurrence_unresolved(tmp_path):
    native, plan = plan_for(document(tmp_path))
    result = execute_visual_jobs(
        plan, lambda p, n: response(p, n, outcome="partial"), tmp_path / "cache"
    )
    effective = rebuild(native, plan, result)
    assert len([r for r in effective.regions if r.kind == "visual-transcription"]) == 2
    assert effective.unverified_content == native.unverified_content


def test_raw_receipt_is_durable_before_parse_and_resume_avoids_model(
    tmp_path, monkeypatch
):
    import bridge_evidence.contest.visual_evidence as module

    native, plan = plan_for(document(tmp_path))
    cache = tmp_path / "cache"
    parse = module._parse_response

    def interrupt(*args, **kwargs):
        assert list(cache.glob("raw/*.json"))
        raise KeyboardInterrupt

    monkeypatch.setattr(module, "_parse_response", interrupt)
    with pytest.raises(KeyboardInterrupt):
        execute_visual_jobs(plan, response, cache)
    monkeypatch.setattr(module, "_parse_response", parse)

    def no_call(*args):
        raise AssertionError("completed model request repeated")

    recovered = execute_visual_jobs(plan, no_call, cache, no_call)
    assert not rebuild(native, plan, recovered).unverified_content


@pytest.mark.parametrize("change", ["trial", "model", "config", "report"])
def test_cache_identity_rejects_different_trial_model_config_or_report(
    tmp_path, change
):
    path = document(tmp_path)
    native, plan = plan_for(path)
    cache = tmp_path / "cache"
    execute_visual_jobs(plan, response, cache)
    cfg = {"trial_id": "trial-2" if change == "trial" else "trial-1"}
    if change == "config":
        cfg["max_tokens"] = 1024
    if change == "report":
        other = tmp_path / "other.docx"
        other.write_bytes(path.read_bytes())
        path = other
    changed = plan_visual_jobs(
        path, native, "b" * 64 if change == "model" else "a" * 64, cfg
    )
    with pytest.raises(ValueError, match="identity"):
        execute_visual_jobs(changed, response, cache)


@pytest.mark.parametrize("field", ["evidence", "calls", "raw", "plan"])
def test_independent_rebuild_rejects_modified_evidence_and_identity(tmp_path, field):
    native, plan = plan_for(document(tmp_path))
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    if field == "evidence":
        result["evidence"][0]["response"]["text_blocks"][0]["text"] = "伪造"
    if field == "calls":
        result["calls"][0]["prompt_sha256"] = "b" * 64
    if field == "raw":
        result["raw"][0]["response"]["text"] = "{}"
    if field == "plan":
        plan["jobs"][0]["view_sha256"] = "b" * 64
    with pytest.raises(ValueError):
        rebuild(native, plan, result)


def test_nonzero_crop_rotation_flip_and_symbols_remain_unresolved(tmp_path):
    path = document(tmp_path, repeat=4)

    def edits(values):
        root = etree.fromstring(values["word/document.xml"])
        pics = root.findall(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"
        )
        etree.SubElement(
            pics[1].find(
                "{http://schemas.openxmlformats.org/drawingml/2006/picture}blipFill"
            ),
            f"{{{A}}}srcRect",
            l="100",
        )
        pics[2].find(f".//{{{A}}}xfrm").set("rot", "60000")
        pics[3].find(f".//{{{A}}}xfrm").set("flipH", "1")
        run = etree.SubElement(root.find(f"{{{W}}}body/{{{W}}}p"), f"{{{W}}}r")
        etree.SubElement(run, f"{{{W}}}sym", {f"{{{W}}}char": "F0B7"})
        values["word/document.xml"] = etree.tostring(root)

    mutate(path, edits)
    native, plan = plan_for(path)
    assert len(plan["jobs"][0]["occurrence_ids"]) == 1
    assert len(plan["unresolved"]) == 4
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    effective = rebuild(native, plan, result)
    assert len(effective.unverified_content) == 4
    assert any(r.get("kind") == "sym" for r in effective.unverified_content)


def test_decode_limits_and_bad_raster_fail_closed_without_whole_plan_failure(tmp_path):
    native, plan = plan_for(document(tmp_path), max_pixels=100)
    assert not plan["jobs"] and len(plan["unresolved"]) == 2
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    assert rebuild(native, plan, result).unverified_content == native.unverified_content


def test_runtime_failure_is_preserved_and_budget_runs_before_callback(tmp_path):
    native, plan = plan_for(document(tmp_path))
    seen = []

    def generate(*args):
        from bridge_evidence.contest.rsi_runtime import REQUEST_CALL_ID

        assert seen
        error = TimeoutError("backend timed out")
        error.call_id, error.latency_seconds = REQUEST_CALL_ID.get(), 1.0
        raise error

    result = execute_visual_jobs(plan, generate, tmp_path / "cache", seen.append)
    assert result["calls"][0]["status"] == "failed"
    assert result["raw"][0]["error_type"] == "TimeoutError"
    assert rebuild(native, plan, result).unverified_content == native.unverified_content


def test_rebuild_requires_the_actual_original_docx_and_native_snapshot(tmp_path):
    path = document(tmp_path)
    native, plan = plan_for(path)
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    changed = deepcopy(native)
    changed.regions.pop()
    with pytest.raises(ValueError):
        rebuild(changed, plan, result)
    path.write_bytes(b"different document")
    with pytest.raises(ValueError):
        rebuild(native, plan, result)


def test_budget_rejection_preserves_earlier_calls_without_fake_receipt(tmp_path):
    path = document(tmp_path, repeat=1)
    other = tmp_path / "other.png"
    Image.new("RGB", (16, 16), "red").save(other)
    doc = Document(path)
    doc.add_picture(str(other))
    doc.save(path)
    native, plan = plan_for(path)
    seen = []

    def budget(job):
        assert isinstance(job["prior_calls"], list)
        seen.append(job)
        if len(seen) == 2:
            assert len(job["prior_calls"]) == 1
            raise TimeoutError("report budget expired")

    result = execute_visual_jobs(plan, response, tmp_path / "cache", budget)
    assert len(result["calls"]) == len(result["raw"]) == 1
    assert result["errors"][-1]["stage"] == "budget"
    assert result["spent_seconds"] >= 0
    assert len(rebuild(native, plan, result).unverified_content) == 1


def test_pending_request_never_reissued_after_missing_raw_crash(tmp_path, monkeypatch):
    import bridge_evidence.contest.visual_evidence as module

    native, plan = plan_for(document(tmp_path))
    cache = tmp_path / "cache"
    durable = module._durable

    def crash(path, value):
        if path.parent.name == "raw":
            raise KeyboardInterrupt
        return durable(path, value)

    monkeypatch.setattr(module, "_durable", crash)
    with pytest.raises(KeyboardInterrupt):
        execute_visual_jobs(plan, response, cache)
    assert list(cache.glob("pending/*.json")) and not list(cache.glob("raw/*.json"))
    monkeypatch.setattr(module, "_durable", durable)

    def no_model(*args):
        raise AssertionError("unknown outcome reissued")

    result = execute_visual_jobs(plan, no_model, cache)
    assert result["calls"][0]["status"] == "unknown"
    assert result["calls"][0]["call_id"]
    assert result["errors"] and rebuild(native, plan, result).unverified_content


def test_vml_oval_is_not_interpreted_as_uncropped_raster(tmp_path):
    from bridge_evidence.contest.visual_inventory import R, V

    path = document(tmp_path, repeat=1)

    def change(values):
        root = etree.fromstring(values["word/document.xml"])
        image = root.find(f".//{{{A}}}blip")
        rid = image.get(f"{{{R}}}embed")
        paragraph = etree.SubElement(root.find(f"{{{W}}}body"), f"{{{W}}}p")
        oval = etree.SubElement(
            etree.SubElement(paragraph, f"{{{W}}}r"), f"{{{V}}}oval"
        )
        etree.SubElement(oval, f"{{{V}}}imagedata", {f"{{{R}}}id": rid})
        values["word/document.xml"] = etree.tostring(root)

    mutate(path, change)
    native, plan = plan_for(path)
    assert len(plan["unresolved"]) == 1
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    assert len(rebuild(native, plan, result).unverified_content) == 1


def test_pillow_decompression_bomb_keeps_images_unresolved(tmp_path, monkeypatch):
    path = document(tmp_path)
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 1)
    native, plan = plan_for(path)
    assert not plan["jobs"] and len(plan["unresolved"]) == 2


def test_cached_evidence_refuses_changed_original_file(tmp_path):
    path = document(tmp_path)
    native, plan = plan_for(path)
    cache = tmp_path / "cache"
    execute_visual_jobs(plan, response, cache)
    path.write_bytes(b"changed")
    with pytest.raises(ValueError, match="DOCX identity"):
        execute_visual_jobs(plan, response, cache)


def test_sigkill_after_actual_ledger_commit_never_reissues_pending_job(tmp_path):
    from bridge_evidence.contest.timing import TimingLedger

    native, plan = plan_for(document(tmp_path))
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(json.dumps(plan))
    cache = tmp_path / "cache"
    ledger_path = tmp_path / "timing.sqlite"
    script = """
import json, os, signal, sys
from pathlib import Path
from bridge_evidence.contest.backends import AccountedBackend
from bridge_evidence.contest.timing import TimingLedger
import bridge_evidence.contest.visual_evidence as module
plan = json.loads(Path(sys.argv[1]).read_text())
ledger = TimingLedger(sys.argv[3])
def backend(request, maximum):
    value = {"schema_version":1,"job_id":request["job_id"],"outcome":"complete",
             "text_blocks":[{"id":"t1","bbox":[0,0,1,1],"text":"local fixture"}],
             "observations":[],"unresolved_regions":[]}
    return {"text":json.dumps(value),"finish_reason":"stop","output_tokens":50,
            "prompt_tokens":100,"backend_request_id":"local-test-backend"}
durable = module._durable
def kill_before_raw(path, value):
    if path.parent.name == "raw": os.kill(os.getpid(), signal.SIGKILL)
    durable(path, value)
module._durable = kill_before_raw
module.execute_visual_jobs(plan, AccountedBackend(backend, ledger), sys.argv[2])
"""
    child = subprocess.run(
        [sys.executable, "-c", script, str(plan_path), str(cache), str(ledger_path)],
        capture_output=True,
        text=True,
        timeout=20,
    )
    assert child.returncode == -9, child.stderr
    ledger = TimingLedger(ledger_path)
    with ledger.connect() as db:
        original_id, original_status = db.execute(
            "SELECT id,status FROM calls"
        ).fetchone()
    assert original_status == "success"

    def no_model(*args):
        raise AssertionError("committed request repeated")

    result = execute_visual_jobs(plan, no_model, cache)
    assert result["recovery_blocked"]
    assert result["calls"][0]["call_id"] == original_id
    with ledger.connect() as db:
        assert db.execute("SELECT count(*) FROM calls").fetchone()[0] == 1
    assert rebuild(native, plan, result).unverified_content


@pytest.mark.parametrize("what", ["vector", "corrupt"])
def test_nonraster_or_corrupt_media_never_becomes_complete(tmp_path, what):
    path = document(tmp_path)

    def change(values):
        target = next(p for p in values if p.startswith("word/media/"))
        values[target] = b"not a decodable PNG"
        if what == "vector":
            root = etree.fromstring(values["[Content_Types].xml"])
            next(node for node in root if node.get("Extension") == "png").set(
                "ContentType", "image/x-emf"
            )
            values["[Content_Types].xml"] = etree.tostring(root)

    mutate(path, change)
    native, plan = plan_for(path)
    assert not plan["jobs"] and len(plan["unresolved"]) == 2
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    assert rebuild(native, plan, result).unverified_content == native.unverified_content


def test_ole_with_supported_drawing_preview_keeps_ole_unresolved(tmp_path):
    from bridge_evidence.contest.visual_inventory import O

    path = document(tmp_path, repeat=1)

    def change(values):
        root = etree.fromstring(values["word/document.xml"])
        run = root.find(f".//{{{W}}}drawing").getparent()
        obj = etree.SubElement(run, f"{{{W}}}object")
        etree.SubElement(obj, f"{{{O}}}OLEObject")
        etree.SubElement(
            run, "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
        )
        values["word/document.xml"] = etree.tostring(root)

    mutate(path, change)
    native, plan = plan_for(path)
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    unresolved = rebuild(native, plan, result).unverified_content
    assert {item["kind"] for item in unresolved} == {"object", "OLEObject", "oMath"}


@pytest.mark.parametrize("tokens", [8192, 8193])
def test_actual_visual_prompt_usage_is_bounded_without_discarding_receipt(
    tmp_path, tokens
):
    native, plan = plan_for(document(tmp_path))

    def generate(request, maximum):
        value = response(request, maximum)
        value["prompt_tokens"] = tokens
        return value

    cache = tmp_path / "cache"
    result = execute_visual_jobs(plan, generate, cache)
    assert result["calls"][0]["prompt_tokens"] == tokens
    assert list(cache.glob("raw/*.json"))
    effective = rebuild(native, plan, result)
    assert bool(result["evidence"]) is (tokens <= 8192)
    assert bool(effective.unverified_content) is (tokens > 8192)


_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
_LOCAL_DPI_URI = "{28A0092B-C50C-407E-A947-70E740481C1C}"


def local_dpi_extension(path, val="0", uri=_LOCAL_DPI_URI, tamper=None):
    def change(values):
        root = etree.fromstring(values["word/document.xml"])
        blip = root.find(f".//{{{A}}}blip")
        extension_list = etree.SubElement(blip, f"{{{A}}}extLst")
        extension = etree.SubElement(extension_list, f"{{{A}}}ext", uri=uri)
        leaf = etree.SubElement(extension, f"{{{_A14}}}useLocalDpi")
        if val is not None:
            leaf.set("val", val)
        if tamper:
            tamper(root, blip, extension_list, extension, leaf)
        values["word/document.xml"] = etree.tostring(root)

    mutate(path, change)


@pytest.mark.parametrize("val", [None, "0", "1", "true", "false"])
@pytest.mark.parametrize("uri", [_LOCAL_DPI_URI, _LOCAL_DPI_URI.lower()])
def test_exact_local_dpi_metadata_preserves_original_pixels_and_occurrences(
    tmp_path, val, uri
):
    path = document(tmp_path)
    _, original = plan_for(path)
    original_view = original["jobs"][0]["view_sha256"]
    local_dpi_extension(path, val, uri)
    native, plan = plan_for(path)
    assert plan["identity"]["view_validation_version"] == "full-raster-view-v2"
    assert not plan["unresolved"] and len(plan["jobs"]) == 1
    job = plan["jobs"][0]
    assert job["view_sha256"] == original_view
    assert len(job["occurrence_ids"]) == 2
    # One marked and one unmarked occurrence share pixels, but retain distinct XML.
    assert len({item["display_xml_sha256"] for item in job["bindings"]}) == 2
    assert {item["display_xml_sha256"] for item in job["bindings"]} == {
        item["display_xml_sha256"] for item in plan["inventory"]["occurrences"]
    }
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    effective = rebuild(native, plan, result)
    assert not effective.unverified_content
    assert (
        len([item for item in effective.regions if item.kind == "visual-transcription"])
        == 2
    )


@pytest.mark.parametrize(
    "case",
    [
        "wrong-leaf-namespace",
        "wrong-ext-namespace",
        "wrong-list-namespace",
        "unknown-guid",
        "leaf-attribute",
        "ext-attribute",
        "list-attribute",
        "invalid-bool",
        "leaf-child",
        "leaf-sibling",
        "ext-sibling",
        "list-sibling",
        "pixel-effect",
        "leaf-text",
        "ext-text",
        "list-text",
        "blip-text",
        "leaf-tail",
        "comment",
        "crop",
        "rotation",
        "geometry",
    ],
)
def test_local_dpi_whitelist_does_not_admit_unknown_structure_or_geometry(
    tmp_path, case
):
    path = document(tmp_path, repeat=1)

    def tamper(root, blip, extension_list, extension, leaf):
        if case == "wrong-leaf-namespace":
            leaf.tag = "{urn:untrusted}useLocalDpi"
        if case == "wrong-ext-namespace":
            extension.tag = "{urn:untrusted}ext"
        if case == "wrong-list-namespace":
            extension_list.tag = "{urn:untrusted}extLst"
        if case == "unknown-guid":
            extension.set("uri", "{00000000-0000-0000-0000-000000000000}")
        if case == "leaf-attribute":
            leaf.set("unrecognized", "true")
        if case == "ext-attribute":
            extension.set("unrecognized", "true")
        if case == "list-attribute":
            extension_list.set("unrecognized", "true")
        if case == "invalid-bool":
            leaf.set("val", "yes")
        if case == "leaf-child":
            etree.SubElement(leaf, f"{{{A}}}alphaInv")
        if case == "leaf-sibling":
            etree.SubElement(extension, f"{{{_A14}}}useLocalDpi")
        if case == "ext-sibling":
            etree.SubElement(extension_list, f"{{{A}}}ext", uri=_LOCAL_DPI_URI)
        if case == "list-sibling":
            etree.SubElement(blip, f"{{{A}}}extLst")
        if case == "pixel-effect":
            etree.SubElement(blip, f"{{{A}}}alphaInv")
        if case == "leaf-text":
            leaf.text = "unexpected"
        if case == "ext-text":
            extension.text = "unexpected"
        if case == "list-text":
            extension_list.text = "unexpected"
        if case == "blip-text":
            blip.text = "unexpected"
        if case == "leaf-tail":
            leaf.tail = "unexpected"
        if case == "comment":
            leaf.append(etree.Comment("unknown markup"))
        if case == "crop":
            etree.SubElement(blip.getparent(), f"{{{A}}}srcRect", l="1")
        if case == "rotation":
            root.find(f".//{{{A}}}xfrm").set("rot", "60000")
        if case == "geometry":
            root.find(f".//{{{A}}}prstGeom").set("prst", "ellipse")

    local_dpi_extension(path, tamper=tamper)
    native, plan = plan_for(path)
    assert not plan["jobs"] and len(plan["unresolved"]) == 1
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    assert rebuild(native, plan, result).unverified_content == native.unverified_content


def test_view_validation_version_invalidates_previous_cache(tmp_path, monkeypatch):
    import bridge_evidence.contest.visual_evidence as module

    path = document(tmp_path)
    native, current = plan_for(path)
    monkeypatch.setattr(module, "VIEW_VALIDATION_VERSION", "legacy-full-raster-v1")
    old = module.plan_visual_jobs(path, native, "a" * 64, {"trial_id": "trial-1"})
    cache = tmp_path / "cache"
    execute_visual_jobs(old, response, cache)
    monkeypatch.setattr(module, "VIEW_VALIDATION_VERSION", "full-raster-view-v2")
    with pytest.raises(ValueError, match="identity"):
        execute_visual_jobs(current, response, cache)


def test_plan_without_view_validation_version_is_rejected_before_dispatch(tmp_path):
    _, plan = plan_for(document(tmp_path))
    plan["identity"].pop("view_validation_version")
    plan["plan_sha256"] = json_digest(
        {key: value for key, value in plan.items() if key != "plan_sha256"}
    )
    with pytest.raises(ValueError, match="view validation version"):
        execute_visual_jobs(plan, response, tmp_path / "cache")
    assert not (tmp_path / "cache").exists()


def test_local_dpi_whitespace_formatting_is_not_pixel_content(tmp_path):
    path = document(tmp_path, repeat=1)

    def whitespace(root, *nodes):
        for node in nodes:
            node.text, node.tail = "\n  ", "\n"

    local_dpi_extension(path, tamper=whitespace)
    _, plan = plan_for(path)
    assert len(plan["jobs"]) == 1 and not plan["unresolved"]


def response_v2(request, maximum, *, outcome="complete", bbox=None):
    value = response(request, maximum, outcome=outcome)
    data = json.loads(value["text"])
    data["schema_version"] = 2
    data["bbox_space"] = "normalized_1000"
    data["text_blocks"][0]["bbox"] = [526, 564, 608, 666] if bbox is None else bbox
    for row in data["unresolved_regions"]:
        row["bbox"] = [0, 0, 500, 500]
    value["text"] = json.dumps(data, ensure_ascii=False)
    return value


@pytest.mark.parametrize("outcome", ["complete", "partial", "unreadable"])
def test_v2_explicit_thousandths_preserves_original_and_canonical_units(
    tmp_path, outcome
):
    native, plan = plan_for(document(tmp_path))
    assert plan["identity"]["response_protocol_version"] == "visual-response-v2"
    instruction = build_visual_request(plan, plan["jobs"][0])["instruction"]
    assert (
        "schema_version:2" in instruction
        and 'bbox_space:"normalized_1000"' in instruction
    )
    assert "[526,564,608,666]" in instruction
    raw_texts = []

    def generate(request, maximum):
        value = response_v2(request, maximum, outcome=outcome)
        raw_texts.append(value["text"])
        return value

    result = execute_visual_jobs(plan, generate, tmp_path / "cache")
    assert not result["errors"]
    assert result["raw"][0]["response"]["text"] == raw_texts[0]
    item = result["evidence"][0]
    assert (
        item["response"]["bbox_space"] == item["source_bbox_space"] == "normalized_1000"
    )
    assert item["response"]["text_blocks"][0]["bbox"] == [526, 564, 608, 666]
    assert item["canonical_bbox_space"] == "normalized_1"
    assert item["canonical_blocks"]["text_blocks"][0]["bbox"] == [
        0.526,
        0.564,
        0.608,
        0.666,
    ]
    effective = rebuild(native, plan, result)
    assert bool(effective.unverified_content) is (outcome != "complete")
    region = next(r for r in effective.regions if r.kind == "visual-transcription")
    assert region.location["bbox"] == [0.526, 0.564, 0.608, 0.666]
    assert region.location["bbox_space"] == "normalized_1"
    assert region.location["source_bbox_space"] == "normalized_1000"
    assert region.location["source_bbox"] == [526, 564, 608, 666]
    if outcome != "complete":
        assert item["canonical_blocks"]["unresolved_regions"][0]["bbox"] == [
            0.0,
            0.0,
            0.5,
            0.5,
        ]
    item["canonical_blocks"]["text_blocks"][0]["bbox"][0] = 0.1
    with pytest.raises(ValueError, match="independently parsed"):
        rebuild(native, plan, result)


@pytest.mark.parametrize(
    "bbox,expected",
    [
        ([0, 0, 1000, 1000], [0.0, 0.0, 1.0, 1.0]),
        ([0, 0, 1, 1], [0.0, 0.0, 0.001, 0.001]),
    ],
)
def test_v2_boundary_and_tiny_box_use_declared_units_without_guessing(
    tmp_path, bbox, expected
):
    native, plan = plan_for(document(tmp_path))
    result = execute_visual_jobs(
        plan, lambda p, n: response_v2(p, n, bbox=bbox), tmp_path / "cache"
    )
    assert (
        result["evidence"][0]["canonical_blocks"]["text_blocks"][0]["bbox"] == expected
    )
    assert not rebuild(native, plan, result).unverified_content


@pytest.mark.parametrize(
    "case",
    [
        "negative",
        "huge-integer",
        "over1000",
        "zero-area",
        "reverse",
        "strings",
        "float-integer",
        "fractional",
        "bool",
        "missing-space",
        "unknown-space",
        "normalized-one-space",
        "string-schema",
        "v1-thousandths",
        "v1-added-space",
        "mixed-blocks",
        "block-space",
    ],
)
def test_undeclared_mixed_or_invalid_visual_coordinate_units_remain_unresolved(
    tmp_path, case
):
    native, plan = plan_for(document(tmp_path))

    def generate(request, maximum):
        value = response_v2(request, maximum)
        data = json.loads(value["text"])
        box_cases = {
            "negative": [-1, 0, 1000, 1000],
            "huge-integer": [0, 0, 10**1000, 1000],
            "over1000": [0, 0, 1001, 1000],
            "zero-area": [0, 0, 0, 1000],
            "reverse": [1000, 0, 0, 1000],
            "strings": ["526", "564", "608", "666"],
            "float-integer": [0.0, 0.0, 1000.0, 1000.0],
            "fractional": [0.526, 0.564, 0.608, 0.666],
            "bool": [False, False, 1000, 1000],
        }
        if case in box_cases:
            data["text_blocks"][0]["bbox"] = box_cases[case]
        if case == "missing-space":
            data.pop("bbox_space")
        if case == "unknown-space":
            data["bbox_space"] = "pixels"
        if case == "normalized-one-space":
            data["bbox_space"] = "normalized_1"
        if case == "string-schema":
            data["schema_version"] = "2"
        if case == "v1-thousandths":
            data["schema_version"] = 1
            data.pop("bbox_space")
        if case == "v1-added-space":
            data["schema_version"] = 1
        if case == "mixed-blocks":
            data["observations"] = [
                {"id": "o1", "bbox": [0.1, 0.2, 0.3, 0.4], "description": "mixed units"}
            ]
        if case == "block-space":
            data["text_blocks"][0]["bbox_space"] = "normalized_1"
        value["text"] = json.dumps(data)
        return value

    result = execute_visual_jobs(plan, generate, tmp_path / "cache")
    assert result["errors"] and result["raw"] and result["calls"]
    assert not result["evidence"]
    assert rebuild(native, plan, result).unverified_content == native.unverified_content


def test_v1_compatible_normalized_response_has_unambiguous_typed_metadata(tmp_path):
    native, plan = plan_for(document(tmp_path))
    result = execute_visual_jobs(plan, response, tmp_path / "cache")
    item = result["evidence"][0]
    assert item["response"]["schema_version"] == 1
    assert "bbox_space" not in item["response"]
    assert item["source_bbox_space"] == item["canonical_bbox_space"] == "normalized_1"
    assert (
        item["response"]["text_blocks"][0]["bbox"]
        == item["canonical_blocks"]["text_blocks"][0]["bbox"]
    )
    assert not rebuild(native, plan, result).unverified_content


def test_response_protocol_version_rejects_pre_v2_plan_before_dispatch(tmp_path):
    _, plan = plan_for(document(tmp_path))
    plan["identity"].pop("response_protocol_version")
    plan["plan_sha256"] = json_digest(
        {key: value for key, value in plan.items() if key != "plan_sha256"}
    )
    with pytest.raises(ValueError, match="response protocol version"):
        execute_visual_jobs(plan, response, tmp_path / "cache")
