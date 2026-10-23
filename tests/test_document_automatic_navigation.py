"""Receipt-owned positions, including recovery; no SDK/model/network is used."""
from copy import deepcopy
import inspect
import json

from docx import Document
import pytest

from bridge_evidence.contest.document_tools import DocumentTools
from bridge_evidence.contest.evidence_manifest import build_manifest, manifest_from_text
from bridge_evidence.contest.tong_react_adapter import wrap_tool_response
from bridge_evidence.contest import tong_react_adapter as adapter


def _tools(text="1 甲章\n甲正文\n甲末尾\n2 乙章\n乙正文", *, budget=3000, **kwargs):
    return DocumentTools(manifest_from_text(text), count_tokens=len, max_tool_tokens=budget,
                         native_response_budget=True, **kwargs)


def _drain(handler, **arguments):
    pages = []
    for _ in range(1000):
        result = handler(**arguments)
        assert result["ok"], result
        pages.append(result)
        if result["continuation"] is None:
            return pages
    raise AssertionError("Automatic navigation did not complete")


def _spans(pages):
    return [span for page in pages for span in page.get("spans", [])]


def test_model_signatures_do_not_require_positions_and_precise_apis_are_available():
    tools = _tools()
    handlers = tools.automatic_handlers()
    for name in ("read_section", "read_source", "read_table_rows", "read_outline", "search_in_document"):
        parameters = inspect.signature(handlers[name]).parameters
        assert not {"cursor", "start", "end"}.intersection(parameters)
        assert parameters["restart"].default is False
    assert {"read_source_range", "read_section_range", "read_table_range"}.issubset(handlers)
    # The legacy API is genuinely unchanged.
    assert "cursor" in inspect.signature(tools.handlers()["read_section"]).parameters


def test_native_dispatch_uses_cursor_free_signatures_and_precise_aliases():
    tools = _tools()
    session = {"handlers": tools.automatic_handlers(), "tool_calls": [], "max_calls": 100,
               "tool_budget_exceeded": False, "count_tokens": len, "max_tool_tokens": 3000}
    for name, arguments in [
        ("read_section", {"section_id": "E000001", "max_units": 1}),
        ("read_outline", {"limit": 1}),
        ("read_source", {"evidence_id": "E000002"}),
        ("search_in_document", {"query": ["正文"], "max_matches": 1}),
        ("read_source_range", {"evidence_id": "E000002", "start": 0, "end": 2}),
        ("read_section_range", {"section_id": "E000004"}),
        ("read_section", {"section_id": "E000001", "max_units": 1}),
    ]:
        result = adapter._invoke_active(session, name, arguments)
        assert result["ok"], (name, result)
        assert tools.receipts[-1]["tool"] == name
        assert tools.receipts[-1]["result_tokens"] == len(adapter._json_bytes(result).decode())
    assert result["result"]["spans"][0]["evidence_id"] == "E000002"


def test_chapter_switches_have_independent_positions_and_finish_requires_restart():
    tools = _tools()
    read = tools.automatic_handlers()["read_section"]
    first = read("E000001", max_units=1)
    second = read("E000004", max_units=1)
    resume = read("E000001", max_units=1)
    assert first["spans"][0]["evidence_id"] == "E000001"
    assert second["spans"][0]["evidence_id"] == "E000004"
    assert resume["spans"][0]["evidence_id"] == "E000002"
    assert resume["continuation"] == {"tool": "read_section", "arguments": {"section_id": "E000001", "max_units": 1}}
    last = read("E000001")
    assert last["spans"][0]["evidence_id"] == "E000003" and last["continuation"] is None
    again = read("E000001")
    assert again["status"] == "already_delivered" and again["spans"] == []
    assert again["automatic_navigation"]["delivery_complete"] is True
    assert "not_understanding_or_extraction" in again["automatic_navigation"]["progress_meaning"]
    assert read("E000001", max_units=1, restart=True)["spans"][0]["evidence_id"] == "E000001"


def test_long_source_pages_resume_the_original_section_without_skip_or_duplicate():
    source = '甲😀带"引号"与\\路径' * 500 + "\n最后一段"
    delivered = []
    tools = _tools(source, budget=1200, on_read=lambda *span: delivered.append(span))
    pages = _drain(tools.automatic_handlers()["read_section"], section_id="root", max_units=8)
    spans = _spans(pages)
    assert "".join(span["text"] for span in spans) == source
    assert len(pages) > 3
    assert delivered == [(span["evidence_id"], span["start"], span["end"]) for span in spans]
    assert len(delivered) == len(set(delivered))
    assert all("resume_section" not in page for page in pages)
    assert all(page["continuation"] is None or page["continuation"]["tool"] == "read_section" for page in pages)
    assert all(row["tool"] == "read_section" and row["result_tokens"] <= 1200 for row in tools.receipts)
    assert tools._response_projection is None and tools._response_tool is None


def test_source_object_progress_and_explicit_reread_do_not_change_chapter_position():
    tools = _tools("很长正文" * 500 + "\n尾段", budget=1200)
    handlers = tools.automatic_handlers()
    source_pages = _drain(handlers["read_source"], evidence_id="E000001")
    assert "".join(span["text"] for span in _spans(source_pages)) == tools._index["E000001"]["text"]
    before = tools.navigation_state()
    exact = handlers["read_source_range"]("E000001", 2, 7)
    assert exact["spans"][0]["start"] == 2 and exact["spans"][0]["end"] == 7
    assert tools.navigation_state() == before
    chapter = handlers["read_section"]("root", max_units=1)
    assert chapter["spans"][0]["start"] == 0
    assert tools.receipts[-2]["tool"] == "read_source_range"


def test_snapshot_resumes_successful_receipts_and_rejects_other_documents():
    saved = []
    first = _tools("长原文" * 800 + "\n尾段", budget=1000)
    handlers = first.automatic_handlers(on_progress=lambda state: saved.append(deepcopy(state)))
    initial = handlers["read_section"]("root")
    assert saved[-1] == first.navigation_state()
    restored = _tools("长原文" * 800 + "\n尾段", budget=1000)
    pages = _drain(restored.automatic_handlers(state=saved[-1])["read_section"], section_id="root")
    assert "".join(span["text"] for span in _spans([initial, *pages])) == "长原文" * 800 + "\n尾段"
    with pytest.raises(ValueError, match="immutable document"):
        _tools("另一个文档").automatic_handlers(state=saved[-1])


def test_two_tools_instances_never_share_positions_even_for_identical_source():
    first, second = _tools(), _tools()
    first.automatic_handlers()["read_section"]("root", max_units=1)
    a = first.automatic_handlers()["read_section"]("root", max_units=1)
    b = second.automatic_handlers()["read_section"]("root", max_units=1)
    assert a["spans"][0]["evidence_id"] == "E000002"
    assert b["spans"][0]["evidence_id"] == "E000001"


def test_read_callback_failure_and_bad_arguments_never_advance_navigation():
    failed = True
    def persist_read(*_):
        if failed:
            raise RuntimeError("durable read store unavailable")
    tools = _tools(on_read=persist_read)
    read = tools.automatic_handlers()["read_section"]
    before = tools.navigation_state()
    assert not read("root", max_units=1)["ok"]
    assert tools.navigation_state() == before
    failed = False
    assert read("root", max_units=1)["spans"][0]["evidence_id"] == "E000001"
    before = tools.navigation_state()
    assert not read("root", max_units=0)["ok"]
    assert not read("missing")["ok"]
    assert tools.navigation_state() == before


def test_navigation_persistence_failure_does_not_publish_position_or_success_receipt():
    failed = True
    def persist(state):
        if failed:
            raise OSError("navigation disk full")
    tools = _tools()
    read = tools.automatic_handlers(on_progress=persist)["read_section"]
    before = tools.navigation_state()
    result = read("root", max_units=1)
    assert result["ok"] is False and result["error"]["type"] == "OSError"
    assert tools.navigation_state() == before
    assert tools.receipts[-1]["status"] == "error" and tools.receipts[-1]["ranges"] == []
    failed = False
    assert read("root", max_units=1)["spans"][0]["evidence_id"] == "E000001"


def test_final_budget_failure_does_not_advance_or_grant_read_coverage():
    observed = []
    tools = _tools(on_read=lambda *span: observed.append(span))
    read = tools.automatic_handlers()["read_section"]
    original = tools._response_tokens
    # A deliberately unstable tokenizer makes preliminary fit checks pass
    # while the final result fails. The common receipt gate must still win.
    def count(result):
        if result.get("automatic_navigation") and "spans" in result:
            return 99999
        return original(result)
    tools._response_tokens = count
    before = tools.navigation_state()
    result = read("root", max_units=1)
    assert not result["ok"] and not observed
    assert tools.navigation_state() == before


def test_outline_search_and_source_state_are_independent_and_query_scoped():
    tools = _tools("1 检测\n裂缝 裂缝 剥落\n2 历史\n裂缝 剥落", budget=2200)
    handlers = tools.automatic_handlers()
    handlers["read_section"]("E000001", max_units=1)
    outline_first = handlers["read_outline"](limit=1)
    a1 = handlers["search_in_document"]("裂缝", section_id="E000001", max_matches=1)
    b1 = handlers["search_in_document"]("裂缝", section_id="E000003", max_matches=1)
    c1 = handlers["search_in_document"]("剥落", section_id="E000001", max_matches=1)
    a2 = handlers["search_in_document"]("裂缝", section_id="E000001", max_matches=1)
    outline_second = handlers["read_outline"](limit=1)
    assert a1["matches"][0]["match_start"] == 0 and a2["matches"][0]["match_start"] == 3
    assert b1["matches"][0]["evidence_id"] == "E000004"
    assert c1["matches"][0]["matched_query"] == "剥落"
    assert outline_first["sections"][0]["id"] == "root"
    assert outline_second["sections"][0]["id"] == "E000001"
    assert handlers["read_section"]("E000001", max_units=1)["spans"][0]["evidence_id"] == "E000002"
    assert handlers["search_in_document"]("裂缝", section_id="E000001")["status"] == "already_delivered"
    assert handlers["search_in_document"]("裂缝", section_id="E000001", restart=True)["matches"][0]["match_start"] == 0
    # Model references choose the source object, never manual offsets.
    assert a1["matches"][0]["read"] == {"tool": "read_source", "arguments": {
        "evidence_id": "E000002", "restart": True}}


@pytest.mark.parametrize("kind", ["read_section", "read_source", "read_table_rows"])
def test_large_table_context_automatically_pages_then_delivers_primary_without_loss(tmp_path, kind):
    path = tmp_path / "large-table.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=1)
    table.cell(0, 0).text = "表头单位毫米" * 450
    table.cell(1, 0).text = "第一行正文" * 180
    table.cell(2, 0).text = "第二行正文"
    doc.save(path)
    tools = DocumentTools(build_manifest(path), count_tokens=len, max_tool_tokens=1200, native_response_budget=True)
    handlers = tools.automatic_handlers()
    arguments = {"read_section": {"section_id": "root", "max_units": 1},
                 "read_source": {"evidence_id": "E000002"},
                 "read_table_rows": {"table_id": "T00001", "max_rows": 1}}[kind]
    pages = _drain(handlers[kind], **arguments)
    spans = _spans(pages)
    expected = (tools._units if kind in {"read_section", "read_table_rows"} else [tools._index["E000002"]])
    for unit in expected:
        # Recovery may legitimately deliver the same original as table context
        # and as primary; every requested character must nevertheless arrive.
        covered = {index for span in spans if span["evidence_id"] == unit["id"]
                   for index in range(span["start"], span["end"])}
        assert covered == set(range(len(unit["text"])))
    assert any(page.get("context_part") for page in pages)
    assert all(page["continuation"] is None or page["continuation"]["tool"] == kind for page in pages)
    assert all(receipt["result_tokens"] <= 1200 for receipt in tools.receipts)
    assert all(len(json.dumps(wrap_tool_response(kind, page), ensure_ascii=False, sort_keys=True,
                              separators=(",", ":"))) <= 1200 for page in pages)


def test_table_row_pages_and_precise_table_reread_are_separate(tmp_path):
    path = tmp_path / "rows.docx"
    doc = Document()
    table = doc.add_table(rows=5, cols=1)
    for index in range(5):
        table.cell(index, 0).text = f"第 {index} 行"
    doc.save(path)
    tools = DocumentTools(build_manifest(path), count_tokens=len, max_tool_tokens=5000, native_response_budget=True)
    handlers = tools.automatic_handlers()
    first = handlers["read_table_rows"]("T00001", max_rows=2)
    assert [span["row_index"] for span in first["spans"]] == [0, 1]
    before = tools.navigation_state()
    assert handlers["read_table_range"]("T00001", 0, 1)["spans"][0]["row_index"] == 0
    assert tools.navigation_state() == before
    tail = _drain(handlers["read_table_rows"], table_id="T00001", max_rows=2)
    assert [span["row_index"] for span in _spans(tail)] == [2, 3, 4]
    assert handlers["read_table_rows"]("T00001")["status"] == "already_delivered"
    assert handlers["read_table_rows"]("T00001", max_rows=1, restart=True)["spans"][0]["row_index"] == 0


def test_explicit_alias_pagination_counts_actual_tool_name_and_is_callable():
    source = "不同字符病害描述" * 400
    tools = _tools(source, budget=1200)
    handlers = tools.automatic_handlers()
    page = handlers["read_source_range"]("E000001")
    spans = []
    while True:
        assert page["ok"], page
        spans.extend(page["spans"])
        assert tools.receipts[-1]["tool"] == "read_source_range"
        assert tools.receipts[-1]["result_tokens"] <= 1200
        call = page["continuation"]
        if call is None:
            break
        assert call == {"tool": "read_source", "arguments": {"evidence_id": "E000001", "restart": True}}
        # The explicit Python API remains available to programmatic clients;
        # returned model recommendations do not ask it to calculate offsets.
        page = handlers["read_source_range"]("E000001", start=spans[-1]["end"])
    assert "".join(span["text"] for span in spans) == source
    assert tools.navigation_state()["objects"] == {}


def test_neighbor_budget_recovery_recommends_an_ordinary_source_read():
    tools = _tools("邻居很长" * 1000 + "\n正文", budget=1000)
    result = tools.automatic_handlers()["get_neighbors"]("E000002")
    assert not result["ok"] and result["status"] == "budget_exceeded"
    assert all(call["tool"] == "read_source" and set(call["arguments"]) == {"evidence_id", "restart"}
               for call in result["recovery"])
    assert tools.navigation_state()["objects"] == {}


def test_precise_range_error_recovery_uses_the_exposed_alias_and_legal_bounds():
    tools = _tools()
    handlers = tools.automatic_handlers()
    error = handlers["read_source_range"]("E000002", 0, 9999)
    assert not error["ok"] and error["error"]["code"] == "SOURCE_RANGE"
    call = error["next_read"]
    assert call["tool"] == "read_source"
    assert call["arguments"] == {"evidence_id": "E000002", "restart": True}
    read = handlers[call["tool"]](**call["arguments"])
    assert read["ok"] and read["spans"][0]["text"] == "甲正文\n"
    assert len(tools.navigation_state()["objects"]) == 1


@pytest.mark.parametrize("name,arguments", [
    ("read_section", {"section_id": "root", "restart": "true"}),
    ("read_outline", {"limit": 0}),
    ("search_in_document", {"query": "裂缝", "max_matches": True}),
    ("search_in_document", {"query": [], "scope": "document"}),
])
def test_bad_navigation_arguments_leave_state_unchanged(name, arguments):
    tools = _tools()
    handlers = tools.automatic_handlers()
    before = tools.navigation_state()
    assert not handlers[name](**arguments)["ok"]
    assert tools.navigation_state() == before
