#!/usr/bin/env python3
"""Fail-closed structural regression gate for IAIC result releases.

The gate does not pretend to predict the hidden leaderboard.  It answers the
safer question: did a candidate preserve the known-good v12 result structure,
history evidence coverage, and a bounded change surface?
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document


def _files(directory: Path, suffix: str) -> dict[str, Path]:
    paths = sorted(directory.rglob(f"*{suffix}"))
    result: dict[str, Path] = {}
    for path in paths:
        if path.name in result:
            raise ValueError(f"目录中存在重名文件: {path.name}")
        result[path.name] = path
    return result


def _tables(path: Path) -> list[list[list[str]]]:
    document = Document(path)
    return [
        [[cell.text.strip() for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]


def _state_coverage(directory: Path | None) -> dict[str, int]:
    if directory is None:
        return {"files": 0, "with_history_evidence": 0, "known_is_new": 0, "unknown_is_new": 0}
    states = _files(directory, ".json")
    coverage = Counter(files=len(states))
    for path in states.values():
        payload = json.loads(path.read_text(encoding="utf-8"))
        evidence = payload.get("evidence") or {}
        if any(str(key).startswith("HISTORY:") for key in evidence):
            coverage["with_history_evidence"] += 1
        for disease in payload.get("diseases") or []:
            if disease.get("is_new") is None:
                coverage["unknown_is_new"] += 1
            else:
                coverage["known_is_new"] += 1
    return {
        "files": coverage["files"],
        "with_history_evidence": coverage["with_history_evidence"],
        "known_is_new": coverage["known_is_new"],
        "unknown_is_new": coverage["unknown_is_new"],
    }


def _state_payloads(directory: Path | None) -> dict[str, dict[str, Any]]:
    if directory is None:
        return {}
    return {
        path.name: json.loads(path.read_text(encoding="utf-8"))
        for path in _files(directory, ".json").values()
    }


def _current_anchor(disease: dict[str, Any]) -> str | None:
    return next(
        (
            str(item)
            for item in disease.get("evidence") or []
            if re.fullmatch(r"T\d+R\d+", str(item))
        ),
        None,
    )


def _inactive_placeholder(disease: dict[str, Any]) -> bool:
    """Return true only for non-disease photo/repair ledger placeholders."""
    disease_type = re.sub(r"[\s、，,]", "", str(disease.get("type") or ""))
    description = re.sub(r"\s", "", str(disease.get("description") or ""))
    if disease_type in {"现状", "工作照", "桥涵标现状", "修复现状"}:
        return True
    if disease_type in {"修补", "修复"} and re.search(r"修补现状|修复现状|修补痕迹", description):
        return True
    return False


def _allowed_inactive_disease_prune(
    *,
    name: str,
    table_index: int,
    baseline_rows: list[list[str]],
    candidate_rows: list[list[str]],
    baseline_states: dict[str, dict[str, Any]],
    candidate_states: dict[str, dict[str, Any]],
) -> tuple[bool, dict[str, Any] | None]:
    """Allow only evidence-audited removal of inactive rows from table 3."""
    if table_index != 2 or not baseline_rows or not candidate_rows:
        return False, None
    if baseline_rows[0] != candidate_rows[0]:
        return False, None
    if any(len(row) != len(baseline_rows[0]) for row in baseline_rows + candidate_rows):
        return False, None
    state_name = str(Path(name).with_suffix(".json").name)
    baseline = baseline_states.get(state_name)
    candidate = candidate_states.get(state_name)
    if baseline is None or candidate is None:
        return False, None
    baseline_diseases = baseline.get("diseases") or []
    candidate_diseases = candidate.get("diseases") or []
    if len(baseline_rows) != len(baseline_diseases) + 1:
        return False, None
    if len(candidate_rows) != len(candidate_diseases) + 1:
        return False, None
    baseline_by_anchor = {
        anchor: disease
        for disease in baseline_diseases
        if (anchor := _current_anchor(disease)) is not None
    }
    candidate_anchors = {
        anchor
        for disease in candidate_diseases
        if (anchor := _current_anchor(disease)) is not None
    }
    if len(baseline_by_anchor) != len(baseline_diseases):
        return False, None
    if len(candidate_anchors) != len(candidate_diseases):
        return False, None
    if not candidate_anchors.issubset(baseline_by_anchor):
        return False, None
    removed = [
        disease
        for anchor, disease in baseline_by_anchor.items()
        if anchor not in candidate_anchors
    ]
    row_delta = len(baseline_rows) - len(candidate_rows)
    if row_delta <= 0 or row_delta != len(removed):
        return False, None
    if not all(_inactive_placeholder(disease) for disease in removed):
        return False, None
    return True, {
        "file": name,
        "table": table_index,
        "removed": len(removed),
        "anchors": [_current_anchor(disease) for disease in removed],
        "types": [str(disease.get("type") or "") for disease in removed],
    }


def _aligned_row_change_cost(
    baseline_rows: list[list[str]],
    candidate_rows: list[list[str]],
) -> int:
    """Cell edit cost without treating every post-deletion row as shifted."""
    previous = [0]
    for row in candidate_rows:
        previous.append(previous[-1] + len(row))
    for baseline_row in baseline_rows:
        current = [previous[0] + len(baseline_row)]
        for column, candidate_row in enumerate(candidate_rows, 1):
            if len(baseline_row) == len(candidate_row):
                substitution = sum(left != right for left, right in zip(baseline_row, candidate_row))
            else:
                substitution = max(len(baseline_row), len(candidate_row))
            current.append(min(
                previous[column - 1] + substitution,
                previous[column] + len(baseline_row),
                current[column - 1] + len(candidate_row),
            ))
        previous = current
    return previous[-1]


def compare(
    baseline_result: Path,
    candidate_result: Path,
    baseline_state: Path | None,
    candidate_state: Path | None,
    max_changed_cell_ratio: float,
) -> dict[str, Any]:
    baseline_files = _files(baseline_result, ".docx")
    candidate_files = _files(candidate_result, ".docx")
    baseline_names = set(baseline_files)
    candidate_names = set(candidate_files)
    failures: list[str] = []
    if baseline_names != candidate_names:
        failures.append("候选结果文件集合与 v12.6 不一致")

    changed_cells = 0
    total_cells = 0
    changed_files = 0
    schema_mismatches: list[str] = []
    samples: list[dict[str, Any]] = []
    changed_by_table: Counter[int] = Counter()
    audited_inactive_prunes: list[dict[str, Any]] = []
    baseline_states = _state_payloads(baseline_state)
    candidate_states = _state_payloads(candidate_state)
    for name in sorted(baseline_names & candidate_names):
        baseline_tables = _tables(baseline_files[name])
        candidate_tables = _tables(candidate_files[name])
        file_changed = False
        if len(baseline_tables) != len(candidate_tables):
            schema_mismatches.append(f"{name}: table {len(baseline_tables)} -> {len(candidate_tables)}")
            file_changed = True
        for table_index in range(min(len(baseline_tables), len(candidate_tables))):
            base_rows = baseline_tables[table_index]
            cand_rows = candidate_tables[table_index]
            if [len(row) for row in base_rows] != [len(row) for row in cand_rows]:
                allowed, prune = _allowed_inactive_disease_prune(
                    name=name,
                    table_index=table_index,
                    baseline_rows=base_rows,
                    candidate_rows=cand_rows,
                    baseline_states=baseline_states,
                    candidate_states=candidate_states,
                )
                if not allowed:
                    schema_mismatches.append(f"{name}: table {table_index} row/column shape changed")
                    file_changed = True
                else:
                    assert prune is not None
                    audited_inactive_prunes.append(prune)
                    cost = _aligned_row_change_cost(base_rows, cand_rows)
                    total_cells += max(
                        sum(len(row) for row in base_rows),
                        sum(len(row) for row in cand_rows),
                    )
                    changed_cells += cost
                    changed_by_table[table_index] += cost
                    file_changed = file_changed or cost > 0
                    if len(samples) < 50:
                        samples.append({
                            "file": name,
                            "table": table_index,
                            "audited_inactive_rows_removed": prune["removed"],
                            "anchors": prune["anchors"],
                        })
                    continue
            for row_index in range(min(len(base_rows), len(cand_rows))):
                for column_index in range(min(len(base_rows[row_index]), len(cand_rows[row_index]))):
                    total_cells += 1
                    before = base_rows[row_index][column_index]
                    after = cand_rows[row_index][column_index]
                    if before == after:
                        continue
                    file_changed = True
                    changed_cells += 1
                    changed_by_table[table_index] += 1
                    if len(samples) < 50:
                        samples.append({
                            "file": name,
                            "table": table_index,
                            "row": row_index,
                            "column": column_index,
                            "before": before[:300],
                            "after": after[:300],
                        })
        changed_files += int(file_changed)

    if schema_mismatches:
        failures.append("候选结果存在表格结构变化")
    changed_ratio = changed_cells / total_cells if total_cells else 1.0
    if changed_ratio > max_changed_cell_ratio:
        failures.append(
            f"改单元格比例 {changed_ratio:.2%} 超过门限 {max_changed_cell_ratio:.2%}"
        )

    baseline_history = _state_coverage(baseline_state)
    candidate_history = _state_coverage(candidate_state)
    if baseline_state is not None and candidate_state is not None:
        if candidate_history["with_history_evidence"] < baseline_history["with_history_evidence"]:
            failures.append("历史证据覆盖低于 v12.6")
        if candidate_history["known_is_new"] < baseline_history["known_is_new"]:
            failures.append("新增/非新增确定值数量低于 v12.6")

    return {
        "baseline": str(baseline_result.resolve()),
        "candidate": str(candidate_result.resolve()),
        "passed": not failures,
        "verdict": "eligible_for_proxy_evaluation" if not failures else "rejected_before_scoring",
        "failures": failures,
        "files": {
            "baseline": len(baseline_files),
            "candidate": len(candidate_files),
            "changed": changed_files,
            "missing": sorted(baseline_names - candidate_names),
            "extra": sorted(candidate_names - baseline_names),
        },
        "cells": {
            "compared": total_cells,
            "changed": changed_cells,
            "changed_ratio": changed_ratio,
            "changed_by_table": dict(sorted(changed_by_table.items())),
        },
        "schema_mismatches": schema_mismatches[:100],
        "audited_inactive_disease_prunes": audited_inactive_prunes,
        "history": {"baseline": baseline_history, "candidate": candidate_history},
        "change_samples": samples,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="将 IAIC 候选结果逐单元格与 v12.6 比较")
    parser.add_argument("--baseline-result", type=Path, required=True)
    parser.add_argument("--candidate-result", type=Path, required=True)
    parser.add_argument("--baseline-state", type=Path)
    parser.add_argument("--candidate-state", type=Path)
    parser.add_argument("--max-changed-cell-ratio", type=float, default=0.05)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    report = compare(
        args.baseline_result,
        args.candidate_result,
        args.baseline_state,
        args.candidate_state,
        args.max_changed_cell_ratio,
    )
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    raise SystemExit(0 if report["passed"] else 2)


if __name__ == "__main__":
    main()
