"""文档结构化切片器（Document Splitter）。

按工程结构物理部位与功能模块切分桥梁定检报告：
1. overview: 工程概况、评定表、综合评价与总体结论
2. advice: 养护维修建议章节
3. deck: 桥面系检查记录与表格
4. superstructure: 上部结构检查记录与表格
5. substructure: 下部结构与支座系统检查记录与表格
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any
from docx import Document


def _clean(s: str | None) -> str:
    if not s:
        return ""
    return re.sub(r"\s+", " ", str(s)).strip()


def split_document_sections(doc_path: Path | str) -> dict[str, str]:
    """将 DOCX 文档按部位和功能拆分为不同上下文切片。"""
    path = Path(doc_path)
    if path.suffix.lower() == ".doc":
        raise ValueError(f"doc_splitter 仅支持 .docx 文件，请先转换为 .docx: {path}")

    doc = Document(str(path))

    overview_paras: list[str] = []
    advice_paras: list[str] = []
    deck_items: list[str] = []
    super_items: list[str] = []
    sub_items: list[str] = []

    # 1. 遍历段落，根据标题与章节归类
    current_sec = "overview"

    for p in doc.paragraphs:
        text = _clean(p.text)
        if not text:
            continue

        if len(text) < 40:
            if re.search(r"^(?:第?[一二三四五六七八九十\d]+[\.、\s]*)?[0-9\.]*\s*(?:工程概况|基本情况|桥梁概况|检测评估报告综述|检测结论|总体评价|结论)", text):
                current_sec = "overview"
            elif re.search(r"^(?:第?[一二三四五六七八九十\d]+[\.、\s]*)?[0-9\.]*\s*(?:养护建议|维修建议|养护对策|处置建议)\b", text) or re.search(r"^\d+[\.、\s]*建议", text):
                current_sec = "advice"
            elif re.search(r"^(?:第?[一二三四五六七八九十\d]+[\.、\s]*)?[0-9\.]*\s*桥面系", text):
                current_sec = "deck"
            elif re.search(r"^(?:第?[一二三四五六七八九十\d]+[\.、\s]*)?[0-9\.]*\s*上部结构", text):
                current_sec = "super"
            elif re.search(r"^(?:第?[一二三四五六七八九十\d]+[\.、\s]*)?[0-9\.]*\s*(?:下部结构|支座)", text):
                current_sec = "sub"
            elif re.search(r"^(?:第?[一二三四五六七八九十\d]+[\.、\s]*)?[0-9\.]*\s*(?:技术状况评定|荷载试验|桥梁结构检算|旧桥检算|试验结果)", text):
                current_sec = "calc_eval"
            elif re.search(r"^(?:附录|表C-|附件)", text):
                current_sec = "appendix"

        if current_sec == "overview":
            overview_paras.append(text)
        elif current_sec == "advice":
            advice_paras.append(text)
        elif current_sec == "deck":
            deck_items.append(text)
        elif current_sec == "super":
            super_items.append(text)
        elif current_sec == "sub":
            sub_items.append(text)

    # 2. 遍历表格，根据表前标题或表头归类
    last_defect_sec = "super"
    for t_idx, table in enumerate(doc.tables):
        if len(table.rows) < 2:
            continue

        prev_elem = table._element.getprevious()
        caption = _clean(prev_elem.text) if prev_elem is not None and hasattr(prev_elem, "text") else ""
        header_cells = [_clean(c.text) for c in table.rows[0].cells]
        header_str = caption + " " + " | ".join(header_cells)

        # 忽略计算、荷载试验、材料物理试验等无关大表
        if any(ign in header_str for ign in ("荷载试验", "静载", "动载", "有限元", "检算", "承载能力", "电位", "电阻率", "氯离子", "永久观测点", "交通量", "设备编号")):
            continue

        # 过滤与转录表格文本
        table_lines = [f"[表格: {caption or f'表{t_idx}'}] 表头: " + " | ".join(header_cells)]
        for row in table.rows[1:]:
            row_cells = [_clean(c.text) for c in row.cells]
            row_text = " | ".join(c for c in row_cells if c)
            if row_text and not re.search(r"^(?:照片|示意图)", row_text):
                table_lines.append(row_text)
        table_text = "\n".join(table_lines)

        # 分类归入对应切片
        if any(k in header_str for k in ("养护建议", "维修建议", "建议措施")):
            advice_paras.append(table_text)
        elif any(k in header_str for k in ("桥面系", "桥面铺装", "伸缩缝", "护栏", "人行道")):
            deck_items.append(table_text)
            last_defect_sec = "deck"
        elif any(k in header_str for k in ("上部结构", "主梁", "T梁", "箱梁", "横隔板", "翼缘板")):
            super_items.append(table_text)
            last_defect_sec = "super"
        elif any(k in header_str for k in ("下部结构", "桥墩", "桥台", "盖梁", "墩台", "支座")):
            sub_items.append(table_text)
            last_defect_sec = "sub"
        elif any(k in header_str for k in ("评定表", "综合评定", "评分", "扣分", "概况", "资料卡", "综述")):
            overview_paras.append(table_text)
        elif any(k in header_str for k in ("病害", "缺损", "损坏", "裂缝", "外观")):
            if last_defect_sec == "deck":
                deck_items.append(table_text)
            elif last_defect_sec == "sub":
                sub_items.append(table_text)
            else:
                super_items.append(table_text)
        elif caption == "" and len(header_cells) > 1 and ("序号" in header_str or "位置" in header_str or "病害" in header_str or header_cells[0] == ""):
            if last_defect_sec == "deck":
                deck_items.append(table_text)
            elif last_defect_sec == "sub":
                sub_items.append(table_text)
            else:
                super_items.append(table_text)

    return {
        "overview": "\n\n".join(overview_paras[:120]),
        "advice": "\n\n".join(advice_paras[:80]),
        "deck": "\n\n".join(deck_items[:150]),
        "superstructure": "\n\n".join(super_items[:250]),
        "substructure": "\n\n".join(sub_items[:250]),
    }
