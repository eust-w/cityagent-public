"""Export saved drafts and page the actual DOCX back without adding facts.

Preview generation does not complete a report or perform a semantic review.
The workflow owns the native finish and same-agent file-check evidence. Failed
runs publish a separate, explicitly incomplete artifact, even with no draft.
"""
from __future__ import annotations

from copy import deepcopy
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import uuid

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .docx_writer import (
    ADVICE_FIELDS, BASIC_FIELDS, DEFECT_FIELDS,
    write_numbered_extraction_docx as write_extraction_docx,
)
from .document_delivery_validation import (
    VALIDATOR_VERSION, _expected, _validate_table, validate_delivery_data)
from .run_state import atomic_json, json_digest


SCOPE_NOTE = "本文件处理范围为已解析文本与表格；不声明已检查全部图片或完整物理报告。"
_VIEWS = ("basic", "advice", "defects")
_SECTIONS = {"summary", "text", *_VIEWS}
CONTENT_IDENTITY_VERSION = "document-word-content-v2"


class DeliveryArtifactError(ValueError):
    """The exported file no longer matches the saved draft or receipt."""


def _sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _version(snapshot: dict):
    version = snapshot.get("workspace_version", snapshot.get("version"))
    if type(version) is not int or version < 0:
        raise DeliveryArtifactError("A saved workspace version is required")
    if "version" in snapshot and snapshot["version"] != version:
        raise DeliveryArtifactError("Workspace version aliases disagree")
    return version


def _binding(snapshot: dict) -> tuple[int, str, object]:
    if not isinstance(snapshot, dict):
        raise DeliveryArtifactError("Saved workspace snapshot must be an object")
    version = _version(snapshot)
    draft = deepcopy(snapshot.get("draft"))
    digest = json_digest(draft)
    declared = snapshot.get("draft_sha256")
    # An initial empty workspace deliberately declares no draft hash.
    if declared != digest and not (draft is None and declared is None):
        raise DeliveryArtifactError("Saved draft hash does not match saved draft")
    return version, digest, draft


def artifact_content_identity(snapshot: dict, document_id: str, scope: dict) -> dict:
    """Immutable Word inputs, separate from mutable reading/audit counters."""
    _version_value, digest, _draft = _binding(snapshot)
    return {"schema": CONTENT_IDENTITY_VERSION, "document_id": document_id,
            "document_key": _sha(document_id.encode("utf-8")),
            "workspace_identity_sha256": json_digest(snapshot.get("identity")),
            "draft_sha256": digest, "scope_sha256": json_digest(scope),
            "export_contract_sha256": json_digest({"validator": VALIDATOR_VERSION,
                "basic_fields": BASIC_FIELDS, "advice_fields": ADVICE_FIELDS,
                "defect_fields": DEFECT_FIELDS, "scope_note": SCOPE_NOTE,
                "layout": "submission_clean_v1" if scope.get("submission_format") else "processing_note_after_title"})}


def artifact_matches_snapshot(snapshot: dict, artifact: dict | None, *,
                              document_id: str | None = None, scope: dict | None = None) -> bool:
    """Pure identity check; callers must still validate actual file bytes.

    Legacy receipts retain their original strict version contract. New receipts
    may survive a later reading-only workspace version, never a different draft,
    document, processing scope or export contract.
    """
    if not isinstance(snapshot, dict) or not isinstance(artifact, dict):
        return False
    if document_id is not None and artifact.get("document_id") != document_id:
        return False
    if scope is not None and artifact.get("scope") != scope:
        return False
    if "content_identity_sha256" not in artifact:
        return (artifact.get("workspace_version") == snapshot.get("version", snapshot.get("workspace_version"))
                and artifact.get("draft_sha256") == snapshot.get("draft_sha256"))
    try:
        version = _version(snapshot)
        exported_version = artifact.get("workspace_version")
        expected = artifact_content_identity(snapshot, document_id or artifact["document_id"],
            scope if scope is not None else artifact["scope"])
        return (type(exported_version) is int and 0 <= exported_version <= version
                and artifact.get("document_key") == expected["document_key"]
                and artifact.get("draft_sha256") == expected["draft_sha256"]
                and artifact.get("content_identity") == expected
                and artifact["content_identity_sha256"] == json_digest(expected))
    except (KeyError, TypeError, ValueError):
        return False


def artifact_page_matches(page: dict, artifact: dict | None) -> bool:
    """Match the immutable exported file; never rewrite a page's old proof."""
    if not isinstance(page, dict) or not isinstance(artifact, dict):
        return False
    if any(page.get(key) != artifact.get(key) for key in
           ("workspace_version", "draft_sha256", "file_sha256")):
        return False
    return (page.get("content_identity_sha256") == artifact["content_identity_sha256"]
            if "content_identity_sha256" in artifact else "content_identity_sha256" not in page)


def _has_content(data: dict) -> bool:
    return any(value not in (None, "") for value in data["基本信息"].values()) or bool(
        data["维修建议"] or data["病害明细"])


def _field_issues(data) -> list[dict]:
    """Name unmapped fields without copying their values into tool metadata."""
    if not isinstance(data, dict):
        return [{"path": "draft", "issue": "not_a_three_section_object"}]
    issues = []
    expected = {"基本信息", "维修建议", "病害明细"}
    for field in sorted(set(data) - expected):
        issues.append({"path": field, "issue": "unsupported_section"})
    for field in sorted(expected - set(data)):
        issues.append({"path": field, "issue": "missing_section"})
    basic = data.get("基本信息")
    if isinstance(basic, dict):
        for field in sorted(set(basic) - {name for name, _ in BASIC_FIELDS}):
            issues.append({"path": f"基本信息.{field}", "issue": "unsupported_field"})
    for section, fields in (("维修建议", ADVICE_FIELDS), ("病害明细", DEFECT_FIELDS)):
        rows = data.get(section)
        if isinstance(rows, list):
            for index, row in enumerate(rows):
                if isinstance(row, dict):
                    for field in sorted(set(row) - set(fields)):
                        issues.append({"path": f"{section}[{index}].{field}", "issue": "unsupported_field"})
    return issues


def _safe_note(text: str) -> str:
    # Status explanations are administrative text, not report facts. XML cannot
    # represent all Unicode characters, so invalid explanatory text is omitted.
    return "".join(char for char in text if char in "\t\n\r" or
                   "\u0020" <= char <= "\ud7ff" or "\ue000" <= char <= "\ufffd" or
                   "\U00010000" <= char <= "\U0010ffff")


def _front_processing_note(path: Path, note: str) -> None:
    """Move the writer's two actual note paragraphs after its unchanged title."""
    document = Document(BytesIO(path.read_bytes()))
    paragraphs = document.paragraphs
    if (len(paragraphs) < 3 or paragraphs[-2].text != "处理范围说明"
            or paragraphs[-1].text != note):
        raise DeliveryArtifactError("Writer processing note is missing or different")
    first, heading, explanation = paragraphs[0], paragraphs[-2], paragraphs[-1]
    first._p.addnext(heading._p)
    heading._p.addnext(explanation._p)
    document.save(str(path))


def validate_delivery_docx(data: dict, path: str | Path, *, processing_note: str) -> dict:
    """Use the shared exact values/table checks, with only note placement changed.

    This validates the actual exported bytes, including every original field,
    header, duplicate paragraph and row; it does not validate a reconstructed
    surrogate file or alter the shared writer/validator's contract.
    """
    validate_delivery_data(data)
    content = Path(path).read_bytes()
    document = Document(BytesIO(content))
    expected = _expected(data, processing_note)
    # The existing writer's final heading/note are now directly after title.
    expected = expected[:1] + expected[-2:] + expected[1:-2]
    actual = []
    for element in document.element.body:
        if element.tag == qn("w:p"):
            actual.append(("paragraph", Paragraph(element, document)))
        elif element.tag == qn("w:tbl"):
            actual.append(("table", Table(element, document)))
        elif element.tag != qn("w:sectPr"):
            raise DeliveryArtifactError("DOCX contains an unsupported body block")
    if len(actual) != len(expected):
        raise DeliveryArtifactError("DOCX block count differs from saved delivery")
    table_rows, cells_checked, paragraphs_checked = [], 0, 0
    for index, ((kind, block), (expected_kind, value)) in enumerate(zip(actual, expected)):
        if kind != expected_kind:
            raise DeliveryArtifactError(f"DOCX block {index} placement differs from delivery layout")
        if kind == "paragraph":
            if block.text != value:
                raise DeliveryArtifactError(f"DOCX paragraph {index} differs from saved delivery")
            paragraphs_checked += 1
        else:
            cells_checked += _validate_table(block, value, len(table_rows))
            table_rows.append(len(value) - 1)
    return {"schema_version": 1, "validator": VALIDATOR_VERSION,
            "layout": "processing_note_after_title", "report_sha256": json_digest(data),
            "docx_sha256": _sha(content), "table_data_rows": table_rows,
            "cells_checked": cells_checked, "paragraphs_checked": paragraphs_checked,
            "processing_note_sha256": json_digest(processing_note)}


class DocumentDeliveryArtifacts:
    """One document's immutable Word/JSON exports and actual-file readback.

    ``incomplete=False`` means a preview without a failure marker, never an
    assertion that work or model checking has completed. Every status starts
    ``complete=False``. A workflow may update status.json after actual finish,
    while leaving the bound Word and JSON bytes unchanged.
    """

    def __init__(self, directory: str | Path, document_id: str, scope: dict | None = None):
        if not isinstance(document_id, str) or not document_id:
            raise ValueError("A nonempty document ID is required")
        self.document_id = document_id
        self.document_key = _sha(document_id.encode("utf-8"))
        self.directory = Path(directory).resolve() / ("document-" + self.document_key)
        self.scope = deepcopy(scope) if scope is not None else {
            "stage": "parsed_text_and_tables", "physical_report_complete": False}
        # Validate metadata before any file publication.
        json_digest(self.scope)
        self._receipt: dict | None = None

    def export(self, snapshot: dict, *, incomplete: bool = True, reason: str = "",
               structure_status: dict | None = None) -> dict:
        """Save the exact draft JSON, export Word, then validate its readback.

        Invalid/empty drafts and writer failures produce notice-only Word plus
        the unfiltered original draft JSON. No unknown fields are dropped into
        an apparently successful report. Filesystem errors are raised rather
        than claiming a file exists when it could not be written.
        """
        if type(incomplete) is not bool or not isinstance(reason, str):
            raise ValueError("incomplete must be boolean and reason must be text")
        raw_draft = deepcopy(snapshot.get("draft")) if isinstance(snapshot, dict) else None
        # Workspace snapshots are JSON values; a serialization error remains an
        # explicit failure and is never replaced by invented report values.
        draft_digest = json_digest(raw_draft)
        declared_digest = snapshot.get("draft_sha256") if isinstance(snapshot, dict) else None
        version = None
        export_error = None
        try:
            version, draft_digest, raw_draft = _binding(snapshot)
        except DeliveryArtifactError as error:
            export_error = type(error).__name__
        status_input = deepcopy(structure_status) if structure_status is not None else {}
        json_digest(status_input)
        category = "partial" if incomplete else "preview"
        export_dir = self.directory / category / (
            f"v{version if version is not None else 'unknown'}-{draft_digest[:16]}-{uuid.uuid4().hex}")
        export_dir.mkdir(parents=True, exist_ok=False)
        json_path = export_dir / ("partial.json" if incomplete else "report.json")
        docx_path = export_dir / ("partial.docx" if incomplete else "report.docx")
        status_path = export_dir / "status.json"
        # Preserve the actual raw draft even if schema/readback validation fails.
        atomic_json(json_path, raw_draft)
        receipt = {
            "schema_version": 1, "document_id": self.document_id,
            "document_key": self.document_key, "workspace_version": version,
            "draft_sha256": draft_digest, "saved_draft_sha256": declared_digest,
            "draft_json_sha256": _sha(json_path.read_bytes()),
            "json_path": str(json_path), "docx_path": str(docx_path),
            "status_path": str(status_path), "scope": deepcopy(self.scope),
            "processing_scope": "parsed_text_and_tables_only",
            "physical_report_complete": False,
            "complete": False, "incomplete": True,
            "requested_failure_marker": incomplete,
            "phase": "incomplete" if incomplete else "pending_agent_self_check",
            "reason": reason, "structure_status": status_input,
            "semantic_review": {"performed": False, "passed": None,
                                "status": "not_performed_by_artifact_module"},
            "agent_file_self_check": {"performed": False, "passed": None,
                                      "status": "pending_workflow_evidence"},
            "export_error": export_error, "artifact_kind": "notice_only",
            "export_issues": _field_issues(raw_draft),
            "structure_valid": False, "content_validated": False,
            "required_views": [],
        }
        if export_error is None:
            receipt["content_identity"] = artifact_content_identity(snapshot, self.document_id, self.scope)
            receipt["content_identity_sha256"] = json_digest(receipt["content_identity"])
        atomic_json(status_path, receipt)
        note = SCOPE_NOTE + (
            " 当前交付未完成；文件只保留已保存草稿，不代表完整结果。" if incomplete else
            " 当前文件为待自查预览；是否完成交付以同目录 status.json 为准。")
        pending = export_dir / (".word-" + uuid.uuid4().hex + ".docx")
        try:
            if export_error is not None:
                raise DeliveryArtifactError("Saved workspace binding is invalid")
            validate_delivery_data(raw_draft)
            if not _has_content(raw_draft):
                raise DeliveryArtifactError("Saved draft is empty")
            if self.scope.get("submission_format"):
                from .document_delivery_validation import validate_delivery_docx as validate_clean
                write_extraction_docx(raw_draft, str(pending), submission_format=True)
                validation = validate_clean(raw_draft, pending, submission_format=True)
            else:
                write_extraction_docx(raw_draft, str(pending), processing_note=note)
                _front_processing_note(pending, note)
                validation = validate_delivery_docx(raw_draft, pending, processing_note=note)
            os.replace(pending, docx_path)
            receipt.update(artifact_kind="three_table_report", structure_valid=True,
                           content_validated=True, required_views=list(_VIEWS),
                           readback_validation=validation)
        except Exception as error:
            # Failure produces a new administrative notice, without inferring,
            # filtering, stringifying, or silently completing report facts.
            pending.unlink(missing_ok=True)
            receipt.update(export_error=export_error or type(error).__name__,
                           phase="incomplete", incomplete=True)
            document = Document()
            document.add_heading("未完成的信息提取交付", level=0)
            document.add_paragraph("尚未生成可完整导出的有效草稿。原始保存稿保留在同目录 JSON 文件中。")
            document.add_paragraph(SCOPE_NOTE)
            if reason:
                document.add_paragraph("未完成说明：" + _safe_note(reason))
            document.save(str(pending))
            notice = Document(BytesIO(pending.read_bytes()))
            if notice.tables or not notice.paragraphs or notice.paragraphs[0].text != "未完成的信息提取交付":
                raise DeliveryArtifactError("Incomplete notice readback failed")
            os.replace(pending, docx_path)
        finally:
            pending.unlink(missing_ok=True)
        receipt["file_sha256"] = receipt["docx_sha256"] = _sha(docx_path.read_bytes())
        receipt["files_available"] = True
        atomic_json(status_path, receipt)
        self._receipt = deepcopy(receipt)
        return deepcopy(receipt)

    def matches_snapshot(self, snapshot: dict, receipt: dict | None) -> bool:
        return artifact_matches_snapshot(snapshot, receipt, document_id=self.document_id, scope=self.scope)

    def _read_bound(self, snapshot: dict, receipt: dict | None,
                    expected_file_sha256: str | None):
        record = deepcopy(receipt if receipt is not None else self._receipt)
        if not isinstance(record, dict):
            raise DeliveryArtifactError("No exported artifact receipt is available")
        version, digest, _draft = _binding(snapshot)
        if (record.get("document_key") != self.document_key
                or not self.matches_snapshot(snapshot, record)):
            raise DeliveryArtifactError("Artifact does not match the current saved workspace")
        for field in ("json_path", "docx_path", "status_path"):
            try:
                Path(record[field]).resolve().relative_to(self.directory)
            except (KeyError, TypeError, ValueError) as error:
                raise DeliveryArtifactError("Artifact path is outside this document") from error
        json_content = Path(record["json_path"]).read_bytes()
        if _sha(json_content) != record.get("draft_json_sha256") or json_digest(json.loads(json_content)) != digest:
            raise DeliveryArtifactError("Saved draft artifact has changed")
        content = Path(record["docx_path"]).read_bytes()
        file_sha = _sha(content)
        if file_sha != record.get("file_sha256") or file_sha != record.get("docx_sha256"):
            raise DeliveryArtifactError("Exported Word has changed")
        if expected_file_sha256 is not None and file_sha != expected_file_sha256:
            raise DeliveryArtifactError("Expected Word hash does not match")
        try:
            document = Document(BytesIO(content))
        except Exception as error:
            raise DeliveryArtifactError("Cannot reopen exported Word") from error
        if record.get("artifact_kind") == "three_table_report":
            if len(document.tables) != 3:
                raise DeliveryArtifactError("The actual Word does not have three delivery tables")
            # Receipts can be restored by a caller, but may not replace actual
            # content validation. Recheck bytes against the saved JSON on read.
            note = SCOPE_NOTE + (
                " 当前交付未完成；文件只保留已保存草稿，不代表完整结果。"
                if record.get("requested_failure_marker") else
                " 当前文件为待自查预览；是否完成交付以同目录 status.json 为准。")
            if self.scope.get("submission_format"):
                from .document_delivery_validation import validate_delivery_docx as validate_clean
                validate_clean(_draft, record["docx_path"], submission_format=True)
            else:
                validate_delivery_docx(_draft, record["docx_path"], processing_note=note)
            if _sha(Path(record["docx_path"]).read_bytes()) != file_sha:
                raise DeliveryArtifactError("Exported Word changed during readback")
        return record, document

    @staticmethod
    def _views(record: dict, document) -> tuple[dict[str, str], dict]:
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        sections = {name: "" for name in _VIEWS}
        table_metadata = {}
        if record.get("artifact_kind") == "three_table_report":
            for name, table in zip(_VIEWS, document.tables):
                headers = [cell.text for cell in table.rows[0].cells]
                rows = [[cell.text for cell in row.cells] for row in table.rows[1:]]
                sections[name] = "\n".join(json.dumps(dict(zip(headers, row)), ensure_ascii=False)
                                           for row in rows)
                table_metadata[name] = {"headers": headers, "row_count": len(rows),
                                        "total_chars": len(sections[name])}
        else:
            table_metadata = {name: {"headers": [], "row_count": 0, "total_chars": 0}
                              for name in _VIEWS}
        body_text = []
        for element in document.element.body:
            if element.tag == qn("w:p"):
                body_text.append(Paragraph(element, document).text)
            elif element.tag == qn("w:tbl"):
                for row in Table(element, document).rows:
                    body_text.append(json.dumps([cell.text for cell in row.cells], ensure_ascii=False))
        sections["text"] = "\n".join(body_text)
        summary = {"artifact_kind": record["artifact_kind"], "complete": False,
                   "phase": record["phase"], "structure_valid": record["structure_valid"],
                   "processing_scope": record["processing_scope"],
                   "physical_report_complete": False,
                   "required_views": record["required_views"], "sections": table_metadata,
                   "export_error": record.get("export_error"),
                   "export_issues": record.get("export_issues", []),
                   "actual_word_paragraphs": paragraphs}
        sections["summary"] = json.dumps(summary, ensure_ascii=False, indent=2)
        return sections, table_metadata

    def read_page(self, snapshot: dict, *, section: str = "summary", cursor: int = 0,
                  max_chars: int = 4096, expected_file_sha256: str | None = None,
                  receipt: dict | None = None) -> dict:
        """Read an exact character range of a view built from reopened Word.

        Long individual fields can be read over multiple pages. Reading an
        empty actual table returns a zero-length page with its real headers;
        neither draft text nor model notes stand in for file contents.
        """
        if section not in _SECTIONS or type(cursor) is not int or cursor < 0:
            raise ValueError("Unsupported section or invalid cursor")
        if type(max_chars) is not int or not 1 <= max_chars <= 65536:
            raise ValueError("max_chars must be between 1 and 65536")
        record, document = self._read_bound(snapshot, receipt, expected_file_sha256)
        sections, tables = self._views(record, document)
        text = sections[section]
        if cursor > len(text):
            raise ValueError("Cursor exceeds actual file view")
        end = min(cursor + max_chars, len(text))
        # Tool responses expose binding metadata, never the potentially large
        # workspace progress/status snapshot retained in the disk receipt.
        public_receipt = {key: deepcopy(record[key]) for key in (
            "schema_version", "document_id", "document_key", "workspace_version", "draft_sha256",
            "saved_draft_sha256", "draft_json_sha256", "file_sha256", "docx_sha256",
            "json_path", "docx_path", "status_path", "processing_scope", "physical_report_complete",
            "complete", "incomplete", "phase", "requested_failure_marker", "artifact_kind",
            "structure_valid", "content_validated", "required_views")}
        if "content_identity_sha256" in record:
            public_receipt["content_identity_sha256"] = record["content_identity_sha256"]
        result = {"section": section, "text": text[cursor:end], "cursor": cursor,
                  "end": end, "next_cursor": end if end < len(text) else None,
                  "total_chars": len(text), "complete_page": end == len(text),
                  "receipt": public_receipt}
        if section in _VIEWS:
            result.update(tables[section])
        if section == "summary":
            result.update(sections={name: {"total_chars": len(value)} for name, value in sections.items()},
                          required_views=record["required_views"], tables=tables)
        return result
