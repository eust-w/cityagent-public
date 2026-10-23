from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "scripts/benchmark_conversion_cache.py"
spec = importlib.util.spec_from_file_location("conversion_cache_benchmark", SCRIPT)
benchmark = importlib.util.module_from_spec(spec)
spec.loader.exec_module(benchmark)


def inputs(tmp_path):
    root = tmp_path / "training"
    root.mkdir()
    source = root / "report.docx"
    doc = Document()
    doc.add_paragraph("同一位置的重复证据")
    doc.add_paragraph("同一位置的重复证据")
    doc.save(source)
    broken = root / "broken.docx"
    broken.write_bytes(b"invalid OPC archive")
    manifest = tmp_path / "inputs.json"
    benchmark.prep.save(manifest, {"scope": "official-training", "training_root": str(root),
        "reports": [{"path": str(path), "sha256": benchmark.prep.digest(path)}
                    for path in (source, broken)]})
    return manifest


def test_actual_native_preparation_binds_source_and_retains_all_failures(tmp_path):
    manifest = inputs(tmp_path)
    output = tmp_path / "run"
    result = benchmark.benchmark(manifest, output, repeats=3, timeout=10)
    assert result["model_requests"] == 0
    assert result["paired_reports"] == 6
    assert result["failures"] == {"cold": 3, "hit": 3}
    assert not result["all_pairs_qualified"]
    assert result["qualified_preparation_wall_reduction"] is None
    rounds = json.loads((output / "rounds.json").read_text())
    assert len(rounds) == 6
    for item in rounds:
        good, bad = item["reports"]
        assert good["status"] == "success"
        assert good["source_binding"]["source_sha256"] == good["source"]["sha256"]
        assert good["source_binding"]["prepared_docx_sha256"] == good["target_sha256"]
        assert good["source_binding"]["conversion"] is None
        assert good["counts"]["units"] == 2
        assert good["subprocess_counts"] == {"version": 0, "convert": 0, "other": 0}
        assert good["loaded_model_or_sdk_modules"] == []
        assert bad["status"] == "failed"
        assert bad["wall_seconds"] > 0
    public = (output / "public-results.json").read_text()
    assert str(tmp_path) not in public
    assert "同一位置" not in public
    assert "invalid OPC archive" not in public
    for relative, expected in json.loads((output / "identity.json").read_text())["code"].items():
        assert benchmark.prep.digest(output / "executed-source" / relative) == expected


def record(*, suffix=".doc", phase="cold"):
    return {"status": "success", "source": {"path": "report" + suffix, "sha256": "source"},
            "target_sha256": "target", "content_sha256": "content",
            "package_parts_sha256": "parts", "source_binding_sha256": "binding",
            "subprocess_counts": {"version": int(suffix == ".doc"),
                "convert": int(suffix == ".doc" and phase == "cold"), "other": 0}}


@pytest.mark.parametrize("field", ["target_sha256", "content_sha256", "package_parts_sha256",
                                   "source_binding_sha256"])
def test_pair_rejects_any_byte_or_binding_change(field):
    cold, hit = record(), record(phase="hit")
    hit[field] = "changed"
    assert not benchmark.compare_pair(cold, hit)["qualified"]


def test_pair_requires_real_cache_hit_and_complete_denominator():
    cold, hit = record(), record(phase="hit")
    assert benchmark.compare_pair(cold, hit)["qualified"]
    hit["subprocess_counts"]["convert"] = 1
    assert not benchmark.compare_pair(cold, hit)["qualified"]
    hit["subprocess_counts"]["convert"] = 0
    hit["status"] = "failed"
    assert not benchmark.compare_pair(cold, hit)["qualified"]


def test_r7_quantile_uses_all_observed_service_latencies():
    assert benchmark.quantile([1, 2, 3, 4], .5) == 2.5
    assert benchmark.quantile([1, 2, 3, 4], .95) == pytest.approx(3.85)


def test_actual_subprocess_audit_records_command_without_environment(tmp_path):
    events = tmp_path / "events.jsonl"
    code = ("import runpy,subprocess,sys; m=runpy.run_path(" + repr(str(SCRIPT)) + "); "
            "m['install_process_audit'](" + repr(str(events)) + "); "
            "subprocess.run([sys.executable,'-c','pass'],check=True)")
    subprocess.run([sys.executable, "-B", "-c", code], check=True, timeout=10)
    actual = json.loads(events.read_text())
    assert actual["kind"] == "other"
    assert actual["argv"][1:] == ["-c", "pass"]
    assert "env" not in actual
    assert actual["event"] == "subprocess.Popen"


def test_worker_rejects_changed_frozen_source(tmp_path):
    manifest = inputs(tmp_path)
    source = benchmark.prep.load_inputs(manifest)["reports"][0]
    Path(source["path"]).write_bytes(b"changed after freezing")
    result = benchmark.run_report(source, tmp_path / "attempt", tmp_path / "cache", 10)
    assert result["status"] == "failed"
    assert "Source changed" in result["error"]
    assert result["subprocess_counts"]["convert"] == 0


def test_actual_worker_launch_failure_remains_a_failed_receipt(tmp_path, monkeypatch):
    manifest = inputs(tmp_path)
    source = benchmark.prep.load_inputs(manifest)["reports"][0]
    # Exercise a real failed OS launch; no fake elapsed-time or model receipt.
    monkeypatch.setattr(benchmark.sys, "executable", str(tmp_path / "missing-python"))
    result = benchmark.run_report(source, tmp_path / "attempt", tmp_path / "cache", 10)
    assert result["status"] == "failed"
    assert result["error_type"] == "FileNotFoundError"
    assert not result["audit_complete"]
    assert result["wall_seconds"] > 0
    assert (tmp_path / "attempt/observed-receipt.json").is_file()


def test_worker_rejects_source_change_between_initial_hash_and_binding(tmp_path):
    manifest = inputs(tmp_path)
    source = benchmark.prep.load_inputs(manifest)["reports"][0]
    replacement = tmp_path / "replacement.docx"
    doc = Document()
    doc.add_paragraph("Changed after the initial source hash")
    doc.save(replacement)
    directory = tmp_path / "attempt"
    directory.mkdir()
    request = directory / "request.json"
    benchmark.prep.save(request, {"source": source, "directory": str(directory),
                                 "cache_directory": str(tmp_path / "cache")})
    # Trigger a real file mutation at the entry import boundary, after the first
    # frozen hash check. This race fixture is separate from timing campaigns.
    code = """
import runpy, sys
from pathlib import Path
module = runpy.run_path(sys.argv[1])
replacement = Path(sys.argv[4]).read_bytes()
changed = False
def mutate(event, args):
    global changed
    if event == 'open' and isinstance(args[0], str) and not changed:
        name = Path(args[0]).name
        if name == 'run.py' or (name.startswith('run.') and name.endswith('.pyc')):
            changed = True
            Path(sys.argv[3]).write_bytes(replacement)
sys.addaudithook(mutate)
module['prepare_worker'](sys.argv[2])
assert changed
"""
    subprocess.run([sys.executable, "-B", "-c", code, str(SCRIPT), str(request), source["path"], str(replacement)],
                   check=True, timeout=10)
    record = json.loads((directory / "receipt.json").read_text())
    assert record["status"] == "failed"
    assert "changed during preparation" in record["error"]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="Real LibreOffice version probe required")
def test_real_cache_hit_still_rejects_unbound_conversion_metadata(tmp_path):
    # Deliberately crafted cache, not a performance observation. Version lookup,
    # cache lookup and source binding all run through the real production entry.
    source = tmp_path / "report.doc"
    source.write_bytes(b"source bytes for cache integrity fixture")
    source_sha = benchmark.prep.digest(source)
    version = subprocess.check_output(["soffice", "--version"], text=True, timeout=10).strip()
    cache = tmp_path / "cache"
    key = benchmark.prep.json_digest({"source_sha256": source_sha, "converter": version, "conversion": "docx"})
    target = cache / key / "report.docx"
    target.parent.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph("Cached output has a corrupted source binding")
    doc.save(target)
    benchmark.prep.save(target.parent / "conversion.json", {"source_sha256": "wrong-source",
        "converter": version, "conversion": "docx", "output_sha256": benchmark.prep.digest(target)})
    result = benchmark.run_report({"path": str(source), "sha256": source_sha}, tmp_path / "attempt", cache, 10)
    assert result["status"] == "failed"
    assert result["error_type"] == "ValueError"
    assert "DOC conversion evidence does not match source" in result["error"]
    assert result["subprocess_counts"] == {"version": 1, "convert": 0, "other": 0}
