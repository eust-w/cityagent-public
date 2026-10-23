#!/usr/bin/env python3
"""Two real SDK workflows, actual model requests and isolated output receipts."""

import argparse
import importlib.util
import json
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from bridge_evidence.contest.backends import (
    AccountedBackend,
    APIBackend,
    ModelManifest,
    normalize_api_url,
)
from bridge_evidence.contest.rsi_pipeline import parallel_runtime_digest
from bridge_evidence.contest.rsi_runtime import Dispatcher, WorkLock
from bridge_evidence.contest.run_state import (
    atomic_json,
    bind_work_directory,
    build_run_identity,
    file_digest,
)
from bridge_evidence.contest.timing import TimingLedger


def create_backend(args, model, tokenizer):
    return APIBackend(args.url, model.served_model, tokenizer, timeout=180)


def build_parser():
    parser = argparse.ArgumentParser()
    parser.add_argument("--model-manifest", type=Path, required=True)
    parser.add_argument("--sdk-acceptance", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--url", type=normalize_api_url, default="http://127.0.0.1:8000/v1")
    parser.add_argument("--workers", type=int, choices=(2, 4), default=2)
    return parser


def create_workflow_runtime(entry, args, inputs, work, backend):
    runtime = entry.build_parser().parse_args([
        "--input-dir", str(inputs),
        "--result-dir", str(args.output_dir / "result"),
        "--work-dir", str(work),
        "--engine", "tongagents",
        "--max-new-tokens", "1024",
        "--backend-url", backend.url,
    ])
    runtime.sdk_acceptance = args.sdk_acceptance
    return runtime


def main():
    parser = build_parser()
    args = parser.parse_args()
    if not str(args.output_dir.resolve()).startswith("/workspace/"):
        parser.error("official /workspace required")
    args.output_dir.mkdir(parents=True, exist_ok=False)
    from concurrent.futures import ThreadPoolExecutor

    from docx import Document
    from transformers import AutoTokenizer

    model = ModelManifest.load(args.model_manifest).validate()
    model_identity = model.identity()
    inputs = args.output_dir / "input"
    inputs.mkdir()
    paths = []
    for i in range(args.workers):
        document = Document()
        document.add_paragraph(
            f"并行验收桥{i}，2018年定期检测。总体评分90分，技术状况等级2类。"
        )
        document.add_paragraph("主梁存在裂缝，建议封闭处理。")
        path = inputs / f"parallel-{i}.docx"
        document.save(path)
        paths.append(path)
    spec = importlib.util.spec_from_file_location(
        "contest_parallel_probe", ROOT / "run.py"
    )
    entry = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(entry)
    tokenizer = AutoTokenizer.from_pretrained(
        model.tokenizer, local_files_only=True, trust_remote_code=True
    )
    backend = create_backend(args, model, tokenizer)
    identity = build_run_identity(
        Path(model.base),
        ROOT,
        {"purpose": "parallel-sdk-acceptance", "model": model_identity, "url": backend.url},
        model_files=model_identity["files"]["base"],
    )
    work = args.output_dir / "work"
    with WorkLock(work):
        fingerprint = bind_work_directory(work, identity, resume=False)
        ledger = TimingLedger(work / "timing.sqlite")
        generate = AccountedBackend(backend, ledger)
        runtime = create_workflow_runtime(entry, args, inputs, work, backend)
        runtime.result_dir.mkdir()
        intervals = {}
        with (
            Dispatcher(generate, capacity=args.workers) as dispatcher,
            ThreadPoolExecutor(max_workers=args.workers) as pool,
        ):
            dispatcher.begin_foreground()

            def execute(path):
                start = time.monotonic()
                record = entry.process_document(
                    runtime, path, tokenizer, dispatcher, fingerprint=fingerprint
                )
                intervals[path.name] = (start, time.monotonic())
                return record

            records = list(pool.map(execute, paths))
            dispatcher.end_foreground()
        overlap = max(v[0] for v in intervals.values()) < min(
            v[1] for v in intervals.values()
        )
        if not overlap:
            raise ValueError("SDK workflows did not overlap")
    report = {
        "schema_version": 1,
        "url": backend.url,
        "sdk_acceptance_sha256": file_digest(args.sdk_acceptance),
        "runtime_sha256": file_digest(ROOT / "run.py"),
        "overlap": overlap,
        "protected_sha256": parallel_runtime_digest(),
        "receipts": [
            {
                "document_sha256": r["document"]["sha256"],
                "workflow": r["workflow"],
                "model_call_ids": [c["call_id"] for c in r["model_calls"]],
                "output_path": str(runtime.result_dir / r["output"]["name"]),
                "output_sha256": r["output"]["sha256"],
            }
            for r in records
        ],
    }
    atomic_json(args.output_dir / "parallel-acceptance.json", report)
    print(
        json.dumps(
            {
                "passed": True,
                "actual_requests": sum(len(r["model_calls"]) for r in records),
            }
        )
    )


if __name__ == "__main__":
    main()
