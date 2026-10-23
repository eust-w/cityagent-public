from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace
import pytest
from docx import Document

from bridge_evidence.contest import contest_batch_cli
from bridge_evidence.contest.contest_batch_cli import source_files, publish_report, prepare_inputs, validate_source_scope
from bridge_evidence.contest.document_delivery_artifacts import DocumentDeliveryArtifacts
from bridge_evidence.contest.document_delivery_validation import validate_delivery_docx
from bridge_evidence.contest.docx_writer import write_extraction_docx
from bridge_evidence.contest.run_state import json_digest, atomic_json
from bridge_evidence.contest.document_multiagent_batch import prepare_batch, file_sha
from test_document_multiagent_batch import runtime


def draft():
    return {"基本信息": {"桥梁名称": "合成桥", "上一次总体评分": "90", "总体结论": "原文结论"},
        "病害明细": [{"序号": "model-id", "病害描述": "左幅1号支座锈蚀"}],
        "维修建议": [{"序号": "", "维修建议": "原文建议"}]}


def test_clean_export_keeps_values_and_program_numbers_without_internal_notices(tmp_path):
    # Discover canonical column keys from the actual contract.
    from bridge_evidence.contest.docx_writer import DEFECT_FIELDS, ADVICE_FIELDS
    data = draft()
    data["病害明细"] = [{"序号": "model-id", DEFECT_FIELDS[-1]: "病害原文"}]
    data["维修建议"] = [{"序号": "", ADVICE_FIELDS[-1]: "建议原文"}]
    original = deepcopy(data)
    snapshot = {"version": 1, "draft": data, "draft_sha256": json_digest(data), "identity": {}}
    manager = DocumentDeliveryArtifacts(tmp_path / "exports", "doc-one", {"submission_format": True})
    receipt = manager.export(snapshot, incomplete=False)
    assert receipt["artifact_kind"] == "three_table_report"
    assert receipt["phase"] == "pending_agent_self_check"
    assert receipt["complete"] is False
    doc = Document(receipt["docx_path"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "原文结论" in text
    for forbidden in ["处理范围说明", "待自查预览", "status.json", "（有对比年度）"]:
        assert forbidden not in text
    assert [t.rows[1].cells[0].text for t in doc.tables[1:]] == ["1", "1"]
    assert data == original
    for view in ["basic", "advice", "defects"]:
        manager.read_page(snapshot, section=view, receipt=receipt)
    assert manager.matches_snapshot(snapshot, receipt)
    other = DocumentDeliveryArtifacts(tmp_path / "exports", "doc-one", {})
    assert not other.matches_snapshot(snapshot, receipt)
    doc.tables[0].rows[1].cells[1].text = "篡改"
    doc.save(receipt["docx_path"])
    with pytest.raises(ValueError):
        manager.read_page(snapshot, section="basic", receipt=receipt)


def test_input_name_collision_and_symlink_escape_rejected(tmp_path):
    root = tmp_path / "inputs"
    root.mkdir()
    (root / "原始名字（左幅）.doc").write_bytes(b"doc")
    assert source_files(root)[0].stem + ".docx" == "原始名字（左幅）.docx"
    (root / "原始名字（左幅）.docx").write_bytes(b"docx")
    with pytest.raises(ValueError, match="collide"):
        source_files(root)
    (root / "原始名字（左幅）.docx").unlink()
    outside = tmp_path / "outside.docx"
    outside.write_bytes(b"docx")
    (root / "alias.docx").symlink_to(outside)
    with pytest.raises(ValueError, match="symlink"):
        source_files(root)


def test_test_scope_needs_explicit_inference_declaration(tmp_path):
    path = runtime(tmp_path, count=1)
    value = json.loads(path.read_bytes())
    value["scope"] = "official-test"
    value["documents"][0]["scope"] = "official-test"
    atomic_json(path, value)
    with pytest.raises(ValueError):
        prepare_batch(path, file_sha(path))
    value["test_inference_only"] = True
    atomic_json(path, value)
    prepared = prepare_batch(path, file_sha(path))
    assert prepared["scope"] == prepared["documents"][0]["scope"] == "official-test"
    assert prepared["model_calls"] == 0


def test_failed_document_is_never_published_as_final(tmp_path):
    result = publish_report({}, {"complete": False, "status": "failed"}, tmp_path)
    assert result["published"] is False
    assert not list(tmp_path.iterdir())


def test_unsettled_success_claim_cannot_publish(tmp_path):
    with pytest.raises(ValueError, match="settled document generation"):
        publish_report({}, {"complete": True}, tmp_path)
    assert not list(tmp_path.iterdir())


@pytest.fixture
def publication_case(tmp_path, monkeypatch):
    # Exercise filesystem publication of actual DOCX bytes; native workflow
    # receipt validation has separate end-to-end batch coverage.
    source = tmp_path / "inspected.docx"
    doc = Document()
    doc.add_paragraph("已检查的合成报告")
    doc.save(source)
    receipt = {"scope": {"submission_format": True}, "docx_path": str(source),
               "file_sha256": file_sha(source)}
    monkeypatch.setattr(contest_batch_cli, "verify_current_delivery", lambda *a, **kw: deepcopy(receipt))
    document = {"official_original": {"path": str(tmp_path / "input" / "桥梁报告.doc")},
                "manifest_identity_sha256": "a" * 64}
    record = {"complete": True, "generation_settlement": {"settled": True},
              "run_directory": str(tmp_path / "run"), "result": {}}
    result_dir = tmp_path / "results"
    result_dir.mkdir()
    foreign = result_dir / ".pending-another-worker.docx"
    foreign.write_bytes(b"another worker's temporary file")
    return document, record, result_dir, source, foreign


@pytest.mark.parametrize("failure", ["write", "short-write", "close"])
def test_publication_write_or_close_failure_leaves_no_final_or_own_temporary(publication_case, monkeypatch, failure):
    document, record, result_dir, source, foreign = publication_case
    before = source.read_bytes()
    real_open = Path.open

    class BrokenStream:
        def __init__(self, stream):
            self.stream = stream
        def __enter__(self):
            return self
        def __getattr__(self, key):
            return getattr(self.stream, key)
        def write(self, raw):
            if failure in {"write", "short-write"}:
                written = self.stream.write(raw[:len(raw) // 2])
                if failure == "write":
                    raise OSError("synthetic interrupted write")
                return written
            return self.stream.write(raw)
        def __exit__(self, *_):
            self.stream.close()
            if failure == "close":
                raise OSError("synthetic close failure")
            return False

    def broken_open(path, mode="r", *args, **kwargs):
        stream = real_open(path, mode, *args, **kwargs)
        if path.parent == result_dir and path.name.startswith(".pending-") and mode == "xb":
            return BrokenStream(stream)
        return stream
    monkeypatch.setattr(Path, "open", broken_open)
    monkeypatch.setattr(contest_batch_cli, "publish_no_replace", lambda *args: pytest.fail("incomplete file published"))
    with pytest.raises(OSError):
        publish_report(document, record, result_dir)
    assert list(result_dir.iterdir()) == [foreign]
    assert foreign.read_bytes() == b"another worker's temporary file"
    assert source.read_bytes() == before


@pytest.mark.parametrize("kind", ["file", "symlink"])
def test_publication_never_overwrites_an_existing_target(publication_case, tmp_path, kind):
    document, record, result_dir, source, foreign = publication_case
    target = result_dir / "桥梁报告.docx"
    if kind == "file":
        target.write_bytes(b"previous result")
    else:
        other = tmp_path / "other-result.docx"
        other.write_bytes(b"previous result")
        target.symlink_to(other)
    with pytest.raises(FileExistsError):
        publish_report(document, record, result_dir)
    assert target.read_bytes() == b"previous result"
    assert target.is_symlink() is (kind == "symlink")
    assert set(result_dir.iterdir()) == {target, foreign}


def test_publication_commits_exact_closed_bytes_and_cleans_only_own_temporary(publication_case):
    document, record, result_dir, source, foreign = publication_case
    published = publish_report(document, record, result_dir)
    target = result_dir / "桥梁报告.docx"
    assert published == {"published": True, "output_name": target.name, "sha256": file_sha(source)}
    assert target.read_bytes() == source.read_bytes()
    assert set(result_dir.iterdir()) == {target, foreign}
    assert (result_dir.parent / ".cityagent-publish-locks").is_dir()


def test_temporary_name_collision_does_not_delete_unowned_file(publication_case, monkeypatch):
    document, record, result_dir, source, foreign = publication_case
    existing = result_dir / ".pending-collision.docx"
    existing.write_bytes(b"other process owns this exact temporary path")
    monkeypatch.setattr(contest_batch_cli, "uuid4", lambda: SimpleNamespace(hex="collision"))
    with pytest.raises(FileExistsError):
        publish_report(document, record, result_dir)
    assert existing.read_bytes() == b"other process owns this exact temporary path"
    assert set(result_dir.iterdir()) == {existing, foreign}


@pytest.mark.parametrize("projection", [False, True])
def test_input_preparation_preserves_images_unless_projection_is_explicit(tmp_path, monkeypatch, projection):
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    source = source_dir / "桥梁报告.docx"
    doc = Document()
    doc.add_paragraph("原始正文")
    doc.save(source)
    full = {"document_sha256": file_sha(source), "units": [{"id": "E1", "text": "原始正文"}],
            "images": [{"id": "I1", "source_path": "word/media/image1.png"}],
            "unverified_content": [{"id": "O1", "kind": "embedded_object"}]}
    def fake_parser(converted, output, **kwargs):
        atomic_json(output / "materials.json", {"docx_sha256": file_sha(converted)})
    monkeypatch.setattr(contest_batch_cli, "prepare_document", fake_parser)
    monkeypatch.setattr(contest_batch_cli, "build_manifest", lambda converted: deepcopy(full))
    kwargs = {"text_table_projection": True} if projection else {}
    runtime_path = prepare_inputs(source_dir, tmp_path / "prepared", "official-training", "unused", **kwargs)
    value = json.loads(runtime_path.read_bytes())
    document = value["documents"][0]
    actual = json.loads(Path(document["manifest"]["path"]).read_bytes())
    original = json.loads(Path(document["full_source_inventory"]["path"]).read_bytes())
    assert original == full
    assert actual == ({**full, "images": [], "unverified_content": []} if projection else full)
    assert value["text_table_projection"] is projection
    assert document["source_scope"]["projection_explicitly_requested"] is projection
    scopes = validate_source_scope(value, text_table_projection=projection)
    assert scopes[document["document_id"]]["excluded_images"] == int(projection)
    assert scopes[document["document_id"]]["whole_original_report_verified"] is False
    if projection:
        # A prepared runtime cannot silently carry authorization to a new run.
        with pytest.raises(ValueError, match="--text-table-projection"):
            validate_source_scope(value)


def test_projection_switch_cannot_authorize_text_changes_or_false_exclusion_receipt(tmp_path):
    source = runtime(tmp_path, 1)
    prepared = prepare_batch(source, file_sha(source))
    document = prepared["documents"][0]
    scopes = validate_source_scope(prepared)
    assert scopes[document["document_id"]]["mode"] == "full_document"
    manifest_path = Path(document["manifest"]["path"])
    original = json.loads(manifest_path.read_bytes())
    changed = deepcopy(original)
    changed["units"][0]["text"] = "改写原文"
    atomic_json(manifest_path, changed)
    with pytest.raises(ValueError, match="excludes or changes"):
        validate_source_scope(prepared, text_table_projection=True)
    atomic_json(manifest_path, original)
    atomic_json(Path(document["excluded_content"]["path"]), {"images": [{"id": "invented"}]})
    with pytest.raises(ValueError, match="Excluded-content receipt differs"):
        validate_source_scope(prepared, text_table_projection=True)
