from __future__ import annotations

from tempfile import TemporaryDirectory
from pathlib import Path
import unittest

from docx import Document
from docx.oxml.ns import qn

from bridge_evidence.models.schema import Disease, DocumentState, Recommendation
from bridge_evidence.rendering.docx_report import DocxReportRenderer


class DocxReportRendererTest(unittest.TestCase):
    def test_renderer_declares_system_cjk_font(self) -> None:
        state = DocumentState(document_id="sha256:font", source_path="test.docx")
        state.bridge.name.value = "中文桥梁"
        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "report.docx"
            DocxReportRenderer().render(state, output)
            document = Document(output)

        normal = document.styles["Normal"]
        self.assertEqual(normal._element.rPr.rFonts.get(qn("w:eastAsia")), "Arial Unicode MS")

    def test_merged_ledgers_use_continuous_display_sequence(self) -> None:
        state = DocumentState(document_id="sha256:test", source_path="test.docx")
        state.recommendations = [
            Recommendation(sequence=7, category="尽快维修", content="修补裂缝"),
            Recommendation(sequence=1, category="预防性养护", content="定期巡查"),
        ]
        state.diseases = [
            Disease(sequence="95", component="主梁", type="裂缝", description="左幅裂缝"),
            Disease(sequence="1", component="横向联系", type="裂缝", description="右幅裂缝"),
        ]

        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "report.docx"
            DocxReportRenderer().render(state, output)
            document = Document(output)

        recommendation_numbers = [row.cells[0].text for row in document.tables[1].rows[1:]]
        disease_numbers = [row.cells[0].text for row in document.tables[2].rows[1:]]
        self.assertEqual(recommendation_numbers, ["1", "2"])
        self.assertEqual(disease_numbers, ["1", "2"])

    def test_disease_table_keeps_exact_location_and_repeated_history(self) -> None:
        state = DocumentState(document_id="sha256:location", source_path="test.docx")
        state.diseases = [
            Disease(
                sequence="1",
                component="主梁",
                location="右幅4#跨4#板距4#墩0.5m处梁底",
                type="剥落",
                description="梁底混凝土剥落",
                previous_state="上一次已发现",
                development="无明显发展",
            ),
            Disease(
                sequence="2",
                component="主梁",
                location="右幅4#跨5#板板间",
                type="渗水泛碱",
                description="板间渗水泛碱",
                previous_state="上一次已发现",
                development="无明显发展",
            ),
        ]

        with TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "report.docx"
            DocxReportRenderer().render(state, output)
            document = Document(output)

        rows = document.tables[2].rows[1:]
        self.assertEqual(rows[0].cells[1].text, "右幅4#跨4#板距4#墩0.5m处梁底")
        self.assertEqual(rows[1].cells[5].text, "上一次已发现")
        self.assertEqual(rows[1].cells[6].text, "无明显发展")


if __name__ == "__main__":
    unittest.main()
