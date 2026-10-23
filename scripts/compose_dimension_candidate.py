#!/usr/bin/env python3
"""Compose a conservative IAIC candidate along independent score dimensions.

The official feedback shows that summary reasoning and detail temporal logic can
improve while free rewriting of already extracted detail facts regresses text
consistency.  This release helper therefore preserves the candidate document as
the base, keeps its summary and temporal decisions, and restores only the
evidence-aligned detail wording and recommendation ledger from a stronger
baseline.  Disease rows are joined by their source table/row evidence anchor,
never by bridge name, filename pattern, or row position alone.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from docx import Document


CURRENT_ANCHOR = re.compile(r"T\d+R\d+")


def _files(directory: Path, suffix: str) -> dict[str, Path]:
    files = {path.name: path for path in sorted(directory.glob(f"*{suffix}"))}
    if len(files) != len(list(directory.glob(f"*{suffix}"))):
        raise ValueError(f"目录中存在重名文件: {directory}")
    return files


def _anchor(disease: dict[str, Any]) -> str | None:
    return next(
        (
            str(item)
            for item in disease.get("evidence") or []
            if CURRENT_ANCHOR.fullmatch(str(item))
        ),
        None,
    )


def _copy_row(source, target, columns: range | None = None) -> int:
    indices = columns if columns is not None else range(len(source.cells))
    changed = 0
    for index in indices:
        before = target.cells[index].text
        after = source.cells[index].text
        if before != after:
            target.cells[index].text = after
            changed += 1
    return changed


def compose(
    *,
    baseline_result: Path,
    baseline_state: Path,
    candidate_result: Path,
    candidate_state: Path,
    output_result: Path,
) -> dict[str, Any]:
    baseline_documents = _files(baseline_result, ".docx")
    candidate_documents = _files(candidate_result, ".docx")
    baseline_states = _files(baseline_state, ".json")
    candidate_states = _files(candidate_state, ".json")
    expected = set(candidate_documents)
    if set(baseline_documents) != expected:
        raise ValueError("基线与候选 DOCX 文件集合不一致")
    expected_states = {str(Path(name).with_suffix(".json")) for name in expected}
    if set(baseline_states) != expected_states or set(candidate_states) != expected_states:
        raise ValueError("基线或候选状态文件集合与 DOCX 不一致")

    output_result.mkdir(parents=True, exist_ok=True)
    report: dict[str, Any] = {
        "documents": 0,
        "summary_recommendation_count_cells_restored": 0,
        "recommendation_cells_restored": 0,
        "detail_fact_cells_restored": 0,
        "candidate_temporal_cells_preserved": 0,
        "disease_rows_joined_by_source_anchor": 0,
        "candidate_only_disease_rows": 0,
        "baseline_only_inactive_rows": 0,
        "errors": [],
    }

    for name in sorted(expected):
        state_name = str(Path(name).with_suffix(".json"))
        baseline_payload = json.loads(baseline_states[state_name].read_text(encoding="utf-8"))
        candidate_payload = json.loads(candidate_states[state_name].read_text(encoding="utf-8"))
        baseline_diseases = baseline_payload.get("diseases") or []
        candidate_diseases = candidate_payload.get("diseases") or []
        baseline_by_anchor = {
            anchor: index
            for index, disease in enumerate(baseline_diseases)
            if (anchor := _anchor(disease)) is not None
        }
        candidate_anchors = [_anchor(disease) for disease in candidate_diseases]
        if len(baseline_by_anchor) != len(baseline_diseases):
            report["errors"].append(f"{name}: 基线病害缺少或重复当前证据锚点")
            continue
        if any(anchor is None for anchor in candidate_anchors) or len(set(candidate_anchors)) != len(candidate_anchors):
            report["errors"].append(f"{name}: 候选病害缺少或重复当前证据锚点")
            continue

        baseline_doc = Document(baseline_documents[name])
        candidate_doc = Document(candidate_documents[name])
        if len(baseline_doc.tables) != 3 or len(candidate_doc.tables) != 3:
            report["errors"].append(f"{name}: 结果文档不是三表结构")
            continue

        # Row 16 is the recommendation-category count and must follow the
        # restored baseline recommendation ledger.  All other summary rows,
        # including candidate LLM narratives, remain untouched.
        baseline_summary = baseline_doc.tables[0]
        candidate_summary = candidate_doc.tables[0]
        if len(baseline_summary.rows) != 17 or len(candidate_summary.rows) != 17:
            report["errors"].append(f"{name}: 简要信息表结构异常")
            continue
        report["summary_recommendation_count_cells_restored"] += _copy_row(
            baseline_summary.rows[16], candidate_summary.rows[16], range(3)
        )

        baseline_recommendations = baseline_doc.tables[1]
        candidate_recommendations = candidate_doc.tables[1]
        if (
            len(baseline_recommendations.rows) != len(candidate_recommendations.rows)
            or any(
                len(left.cells) != len(right.cells)
                for left, right in zip(baseline_recommendations.rows, candidate_recommendations.rows)
            )
        ):
            report["errors"].append(f"{name}: 建议表结构与基线不一致")
            continue
        for source_row, target_row in zip(
            baseline_recommendations.rows[1:], candidate_recommendations.rows[1:]
        ):
            report["recommendation_cells_restored"] += _copy_row(source_row, target_row)

        baseline_details = baseline_doc.tables[2]
        candidate_details = candidate_doc.tables[2]
        if len(baseline_details.rows) != len(baseline_diseases) + 1:
            report["errors"].append(f"{name}: 基线状态与病害表行数不一致")
            continue
        if len(candidate_details.rows) != len(candidate_diseases) + 1:
            report["errors"].append(f"{name}: 候选状态与病害表行数不一致")
            continue

        for candidate_index, anchor in enumerate(candidate_anchors):
            assert anchor is not None
            baseline_index = baseline_by_anchor.get(anchor)
            if baseline_index is None:
                report["candidate_only_disease_rows"] += 1
                continue
            source_row = baseline_details.rows[baseline_index + 1]
            target_row = candidate_details.rows[candidate_index + 1]
            # Preserve the candidate's continuous display sequence and all
            # three temporal columns.  Restore only location/type/description.
            report["detail_fact_cells_restored"] += _copy_row(
                source_row, target_row, range(1, 4)
            )
            report["candidate_temporal_cells_preserved"] += 3
            report["disease_rows_joined_by_source_anchor"] += 1

        report["baseline_only_inactive_rows"] += len(
            set(baseline_by_anchor) - {str(item) for item in candidate_anchors}
        )
        candidate_doc.save(output_result / name)
        report["documents"] += 1

    report["passed"] = not report["errors"] and report["documents"] == len(expected)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="按官方五维指标组合 IAIC 稳健候选")
    parser.add_argument("--baseline-result", type=Path, required=True)
    parser.add_argument("--baseline-state", type=Path, required=True)
    parser.add_argument("--candidate-result", type=Path, required=True)
    parser.add_argument("--candidate-state", type=Path, required=True)
    parser.add_argument("--output-result", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = compose(
        baseline_result=args.baseline_result,
        baseline_state=args.baseline_state,
        candidate_result=args.candidate_result,
        candidate_state=args.candidate_state,
        output_result=args.output_result,
    )
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    raise SystemExit(0 if report["passed"] else 2)


if __name__ == "__main__":
    main()
