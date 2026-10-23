from __future__ import annotations

import importlib.util
from pathlib import Path

from docx import Document


SCRIPT = Path(__file__).parents[1] / "scripts" / "compose_table_ensemble.py"
SPEC = importlib.util.spec_from_file_location("compose_table_ensemble", SCRIPT)
assert SPEC is not None and SPEC.loader is not None
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)
compose = MODULE.compose


def _write_document(path: Path, prefix: str) -> None:
    document = Document()
    document.add_paragraph(f"{prefix}-opening")
    for table_index in range(3):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = f"{prefix}-table-{table_index}-header"
        table.cell(1, 1).text = f"{prefix}-table-{table_index}-body"
        document.add_paragraph(f"{prefix}-after-{table_index}")
    document.save(path)


def test_compose_uses_summary_table_and_complete_detail_body(tmp_path: Path) -> None:
    summary_dir = tmp_path / "summary"
    detail_dir = tmp_path / "detail"
    output_dir = tmp_path / "output"
    summary_dir.mkdir()
    detail_dir.mkdir()
    _write_document(summary_dir / "sample.docx", "summary")
    _write_document(detail_dir / "sample.docx", "detail")

    report = compose(
        summary_result=summary_dir,
        detail_result=detail_dir,
        output_result=output_dir,
    )

    assert report["passed"] is True
    assert report["summary_tables_copied"] == 1
    assert report["detail_tables_preserved"] == 2
    output = Document(output_dir / "sample.docx")
    assert output.tables[0].cell(1, 1).text == "summary-table-0-body"
    assert output.tables[1].cell(1, 1).text == "detail-table-1-body"
    assert output.tables[2].cell(1, 1).text == "detail-table-2-body"
    assert [paragraph.text for paragraph in output.paragraphs] == [
        "detail-opening",
        "detail-after-0",
        "detail-after-1",
        "detail-after-2",
    ]
