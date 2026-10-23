from __future__ import annotations

import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
import zipfile
from dataclasses import asdict, dataclass, field, replace
from pathlib import Path
from xml.etree import ElementTree

from docx import Document

from .word_quality import (
    DeepAuditResult,
    LegacyDocPreflight,
    WordQualityGate,
    fingerprint_docx,
)

CONVERSION_POLICY_VERSION = "5"
MIN_REFERENCE_CHARS = 64
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CONTENT_PART_RE = re.compile(
    r"^word/(?:document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$"
)
NUMBER_RE = re.compile(r"\d+(?:[.+-]\d+)*%?")


@dataclass(frozen=True)
class DocxStats:
    text: str
    paragraphs: int
    tables: int
    images: int


@dataclass(frozen=True)
class ConversionAudit:
    policy_version: str
    converter: str
    status: str
    accepted: bool
    confidence: float
    requires_review: bool
    source_reference: str | None
    source_text_chars: int | None
    converted_text_chars: int
    text_coverage: float | None
    numeric_coverage: float | None
    length_ratio: float | None
    paragraphs: int
    tables: int
    images: int
    warnings: tuple[str, ...] = ()
    preflight: dict[str, object] = field(default_factory=dict)
    environment_fingerprint: str | None = None
    environment: dict[str, str] = field(default_factory=dict)
    quality_axes: dict[str, object] = field(default_factory=dict)

    def to_dict(self) -> dict[str, object]:
        return asdict(self)

    @classmethod
    def from_dict(cls, value: dict[str, object]) -> ConversionAudit:
        data = dict(value)
        data["warnings"] = tuple(str(item) for item in data.get("warnings", ()))
        return cls(**data)  # type: ignore[arg-type]


@dataclass(frozen=True)
class AdaptedDocument:
    source_path: Path
    docx_path: Path
    sha256: str
    converted: bool
    audit: ConversionAudit


class FormatAdapter:
    """Adapt Word inputs while refusing silent DOC conversion loss."""

    def __init__(
        self,
        work_dir: Path,
        *,
        quality_gate: WordQualityGate | None = None,
    ):
        self.work_dir = work_dir
        self.convert_dir = work_dir / "converted"
        self.convert_dir.mkdir(parents=True, exist_ok=True)
        self._quality_gate = quality_gate

    @staticmethod
    def digest(path: Path) -> str:
        h = hashlib.sha256()
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()

    def adapt(self, path: Path) -> AdaptedDocument:
        path = path.resolve()
        digest = self.digest(path)
        suffix = path.suffix.lower()
        if suffix == ".docx":
            stats = self._inspect_docx(path)
            audit = self._native_audit(stats)
            return AdaptedDocument(path, path, digest, False, audit)
        if suffix != ".doc":
            raise ValueError(f"Unsupported input format: {path}")

        quality_gate = self._quality_gate
        if quality_gate is None:
            # Validate the binary before discovering or invoking optional tools.
            # Malformed or encrypted inputs must fail even without LibreOffice.
            preflight = WordQualityGate.preflight(path)
            quality_gate = WordQualityGate()
            self._quality_gate = quality_gate
        else:
            preflight = quality_gate.preflight(path)
        environment_fingerprint = quality_gate.cache_fingerprint
        # The policy version prevents an older text-only wvWare result from
        # being mistaken for a newly audited, high-fidelity conversion.
        target = self.convert_dir / (
            f"v{CONVERSION_POLICY_VERSION}-{digest[:16]}-"
            f"{environment_fingerprint[:12]}-{path.stem}.docx"
        )
        cached = self._load_cached(target, digest, environment_fingerprint)
        if cached is not None:
            return AdaptedDocument(path, target, digest, True, cached)

        attempts: list[str] = []
        candidates: list[tuple[Path, ConversionAudit]] = []
        with tempfile.TemporaryDirectory(prefix="bridge-evidence-") as raw_temp_dir:
            temp_dir = Path(raw_temp_dir)
            source_text, source_reference = self._extract_legacy_text(path, temp_dir)

            try:
                converted = self._convert_with_libreoffice(path, temp_dir)
                audit = self._audit_candidate(
                    converted,
                    converter="libreoffice",
                    source_text=source_text,
                    source_reference=source_reference,
                )
                if audit.accepted:
                    audit = self._apply_deep_audit(
                        audit,
                        quality_gate.assess(
                            path,
                            converted,
                            temp_dir / "libreoffice-quality",
                        ),
                        preflight,
                        quality_gate,
                    )
                if audit.accepted:
                    candidates.append((converted, audit))
                else:
                    attempts.append(self._rejection_summary("LibreOffice", audit))
            except (OSError, RuntimeError, ValueError, zipfile.BadZipFile, subprocess.TimeoutExpired) as exc:
                attempts.append(f"LibreOffice: {exc}")

            # A verified LibreOffice result is the canonical path. Run the
            # independent fallback only when conversion failed or its audit is
            # weak enough that another candidate can materially improve it.
            if not candidates or candidates[0][1].confidence < 0.85:
                try:
                    converted = self._convert_with_wv(path, temp_dir)
                    if converted is None:
                        attempts.append("wvWare/Pandoc: unavailable")
                    else:
                        audit = self._audit_candidate(
                            converted,
                            converter="wvware-pandoc",
                            source_text=source_text,
                            source_reference=source_reference,
                        )
                        if audit.accepted:
                            audit = self._apply_deep_audit(
                                audit,
                                quality_gate.assess(
                                    path,
                                    converted,
                                    temp_dir / "wvware-quality",
                                ),
                                preflight,
                                quality_gate,
                            )
                        if audit.accepted:
                            candidates.append((converted, audit))
                        else:
                            attempts.append(self._rejection_summary("wvWare/Pandoc", audit))
                except (OSError, RuntimeError, ValueError, zipfile.BadZipFile, subprocess.TimeoutExpired) as exc:
                    attempts.append(f"wvWare/Pandoc: {exc}")

            if not candidates:
                detail = "; ".join(attempts) or "no converter produced a candidate"
                raise RuntimeError(
                    f"DOC conversion failed fidelity gate for {path.name}: {detail}"
                )

            converted, audit = max(candidates, key=lambda item: item[1].confidence)
            self._store_cached(converted, target, digest, audit)

        return AdaptedDocument(path, target, digest, True, audit)

    def _convert_with_libreoffice(self, path: Path, temp_dir: Path) -> Path:
        soffice = self._soffice()
        profile_dir = temp_dir / "libreoffice-profile"
        proc = subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--headless",
                "--convert-to",
                "docx:Office Open XML Text",
                "--outdir",
                str(temp_dir),
                str(path),
            ],
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=1800,
            check=False,
        )
        converted = temp_dir / f"{path.stem}.docx"
        if proc.returncode or not converted.exists():
            raise RuntimeError(proc.stdout.strip() or f"exit code {proc.returncode}")
        return converted

    @staticmethod
    def _convert_with_wv(path: Path, temp_dir: Path) -> Path | None:
        wvhtml = shutil.which("wvHtml")
        pandoc = shutil.which("pandoc")
        if not wvhtml or not pandoc:
            return None
        html_name = "wv-converted.html"
        docx_name = "wv-converted.docx"
        extract = subprocess.run(
            [wvhtml, str(path), html_name],
            cwd=temp_dir,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=600,
            check=False,
        )
        html_path = temp_dir / html_name
        if extract.returncode or not html_path.exists():
            raise RuntimeError(extract.stdout.strip() or f"wvHtml exit code {extract.returncode}")

        # Do not strip images here. Scanned rating tables and annotated defect
        # photos are evidence; a text-only DOCX must never silently become the
        # canonical document merely because it converts faster.
        render = subprocess.run(
            [
                pandoc,
                html_name,
                "-f",
                "html",
                "-t",
                "docx",
                "--resource-path=.",
                "-o",
                docx_name,
            ],
            cwd=temp_dir,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=900,
            check=False,
        )
        docx_path = temp_dir / docx_name
        if render.returncode or not docx_path.exists():
            raise RuntimeError(render.stdout.strip() or f"Pandoc exit code {render.returncode}")
        return docx_path

    @staticmethod
    def _extract_legacy_text(path: Path, temp_dir: Path) -> tuple[str | None, str | None]:
        """Read DOC text through an independent parser for fidelity checks."""
        wvtext = shutil.which("wvText")
        if wvtext:
            output_path = temp_dir / "source-reference.txt"
            try:
                proc = subprocess.run(
                    [wvtext, str(path), str(output_path)],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    timeout=300,
                    check=False,
                )
                if proc.returncode == 0 and output_path.exists():
                    text = FormatAdapter._clean_reference_text(
                        output_path.read_text(encoding="utf-8", errors="replace")
                    )
                    if FormatAdapter._normalize_text(text):
                        return text, "wvText"
            except (OSError, subprocess.TimeoutExpired):
                pass

        antiword = shutil.which("antiword")
        if antiword:
            try:
                proc = subprocess.run(
                    [antiword, str(path)],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    timeout=300,
                    check=False,
                )
                if proc.returncode == 0:
                    text = FormatAdapter._clean_reference_text(
                        proc.stdout.decode("utf-8", errors="replace")
                    )
                    if FormatAdapter._normalize_text(text):
                        return text, "antiword"
            except (OSError, subprocess.TimeoutExpired):
                pass
        return None, None

    @classmethod
    def _audit_candidate(
        cls,
        path: Path,
        *,
        converter: str,
        source_text: str | None,
        source_reference: str | None,
    ) -> ConversionAudit:
        stats = cls._inspect_docx(path)
        converted = cls._normalize_text(stats.text)
        source = cls._normalize_text(source_text or "")
        warnings: list[str] = []
        accepted = True
        coverage: float | None = None
        numeric_coverage: float | None = None
        length_ratio: float | None = None

        if len(source) >= MIN_REFERENCE_CHARS:
            coverage = cls._shingle_coverage(source, converted)
            length_ratio = len(converted) / len(source) if source else None
            source_numbers = cls._number_tokens(source_text or "")
            converted_numbers = cls._number_tokens(stats.text)
            numeric_coverage = (
                len(source_numbers & converted_numbers) / len(source_numbers)
                if source_numbers
                else 1.0
            )

            if length_ratio < 0.8:
                warnings.append("converted_text_shorter_than_source")
            if coverage < 0.8:
                warnings.append("source_text_coverage_low")
            if len(source_numbers) >= 5 and numeric_coverage < 0.95:
                warnings.append("numeric_coverage_low")

            if length_ratio < 0.45 or coverage < 0.5:
                accepted = False
            if len(source_numbers) >= 5 and numeric_coverage < 0.75:
                accepted = False

            fidelity = (
                0.55 * coverage
                + 0.30 * numeric_coverage
                + 0.15 * min(length_ratio, 1.0)
            )
            status = "verified" if not warnings else "warning"
        else:
            warnings.append("independent_source_reference_unavailable")
            fidelity = 0.72
            status = "structural_only"

        if converter == "wvware-pandoc":
            warnings.append("fallback_converter_may_change_layout_or_media")
            fidelity *= 0.82

        confidence = round(max(0.0, min(1.0, fidelity)), 4)
        requires_review = bool(warnings) or confidence < 0.85
        return ConversionAudit(
            policy_version=CONVERSION_POLICY_VERSION,
            converter=converter,
            status=status if accepted else "rejected",
            accepted=accepted,
            confidence=confidence,
            requires_review=requires_review,
            source_reference=source_reference,
            source_text_chars=len(source) if source else None,
            converted_text_chars=len(converted),
            text_coverage=round(coverage, 4) if coverage is not None else None,
            numeric_coverage=(
                round(numeric_coverage, 4) if numeric_coverage is not None else None
            ),
            length_ratio=round(length_ratio, 4) if length_ratio is not None else None,
            paragraphs=stats.paragraphs,
            tables=stats.tables,
            images=stats.images,
            warnings=tuple(warnings),
        )

    @staticmethod
    def _native_audit(stats: DocxStats) -> ConversionAudit:
        return ConversionAudit(
            policy_version=CONVERSION_POLICY_VERSION,
            converter="native-ooxml",
            status="native",
            accepted=True,
            confidence=1.0,
            requires_review=False,
            source_reference=None,
            source_text_chars=None,
            converted_text_chars=len(FormatAdapter._normalize_text(stats.text)),
            text_coverage=None,
            numeric_coverage=None,
            length_ratio=None,
            paragraphs=stats.paragraphs,
            tables=stats.tables,
            images=stats.images,
        )

    @staticmethod
    def _apply_deep_audit(
        audit: ConversionAudit,
        deep: DeepAuditResult,
        preflight: LegacyDocPreflight,
        quality_gate: WordQualityGate,
    ) -> ConversionAudit:
        weighted_scores = [(audit.confidence, 0.50)]
        if deep.structure.score is not None:
            weighted_scores.append((deep.structure.score, 0.20))
        if deep.visual.score is not None:
            weighted_scores.append((deep.visual.score, 0.30))
        confidence = round(
            sum(score * weight for score, weight in weighted_scores)
            / sum(weight for _, weight in weighted_scores),
            4,
        )
        accepted = audit.accepted and deep.accepted
        warnings = tuple(dict.fromkeys((*audit.warnings, *deep.warnings)))
        if not accepted:
            status = "rejected"
        elif warnings:
            status = "warning"
        else:
            status = "verified"
        return replace(
            audit,
            status=status,
            accepted=accepted,
            confidence=confidence,
            requires_review=(
                audit.requires_review
                or deep.requires_review
                or confidence < 0.85
            ),
            warnings=warnings,
            preflight=preflight.to_dict(),
            environment_fingerprint=quality_gate.cache_fingerprint,
            environment=dict(getattr(quality_gate, "environment", {})),
            quality_axes=deep.to_dict(),
        )

    @staticmethod
    def _inspect_docx(path: Path) -> DocxStats:
        # Enforce bounded ZIP/XML parsing before python-docx resolves package
        # relationships. The returned fingerprint is reused by the deep gate
        # for converted DOC inputs and acts as a security preflight for native
        # DOCX inputs.
        fingerprint_docx(path)
        required_parts = {"[Content_Types].xml", "word/document.xml"}
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
            if bad_member:
                raise ValueError(f"DOCX ZIP member failed CRC: {bad_member}")
            names = set(archive.namelist())
            missing = required_parts - names
            if missing:
                raise ValueError(f"DOCX is missing required OOXML parts: {sorted(missing)}")

            text_parts: list[str] = []
            paragraphs = 0
            tables = 0
            for name in sorted(names):
                if not CONTENT_PART_RE.match(name):
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError as exc:
                    raise ValueError(f"Invalid OOXML in {name}: {exc}") from exc
                text_parts.extend(
                    element.text or ""
                    for element in root.iter(f"{{{WORD_NAMESPACE}}}t")
                )
                paragraphs += sum(1 for _ in root.iter(f"{{{WORD_NAMESPACE}}}p"))
                tables += sum(1 for _ in root.iter(f"{{{WORD_NAMESPACE}}}tbl"))
            images = sum(
                1 for name in names if name.startswith("word/media/") and not name.endswith("/")
            )

        # ZIP/XML validity is insufficient: python-docx must also be able to
        # resolve the package relationships used by the downstream parser.
        Document(path)
        if not paragraphs and not tables and not images:
            raise ValueError(f"Converted document is empty: {path}")
        # Keep a boundary between OOXML text nodes. This prevents adjacent
        # table cells such as 15 | 02 from becoming the artificial number 1502
        # during the independent numeric-fidelity check.
        return DocxStats("\n".join(text_parts), paragraphs, tables, images)

    @staticmethod
    def _clean_reference_text(text: str) -> str:
        """Remove extractor markup that is not visible Word document text."""
        cleaned = html.unescape(text)
        cleaned = re.sub(r"</?[A-Za-z][^>]*>", " ", cleaned)
        return "".join(
            " " if unicodedata.category(character) in {"Cc", "Cf"} else character
            for character in cleaned
        )

    @staticmethod
    def _normalize_text(text: str) -> str:
        normalized = unicodedata.normalize("NFKC", text)
        return "".join(
            character
            for character in normalized
            if not character.isspace()
            and unicodedata.category(character) not in {"Cc", "Cf"}
        )

    @staticmethod
    def _number_tokens(text: str) -> set[str]:
        normalized = unicodedata.normalize("NFKC", text)
        return set(NUMBER_RE.findall(normalized))

    @staticmethod
    def _shingle_coverage(source: str, converted: str, width: int = 3) -> float:
        if not source:
            return 1.0
        if len(source) < width:
            return 1.0 if source in converted else 0.0
        source_shingles = {source[index : index + width] for index in range(len(source) - width + 1)}
        if not source_shingles:
            return 1.0
        converted_shingles = {
            converted[index : index + width]
            for index in range(max(0, len(converted) - width + 1))
        }
        return len(source_shingles & converted_shingles) / len(source_shingles)

    @staticmethod
    def _soffice() -> str:
        """Return a working LibreOffice executable, not a stale command wrapper."""
        candidates = [
            shutil.which("soffice"),
            shutil.which("libreoffice"),
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
            "/usr/bin/soffice",
            "/usr/bin/libreoffice",
            "/usr/local/bin/soffice",
            "/usr/local/bin/libreoffice",
        ]
        for candidate in candidates:
            if not candidate or not Path(candidate).exists():
                continue
            try:
                probe = subprocess.run(
                    [candidate, "--version"],
                    text=True,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    timeout=15,
                    check=False,
                )
            except (OSError, subprocess.TimeoutExpired):
                continue
            if probe.returncode == 0 and "LibreOffice" in probe.stdout:
                return candidate
        raise RuntimeError(
            "A working LibreOffice/soffice installation is required to convert .doc files "
            "without losing table structure"
        )

    def _load_cached(
        self,
        target: Path,
        source_digest: str,
        environment_fingerprint: str,
    ) -> ConversionAudit | None:
        manifest_path = self._manifest_path(target)
        if not target.exists() or not manifest_path.exists():
            return None
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            if manifest.get("policy_version") != CONVERSION_POLICY_VERSION:
                return None
            if manifest.get("source_sha256") != source_digest:
                return None
            if manifest.get("docx_sha256") != self.digest(target):
                return None
            audit = ConversionAudit.from_dict(manifest["audit"])
            if not audit.accepted:
                return None
            if audit.environment_fingerprint != environment_fingerprint:
                return None
            self._inspect_docx(target)
            return audit
        except (KeyError, TypeError, ValueError, OSError, json.JSONDecodeError, zipfile.BadZipFile):
            return None

    def _store_cached(
        self,
        converted: Path,
        target: Path,
        source_digest: str,
        audit: ConversionAudit,
    ) -> None:
        fd, raw_staged_path = tempfile.mkstemp(
            prefix=f".{target.stem}-",
            suffix=".docx",
            dir=self.convert_dir,
        )
        os.close(fd)
        staged_path = Path(raw_staged_path)
        try:
            shutil.copy2(converted, staged_path)
            self._inspect_docx(staged_path)
            docx_digest = self.digest(staged_path)
            os.replace(staged_path, target)
        finally:
            staged_path.unlink(missing_ok=True)

        manifest = {
            "policy_version": CONVERSION_POLICY_VERSION,
            "source_sha256": source_digest,
            "docx_sha256": docx_digest,
            "audit": audit.to_dict(),
        }
        manifest_path = self._manifest_path(target)
        fd, raw_manifest_path = tempfile.mkstemp(
            prefix=f".{manifest_path.stem}-",
            suffix=".json",
            dir=self.convert_dir,
        )
        os.close(fd)
        staged_manifest = Path(raw_manifest_path)
        try:
            staged_manifest.write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            os.replace(staged_manifest, manifest_path)
        finally:
            staged_manifest.unlink(missing_ok=True)

    @staticmethod
    def _manifest_path(target: Path) -> Path:
        return target.with_suffix(".conversion.json")

    @staticmethod
    def _rejection_summary(label: str, audit: ConversionAudit) -> str:
        warnings = ",".join(audit.warnings) or "quality below threshold"
        return f"{label}: rejected ({warnings}, confidence={audit.confidence:.4f})"
