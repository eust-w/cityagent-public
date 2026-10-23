"""Optional same-model raster stages share deadlines and retain actual evidence."""

import importlib.util
import json
from pathlib import Path

import pytest
from docx import Document
from test_chunked_integration import CharTokenizer, FragmentBackend

import run
from bridge_evidence.contest.candidate import validate_config
from bridge_evidence.contest.run_state import file_digest, json_digest


@pytest.fixture
def options(tmp_path):
    args = run.build_parser().parse_args(
        [
            "--engine",
            "direct",
            "--workflow",
            "chunked",
            "--backend",
            "vllm-api",
            "--work-dir",
            str(tmp_path / "work"),
            "--input-dir",
            str(tmp_path),
            "--result-dir",
            str(tmp_path / "result"),
        ]
    )
    args.candidate = validate_config(
        {"workflow": "chunked", "input_budget": 3000, "output_budget": 500}
    )
    args.model_identity = {
        "manifest": {"served_model": "fixed"},
        "files": {"base": {"config.json": "f" * 64}},
    }
    args.visual_model_identity_sha256 = json_digest(args.model_identity)
    args.visual_trial_id = "a" * 64
    args.visual_capability_binding = {
        "mode": "same-model",
        "model_identity_sha256": args.visual_model_identity_sha256,
    }
    args.visual_capability = tmp_path / "capability.json"
    args.visual_capability.write_text("{}")
    return args


def make_report(tmp_path):
    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("测试桥")
    doc.save(path)
    return path


def test_visual_is_off_by_default_and_native_prepare_unchanged(
    options, tmp_path, monkeypatch
):
    from bridge_evidence.contest import visual_evidence

    assert options.visual_evidence == "off"
    monkeypatch.setattr(
        visual_evidence,
        "plan_visual_jobs",
        lambda *a: pytest.fail("native path invoked visual planner"),
    )
    prepared = run.prepare_document(options, make_report(tmp_path), CharTokenizer())
    assert "visual_plan" not in prepared
    assert prepared["chunk_plan"].coverage(prepared["snapshot"])["complete"]


@pytest.mark.parametrize(
    "field,value",
    [
        ("workflow", "whole"),
        ("workflow", "divide-fixed"),
        ("backend", "local-vllm"),
        ("visual_capability", None),
    ],
)
def test_visual_invalid_route_is_rejected_before_model_work(options, field, value):
    options.visual_evidence = "same-model"
    setattr(options, field, value)
    with pytest.raises(ValueError):
        run.validate_visual_options(options)


def test_capability_is_revalidated_against_actual_fixed_model(options, monkeypatch):
    from bridge_evidence.contest import vision_capability

    calls = []
    options.visual_evidence = "same-model"

    def validate(path, identity):
        calls.append((path, identity))
        return {
            "mode": "same-model",
            "model_identity_sha256": json_digest(identity),
            "probe_sha256": "d" * 64,
        }

    monkeypatch.setattr(vision_capability, "validate_vision_capability", validate)
    binding = run.configure_visual_runtime(options, options.model_identity)
    assert calls == [(options.visual_capability, options.model_identity)]
    assert binding["mode"] == "same-model"
    assert binding["config"] == {"max_tokens": 2048}
    assert "trial_id" not in binding
    assert options.visual_model_identity_sha256 == json_digest(options.model_identity)


def test_cpu_preparation_binds_plan_to_trial_model_and_native_sources(
    options, tmp_path, monkeypatch
):
    from bridge_evidence.contest import visual_evidence

    options.visual_evidence = "same-model"
    path = make_report(tmp_path)
    seen = []

    def plan(source, native, identity, config):
        seen.append((source, native, identity, config))
        return {"plan_sha256": "v" * 64, "jobs": []}

    monkeypatch.setattr(visual_evidence, "plan_visual_jobs", plan)
    prepared = run.prepare_document(options, path, CharTokenizer())
    assert seen[0][0] == path
    assert seen[0][1] is prepared["snapshot"]
    assert seen[0][2] == options.visual_model_identity_sha256
    assert seen[0][3] == {"trial_id": options.visual_trial_id, "max_tokens": 2048}
    assert prepared["visual_plan"]["plan_sha256"] == "v" * 64
    assert prepared["chunk_plan"].coverage(prepared["snapshot"])["complete"]


def test_evaluator_parser_preserves_default_off():
    # The evaluator's explicit parser is also used by the diagnostic driver.
    spec = importlib.util.spec_from_file_location(
        "visual_evaluator", Path(run.ROOT) / "scripts/evaluate_multitask.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    args = module.build_parser().parse_args(
        [
            "--manifest",
            "m.json",
            "--model-manifest",
            "model.json",
            "--candidate-config",
            "c.json",
            "--output-dir",
            "out",
            "--hardware",
            "fixture",
        ]
    )
    assert args.visual_evidence == "off" and args.visual_capability is None


def mock_visual_stage(monkeypatch, prepared, *, fail_text=False, spent=3.0):
    from bridge_evidence.contest import (
        chunked_inference,
        vision_capability,
        visual_evidence,
    )

    monkeypatch.setattr(
        vision_capability,
        "validate_vision_capability",
        lambda path, identity: {
            "mode": "same-model",
            "model_identity_sha256": json_digest(identity),
        },
    )
    visual_call = {
        "call_id": "visual-1",
        "role": "visual-evidence",
        "status": "success",
        "seconds": 2.0,
    }
    raw = {"response": {"call_id": "visual-1", "text": "visual raw"}}
    seen = {}

    def execute(plan, generate, cache_dir, budget):
        seen["visual_cache"] = cache_dir
        budget({"prior_calls": [], "spent_seconds": 0.0})
        return {
            "calls": [visual_call],
            "raw": [raw],
            "evidence": [{"call_id": "visual-1"}],
            "errors": [],
            "spent_seconds": spent,
        }

    def text(snapshot, tokenizer, generate, **kwargs):
        from bridge_evidence.contest.rsi_runtime import REPORT_DEADLINE

        seen.update(kwargs)
        seen["deadline"] = REPORT_DEADLINE.get()
        if fail_text:
            from bridge_evidence.contest.inference_core import InferenceFailure

            raise InferenceFailure(
                [{"call_id": "text-1", "seconds": 4.0}], raw_output="", wall_seconds=4.0
            )
        return {
            "model_calls": [{"call_id": "text-1", "seconds": 4.0}],
            "seconds": 4.0,
            "generation_evidence": {
                "schema_version": 1,
                "responses": {"text-1": "text raw"},
            },
            "shared_budget": {"prefix": kwargs["prefix_budget"]},
        }

    monkeypatch.setattr(visual_evidence, "execute_visual_jobs", execute)
    monkeypatch.setattr(
        visual_evidence, "rebuild_visual_snapshot", lambda *args: prepared["snapshot"]
    )
    monkeypatch.setattr(chunked_inference, "infer_chunked", text)
    return seen


def visual_prepared(options, tmp_path):
    path = make_report(tmp_path)
    prepared = run.prepare_document(options, path, CharTokenizer())
    options.visual_evidence = "same-model"
    prepared["visual_plan"] = {
        "plan_sha256": "b" * 64,
        "jobs": [{"job_id": "j1"}],
        "identity": {
            "model_identity_sha256": options.visual_model_identity_sha256,
            "config": {"trial_id": options.visual_trial_id, "max_tokens": 2048},
        },
    }
    return path, prepared


def test_visual_and_text_share_original_deadline_call_cap_and_cache_identity(
    options, tmp_path, monkeypatch
):
    from bridge_evidence.contest.rsi_runtime import REPORT_DEADLINE

    path, prepared = visual_prepared(options, tmp_path)
    seen = mock_visual_stage(monkeypatch, prepared)
    record = {}
    result = run.infer_with_visual_evidence(
        options,
        prepared,
        CharTokenizer(),
        lambda *a: None,
        fingerprint=options.visual_trial_id,
        document={"sha256": file_digest(path), "relative_path": path.name},
        record=record,
    )
    assert result["seconds"] == 6.0
    assert [call["call_id"] for call in result["model_calls"]] == ["visual-1", "text-1"]
    assert result["generation_evidence"]["responses"] == {
        "visual-1": "visual raw",
        "text-1": "text raw",
    }
    assert result["visual_contributing_call_ids"] == ["visual-1"]
    prefix = seen["prefix_budget"]
    assert prefix["initial_text_tasks"] == len(prepared["chunk_plan"].tasks)
    assert (
        record["visual_evidence"]["native_plan_sha256"] == prepared["chunk_plan"].digest
    )
    assert (
        len(record["visual_evidence"]["native_chunk_plan"]["tasks"])
        == prefix["initial_text_tasks"]
    )
    assert prefix["initial_visual_jobs"] == 1
    assert prefix["total_call_limit"] == 4 * (len(prepared["chunk_plan"].tasks) + 1) + 2
    assert prefix["used_calls"] == 1 and prefix["elapsed_seconds"] >= 3.0
    assert seen["deadline"] is not None
    assert REPORT_DEADLINE.get() is None
    assert options.visual_trial_id in seen["visual_cache"].parts
    assert record["visual_evidence"]["capability_path"] == str(
        options.visual_capability.resolve()
    )


def test_text_failure_does_not_erase_prior_visual_calls_or_raw(
    options, tmp_path, monkeypatch
):
    path, prepared = visual_prepared(options, tmp_path)
    mock_visual_stage(monkeypatch, prepared, fail_text=True)
    record = {}
    with pytest.raises(RuntimeError):
        run.infer_with_visual_evidence(
            options,
            prepared,
            CharTokenizer(),
            lambda *a: None,
            fingerprint=options.visual_trial_id,
            document={"sha256": file_digest(path), "relative_path": path.name},
            record=record,
        )
    assert record["model_calls"][0]["call_id"] == "visual-1"
    assert record["generation_evidence"]["responses"]["visual-1"] == "visual raw"
    assert record["visual_evidence"]["calls"][0]["seconds"] == 2.0


def test_exhausted_visual_history_cannot_get_fresh_text_deadline(
    options, tmp_path, monkeypatch
):
    path, prepared = visual_prepared(options, tmp_path)
    seen = mock_visual_stage(monkeypatch, prepared, spent=1201.0)
    with pytest.raises(TimeoutError):
        run.infer_with_visual_evidence(
            options,
            prepared,
            CharTokenizer(),
            lambda *a: None,
            fingerprint=options.visual_trial_id,
            document={"sha256": file_digest(path), "relative_path": path.name},
            record={},
        )
    assert "deadline" not in seen


def test_actual_visual_plan_then_text_and_same_trial_resume(
    options, tmp_path, monkeypatch
):
    from test_contest_visual_evidence import document, response

    from bridge_evidence.contest import vision_capability

    monkeypatch.setattr(
        vision_capability,
        "validate_vision_capability",
        lambda path, identity: {
            "mode": "same-model",
            "model_identity_sha256": json_digest(identity),
        },
    )
    path = document(tmp_path, repeat=2)
    options.visual_evidence = "same-model"
    options.candidate = validate_config(
        {"workflow": "chunked", "input_budget": 8192, "output_budget": 500}
    )
    prepared = run.prepare_document(options, path, CharTokenizer())
    text = FragmentBackend()
    requests = []

    def generate(prompt, maximum):
        requests.append(prompt)
        return (
            response(prompt, maximum)
            if isinstance(prompt, dict)
            else text(prompt, maximum)
        )

    document_id = {"sha256": file_digest(path), "relative_path": path.name}
    first = run.infer_with_visual_evidence(
        options,
        prepared,
        CharTokenizer(),
        generate,
        fingerprint=options.visual_trial_id,
        document=document_id,
        record={},
    )
    assert len(first["visual_evidence"]["calls"]) == 1
    assert (
        len(
            [
                r
                for r in first["snapshot"]["regions"]
                if r["kind"] == "visual-transcription"
            ]
        )
        == 2
    )
    assert first["shared_budget"]["prefix"]["initial_visual_jobs"] == 1
    before = len(requests)
    second = run.infer_with_visual_evidence(
        options,
        prepared,
        CharTokenizer(),
        generate,
        fingerprint=options.visual_trial_id,
        document=document_id,
        record={},
    )
    assert len(requests) == before
    assert first["model_calls"] == second["model_calls"]
    assert first["seconds"] == second["seconds"]


def test_visual_rebuild_overhead_is_durable_and_never_refunded(
    options, tmp_path, monkeypatch
):
    from bridge_evidence.contest import visual_evidence

    path, prepared = visual_prepared(options, tmp_path)
    seen = mock_visual_stage(monkeypatch, prepared, spent=3.0)
    clock = [100.0]
    monkeypatch.setattr(run.time, "monotonic", lambda: clock[0])
    rebuilds = []

    def rebuild(*args):
        if not rebuilds:
            clock[0] += 10.0
        rebuilds.append(True)
        return prepared["snapshot"]

    monkeypatch.setattr(visual_evidence, "rebuild_visual_snapshot", rebuild)
    document_id = {"sha256": file_digest(path), "relative_path": path.name}
    first = run.infer_with_visual_evidence(
        options,
        prepared,
        CharTokenizer(),
        lambda *a: None,
        fingerprint=options.visual_trial_id,
        document=document_id,
        record={},
    )
    previous = first["shared_budget"]["prefix"]["elapsed_seconds"]
    run.infer_with_visual_evidence(
        options,
        prepared,
        CharTokenizer(),
        lambda *a: None,
        fingerprint=options.visual_trial_id,
        document=document_id,
        record={},
    )
    assert seen["prefix_budget"]["elapsed_seconds"] >= previous


def test_prepared_visual_plan_from_other_model_is_rejected(
    options, tmp_path, monkeypatch
):
    path, prepared = visual_prepared(options, tmp_path)
    mock_visual_stage(monkeypatch, prepared)
    prepared["visual_plan"]["identity"]["model_identity_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="model"):
        run.infer_with_visual_evidence(
            options,
            prepared,
            CharTokenizer(),
            lambda *a: pytest.fail("wrong model plan dispatched"),
            fingerprint=options.visual_trial_id,
            document={"sha256": file_digest(path), "relative_path": path.name},
            record={},
        )


def test_process_document_retains_visual_provenance_through_publication(
    options, tmp_path, monkeypatch
):
    from PIL import Image
    from test_contest_visual_evidence import response

    from bridge_evidence.contest import vision_capability

    image = tmp_path / "picture.png"
    Image.new("RGB", (20, 16), "green").save(image)
    path = tmp_path / "report.docx"
    doc = Document()
    doc.add_paragraph("测试桥")
    doc.add_paragraph().add_run().add_picture(str(image))
    doc.save(path)
    options.visual_evidence = "same-model"
    options.candidate = validate_config(
        {"workflow": "chunked", "input_budget": 8192, "output_budget": 500}
    )
    options.visual_capability.write_text(
        json.dumps({"model_identity": options.model_identity})
    )
    monkeypatch.setattr(
        vision_capability,
        "validate_vision_capability",
        lambda path, identity: {
            "mode": "same-model",
            "model_identity_sha256": json_digest(identity),
        },
    )
    options.result_dir.mkdir()
    text = FragmentBackend()

    def generate(prompt, maximum):
        return (
            response(prompt, maximum)
            if isinstance(prompt, dict)
            else text(prompt, maximum)
        )

    from bridge_evidence.contest.backends import AccountedBackend
    from bridge_evidence.contest.timing import TimingLedger

    # The fixture's model response helper declares a nominal latency. Use the
    # real accountant to measure this invocation and persist its actual receipt.
    accounted = AccountedBackend(
        generate, TimingLedger(options.work_dir / "timing.sqlite")
    )
    from bridge_evidence.contest.rsi_runtime import Dispatcher

    with Dispatcher(accounted, capacity=1) as dispatcher:
        dispatcher.begin_foreground()
        try:
            record = run.process_document(
                options,
                path,
                CharTokenizer(),
                dispatcher,
                fingerprint=options.visual_trial_id,
            )
        finally:
            dispatcher.end_foreground()
    assert record["status"] == "success"
    assert record["visual_contributing_call_ids"]
    assert record["shared_budget"]["prefix"]["used_calls"] == 1
    assert record["facts"]["evidence"]
    assert (options.result_dir / "report.docx").is_file()


def test_visual_nonce_is_fresh_but_resume_reuses_bound_nonce(options):
    options.visual_evidence = "same-model"
    first = run.visual_run_nonce(options)
    second = run.visual_run_nonce(options)
    assert first != second and len(first) == 32
    from bridge_evidence.contest.run_state import atomic_json

    atomic_json(
        options.work_dir / "run_manifest.json",
        {"identity": {"options": {"visual_trial_nonce": first}}},
    )
    options.resume = True
    assert run.visual_run_nonce(options) == first
