from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pytest
from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "scripts/benchmark_document_preparation.py"
spec = importlib.util.spec_from_file_location("preparation_benchmark", SCRIPT)
benchmark = importlib.util.module_from_spec(spec)
spec.loader.exec_module(benchmark)


def inputs(tmp_path):
    root = tmp_path / "training"
    root.mkdir()
    source = root / "report.docx"
    doc = Document()
    doc.add_paragraph("同一位置的真实重复原文")
    doc.add_paragraph("同一位置的真实重复原文")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "位置"
    table.cell(0, 1).text = "病害"
    table.cell(1, 0).text = "左幅 1 号梁"
    table.cell(1, 1).text = "裂缝"
    doc.save(source)
    manifest = tmp_path / "inputs.json"
    benchmark.save(manifest, {"scope": "official-training", "training_root": str(root),
                             "reports": [{"path": str(source), "sha256": benchmark.digest(source)}]})
    return source, manifest


def test_pinned_inputs_reject_change_and_test_directory(tmp_path):
    source, manifest = inputs(tmp_path)
    source.write_bytes(b"changed")
    with pytest.raises(ValueError, match="changed"):
        benchmark.load_inputs(manifest)
    payload = json.loads(manifest.read_text())
    test_root = tmp_path / "test"
    test_root.mkdir()
    payload["training_root"] = str(test_root)
    benchmark.save(manifest, payload)
    with pytest.raises(ValueError, match="test or label"):
        benchmark.load_inputs(manifest)


def test_real_parser_rounds_preserve_duplicate_rows_and_count_failure(tmp_path):
    _, manifest = inputs(tmp_path)
    payload = json.loads(manifest.read_text())
    invalid = Path(payload["training_root"]) / "broken.docx"
    invalid.write_bytes(b"not an OOXML archive")
    payload["reports"].append({"path": str(invalid), "sha256": benchmark.digest(invalid)})
    benchmark.save(manifest, payload)
    result = benchmark.benchmark(manifest, tmp_path / "runs", repeats=3, timeout=10)
    assert result["arms"]["1"]["attempted_reports"] == 6
    assert result["arms"]["2"]["failures"] == 3
    assert result["arms"]["1"]["failures"] == 3
    assert result["validated_cpu_speedup_ratio"] is None
    assert result["model_requests"] == 0
    assert not result["content_mismatches"]
    rounds = json.loads((tmp_path / "runs/rounds.json").read_text())
    for item in rounds:
        good, bad = item["reports"]
        assert good["counts"]["units"] == 4
        assert good["loaded_model_or_sdk_modules"] == []
        assert bad["status"] == "failed"
        assert bad["wall_seconds"] > 0


def test_content_digest_ignores_only_root_path_and_container_bytes():
    original = {"source_path": "/one/run.docx", "document_sha256": "zip-one",
                "units": [{"source_path": "word/document.xml#p[1]", "text": "完整原文"}],
                "images": [{"id": "I1", "verification": "unverified"}]}
    equivalent = {**original, "source_path": "/other/run.docx", "document_sha256": "zip-two"}
    parts = {"word/document.xml": "pinned-part-bytes"}
    assert benchmark.content_digest(original, "source", parts) == benchmark.content_digest(equivalent, "source", parts)
    assert benchmark.content_digest(original, "source", parts) != benchmark.content_digest({**original, "images": []}, "source", parts)


def test_content_mismatch_suppresses_speedup_claim():
    rounds = []
    for workers, content in [(1, "one"), (2, "different")]:
        rounds.append({"workers": workers, "round": 1, "wall_seconds": 1 / workers,
                       "reports": [{"status": "success", "source": {"sha256": "same"},
                                    "content_sha256": content}]})
    result = benchmark.summarize(rounds, 1)
    assert result["content_mismatches"]
    assert result["validated_cpu_speedup_ratio"] is None


def test_cpu_guard_rejects_real_network_attempt(tmp_path):
    import subprocess
    import sys
    code = ("import runpy,socket; m=runpy.run_path(" + repr(str(SCRIPT)) + "); "
            "m['install_cpu_guard'](); socket.create_connection(('127.0.0.1',1))")
    result = subprocess.run([sys.executable, "-B", "-c", code], capture_output=True, text=True, timeout=10, check=False)
    assert result.returncode != 0
    assert "Network access is forbidden" in result.stderr


@pytest.mark.parametrize("changed_part", ["word/media/image1.png", "word/embeddings/data.bin"])
def test_real_package_same_named_media_and_embedding_changes_break_equivalence(tmp_path, changed_part):
    import io
    import zipfile

    from lxml import etree
    from PIL import Image

    from bridge_evidence.contest.evidence_manifest import build_manifest

    source = tmp_path / "source.docx"
    picture = tmp_path / "image.png"
    Image.new("RGB", (8, 8), "red").save(picture)
    doc = Document()
    doc.add_paragraph("相同原文与相同图片路径")
    doc.add_picture(str(picture))
    doc.save(source)
    with zipfile.ZipFile(source) as archive:
        original_parts = {name: archive.read(name) for name in archive.namelist()}
    original_parts["word/embeddings/data.bin"] = b"original embedded bytes"
    types = etree.fromstring(original_parts["[Content_Types].xml"])
    etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                     PartName="/word/embeddings/data.bin", ContentType="application/vnd.openxmlformats-officedocument.oleObject")
    original_parts["[Content_Types].xml"] = etree.tostring(types)
    relations = etree.fromstring(original_parts["word/_rels/document.xml.rels"])
    etree.SubElement(relations, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                     Id="rIdEmbeddedFixture", Target="embeddings/data.bin",
                     Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject")
    original_parts["word/_rels/document.xml.rels"] = etree.tostring(relations)
    document = etree.fromstring(original_parts["word/document.xml"])
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = document.find(w + "body").find(w + "p")
    embedded = etree.SubElement(etree.SubElement(paragraph, w + "r"), w + "object")
    ole = etree.SubElement(embedded, "{urn:schemas-microsoft-com:office:office}OLEObject")
    ole.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rIdEmbeddedFixture")
    original_parts["word/document.xml"] = etree.tostring(document)
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in original_parts.items():
            archive.writestr(name, data)
    changed = tmp_path / "changed.docx"
    changed_parts = dict(original_parts)
    if changed_part.endswith(".png"):
        buffer = io.BytesIO()
        Image.new("RGB", (8, 8), "blue").save(buffer, format="PNG")
        changed_parts[changed_part] = buffer.getvalue()
    else:
        changed_parts[changed_part] = b"changed embedded bytes"
    with zipfile.ZipFile(changed, "w") as archive:
        for name, data in changed_parts.items():
            archive.writestr(name, data)
    before, after = build_manifest(source), build_manifest(changed)
    # The old text/inventory-only digest misses both real byte changes.
    assert {k: v for k, v in before.items() if k not in {"source_path", "document_sha256"}} == {
        k: v for k, v in after.items() if k not in {"source_path", "document_sha256"}}
    assert benchmark.content_digest(before, "fixed-original", benchmark.package_part_hashes(source)) != (
        benchmark.content_digest(after, "fixed-original", benchmark.package_part_hashes(changed)))
