"""Synthetic-only CLI checks for explicitly supplied teacher corpus inputs."""

import json
from pathlib import Path
import subprocess
import sys

from docx import Document
import pytest


ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def corpus(tmp_path: Path) -> dict[str, Path]:
    project = tmp_path / "project"
    source = project / "reports" / "2024" / "合成样本桥.docx"
    label = project / "labels" / "2024" / "合成样本桥.docx"
    teachers = project / "teachers"
    for parent in (source.parent, label.parent, teachers):
        parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("桥梁名称：合成样本桥；总体评分：90；总体等级：A级。")
    document.save(source)
    document = Document()
    summary = document.add_table(rows=1, cols=2)
    summary.rows[0].cells[0].text = "字段"
    summary.rows[0].cells[1].text = "值"
    for key, value in {
        "桥梁名称": "合成样本桥", "总体评分": "90", "总体等级": "A级",
        "上部结构评分": "90", "上部结构等级": "A级",
        "下部结构评分": "90", "下部结构等级": "A级",
        "桥面系评分": "90", "桥面系等级": "A级",
    }.items():
        cells = summary.add_row().cells
        cells[0].text, cells[1].text = key, value
    document.add_table(rows=1, cols=1).cell(0, 0).text = "建议"
    document.add_table(rows=1, cols=1).cell(0, 0).text = "病害部位"
    document.save(label)
    teacher = teachers / "synthetic.teacher.json"
    teacher.write_text(json.dumps({
        "source": {"path": str(source.relative_to(project))},
        "targets": {"scores.overall.score": {"target_value": "89"}},
        "annotation_protocol": {"runtime_import_forbidden": True},
    }), encoding="utf-8")
    manifest = teachers / "manifest.json"
    manifest.write_text(json.dumps({"samples": [{
        "teacher": teacher.name,
        "source": str(source.relative_to(project)),
        "label": str(label.relative_to(project)),
    }]}), encoding="utf-8")
    return {"project": project, "source": source, "label": label,
            "teachers": teachers, "teacher": teacher, "manifest": manifest}


def run_tool(name: str, corpus: dict[str, Path], *options: str) -> subprocess.CompletedProcess[str]:
    selector = (["--manifest", str(corpus["manifest"]), "--force"]
                if name == "build_official_teacher_samples"
                else ["--teacher-root", str(corpus["teachers"])])
    return subprocess.run(
        [sys.executable, str(ROOT / "scripts" / f"{name}.py"), *selector,
         "--project-root", str(corpus["project"]), *options],
        capture_output=True, text=True, check=False,
    )


TOOLS = ("build_official_teacher_samples", "align_teacher_to_official_labels")


@pytest.mark.parametrize("name", TOOLS)
def test_explicit_labels_work_without_sample_specific_adjudications(name: str, corpus: dict[str, Path]) -> None:
    result = run_tool(name, corpus)
    assert result.returncode == 0, result.stderr
    teacher = json.loads(corpus["teacher"].read_text())
    target = teacher["targets"]["scores.overall.score"]
    assert target["target_value"] == "90"
    assert target["status"] == "official_label_resolved"
    assert teacher["benchmark_label"]["path"] == str(corpus["label"].relative_to(corpus["project"]))


@pytest.mark.parametrize("name", TOOLS)
def test_only_explicit_adjudications_mark_a_sample_conflict(name: str, corpus: dict[str, Path]) -> None:
    adjudications = corpus["project"] / "adjudications.json"
    adjudications.write_text(json.dumps({"synthetic.teacher.json": {
        "scores.overall.score": "Synthetic adjudication: source and target use different scopes.",
    }}), encoding="utf-8")
    result = run_tool(name, corpus, "--adjudications", str(adjudications))
    assert result.returncode == 0, result.stderr
    target = json.loads(corpus["teacher"].read_text())["targets"]["scores.overall.score"]
    assert target["status"] == "benchmark_convention_conflict"
    assert target["reason"].startswith("Synthetic adjudication:")


@pytest.mark.parametrize("name", TOOLS)
def test_labels_can_be_paired_under_explicit_year_root(name: str, corpus: dict[str, Path]) -> None:
    manifest = json.loads(corpus["manifest"].read_text())
    del manifest["samples"][0]["label"]
    corpus["manifest"].write_text(json.dumps(manifest), encoding="utf-8")
    result = run_tool(name, corpus, "--label-root", str(corpus["label"].parent.parent))
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("name", TOOLS)
def test_missing_label_location_fails_without_modifying_teacher(name: str, corpus: dict[str, Path]) -> None:
    manifest = json.loads(corpus["manifest"].read_text())
    del manifest["samples"][0]["label"]
    corpus["manifest"].write_text(json.dumps(manifest), encoding="utf-8")
    before = corpus["teacher"].read_bytes()
    result = run_tool(name, corpus)
    assert result.returncode != 0
    assert "--label-root" in result.stderr
    assert corpus["teacher"].read_bytes() == before


@pytest.mark.parametrize("name", TOOLS)
@pytest.mark.parametrize("payload", [[], {"synthetic.teacher.json": {"unknown.field": "Synthetic reason"}},
                                      {"synthetic.teacher.json": {"scores.overall.score": ""}},
                                      {"missing.teacher.json": {"scores.overall.score": "Synthetic reason"}}])
def test_invalid_adjudications_fail_before_writes(name: str, payload: object, corpus: dict[str, Path]) -> None:
    path = corpus["project"] / "adjudications.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    before = corpus["teacher"].read_bytes()
    result = run_tool(name, corpus, "--adjudications", str(path))
    assert result.returncode != 0
    assert "adjudication" in result.stderr.lower()
    assert corpus["teacher"].read_bytes() == before


@pytest.mark.parametrize("name", TOOLS)
def test_explicit_external_label_path_preserves_provenance(name: str, corpus: dict[str, Path]) -> None:
    external = corpus["project"].parent / "synthetic-label.docx"
    external.write_bytes(corpus["label"].read_bytes())
    manifest = json.loads(corpus["manifest"].read_text())
    manifest["samples"][0]["label"] = str(external)
    corpus["manifest"].write_text(json.dumps(manifest), encoding="utf-8")
    result = run_tool(name, corpus)
    assert result.returncode == 0, result.stderr
    teacher = json.loads(corpus["teacher"].read_text())
    assert teacher["benchmark_label"]["path"] == str(external)


@pytest.mark.parametrize("name", TOOLS)
def test_teacher_path_cannot_escape_output_root(name: str, corpus: dict[str, Path]) -> None:
    external = corpus["project"] / "outside.teacher.json"
    external.write_bytes(corpus["teacher"].read_bytes())
    before = external.read_bytes()
    manifest = json.loads(corpus["manifest"].read_text())
    manifest["samples"][0]["teacher"] = "../outside.teacher.json"
    corpus["manifest"].write_text(json.dumps(manifest), encoding="utf-8")
    result = run_tool(name, corpus)
    assert result.returncode != 0
    assert "within" in result.stderr
    assert external.read_bytes() == before
