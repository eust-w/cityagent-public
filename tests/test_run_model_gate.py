"""Synthetic DOCX + explicit fake-model integration tests, never GPU/SDK acceptance.

The direct workflow and real DOCX renderer execute here. The fake backend only
tests control-flow gates; successful fixtures do not certify model semantics.
"""
from __future__ import annotations

from copy import deepcopy
import importlib.util
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("contest_entry_model_gate_tests", ROOT / "run.py")
entry = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(entry)

from bridge_evidence.contest.docx_writer import write_extraction_docx


FINGERPRINT = "a" * 64


class FakeTokenizer:
    """Character tokenizer exclusively for synthetic control-flow assertions."""
    def __init__(self):
        self.messages = []

    def apply_chat_template(self, messages, **kwargs):
        assert kwargs == {"tokenize": False, "add_generation_prompt": True, "enable_thinking": False}
        self.messages.append(messages)
        return "SYSTEM:" + messages[0]["content"] + "|USER:" + messages[1]["content"] + "|ASSISTANT:"

    def __call__(self, text, **kwargs):
        assert kwargs == {"add_special_tokens": False, "truncation": False}
        return {"input_ids": list(map(ord, text))}


class FakeBackend:
    """Deterministic fixture backend; these calls are not actual model evidence."""
    def __init__(self, *responses):
        self.responses = list(responses)
        self.calls = []

    def __call__(self, ids, max_new_tokens):
        self.calls.append({"ids": ids.copy(), "max_new_tokens": max_new_tokens})
        if len(self.calls) > len(self.responses):
            raise AssertionError("Unexpected extra fake-model call")
        result = self.responses[len(self.calls) - 1]
        if isinstance(result, Exception):
            raise result
        return result


def model_response(*, raw=None, finish_reason="stop"):
    data = {
        "基本信息": {
            "桥梁名称": "合成门控测试桥", "报告日期": "2026-10-08", "总体评分": "82.50", "总体等级": "B级",
            "上部结构评分": "81.20", "上部结构等级": "B级",
            "下部结构评分": "83.40", "下部结构等级": "B级",
            "桥面系评分": "82.90", "桥面系等级": "B级",
            "总体结论": "本合成测试报告的评分来自所附人工构造表格，仅用于验证模型调用和文档发布门控。",
            "主要风险点": "本合成测试样本不对应实际桥梁，不构成工程安全判断。",
            "建议统计": "本合成样本未列出额外维护建议。",
        },
        "维修建议": [], "病害明细": [],
    }
    return {"text": json.dumps(data, ensure_ascii=False) if raw is None else raw,
            "finish_reason": finish_reason, "output_tokens": 73}


class RunModelGateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="cityagent-model-gate-test-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.input_dir, self.result_dir, self.work_dir = [self.root / name for name in ("input", "result", "work")]
        for directory in (self.input_dir, self.result_dir, self.work_dir, self.root / "model"):
            directory.mkdir()
        self.source = self.input_dir / "合成门控测试桥.docx"
        doc = Document()
        doc.add_heading("合成门控测试桥定期检测报告", 0)
        doc.add_paragraph("这是人工构造的集成测试文档，不包含真实桥梁数据。")
        info = doc.add_table(rows=0, cols=2)
        for key, value in (("桥梁名称", "合成门控测试桥"), ("报告日期", "2026年10月8日")):
            cells = info.add_row().cells
            cells[0].text, cells[1].text = key, value
        scores = doc.add_table(rows=0, cols=5)
        for values in (("部位名称", "部位技术状况指数", "技术状况等级", "总体评分", "总体等级"),
                       ("上部结构", "81.20", "B级", "82.50", "B级"),
                       ("下部结构", "83.40", "B级", "82.50", "B级"),
                       ("桥面系", "82.90", "B级", "82.50", "B级")):
            for cell, value in zip(scores.add_row().cells, values):
                cell.text = value
        doc.save(self.source)
        self.cli_args = [
            "--input-dir", str(self.input_dir), "--result-dir", str(self.result_dir),
            "--work-dir", str(self.work_dir), "--model", str(self.root / "model"),
            "--engine", "direct", "--max-model-len", "8192", "--max-new-tokens", "512", "--max-attempts", "2",
        ]
        self.args = entry.build_parser().parse_args(self.cli_args)
        self.output = self.result_dir / self.source.name
        self.tokenizer = FakeTokenizer()

    def process(self, backend, **kwargs):
        return entry.process_document(self.args, self.source, self.tokenizer, backend,
                                      fingerprint=kwargs.pop("fingerprint", FINGERPRINT), **kwargs)

    def receipt_path(self):
        document = entry.document_identity(self.source, self.input_dir)
        return self.work_dir / "records" / entry.record_name(document)

    def install_external_text(self, text):
        """Create a self-consistent synthetic text/sidecar pair, never live data."""
        self.args.text_dir = self.root / "external-text"
        self.args.text_dir.mkdir(exist_ok=True)
        path = self.args.text_dir / (self.source.stem + ".txt")
        data = text.encode("utf-8")
        path.write_bytes(data)
        digest = hashlib.sha256(data).hexdigest()
        path.with_suffix(".txt.source.json").write_text(json.dumps({
            "document_sha256": entry.file_digest(self.source), "text_sha256": digest,
            "parser": "contest-libreoffice-utf8-v1",
        }))
        return digest

    def assert_failed_without_output(self, failure, *, calls):
        record = failure.exception.record
        self.assertEqual(record["status"], "failed")
        self.assertFalse(record["submission_ready"])
        self.assertEqual(len(record["model_calls"]), calls)
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.result_dir.glob("*.docx")), [])
        self.assertEqual(list(self.result_dir.glob(".pending-*")), [])
        self.assertEqual(json.loads(self.receipt_path().read_text()), record)
        return record

    def test_valid_model_call_is_rendered_to_pending_then_atomically_published_with_receipt(self):
        backend = FakeBackend(model_response())
        render_paths = []

        def observing_real_writer(data, filename):
            pending = Path(filename)
            render_paths.append(pending)
            self.assertNotEqual(pending, self.output)
            self.assertEqual(pending.parent.parent, self.result_dir)
            self.assertTrue(pending.parent.name.startswith(".pending-"))
            self.assertFalse(self.output.exists())
            write_extraction_docx(data, filename)
            self.assertTrue(pending.is_file())
            self.assertFalse(self.output.exists())

        original_link = entry.os.link
        with patch.object(entry.os, "link", wraps=original_link) as link:
            record = self.process(backend, writer=observing_real_writer)
        link.assert_called_once_with(render_paths[0], self.output)
        self.assertEqual(len(backend.calls), 1)
        self.assertEqual(record["status"], "success")
        self.assertFalse(record["submission_ready"])
        self.assertEqual(record["workflow"]["engine"], "direct")
        self.assertEqual(record["workflow"]["verification"], "non-sdk-baseline")
        self.assertEqual([r["stage"] for r in record["workflow"]["stage_records"]],
                         ["prepare", "infer", "validate", "render"])
        call = record["model_calls"][0]
        self.assertEqual(call["status"], "success")
        self.assertEqual(call["output_tokens"], 73)
        infer_stage = record["workflow"]["stage_records"][1]
        self.assertEqual(infer_stage["model_call_ids"], [call["call_id"]])
        self.assertEqual(infer_stage["accepted_model_call_ids"], [call["call_id"]])
        self.assertEqual(record["output"]["sha256"], entry.file_digest(self.output))
        self.assertEqual(json.loads(self.receipt_path().read_text()), record)
        self.assertEqual(list(self.result_dir.glob(".pending-*")), [])
        rendered = Document(self.output)
        rendered_text = "\n".join(c.text for table in rendered.tables for row in table.rows for c in row.cells)
        self.assertIn("合成门控测试桥", rendered_text)
        self.assertIn("82.50", rendered_text)
        self.assertNotIn("90.00", rendered_text)
        self.assertEqual(entry.validate_execution_record(record,
            expected_document=entry.document_identity(self.source, self.input_dir),
            expected_run_fingerprint=FINGERPRINT, output_path=self.output, profile="legacy"), [])

    def test_default_success_path_preserves_inferred_semantics_without_legacy_candidates(self):
        # The actual JSON parser, direct workflow and DOCX writer run. Only the
        # backend response is synthetic, so this tests integration, not quality.
        predicted = {
            "基本信息": {"桥梁名称": "合成门控测试桥", "总体评分": 0, "总体等级": "E级",
                         "总体结论": "需核查承载能力", "下部结构评分": "",
                         "病害发展趋势与具体说明": "上次无法观测，无法据此判断发展。"},
            "维修建议": [{"序号": "9", "建议类别": "待进一步核查", "建议内容": "复查墩底"}],
            "病害明细": [
                {"序号": "7", "病害描述": "墩底剥落", "是否新增": "无法判断",
                 "上一次定检状态": "上次水位高无法观测", "发展程度": "无法比较"},
                {"序号": "7", "病害描述": "P1裂缝", "是否新增": "否",
                 "上一次定检状态": "上次已存在", "发展程度": "增多"},
                {"病害描述": "P2锈蚀", "是否新增": "", "上一次定检状态": ""},
                {"病害描述": "P2锈蚀", "是否新增": "", "上一次定检状态": ""},
            ],
        }
        raw = json.dumps(predicted, ensure_ascii=False)
        backend = FakeBackend(model_response(raw=raw))
        rendered = []

        def capture_real_writer(data, filename):
            rendered.append(deepcopy(data))
            write_extraction_docx(data, filename)

        self.assertEqual(self.args.fusion_policy, "model-preserving")
        with patch.object(entry, "preserve_model_report", wraps=entry.preserve_model_report) as preserve, \
             patch.object(entry, "_fuse_existing_success_path", side_effect=AssertionError("legacy must not run")) as legacy, \
             patch.object(entry, "extract_defect_rows_from_tables", side_effect=AssertionError("no table replacement")) as candidates:
            record = self.process(backend, writer=capture_real_writer)
        preserve.assert_called_once_with(predicted, label_conventions=False)
        legacy.assert_not_called()
        candidates.assert_not_called()
        self.assertEqual(len(backend.calls), 1)
        self.assertEqual(record["status"], "success")
        self.assertEqual(record["raw_output"], raw)
        self.assertEqual(record["model_json_sha256"], entry.json_digest(predicted))
        self.assertEqual(record["fusion"]["mode"], "model_preserving")
        self.assertEqual(record["fusion"]["result_sha256"], entry.json_digest(rendered[0]))
        self.assertEqual(record["fusion"]["defect_is_new_counts"], {"yes": 0, "no": 1, "unknown": 3})
        self.assertEqual({k: v for k, v in rendered[0]["基本信息"].items() if k != "建议统计"}, predicted["基本信息"])
        for section in ("维修建议", "病害明细"):
            without_numbers = lambda rows: [{key: value for key, value in row.items() if key != "序号"} for row in rows]
            self.assertEqual(without_numbers(rendered[0][section]), without_numbers(predicted[section]))
            self.assertEqual([row["序号"] for row in rendered[0][section]],
                             [str(index) for index in range(1, len(predicted[section]) + 1)])
        # Verify preserved uncertainty also survives actual DOCX rendering.
        output = Document(self.output)
        text = "\n".join(cell.text for table in output.tables for row in table.rows for cell in row.cells)
        self.assertIn("上次水位高无法观测", text)
        self.assertEqual(text.count("P2锈蚀"), 2)
        self.assertNotIn("上一次定检未发现", text)
        self.assertNotIn("结构运行良好", text)
        self.assertNotIn("90.00", text)

    def test_explicit_legacy_cli_runs_legacy_fusion_and_still_publishes_baseline(self):
        self.args = entry.build_parser().parse_args([*self.cli_args, "--fusion-policy", "legacy"])
        backend = FakeBackend(model_response())
        with patch.object(entry, "_fuse_existing_success_path", wraps=entry._fuse_existing_success_path) as legacy, \
             patch.object(entry, "preserve_model_report", side_effect=AssertionError("preserving policy must not run")) as preserve:
            record = self.process(backend)
        legacy.assert_called_once()
        preserve.assert_not_called()
        self.assertEqual(len(backend.calls), 1)
        self.assertEqual(record["status"], "success")
        self.assertFalse(record["submission_ready"])
        self.assertEqual(record["fusion"]["mode"], "legacy_success_path")
        self.assertTrue(self.output.is_file())

    def test_main_binds_fusion_policy_to_run_identity_and_rejects_cross_policy_resume(self):
        # Exercise main's real option construction, hashing and manifest gate,
        # then stop before any Transformers/vLLM import or model initialization.
        (self.root / "model" / "config.json").write_text('{"model_type":"synthetic"}')
        identities = []
        bind = entry.bind_work_directory

        class BeforeModelLoading(RuntimeError):
            pass

        def bind_then_stop(path, identity, *, resume):
            identities.append(deepcopy(identity))
            bind(path, identity, resume=resume)
            raise BeforeModelLoading("synthetic test stops before model loading")

        with patch.object(entry, "preflight", return_value=[]), \
             patch.object(entry, "ensure_ppu_environment"), \
             patch.object(entry, "prepare_model"), \
             patch.object(entry, "bind_work_directory", side_effect=bind_then_stop):
            with patch.object(entry.sys, "argv", [str(ROOT / "run.py"), *self.cli_args]):
                with self.assertRaises(BeforeModelLoading):
                    entry.main()
            manifest_path = self.work_dir / "run_manifest.json"
            before = manifest_path.read_bytes()
            with patch.object(entry.sys, "argv", [str(ROOT / "run.py"), *self.cli_args, "--resume"]):
                with self.assertRaises(BeforeModelLoading):
                    entry.main()
            with patch.object(entry.sys, "argv", [str(ROOT / "run.py"), *self.cli_args,
                                                  "--resume", "--fusion-policy", "legacy"]):
                with self.assertRaisesRegex(ValueError, "Run identity changed"):
                    entry.main()
        self.assertEqual([identity["options"]["fusion"] for identity in identities],
                         ["model-preserving", "model-preserving", "legacy"])
        self.assertEqual(identities[0], identities[1])
        self.assertNotEqual(entry.json_digest(identities[0]), entry.json_digest(identities[2]))
        self.assertEqual(manifest_path.read_bytes(), before)
        self.assertEqual(list(self.result_dir.iterdir()), [])

    def test_invalid_empty_and_length_responses_exhaust_attempts_without_rule_success(self):
        # Separate temporary outputs/receipts keep each case independently observable.
        for case, response in (("invalid", model_response(raw="not JSON")),
                               ("empty", model_response(raw="")),
                               ("length", model_response(finish_reason="length"))):
            with self.subTest(case=case):
                self.args.work_dir = self.root / ("work-" + case)
                self.work_dir = self.args.work_dir
                self.args.result_dir = self.root / ("result-" + case)
                self.result_dir = self.args.result_dir
                self.output = self.result_dir / self.source.name
                backend = FakeBackend(response, response)
                with patch.object(entry, "_fuse_existing_success_path") as fusion:
                    with self.assertRaises(entry.DocumentProcessingError) as failure:
                        self.process(backend)
                fusion.assert_not_called()
                record = self.assert_failed_without_output(failure, calls=2)
                self.assertEqual(len(backend.calls), 2)
                self.assertTrue(all(call["status"] == "failed" and call["errors"] for call in record["model_calls"]))
                self.assertFalse(record["model_validation"]["valid"])

    def test_success_after_failed_attempt_retains_both_attempts_and_only_one_accepted(self):
        backend = FakeBackend(model_response(raw="invalid fixture"), model_response())
        record = self.process(backend)
        calls = record["model_calls"]
        self.assertEqual([call["status"] for call in calls], ["failed", "success"])
        stage = record["workflow"]["stage_records"][1]
        self.assertEqual(stage["model_call_ids"], [call["call_id"] for call in calls])
        self.assertEqual(stage["accepted_model_call_ids"], [calls[-1]["call_id"]])
        self.assertEqual(record["seconds"], sum(call["seconds"] for call in calls))
        self.assertTrue(self.output.is_file())

    def test_backend_errors_are_recorded_without_publishing(self):
        backend = FakeBackend(RuntimeError("synthetic backend unavailable"), OSError("synthetic retry failure"))
        with self.assertRaises(entry.DocumentProcessingError) as failure:
            self.process(backend)
        record = self.assert_failed_without_output(failure, calls=2)
        self.assertEqual(len(backend.calls), 2)
        self.assertEqual([call["errors"] for call in record["model_calls"]],
                         [["backend_exception:RuntimeError"], ["backend_exception:OSError"]])

    def test_overlong_input_fails_before_any_backend_call(self):
        doc = Document(self.source)
        doc.add_paragraph("合成超长全文不可截断" * 2000)
        doc.save(self.source)
        backend = FakeBackend()
        with self.assertRaises(entry.DocumentProcessingError) as failure:
            self.process(backend)
        record = self.assert_failed_without_output(failure, calls=0)
        self.assertEqual(backend.calls, [])
        self.assertFalse(record["model_validation"]["valid"])

    def test_valid_resume_reuses_only_matching_identity_without_backend_or_writes(self):
        first = self.process(FakeBackend(model_response()))
        output_bytes, receipt_bytes = self.output.read_bytes(), self.receipt_path().read_bytes()
        self.args.resume = True
        unused_backend = FakeBackend()
        with patch.object(entry, "atomic_json", wraps=entry.atomic_json) as save:
            resumed = self.process(unused_backend)
        save.assert_not_called()
        self.assertEqual(resumed, first)
        self.assertEqual(unused_backend.calls, [])
        self.assertEqual(self.output.read_bytes(), output_bytes)
        self.assertEqual(self.receipt_path().read_bytes(), receipt_bytes)

    def test_resume_rejects_tampered_output_hash_and_preserves_existing_files(self):
        self.process(FakeBackend(model_response()))
        self.args.resume = True
        receipt_path = self.receipt_path()
        receipt_bytes = receipt_path.read_bytes()
        self.output.write_bytes(self.output.read_bytes() + b"synthetic-output-tampering")
        tampered_bytes = self.output.read_bytes()
        backend = FakeBackend()
        with self.assertRaises(FileExistsError):
            self.process(backend)
        self.assertEqual(backend.calls, [])
        self.assertEqual(self.output.read_bytes(), tampered_bytes)
        self.assertEqual(receipt_path.read_bytes(), receipt_bytes)

    def test_resume_rejects_changed_run_or_source_identity_without_overwriting(self):
        self.process(FakeBackend(model_response()))
        self.args.resume = True
        receipt_path = self.receipt_path()
        output_bytes, receipt_bytes = self.output.read_bytes(), receipt_path.read_bytes()
        backend = FakeBackend()
        with self.assertRaises(FileExistsError):
            self.process(backend, fingerprint="b" * 64)
        doc = Document(self.source)
        doc.add_paragraph("同名原件内容发生了合成测试修改。")
        doc.save(self.source)
        with self.assertRaises(FileExistsError):
            self.process(backend)
        self.assertEqual(backend.calls, [])
        self.assertEqual(self.output.read_bytes(), output_bytes)
        self.assertEqual(receipt_path.read_bytes(), receipt_bytes)
        self.assertFalse(self.receipt_path().exists())

    def test_unverified_sdk_native_evidence_cannot_publish_after_all_stages_render(self):
        self.args.engine = "tongagents"
        backend = FakeBackend(model_response())
        observed = []

        def fake_unverified_sdk_executor(payload, stages, *, engine):
            self.assertEqual(engine, "tongagents")
            # Execute the real stages through the direct workflow. The returned
            # synthetic SDK-shaped events explicitly do not grant verification.
            result = entry.execute_workflow(payload, stages, engine="direct")
            observed.append(result)
            self.assertTrue(any(self.result_dir.glob(".pending-*/*.docx")))
            self.assertFalse(self.output.exists())
            workflow = result["workflow"]
            workflow.update(engine="tongagents", non_sdk=False, sdk_version="FAKE-test-0.0",
                            verification="unverified", native_trace_validated=True,
                            native_events=[{"type": kind, "data": {"node_name": stage}}
                                for stage in ("prepare", "infer", "validate", "render")
                                for kind in ("node_start", "node_end")])
            return result

        with self.assertRaises(entry.DocumentProcessingError) as failure:
            self.process(backend, workflow_executor=fake_unverified_sdk_executor)
        record = self.assert_failed_without_output(failure, calls=1)
        self.assertEqual(len(observed), 1)
        self.assertEqual(record["model_calls"][0]["status"], "success")
        self.assertTrue(record["model_validation"]["valid"])
        self.assertIn("workflow.verification", record["error"])
        self.assertEqual(record["workflow"]["verification"], "unverified")

    def test_publication_race_never_overwrites_another_existing_report(self):
        other_bytes = b"synthetic-other-owner-output"

        def real_writer_with_competing_publish(data, filename):
            write_extraction_docx(data, filename)
            self.output.write_bytes(other_bytes)

        with self.assertRaises(entry.DocumentProcessingError) as failure:
            self.process(FakeBackend(model_response()), writer=real_writer_with_competing_publish)
        self.assertEqual(failure.exception.record["status"], "failed")
        self.assertIn("FileExistsError", failure.exception.record["error"])
        self.assertEqual(self.output.read_bytes(), other_bytes)
        self.assertEqual(list(self.result_dir.glob(".pending-*")), [])

    def test_pending_cleanup_failure_rolls_back_its_published_output(self):
        backend = FakeBackend(model_response())
        observed = []

        def fail_pending_cleanup(path, *args, **kwargs):
            path = Path(path)
            self.assertTrue(path.name.startswith(".pending-"))
            self.assertTrue(self.output.is_file(), "failure must occur after publication")
            observed.append(path)
            raise OSError("synthetic pending cleanup failure")

        with patch.object(entry.shutil, "rmtree", side_effect=fail_pending_cleanup):
            with self.assertRaisesRegex(OSError, "synthetic pending cleanup failure"):
                self.process(backend)
        self.assertEqual(len(backend.calls), 1)
        self.assertEqual(len(observed), 1)
        self.assertFalse(self.output.exists())
        self.assertEqual(list(self.result_dir.glob("*.docx")), [])
        self.assertFalse(self.receipt_path().exists())

    def test_history_move_failure_rolls_back_output_and_preserves_old_receipt(self):
        receipt_path = self.receipt_path()
        receipt_path.parent.mkdir(parents=True)
        old_receipt = b'{"status":"failed","error":"synthetic historical failure"}\n'
        receipt_path.write_bytes(old_receipt)
        real_replace = entry.os.replace
        attempted_moves = []

        def fail_history_move(source, target, *args, **kwargs):
            if Path(source) == receipt_path:
                self.assertTrue(self.output.is_file(), "failure must occur after publication")
                self.assertEqual(Path(target).parent.name, "history")
                attempted_moves.append((Path(source), Path(target)))
                raise OSError("synthetic history move failure")
            return real_replace(source, target, *args, **kwargs)

        backend = FakeBackend(model_response())
        with patch.object(entry.os, "replace", side_effect=fail_history_move):
            with self.assertRaisesRegex(OSError, "synthetic history move failure"):
                self.process(backend)
        self.assertEqual(len(backend.calls), 1)
        self.assertEqual(len(attempted_moves), 1)
        self.assertFalse(self.output.exists())
        self.assertEqual(receipt_path.read_bytes(), old_receipt)
        self.assertEqual(list(self.result_dir.glob(".pending-*")), [])

    def test_receipt_write_failure_rolls_back_its_published_output(self):
        seen_records = []

        def fail_receipt_write(path, record):
            self.assertEqual(Path(path), self.receipt_path())
            self.assertEqual(record["status"], "success")
            self.assertTrue(self.output.is_file(), "failure must occur after publication")
            seen_records.append(record)
            raise OSError("synthetic atomic receipt write failure")

        backend = FakeBackend(model_response())
        with patch.object(entry, "atomic_json", side_effect=fail_receipt_write):
            with self.assertRaisesRegex(OSError, "synthetic atomic receipt write failure"):
                self.process(backend)
        self.assertEqual(len(backend.calls), 1)
        self.assertEqual(len(seen_records), 1)
        self.assertFalse(self.output.exists())
        self.assertFalse(self.receipt_path().exists())
        self.assertEqual(list(self.result_dir.glob(".pending-*")), [])

    def test_failed_receipt_rollback_preserves_another_owner_replacement_inode(self):
        replacement = self.root / "another-owner.docx"
        replacement_bytes = b"synthetic-document-published-by-another-owner"
        replacement.write_bytes(replacement_bytes)
        replacement_identity = (replacement.stat().st_dev, replacement.stat().st_ino)
        observed_original = []

        def replace_output_then_fail_receipt(path, record):
            self.assertTrue(self.output.is_file())
            current = self.output.stat()
            observed_original.append((current.st_dev, current.st_ino))
            self.assertNotEqual(observed_original[-1], replacement_identity)
            entry.os.replace(replacement, self.output)
            raise OSError("synthetic receipt failure after another owner replaced output")

        with patch.object(entry, "atomic_json", side_effect=replace_output_then_fail_receipt):
            with self.assertRaisesRegex(OSError, "another owner replaced output"):
                self.process(FakeBackend(model_response()))
        self.assertEqual(len(observed_original), 1)
        self.assertEqual(self.output.read_bytes(), replacement_bytes)
        current = self.output.stat()
        self.assertEqual((current.st_dev, current.st_ino), replacement_identity)
        self.assertFalse(self.receipt_path().exists())
        self.assertEqual(list(self.result_dir.glob(".pending-*")), [])

    def test_failed_receipt_rollback_compares_symlink_inode_without_following_target(self):
        original_reference = self.root / "original-report-reference.docx"
        replacement_identity = []

        def replace_with_symlink_then_fail_receipt(path, record):
            # Preserve a hard link to our original inode, then simulate another
            # owner's distinct symlink directory entry at the published name.
            entry.os.link(self.output, original_reference)
            self.output.unlink()
            self.output.symlink_to(original_reference)
            link_stat = self.output.lstat()
            replacement_identity.append((link_stat.st_dev, link_stat.st_ino))
            raise OSError("synthetic receipt failure after symlink replacement")

        with patch.object(entry, "atomic_json", side_effect=replace_with_symlink_then_fail_receipt):
            with self.assertRaisesRegex(OSError, "after symlink replacement"):
                self.process(FakeBackend(model_response()))
        self.assertTrue(self.output.is_symlink(), "rollback must not delete a replacement symlink inode")
        current = self.output.lstat()
        self.assertEqual((current.st_dev, current.st_ino), replacement_identity[0])
        self.assertTrue(original_reference.is_file())

    def test_matching_frozen_external_text_is_used_and_valid_resume_does_not_infer(self):
        external = "合成冻结版本全文：合成门控测试桥，2026年10月8日，总体评分82.50。"
        frozen_hash = self.install_external_text(external)
        backend = FakeBackend(model_response())
        with patch.object(entry, "read_bound_text", wraps=entry.read_bound_text) as read_text:
            record = self.process(backend, expected_text_sha256=frozen_hash)
        read_text.assert_called_once_with(self.source, self.args.text_dir)
        self.assertEqual(record["input_text_sha256"], frozen_hash)
        self.assertEqual(self.tokenizer.messages[0][1]["content"], external)
        self.assertEqual(len(backend.calls), 1)
        self.assertTrue(self.output.is_file())
        output_bytes, receipt_bytes = self.output.read_bytes(), self.receipt_path().read_bytes()
        self.args.resume = True
        unused_backend = FakeBackend()
        resumed = self.process(unused_backend, expected_text_sha256=frozen_hash)
        self.assertEqual(resumed, record)
        self.assertEqual(unused_backend.calls, [])
        self.assertEqual(self.output.read_bytes(), output_bytes)
        self.assertEqual(self.receipt_path().read_bytes(), receipt_bytes)

    def test_self_consistent_external_sidecar_cannot_replace_frozen_text(self):
        frozen_hash = self.install_external_text("合成已冻结的原始全文版本。")
        changed = "合成已更换的全文版本；sidecar也已同步更新，但运行身份仍指向旧版本。"
        changed_hash = self.install_external_text(changed)
        self.assertNotEqual(changed_hash, frozen_hash)
        self.assertEqual(entry.read_bound_text(self.source, self.args.text_dir), changed)
        backend = FakeBackend()
        with self.assertRaisesRegex(ValueError, "text changed since run identity"):
            self.process(backend, expected_text_sha256=frozen_hash)
        self.assertEqual(backend.calls, [])
        self.assertFalse(self.output.exists())
        self.assertFalse(self.receipt_path().exists())
        self.assertEqual(list(self.result_dir.glob(".pending-*")), [])

    def test_resume_rejects_replaced_text_even_when_new_sidecar_is_self_consistent(self):
        frozen_hash = self.install_external_text("合成首次运行使用的冻结全文版本。")
        self.process(FakeBackend(model_response()), expected_text_sha256=frozen_hash)
        output_bytes, receipt_bytes = self.output.read_bytes(), self.receipt_path().read_bytes()
        self.args.resume = True
        changed = "合成续跑前被更换的全文版本，sidecar具有合法而更新的哈希。"
        self.assertNotEqual(self.install_external_text(changed), frozen_hash)
        self.assertEqual(entry.read_bound_text(self.source, self.args.text_dir), changed)
        backend = FakeBackend()
        with self.assertRaisesRegex(ValueError, "text changed since run identity"):
            self.process(backend, expected_text_sha256=frozen_hash)
        self.assertEqual(backend.calls, [])
        self.assertEqual(self.output.read_bytes(), output_bytes)
        self.assertEqual(self.receipt_path().read_bytes(), receipt_bytes)


if __name__ == "__main__":
    unittest.main()
