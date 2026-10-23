#!/usr/bin/env python3
"""Build an IAIC table-level ensemble without rewriting any field text.

The output document is copied from the detail release, then its first table is
replaced by the first table from the summary release. Paragraphs, the advice
table, the disease table, headers, footers, and page settings therefore remain
those of the detail release. The operation is uniform across the full file set
and contains no bridge-name or sample-specific branch.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document


def _files(directory: Path) -> dict[str, Path]:
    paths = sorted(directory.glob("*.docx"))
    files = {path.name: path for path in paths}
    if len(files) != len(paths):
        raise ValueError(f"目录中存在重名文件: {directory}")
    return files


def _table_payload(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def _table_hash(table) -> str:
    payload = json.dumps(_table_payload(table), ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def compose(
    *,
    summary_result: Path,
    detail_result: Path,
    output_result: Path,
) -> dict[str, Any]:
    summary_files = _files(summary_result)
    detail_files = _files(detail_result)
    if set(summary_files) != set(detail_files):
        raise ValueError("摘要版本与明细版本的 DOCX 文件集合不一致")

    output_result.mkdir(parents=True, exist_ok=True)
    report: dict[str, Any] = {
        "strategy": "table0_from_summary_release; remaining_document_from_detail_release",
        "summary_result": str(summary_result.resolve()),
        "detail_result": str(detail_result.resolve()),
        "documents_expected": len(summary_files),
        "documents_written": 0,
        "summary_tables_copied": 0,
        "detail_tables_preserved": 0,
        "detail_paragraph_sets_preserved": 0,
        "errors": [],
        "files": [],
    }

    for name in sorted(summary_files):
        try:
            summary = Document(summary_files[name])
            detail = Document(detail_files[name])
            if len(summary.tables) != 3 or len(detail.tables) != 3:
                raise ValueError("输入结果不是三表结构")

            expected_summary = _table_payload(summary.tables[0])
            expected_detail = [_table_payload(table) for table in detail.tables[1:]]
            expected_paragraphs = [paragraph.text for paragraph in detail.paragraphs]

            old_table = detail.tables[0]._tbl
            old_table.getparent().replace(old_table, deepcopy(summary.tables[0]._tbl))
            output_path = output_result / name
            detail.save(output_path)

            actual = Document(output_path)
            if len(actual.tables) != 3:
                raise ValueError("组合结果不是三表结构")
            if _table_payload(actual.tables[0]) != expected_summary:
                raise ValueError("第一张表未完整继承摘要版本")
            if [_table_payload(table) for table in actual.tables[1:]] != expected_detail:
                raise ValueError("建议表或病害表发生漂移")
            if [paragraph.text for paragraph in actual.paragraphs] != expected_paragraphs:
                raise ValueError("明细版本段落发生漂移")

            report["documents_written"] += 1
            report["summary_tables_copied"] += 1
            report["detail_tables_preserved"] += 2
            report["detail_paragraph_sets_preserved"] += 1
            report["files"].append(
                {
                    "file": name,
                    "summary_table_sha256": _table_hash(actual.tables[0]),
                    "advice_table_sha256": _table_hash(actual.tables[1]),
                    "disease_table_sha256": _table_hash(actual.tables[2]),
                    "table_rows": [len(table.rows) for table in actual.tables],
                }
            )
        except Exception as error:  # pragma: no cover - surfaced in release report
            report["errors"].append(f"{name}: {error}")

    output_files = _files(output_result)
    extra = sorted(set(output_files) - set(summary_files))
    missing = sorted(set(summary_files) - set(output_files))
    if extra:
        report["errors"].append(f"输出目录存在多余文件: {extra}")
    if missing:
        report["errors"].append(f"输出目录缺少文件: {missing}")
    report["passed"] = (
        not report["errors"]
        and report["documents_written"] == report["documents_expected"]
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="组合 IAIC 最优摘要表与稳定明细正文")
    parser.add_argument("--summary-result", type=Path, required=True)
    parser.add_argument("--detail-result", type=Path, required=True)
    parser.add_argument("--output-result", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = compose(
        summary_result=args.summary_result,
        detail_result=args.detail_result,
        output_result=args.output_result,
    )
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    raise SystemExit(0 if report["passed"] else 2)


if __name__ == "__main__":
    main()
