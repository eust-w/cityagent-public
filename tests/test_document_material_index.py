"""CPU-only identity, exact source mapping and table-context regression."""
from copy import deepcopy
import json
from pathlib import Path
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from bridge_evidence.contest.document_material_index import DocumentMaterialIndex, MaterialReferenceError
from bridge_evidence.contest.evidence_manifest import build_manifest


def make_material_case(directory):
    """Reusable no-network synthetic DOCX/material pair for integration gates.

    Returns {source_path, manifest, materials, refs}. Table numbers deliberately
    differ from canonical table IDs so matching by ordinal cannot pass.
    """
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    source = directory / 'material-case.docx'
    doc = Document()
    doc.add_heading('1 检测结果', level=1)
    doc.add_paragraph('桥梁甲，2018年检测。')
    doc.add_paragraph('重复说明')
    doc.add_paragraph('重复说明')
    table = doc.add_table(rows=4, cols=3)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row, values in zip(table.rows, [('构件', '本次', '上次'), ('支座', '轻微锈蚀', '轻微锈蚀'), ('桥台', '渗水', '渗水'), ('', '稳定', '稳定')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    table.cell(2, 0).merge(table.cell(3, 0))
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = '父表说明'
    nested = outer.cell(0, 0).add_table(rows=2, cols=2)
    nested.cell(0, 0).text = '项目'
    nested.cell(0, 1).text = '评分'
    nested.cell(1, 0).text = '综合评分'
    nested.cell(1, 1).text = '94.57'
    doc.save(source)
    manifest = build_manifest(source)
    with zipfile.ZipFile(source) as archive:
        tree = etree.fromstring(archive.read('word/document.xml')).getroottree()
    source_map, segments = {}, []
    paragraphs = [u for u in manifest['units'] if u['kind'] == 'paragraph' and u['context']['story'] == 'body']
    for i, unit in enumerate(paragraphs[:2]):
        ref = f'#/texts/{10+i}'
        source_map[ref] = {'xml_paragraph_candidates': [unit['source_path'].split('#', 1)[1]], 'paragraph_match': 'unique_exact'}
        segments.append({'kind': 'text', 'ref': ref, 'key': ref, 'text': unit['text'], 'offset': [0, len(unit['text'])]})
    duplicate_paths = [u['source_path'].split('#', 1)[1] for u in paragraphs if u['text'] == '重复说明']
    source_map['#/texts/99'] = {'xml_paragraph_candidates': duplicate_paths, 'paragraph_match': 'ambiguous_exact'}
    segments.append({'kind': 'text', 'ref': '#/texts/99', 'key': '#/texts/99', 'text': '重复说明', 'offset': [0, 4]})
    table_aliases = {}
    for i, table_info in enumerate(manifest['tables']):
        tref = f'#/tables/{50+i}'
        table_aliases[table_info['id']] = tref
        source_map[tref] = {'xml_table_path': table_info['source_path'].split('#', 1)[1]}
        for unit in manifest['units']:
            if unit.get('table_id') != table_info['id']:
                continue
            for cell in unit['cells']:
                if cell.get('vertical_merge') == 'continue':
                    continue
                cref = f"{tref}/r{unit['row_index']}/c{cell['column_index']}"
                info = {'xml_table_path': source_map[tref]['xml_table_path'], 'xml_cell_path': cell['source_path'].split('#', 1)[1],
                    'row': unit['row_index'], 'column': cell['column_index'], 'row_span': 2 if cell.get('vertical_merge') == 'restart' else 1,
                    'col_span': cell['column_span']}
                source_map[cref] = info
                if cell['text'].strip():
                    text_ref = f'#/texts/{1000+len(source_map)}'
                    node = tree.xpath(info['xml_cell_path'], namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})[0]
                    paragraphs = node.xpath('./w:p', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
                    source_map[text_ref] = {**info, 'xml_paragraph_candidates': [tree.getpath(p) for p in paragraphs]}
                    text = cell['text'].strip()
                    segments.append({'kind': 'text', 'ref': text_ref, 'key': text_ref, 'cell': cref,
                        'text': text, 'offset': [0, len(text)]})
    for tid, tref in table_aliases.items():
        segments.append({'kind': 'table_ref', 'ref': tref, 'key': tref, 'text': '', 'offset': [0, 0]})
    materials = {'docx_sha256': manifest['document_sha256'], 'source': str(source), 'source_map': source_map,
        'packets': [{'packet_id': 'P0001', 'section_id': 'S0001', 'headings': ['1 检测结果'], 'region': 'content',
            'segments': segments, 'docx_sha256': manifest['document_sha256']}],
        'sections': [{'id': 'S0001', 'parent': None, 'title': '1 检测结果', 'headings': ['1 检测结果'],
            'region': 'content', 'level': 1, 'source': {'ref': '#/texts/10'}, 'packet_ids': ['P0001']}]}
    return {'source_path': source, 'manifest': manifest, 'materials': materials,
        'refs': {'heading': '#/texts/10', 'unique': '#/texts/11', 'ambiguous': '#/texts/99',
            'table': table_aliases[manifest['tables'][0]['id']], 'nested_table': table_aliases[manifest['tables'][-1]['id']]}}


@pytest.fixture
def case(tmp_path):
    return make_material_case(tmp_path)


@pytest.fixture
def index(case):
    return DocumentMaterialIndex(case['manifest'], case['materials'], source_path=case['source_path'])


def test_unique_paragraph_keeps_exact_canonical_source(index, case):
    result = index.resolve(case['refs']['unique'])
    assert result['ok'] and len(result['spans']) == 1
    span = result['spans'][0]
    assert span['text'] == '桥梁甲，2018年检测。'
    assert result['semantic_review'] is False
    assert index.validate_manifest(case['manifest'])
    assert index.resolve(index.qualify_ref(case['refs']['unique'])) == result


def test_ambiguous_paragraph_is_not_first_match(index, case):
    result = index.resolve(case['refs']['ambiguous'])
    assert not result['ok'] and result['status'] == 'ambiguous' and not result['spans']
    assert len(result['fallback']['evidence_ids']) == 2
    assert result['unresolved'][0]['code'] == 'ambiguous_xml_text_mapping'


@pytest.mark.parametrize('method', ['resolve', 'ref_exists', 'normalize_ref'])
def test_foreign_namespace_rejected(index, method):
    with pytest.raises(MaterialReferenceError, match='Foreign document'):
        getattr(index, method)('mineru:' + '0'*64 + '#/texts/11')


def test_foreign_manifest_or_original_docx_rejected(case, tmp_path):
    wrong = deepcopy(case['manifest']); wrong['document_sha256'] = 'a'*64
    with pytest.raises(MaterialReferenceError):
        DocumentMaterialIndex(wrong, case['materials'], source_path=case['source_path'])
    wrong = deepcopy(case['manifest']); wrong['units'][0]['text'] += '伪造'
    with pytest.raises(MaterialReferenceError, match='Canonical units'):
        DocumentMaterialIndex(wrong, case['materials'], source_path=case['source_path'])


def test_cell_maps_exact_duplicate_occurrence_by_xml_path(index, case):
    table = case['refs']['table']
    left = index.resolve(table+'/r1/c1'); right = index.resolve(table+'/r1/c2')
    assert left['spans'][0]['text'] == right['spans'][0]['text'] == '轻微锈蚀'
    assert left['spans'][0]['start'] != right['spans'][0]['start']
    assert left['metadata']['canonical_table_id'] == 'T00001'
    assert any(s['role']=='same_row' and '支座' in s['text'] for s in left['context'])
    assert any(s['role']=='header' and '上次' in s['text'] for s in left['context'])


def test_row_returns_original_separators_and_merged_anchor(index, case):
    table = case['refs']['table']
    row = index.resolve(table+'/r3')
    assert all(' | ' in s['text'] for s in row['spans'])
    assert any(s['role']=='merged_cell_anchor' and '桥台' in s['text'] for s in row['context'])
    assert index.resolve(table+'/r2/c0')['metadata']['row_span'] == 2
    assert not index.ref_exists(table+'/r3/c0'), 'Merged-away positions must not be invented as cells'


def test_nested_table_retains_actual_parent_context(index, case):
    result = index.resolve(case['refs']['nested_table']+'/r1/c1')
    assert result['spans'][0]['text'] == '94.57'
    assert result['metadata']['parent_table_ref'] is not None
    assert any(s['role']=='parent_cell' and '父表说明' in s['text'] for s in result['context'])
    assert result['metadata']['header_kind'] == 'first_row_candidate'
    assert any(s['role']=='header_candidate' for s in result['context'])


def test_wrong_xml_path_does_not_search_entire_document(case):
    materials = deepcopy(case['materials'])
    materials['source_map']['#/texts/11']['xml_paragraph_candidates'] = materials['source_map']['#/texts/10']['xml_paragraph_candidates']
    index = DocumentMaterialIndex(case['manifest'], materials, source_path=case['source_path'])
    result = index.resolve('#/texts/11')
    assert not result['ok'] and not result['spans']


def test_wrong_cell_grid_coordinates_are_diagnostic(case):
    materials = deepcopy(case['materials']);ref=case['refs']['table']+'/r1/c1'
    materials['source_map'][ref]['column'] = 999
    index = DocumentMaterialIndex(case['manifest'], materials, source_path=case['source_path'])
    assert not index.resolve(ref)['ok'] and not index.resolve(ref)['spans']


def test_forged_vertical_merge_span_is_rejected(case):
    materials = deepcopy(case['materials']); ref=case['refs']['table']+'/r2/c0'
    materials['source_map'][ref]['row_span'] = 99
    index = DocumentMaterialIndex(case['manifest'], materials, source_path=case['source_path'])
    result = index.resolve(ref)
    assert not result['ok'] and not result['spans']
    assert result['metadata']['row_span'] == 2
    assert result['unresolved'][0]['actual_row_span'] == 2


def test_cell_text_anchor_must_match_original_xml(case):
    materials = deepcopy(case['materials']); ref=case['refs']['table']+'/r1/c1'
    segment = next(s for s in materials['packets'][0]['segments'] if s.get('cell') == ref)
    segment['text'] = '虚构内容'; segment['offset'] = [0,4]
    index = DocumentMaterialIndex(case['manifest'], materials, source_path=case['source_path'])
    result = index.resolve(ref)
    assert not result['ok'] and not result['spans']
    assert any(u['code']=='cell_text_anchor_mismatch' for u in result['unresolved'])


def test_full_manifest_persisted_but_outline_has_no_material_context(index, tmp_path):
    p=index.save(tmp_path/'index.json');saved=json.loads(p.read_text())
    assert saved['source_map'] and saved['refs'] and saved['source_validation']['canonical_units_verified']
    outline=index.compact_outline()
    assert 'source_map' not in outline and 'refs' not in outline
    assert outline['sections'][0]['id']=='#/sections/S0001'
    assert index.packet('P0001')['spans'] and index.section('S0001')['unresolved']
    assert index.aliases_for(index.resolve('#/texts/11')['spans'][0]['evidence_id'])


def test_manifest_digest_binding_even_if_document_hash_same(index, case):
    wrong=deepcopy(case['manifest']);wrong['source_path']='other path'
    with pytest.raises(MaterialReferenceError, match='different canonical manifest'):
        index.validate_manifest(wrong)


def test_resolved_copies_do_not_mutate_index(index):
    result=index.resolve('#/texts/11');result['spans'][0]['text']='changed'
    assert index.resolve('#/texts/11')['spans'][0]['text']=='桥梁甲，2018年检测。'


def test_table_alias_cannot_point_to_different_member_cell_table(case):
    materials=deepcopy(case['materials'])
    materials['source_map'][case['refs']['table']]['xml_table_path']=materials['source_map'][case['refs']['nested_table']]['xml_table_path']
    index=DocumentMaterialIndex(case['manifest'],materials,source_path=case['source_path'])
    result=index.resolve(case['refs']['table'])
    assert not result['ok'] and not result['spans']
    assert result['unresolved'][0]['code']=='xml_table_cell_identity_mismatch'


def test_delivery_ledger_only_confirms_actual_canonical_payload(index,case,tmp_path):
    from bridge_evidence.contest.document_source_delivery import SourceDeliveryLedger, RECEIPT_FIELD
    ledger=SourceDeliveryLedger(case['manifest'],tmp_path/'ledger',run_id='r',agent_id='a')
    result=index.resolve(case['refs']['table']+'/r1/c1')
    page=ledger.prepare_page(result,tool='read_source')
    assert not ledger.is_confirmed(page[RECEIPT_FIELD])
    payload={'messages':[{'role':'tool','tool_call_id':'c','content':json.dumps({'result':page},ensure_ascii=False)}]}
    ledger.observe_request(payload,{'request_id':'q'})
    ledger.confirm_request('q',evidence={'generation_attempted':True,'input_delivery_confirmed':True})
    assert ledger.is_confirmed(page[RECEIPT_FIELD])
    assert ledger.confirmed_ranges()


def test_metadata_alias_alone_does_not_grant_delivery(index,case,tmp_path):
    from bridge_evidence.contest.document_source_delivery import SourceDeliveryLedger, RECEIPT_FIELD
    ledger=SourceDeliveryLedger(case['manifest'],tmp_path/'ledger',run_id='r',agent_id='a')
    page=ledger.prepare_page(index.resolve('#/texts/11'),tool='read_source')
    fake={'ref':page['ref'],RECEIPT_FIELD:page[RECEIPT_FIELD]}
    ledger.observe_request({'messages':[{'role':'tool','tool_call_id':'c','content':json.dumps(fake)}]}, {'request_id':'q'})
    ledger.confirm_request('q',evidence={'input_delivery_confirmed':True})
    assert not ledger.confirmed_ranges()
