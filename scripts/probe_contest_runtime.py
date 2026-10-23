#!/usr/bin/env python3
"""Real imports, mount semantics, DOCX conversion, device and short local generation."""

import argparse
import json
import os
import shutil
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from bridge_evidence.contest.backends import APIBackend, dependency_fingerprint
from bridge_evidence.contest.run_state import atomic_json


def filesystem_probe(directory):
    directory.mkdir(parents=True, exist_ok=False)
    first = directory / "a"
    first.write_text("first")
    result = {}
    try:
        os.link(first, directory / "b")
        result["hardlink"] = True
    except OSError as e:
        result["hardlink"] = False
        result["hardlink_error"] = type(e).__name__
    replacement = directory / "replace"
    replacement.write_text("replacement")
    os.replace(replacement, first)
    result["replace"] = first.read_text() == "replacement"
    atomic_json(directory / "durable.json", result)
    result["reopen"] = json.loads((directory / "durable.json").read_text()) == result
    return result


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--output-dir", type=Path, required=True)
    ap.add_argument("--model", required=True)
    ap.add_argument("--url", default="http://127.0.0.1:8000/v1")
    ap.add_argument("--device-only", action="store_true")
    args = ap.parse_args()
    if not str(args.output_dir.resolve()).startswith("/workspace/"):
        ap.error("all probe artifacts must be under /workspace")
    args.output_dir.mkdir(parents=True, exist_ok=False)
    result = {"dependencies": dependency_fingerprint(), "utc": time.time()}
    if not args.device_only:
        from docx import Document

        doc = Document()
        doc.add_paragraph(
            "运行检查桥定期检测报告。桥梁名称：运行检查桥。报告日期：2026年10月10日。"
        )
        doc.add_paragraph("总体评分82.5，技术状况等级B；桥面伸缩缝有杂物，建议清理。")
        doc.save(args.output_dir / "probe.docx")
        result["filesystem"] = filesystem_probe(args.output_dir / "filesystem")
        if shutil.which("soffice"):
            source = args.output_dir / "probe.docx"
            converted = args.output_dir / "converted"
            converted.mkdir()
            subprocess.run(
                [
                    "soffice",
                    "-env:UserInstallation="
                    + (args.output_dir / "lo-profile").as_uri(),
                    "--headless",
                    "--convert-to",
                    "doc",
                    "--outdir",
                    str(converted),
                    str(source),
                ],
                capture_output=True,
                check=True,
                timeout=120,
            )
            subprocess.run(
                [
                    "soffice",
                    "-env:UserInstallation="
                    + (args.output_dir / "lo-profile").as_uri(),
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(converted),
                    str(converted / "probe.doc"),
                ],
                capture_output=True,
                check=True,
                timeout=120,
            )
            result["conversion"] = (converted / "probe.docx").is_file()
        backend = APIBackend(args.url, args.model, None)
        result["backend"] = backend.probe()
        response = backend([151645], 2)
        result["short_generation"] = {k: v for k, v in response.items() if k != "text"}
    atomic_json(args.output_dir / "probe.json", result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
