"""Structural navigation and literal search; no SDK/network/model is used."""
import copy
import json

from docx import Document
import pytest

from bridge_evidence.contest.document_task_plan import build_document_navigation, build_document_task_plan
from bridge_evidence.contest.document_tools import DocumentTools
from bridge_evidence.contest.evidence_manifest import build_manifest, manifest_from_text
from bridge_evidence.contest.tong_react_adapter import wrap_tool_response


def _manifest(*lines):
    return manifest_from_text("\n".join(lines))


def _tools(manifest, **kwargs):
    return DocumentTools(manifest, count_tokens=len, max_tool_tokens=kwargs.pop("max_tool_tokens", 20000), **kwargs)


def _pages(tools, first):
    page = first
    while True:
        assert page["ok"], page
        yield page
        call = page["continuation"]
        if call is None:
            break
        page = tools.handlers()[call["tool"]](**call["arguments"])


def test_regex_planner_outline_and_section_reads_share_original_ids():
    manifest = _manifest("目录", "1 总体情况......1", "1 总体情况", "桥况", "1.1 检测记录", "裂缝", "2 历史对比", "原文", "附录 A", "附录内容", "")
    original = copy.deepcopy(manifest)
    plan = build_document_task_plan(manifest)
    tools = _tools(manifest)
    planned = {heading["evidence_id"] for task in plan["tasks"] for heading in task["headings"]}
    indexed = {heading["evidence_id"] for heading in tools.navigation["headings"]}
    assert planned == indexed == {"E000003", "E000005", "E000007", "E000009"}
    assert {section["id"] for section in tools.read_outline()["sections"]} == {"root", *indexed}
    assert tools._sections["E000003"] == ["E000003", "E000004", "E000005", "E000006"]
    assert tools._sections["E000005"] == ["E000005", "E000006"]
    assert tools._sections["E000007"] == ["E000007", "E000008"]
    assert tools._sections["E000009"] == ["E000009", "E000010"]
    read = tools.read_section("E000005", 0, 10)
    assert [span["evidence_id"] for span in read["spans"]] == tools._sections["E000005"]
    assert manifest == original


def test_styles_toc_duplicates_and_story_boundaries_are_shared():
    manifest = _manifest("章节", "子节", "页眉", "正文", "重复目录", "附录 B")
    for unit, style in zip(manifest["units"], ["Heading1", "Heading2", "Heading1", "", "TOC1", ""]):
        unit["style_id"] = style
    manifest["units"][2]["context"]["story"] = "header"
    nav = build_document_navigation(manifest)
    assert nav["section_evidence_ids"]["E000001"] == ["E000001", "E000002"]
    assert nav["section_evidence_ids"]["E000003"] == ["E000003"]
    assert nav["heading_paths"][3] == []  # Returning to body does not inherit a stale chapter.
    assert "E000005" not in nav["section_evidence_ids"]
    assert nav["section_evidence_ids"]["E000006"] == ["E000006"]


def test_task_search_defaults_to_primary_and_explicit_global_is_available():
    manifest = _manifest("1 检测", "左幅裂缝", "2 历史", "右幅裂缝")
    task = {"id": "task-local", "primary_evidence_ids": ["E000001", "E000002"],
            "context_evidence_ids": ["E000004"]}
    reads = []
    tools = _tools(manifest, task=task, on_read=lambda *span: reads.append(span))
    assert [hit["evidence_id"] for hit in tools.search_in_document("裂缝")["matches"]] == ["E000002"]
    assert tools.search_in_document("裂缝")["scope"] == "task"
    assert tools.search_in_document("裂缝", scope="document")["total_matches"] == 2
    assert tools.search_in_document("裂缝", section_id="E000003")["matches"][0]["evidence_id"] == "E000004"
    assert tools.search_in_document("裂缝", scope="section", section_id="task-local")["total_matches"] == 1
    assert not reads and all(not receipt["ranges"] for receipt in tools.receipts)
    assert tools.read_outline()["sections"][0]["id"] == "task-local"
    assert "root" in {section["id"] for section in tools.read_outline()["sections"]}
    assert [span["evidence_id"] for span in tools.read_section("task-local", 0, 10)["spans"]] == task["primary_evidence_ids"]


def test_empty_assignment_does_not_silently_expand_to_global():
    tools = _tools(_manifest("原文裂缝"), task={"id": "task-empty", "primary_evidence_ids": []})
    assert tools.search_in_document("裂缝")["total_matches"] == 0
    assert tools.search_in_document("裂缝", scope="document")["total_matches"] == 1
    assert tools.read_section("task-empty", 0, 1)["spans"] == []


@pytest.mark.parametrize("task", [
    {"id": "root", "primary_evidence_ids": []},
    {"id": "E000001", "primary_evidence_ids": []},
    {"id": "x", "primary_evidence_ids": ["missing"]},
    {"id": "x", "primary_evidence_ids": ["E000001", "E000001"]},
    {"id": "x", "primary_evidence_ids": [], "document_sha256": "other"},
])
def test_bad_task_binding_rejected(task):
    with pytest.raises(ValueError, match="Task"):
        _tools(_manifest("原文"), task=task)


def test_multi_keyword_or_literal_phrase_case_and_exact_offsets():
    tools = _tools(_manifest("CrAcK 裂缝和剥落 a.b axb 裂缝", "钢筋锈蚀", "ß ss"))
    result = tools.search_in_document(["crack", "裂缝", "剥落", "a.b", "裂缝"])
    assert result["total_matches"] == 5
    assert [row["matched_query"] for row in result["matches"]] == ["crack", "裂缝", "剥落", "a.b", "裂缝"]
    for row in result["matches"]:
        source = tools._index[row["evidence_id"]]["text"]
        assert row["text"] == source[row["start"]:row["end"]]
        assert source[row["match_start"]:row["match_end"]].lower() == row["matched_query"].lower()
    assert tools.search_in_document("裂缝 剥落")["total_matches"] == 0  # One phrase, not implicit OR.
    assert tools.search_in_document("ß")["total_matches"] == 1  # No casefold rewriting of original query.
    assert tools.search_in_document(["crack", "CrAcK"])["total_matches"] == 1  # OR is a union of occurrences.
    assert not result["grants_coverage"]


def test_occurrence_pagination_preserves_scope_and_native_budget():
    tools = _tools(_manifest("裂缝 剥落 " * 25, "其他裂缝"),
                   task={"id": "assigned", "primary_evidence_ids": ["E000001"]},
                   native_response_budget=True, max_tool_tokens=1700)
    pages = list(_pages(tools, tools.search_in_document(["裂缝", "剥落"], max_matches=2)))
    rows = [row for page in pages for row in page["matches"]]
    assert len(rows) == 50
    assert len({(row["evidence_id"], row["match_start"], row["match_end"]) for row in rows}) == 50
    assert all(page["scope"] == "task" and page["total_matches"] == 50 for page in pages)
    assert pages[-1]["omitted_matches"] == 0
    assert all(len(json.dumps(wrap_tool_response("search_in_document", page), ensure_ascii=False,
                              sort_keys=True, separators=(",", ":"))) <= 1700 for page in pages)
    assert all(not row["ranges"] for row in tools.receipts)


@pytest.mark.parametrize("arguments", [
    {"query": []}, {"query": ["x", " "]}, {"query": ["x"] * 17}, {"query": 3},
    {"query": "x", "scope": "section"}, {"query": "x", "scope": "task"},
    {"query": "x", "scope": "document", "section_id": "root"},
    {"query": "x", "section_id": "missing"}, {"query": "x", "cursor": True},
    {"query": "x", "cursor": 3}, {"query": "x", "max_matches": 101},
    {"query": "x", "max_matches": False}, {"query": "x", "scope": "auto"},
])
def test_invalid_search_arguments_are_bounded_and_do_not_read(arguments):
    seen = []
    tools = _tools(_manifest("x"), on_read=lambda *span: seen.append(span))
    result = tools.search_in_document(**arguments)
    assert not result["ok"] and result["status"] == "error"
    assert not seen and not tools.receipts[-1]["ranges"]


def test_search_keeps_table_header_units_blank_separated_notes_and_merged_origin(tmp_path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("1 检测记录")
    doc.add_paragraph("单位:mm")
    table = doc.add_table(rows=4, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "位置", "宽度(mm)"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "左侧"
    table.cell(1, 1).text, table.cell(2, 1).text = "0.1", "裂缝0.2"
    table.cell(3, 0).text, table.cell(3, 1).text = "其他", "0.3"
    doc.add_paragraph("")
    doc.add_paragraph("单位:mm")
    doc.add_paragraph("注:仅左侧")
    doc.add_paragraph("2 后续章节")
    doc.add_paragraph("不属于表格注释")
    doc.save(path)
    manifest = build_manifest(path)
    before = copy.deepcopy(manifest)
    seen = []
    tools = _tools(manifest, on_read=lambda *span: seen.append(span))
    result = tools.search_in_document(["裂缝", "不存在"])
    assert result["ok"] and result["total_matches"] == 1
    context = result["context"]
    texts = [tools._index[item["evidence_id"]]["text"] for item in context]
    assert "位置 | 宽度(mm)" in texts and "注:仅左侧" in texts
    assert texts.count("单位:mm") == 2 and "" in texts
    assert "2 后续章节" not in texts and "不属于表格注释" not in texts
    origins = [item for item in context if item["role"] == "merged_cell_origin"]
    assert origins and "左侧" in tools._index[origins[0]["evidence_id"]]["text"]
    assert all(item["read"]["arguments"]["evidence_id"] == item["evidence_id"] for item in context)
    assert not seen and manifest == before


def test_long_header_is_explicit_reference_not_silently_lost(tmp_path):
    path = tmp_path / "long.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "原表头" * 3000
    table.cell(1, 0).text = "目标"
    doc.save(path)
    tools = _tools(build_manifest(path), max_tool_tokens=2200, native_response_budget=True)
    result = tools.search_in_document("目标")
    assert result["ok"]
    header = next(item for item in result["context"] if item["role"].startswith("table_header:"))
    assert "text" not in header and header["read"]["arguments"]["end"] == 9000
    assert tools.receipts[-1]["result_tokens"] <= 2200
    recovery = tools.read_source(**header["read"]["arguments"])
    assert recovery.get("continuation") is not None


def test_outline_inventory_pages_all_sections_tables_and_images_without_only_root():
    manifest = _manifest("1 总体情况", "正文", "2 检测", "末尾")
    manifest["images"] = [{"id": "I1", "source_path": "image-source", "target": "image.png"}]
    tools = _tools(manifest, max_tool_tokens=1000, native_response_budget=True)
    pages = list(_pages(tools, tools.read_outline(limit=1)))
    assert [item["id"] for page in pages for item in page.get("sections", [])] == ["root", "E000001", "E000003"]
    assert [item["id"] for page in pages for item in page.get("images", [])] == ["I1"]
    assert pages[-1]["outline_complete"]
    assert all(row["result_tokens"] <= 1000 and not row["ranges"] for row in tools.receipts)
    assert _tools(manifest).read_outline()["sections"][1]["id"] == "E000001"


def test_long_styled_heading_keeps_explicit_full_title_source_pointer():
    manifest = _manifest("标题" * 3000, "正文")
    manifest["units"][0]["style_id"] = "Heading1"
    tools = _tools(manifest, max_tool_tokens=1000)
    pages = list(_pages(tools, tools.read_outline(limit=1)))
    heading = pages[-1]["sections"][0]
    assert heading["id"] == "E000001" and heading["metadata_complete"] is False
    assert heading["read"] == {"tool": "read_source", "arguments": {"evidence_id": "E000001"}}


@pytest.mark.parametrize("arguments", [{"cursor": -1}, {"cursor": True}, {"cursor": 100}, {"limit": 0}, {"limit": 101}])
def test_outline_validation_does_not_restart_pagination(arguments):
    tools = _tools(_manifest("1 标题"))
    assert tools.read_outline(**arguments)["ok"] is False


def test_nonheading_evidence_remains_source_not_section():
    tools = _tools(_manifest("1 标题", "正文"))
    result = tools.read_section("E000002", 0, 1)
    assert result["error"]["code"] == "SECTION_EXPECTS_HEADING_ID"
    assert result["next_read"]["arguments"] == {"evidence_id": "E000002"}
