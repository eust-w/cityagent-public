"""Pair training reports with labels without losing identifiers or guessing.

All documents are synthetic and stay in temporary directories. No model,
document converter, remote host, or competition dataset is used by these tests.
"""
from __future__ import annotations

import importlib.util
import hashlib
import io
import json
from contextlib import redirect_stderr, redirect_stdout
from itertools import permutations
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "build_sft_dataset_for_pairing_test", ROOT / "scripts/build_sft_dataset.py"
)
assert SPEC is not None and SPEC.loader is not None
builder = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(builder)


def label_path(stem: str, suffix: str = "无对比年度的信息提取报告") -> Path:
    return Path(f"{stem}-{suffix}.docx")


class NameNormalizationTests(unittest.TestCase):
    def test_roman_bridge_identifiers_stay_distinct_and_match_their_labels(self):
        identifiers = ["Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ"]
        normalized = [builder.norm_name(Path(f"合成{n}号桥.docx"), False) for n in identifiers]
        self.assertEqual(len(set(normalized)), len(identifiers))
        for numeral, expected_ascii in zip(identifiers, ["ii", "iii", "iv", "v"]):
            with self.subTest(numeral=numeral):
                report = Path(f"合成{numeral}号桥.docx")
                self.assertEqual(builder.norm_name(report, False), f"合成{expected_ascii}号桥")
                self.assertEqual(
                    builder.norm_name(report, False),
                    builder.norm_name(label_path(f"合成{numeral}号桥"), True),
                )

    def test_nfkc_casefold_and_whitespace_normalization(self):
        report = Path(" 合成 ＡＢ１２（北）桥 .docx")
        label = label_path("合成 ab12(北)桥")
        self.assertEqual(builder.norm_name(report, False), "合成ab12(北)桥")
        self.assertEqual(builder.norm_name(report, False), builder.norm_name(label, True))

    def test_label_wrapper_variants_match_without_removing_bridge_numbers(self):
        report = Path("合成12号Ⅱ桥.docx")
        names = [
            "合成12号Ⅱ桥-有对比年度的信息提取报告.docx",
            "合成12号Ⅱ桥.docx-有对比年度的信息提取报告.docx",
            "合成12号Ⅱ桥.docx-有对比年度的信息提取报告(1).docx",
            "合成12号Ⅱ桥.docx-有对比年度的信息提取报告（1）.docx",
            "合成12号Ⅱ桥.docx-有对比年度的信息提取报告2.docx",
        ]
        for name in names:
            with self.subTest(label=name):
                label = Path(name)
                self.assertEqual(builder.norm_name(label, True), "合成12号ii桥")
                pair, = builder.plan_pairs([report], [label])
                self.assertEqual(pair["status"], "exact")
                self.assertEqual(pair["report"], report.name)

    def test_numbers_and_meaningful_punctuation_are_not_discarded(self):
        for left, right in [("桥1", "桥11"), ("A-B桥", "AB桥"), ("A/B桥", "AB桥"), ("A(北)桥", "A北桥")]:
            with self.subTest(left=left, right=right):
                self.assertNotEqual(
                    builder.norm_name(Path(left + ".docx"), False),
                    builder.norm_name(Path(right + ".docx"), False),
                )
        self.assertEqual(builder.norm_name(Path("信息提取报告桥.docx"), True), "信息提取报告桥")


class PairPlanningTests(unittest.TestCase):
    def assert_unassigned(self, record: dict, status: str) -> None:
        self.assertEqual(record["status"], status)
        self.assertIsNone(record["report"])
        self.assertIsInstance(record["reason"], str)
        self.assertTrue(record["reason"])
        self.assertIsInstance(record["candidates"], list)

    def test_roman_numeral_reports_each_match_only_their_own_label(self):
        reports = [Path(f"合成{n}号桥.docx") for n in ["Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ"]]
        labels = [label_path(p.stem) for p in reports]
        pairs = builder.plan_pairs(reports, labels)
        self.assertEqual(len(pairs), 4)
        by_label = {r["label"]: r for r in pairs}
        for report, label in zip(reports, labels):
            record = by_label[label.name]
            self.assertEqual(record["status"], "exact")
            self.assertEqual(record["report"], report.name)
        self.assertEqual(len({r["report"] for r in pairs}), 4)

    def test_duplicate_normalized_report_names_are_rejected_not_overwritten(self):
        reports = [Path("合成Ａ桥.doc"), Path("合成A桥.docx")]
        pair, = builder.plan_pairs(reports, [label_path("合成A桥")])
        self.assert_unassigned(pair, "ambiguous")
        self.assertEqual(len(pair["candidates"]), 2)

    def test_multiple_labels_claiming_one_report_are_all_ambiguous(self):
        report = Path("合成桥.docx")
        labels = [label_path("合成桥"), label_path("合成桥", "含有对比年度的信息提取报告")]
        for order in permutations(labels):
            with self.subTest(order=[p.name for p in order]):
                pairs = builder.plan_pairs([report], list(order))
                self.assertEqual({p["label"] for p in pairs}, {p.name for p in labels})
                for pair in pairs:
                    self.assert_unassigned(pair, "ambiguous")

    def test_collision_does_not_remove_an_unrelated_unique_pair(self):
        reports = [Path("合成Ａ桥.doc"), Path("合成A桥.docx"), Path("独立桥.docx")]
        labels = [label_path("合成A桥"), label_path("独立桥")]
        pairs = {p["label"]: p for p in builder.plan_pairs(reports, labels)}
        self.assert_unassigned(pairs[labels[0].name], "ambiguous")
        self.assertEqual(pairs[labels[1].name]["status"], "exact")
        self.assertEqual(pairs[labels[1].name]["report"], "独立桥.docx")

    def test_nearby_number_is_a_candidate_and_never_an_automatic_pair(self):
        report = Path("合成训练桥12号.docx")
        pair, = builder.plan_pairs([report], [label_path("合成训练桥13号")])
        self.assert_unassigned(pair, "unmatched")
        self.assertTrue(pair["candidates"], "A near match should remain visible for review")

    def test_fuzzy_candidate_cannot_steal_a_report_from_its_exact_label(self):
        report = Path("合成训练桥12号.docx")
        labels = [label_path(report.stem), label_path("合成训练桥13号")]
        pairs = {p["label"]: p for p in builder.plan_pairs([report], labels)}
        self.assertEqual(pairs[labels[0].name]["status"], "exact")
        self.assertEqual(pairs[labels[0].name]["report"], report.name)
        self.assert_unassigned(pairs[labels[1].name], "unmatched")

    def test_different_report_year_is_not_normalized_or_fuzzily_paired(self):
        report = Path("合成桥2024年度定检.docx")
        label = label_path("合成桥2014年度定检")
        self.assertNotEqual(builder.norm_name(report, False), builder.norm_name(label, True))
        pair, = builder.plan_pairs([report], [label])
        self.assert_unassigned(pair, "unmatched")
        self.assertTrue(pair["candidates"])

    def test_plans_are_deterministic_across_file_enumeration_orders(self):
        reports = [Path("合成Ａ桥.doc"), Path("合成A桥.docx"), Path("独立12桥.docx")]
        labels = [label_path("合成A桥"), label_path("独立12桥"), label_path("独立13桥")]
        expected = builder.plan_pairs(reports, labels)
        for report_order in permutations(reports):
            for label_order in permutations(labels):
                self.assertEqual(builder.plan_pairs(list(report_order), list(label_order)), expected)

    def test_no_reports_leaves_every_label_unmatched(self):
        labels = [label_path("合成桥")]
        pair, = builder.plan_pairs([], labels)
        self.assertEqual(pair["label"], labels[0].name)
        self.assert_unassigned(pair, "unmatched")
        self.assertEqual(builder.plan_pairs([Path("合成桥.docx")], []), [])


class BuilderIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.train = self.root / "reports"
        self.labels = self.root / "labels"
        self.year = "2014年"
        (self.train / self.year).mkdir(parents=True)
        (self.labels / self.year).mkdir(parents=True)
        self.out = self.root / "output"

    def run_builder(self, *extra: str) -> None:
        argv = ["build_sft_dataset.py", "--train-root", str(self.train),
                "--label-root", str(self.labels), "--out-dir", str(self.out), *extra]
        with patch.object(builder, "YEARS", [self.year]), patch("sys.argv", argv), redirect_stdout(io.StringIO()):
            builder.main()

    def test_audit_only_writes_metadata_without_reading_or_parsing_documents(self):
        report = self.train / self.year / "合成桥.doc"
        label = self.labels / self.year / label_path("合成桥").name
        report.write_bytes(b"not a real report: audit must not read this")
        label.write_bytes(b"not a real label: audit must not read this")
        original_open = Path.open

        def guarded_open(path, *args, **kwargs):
            if path.is_relative_to(self.train) or path.is_relative_to(self.labels):
                raise AssertionError(f"Audit read document contents: {path.name}")
            return original_open(path, *args, **kwargs)

        with patch.object(Path, "open", guarded_open), \
             patch.object(builder, "report_text", side_effect=AssertionError("report parsed")), \
             patch.object(builder, "label_answer", side_effect=AssertionError("label parsed")), \
             patch("docx.Document", side_effect=AssertionError("document opened")), \
             patch.object(builder, "convert_doc", side_effect=AssertionError("converter called")):
            self.run_builder("--audit-only")

        self.assertEqual({p.name for p in self.out.iterdir()},
                         {"pairing_manifest.json", "stats.json", "anomalies.md"})
        manifest = json.loads((self.out / "pairing_manifest.json").read_text())
        self.assertEqual(len(manifest["pairs"]), 1)
        self.assertEqual(manifest["pairs"][0]["status"], "exact")
        self.assertEqual(manifest["pairs"][0]["year"], self.year)
        stats = json.loads((self.out / "stats.json").read_text())
        self.assertEqual(stats["records"], 0)
        self.assertTrue(stats["audit_only"])
        self.assertEqual(report.read_bytes(), b"not a real report: audit must not read this")
        self.assertEqual(label.read_bytes(), b"not a real label: audit must not read this")

    def test_nonempty_output_is_rejected_and_existing_sft_is_preserved(self):
        self.out.mkdir()
        previous_sft = b'{"previous":"dataset must remain byte-identical"}\n'
        previous_manifest = b'{"previous":"audit must remain byte-identical"}\n'
        (self.out / "sft_train.jsonl").write_bytes(previous_sft)
        (self.out / "pairing_manifest.json").write_bytes(previous_manifest)
        for extra in [(), ("--audit-only",)]:
            with self.subTest(extra=extra), redirect_stderr(io.StringIO()), \
                 patch.object(builder, "report_text", side_effect=AssertionError("report parsed")), \
                 patch.object(builder, "label_answer", side_effect=AssertionError("label parsed")):
                with self.assertRaises(SystemExit) as error:
                    self.run_builder(*extra)
                self.assertEqual(error.exception.code, 2)
            self.assertEqual((self.out / "sft_train.jsonl").read_bytes(), previous_sft)
            self.assertEqual((self.out / "pairing_manifest.json").read_bytes(), previous_manifest)
            self.assertEqual({p.name for p in self.out.iterdir()}, {"sft_train.jsonl", "pairing_manifest.json"})

    def test_output_inside_either_source_tree_is_rejected_without_changing_documents(self):
        report = self.train / self.year / "合成桥.docx"
        label = self.labels / self.year / label_path("合成桥").name
        original_report, original_label = b"original synthetic report", b"original synthetic label"
        report.write_bytes(original_report)
        label.write_bytes(original_label)
        for source_root in [self.train, self.labels]:
            self.out = source_root / "empty-output"
            self.out.mkdir()
            for extra in [(), ("--audit-only",)]:
                with self.subTest(root=source_root.name, extra=extra), redirect_stderr(io.StringIO()), \
                     patch.object(builder, "report_text", side_effect=AssertionError("report parsed")), \
                     patch.object(builder, "label_answer", side_effect=AssertionError("label parsed")):
                    with self.assertRaises(SystemExit) as error:
                        self.run_builder(*extra)
                    self.assertEqual(error.exception.code, 2)
                self.assertEqual(list(self.out.iterdir()), [])
                self.assertEqual(report.read_bytes(), original_report)
                self.assertEqual(label.read_bytes(), original_label)

    def test_builder_outputs_only_exact_pairs_and_never_parses_unresolved_documents(self):
        report = self.train / self.year / "合成Ⅰ号桥.docx"
        label = self.labels / self.year / label_path("合成Ⅰ号桥").name
        doc = Document()
        doc.add_paragraph("合成训练材料，桥墩存在裂缝，尺寸来自本段。" * 40)
        doc.save(report)
        answer = Document()
        tables = [
            [["字段", "值"], ["桥梁名称", "合成Ⅰ号桥"]],
            [["序号", "建议类别", "建议内容", "病害部位"], ["1", "跟踪观测", "跟踪裂缝宽度", "桥墩"]],
            [["序号", "病害部位", "病害类型", "病害描述", "是否新增", "上一次定检状态", "发展程度"],
             ["1", "桥墩", "裂缝", "宽度0.2毫米", "是", "上一次定检未发现", "无"]],
        ]
        for rows in tables:
            table = answer.add_table(rows=len(rows), cols=len(rows[0]))
            for cells, values in zip(table.rows, rows):
                for cell, value in zip(cells.cells, values):
                    cell.text = value
        answer.save(label)

        for name in ["重复Ａ桥.docx", "重复A桥.docx"]:
            (self.train / self.year / name).write_bytes(b"unresolved report must not be parsed")
        for stem in ["重复A桥", "合成Ⅱ号桥"]:
            (self.labels / self.year / label_path(stem).name).write_bytes(b"unresolved label must not be parsed")
        report_reader, label_reader = builder.report_text, builder.label_answer

        def read_report(path, cache):
            self.assertEqual(path, report)
            return report_reader(path, cache)

        def read_label(path, cache, **kwargs):
            self.assertEqual(path, label)
            return label_reader(path, cache, **kwargs)

        with patch.object(builder, "report_text", side_effect=read_report) as report_mock, \
             patch.object(builder, "label_answer", side_effect=read_label) as label_mock, \
             patch.object(builder, "convert_doc", side_effect=AssertionError("converter called")):
            self.run_builder()
        report_mock.assert_called_once()
        label_mock.assert_called_once()
        report_cache = report_mock.call_args.args[1]
        label_cache = label_mock.call_args.args[1]
        self.assertNotEqual(report_cache, label_cache)
        self.assertTrue(report_cache.is_relative_to(self.out))
        self.assertTrue(label_cache.is_relative_to(self.out))
        rows = [json.loads(line) for line in (self.out / "sft_train.jsonl").read_text().splitlines()]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["meta"]["report"], report.name)
        self.assertEqual(rows[0]["meta"]["label"], label.name)
        self.assertEqual(rows[0]["meta"]["pair_confidence"], "exact")
        self.assertEqual(rows[0]["meta"]["pairing_version"], "2-strict-unique")
        self.assertEqual(rows[0]["meta"]["report_sha256"], hashlib.sha256(report.read_bytes()).hexdigest())
        self.assertEqual(rows[0]["meta"]["label_sha256"], hashlib.sha256(label.read_bytes()).hexdigest())
        self.assertEqual(json.loads(rows[0]["output"])["基本信息"]["桥梁名称"], "合成Ⅰ号桥")
        self.assertEqual(json.loads((self.out / "stats.json").read_text())["records"], 1)
        self.assertTrue((self.out / "pairing_manifest.json").is_file())


if __name__ == "__main__":
    unittest.main()
