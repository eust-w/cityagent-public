"""CPU-only scripted save-error recovery through a real Word and finish.

The source, model decisions, gateway acknowledgements and SDK parameter check
are explicit test substitutes. There is no model or native TongAgents process
in this module. Production SpeedMergeWorkflow, DocumentWorkspace, atomic save
transactions, source ledger, DOCX writer/readback and finish binding do execute.
These tests prove mechanical recovery for the prescribed sequence, not that a
real model will choose it or that an official report has semantic coverage.
"""
from copy import deepcopy
import json
from types import SimpleNamespace

from docx import Document
from jsonschema import ValidationError
import pytest

from bridge_evidence.contest import tong_react_adapter as adapter
from bridge_evidence.contest.document_main_review_tools import SCHEMAS
from bridge_evidence.contest.document_main_supplement import SUPPLEMENT_SCHEMAS
from bridge_evidence.contest.document_speed_workflow import SpeedMergeWorkflow
from bridge_evidence.contest.document_workflow import DocumentSettings
from bridge_evidence.contest.evidence_manifest import manifest_from_text
from bridge_evidence.contest.tong_react_adapter import wrap_tool_response
from test_managed_file_self_check import Harness, offline
from test_tong_react_adapter import SchemaParameters, SchemaTool


class ScriptedRecovery(Harness):
    """Prescribe model actions, retaining the real workflow/file machinery."""

    def __init__(self, root, monkeypatch):
        source = ("合成恢复桥\n甲处渗水\n建议甲处疏通排水\n建议乙处清除锈迹\n"
                  "建议丙处更换止水带\n建议丁处清除锈迹后涂漆")
        self.client = SimpleNamespace(model="synthetic-scripted-model",
            base_url="http://unused.invalid", calls=[], max_model_len=262144,
            timeout=900, dynamic_output_budget=True, enable_thinking=False,
            temperature=0, count_tokens=lambda text: max(1, (len(text) + 3) // 4))
        self.wf = SpeedMergeWorkflow(manifest_from_text(source), self.client,
            {"model": self.client.model}, root,
            child_summary=[{"task_id": "synthetic-failed-child", "status": "failed"}],
            settings=DocumentSettings(single_document_agent=True,
                require_semantic_review=False, dynamic_output_budget=True,
                max_model_len=262144, output_tokens=262144, max_model_calls=16,
                report_timeout=900))
        self.history = [{"role": "user", "content": source}]
        units = self.wf.manifest["units"]
        self.wf.source_delivery.register_initial_message(source, [
            {"evidence_id": unit["id"], "start": 0, "end": len(unit["text"])}
            for unit in units])
        # These are already-saved synthetic child results, not model output.
        for unit in units:
            self.wf.workspace.record_read(unit["id"], 0, len(unit["text"]))
        self.wf.workspace.apply_update({"base_version": self.wf.workspace.snapshot()["version"],
            "facts": [
                {"id": "seed-name", "value": units[0]["text"], "status": "explicit",
                 "source_refs": [{"evidence_id": units[0]["id"]}]},
                {"id": "seed-defect", "value": {"病害描述": units[1]["text"]},
                 "status": "explicit", "source_refs": [{"evidence_id": units[1]["id"]}]}],
            "draft": {"基本信息": {"桥梁名称": units[0]["text"]},
                      "病害明细": [{"病害描述": units[1]["text"]}], "维修建议": []},
            "draft_fact_ids": ["seed-name", "seed-defect"],
            "draft_bindings": {"基本信息.桥梁名称": ["seed-name"], "病害明细.0": ["seed-defect"]}})
        self.handlers = self.wf._additional_reading_handlers()
        self.receipts = []
        self.sdk_rejections = []
        monkeypatch.setattr(adapter, "_invoke", self._dispatch)

    def advice(self, index):
        unit = self.wf.manifest["units"][index + 2]
        return {"建议内容": unit["text"], "status": "explicit",
                "source_refs": [{"evidence_id": unit["id"]}]}

    def _dispatch(self, name, payload):
        # The typed production wrapper supplies its actual defaults before
        # dispatch. No missing business field or semantic decision is filled.
        adapter._validate_arguments(name, payload)
        return self.handlers[name](**payload)

    def submit(self, name, payload, *, sdk_check=True):
        message = self.call(name, payload)
        # A synthetic successful gateway acknowledgement drives the real
        # source-delivery ledger. This is not real network-delivery evidence.
        request_id = f"request-{len(self.client.calls)}"
        self.wf.source_delivery.confirm_request(request_id, evidence={
            "request_id": request_id, "input_delivery_confirmed": True})
        schema = SUPPLEMENT_SCHEMAS.get(name, SCHEMAS.get(name))
        if sdk_check:
            fake_tool = type("ScriptedSchemaTool", (SchemaTool,), {
                "parameters": SchemaParameters(schema)})()
            try:
                parsed = fake_tool._normalize_and_validate_params(payload)
            except ValidationError as exc:
                self.sdk_rejections.append({"name": name,
                    "path": list(exc.absolute_path), "validator": exc.validator})
                # The fake SDK rejects before the business handler. Prescribe
                # the next corrected model action; no native SDK retry is
                # claimed by this local test.
                self.history.append(message)
                self.history.append({"role": "tool", "tool_call_id": message["tool_calls"][0]["id"],
                    "content": json.dumps({"ok": False, "error": "synthetic_sdk_schema_rejection"})})
                raise
        else:
            parsed = deepcopy(payload)
        result = adapter._TOOLS[name][0](**parsed)
        envelope = wrap_tool_response(name, result)
        self.wf._observe_native_tool(name, payload, envelope)
        self.history += [message, {"role": "tool", "tool_call_id": message["tool_calls"][0]["id"],
            "content": json.dumps(envelope, ensure_ascii=False)}]
        self.receipts.append((name, deepcopy(result)))
        return result


@pytest.mark.parametrize("fault", ["sdk_array_rejection", "atomic_batch_rejection"])
def test_scripted_parameter_recovery_saves_all_four_measures_and_finishes_current_word(
        tmp_path, monkeypatch, fault):
    h = ScriptedRecovery(tmp_path, monkeypatch)
    deadline = h.wf.deadline
    seed = deepcopy(h.wf.workspace.snapshot()["draft"]["病害明细"])
    first = h.submit("save_recommendations", {"entries": [h.advice(0)]})
    assert first["ok"] and first["save_progress"]["saved_indices"] == [0]
    committed = h.wf.workspace.checkpoint_path.read_bytes()
    if fault == "sdk_array_rejection":
        with pytest.raises(ValidationError):
            h.submit("save_recommendations", {"entries": []})
        assert h.sdk_rejections == [{"name": "save_recommendations",
            "path": ["entries"], "validator": "minItems"}]
        # SDK rejection did not reach or consume a business save attempt.
        assert h.wf._gap_state()["save_attempts"]["save_recommendations"] == 1
    else:
        bad_batch = [h.advice(1), h.advice(2)]
        bad_batch[1]["建议内容"] = " "
        rejected = h.submit("save_recommendations", {"entries": bad_batch})
        assert not rejected["ok"]
        assert rejected["save_progress"]["outcome"] == "not_saved"
        assert rejected["save_progress"]["batch_applied"] is False
        assert rejected["save_progress"]["saved_indices"] == []
        assert rejected["save_progress"]["error"]["path"] == "entries[1].建议内容"
    assert h.wf.workspace.checkpoint_path.read_bytes() == committed

    # The script corrects the rejected input and submits all outstanding rows,
    # including the valid sibling from the rejected whole batch. Third/fourth
    # successful single-entry saves must remain legal after the old two cap.
    successful = [first]
    for index in (1, 2, 3):
        result = h.submit("save_recommendations", {"entries": [h.advice(index)]})
        assert result["ok"] and result["added_count"] == 1
        assert result["save_progress"]["outcome"] == "saved"
        successful.append(result)
    assert successful[-1]["save_progress"]["attempt"] >= 4
    assert h.wf._gap_state()["limits"]["save_calls"] == "shared_main_model_tool_time_budget"
    assert "save_attempts_per_section" not in h.wf._gap_state()["limits"]

    # A repeat is an explicit no-op receipt, not a second row or a state reset.
    complete_bytes = h.wf.workspace.checkpoint_path.read_bytes()
    repeated = h.submit("save_recommendations", {"entries": [h.advice(0)]})
    assert repeated["ok"] and repeated["added_count"] == 0
    assert repeated["save_progress"]["outcome"] == "unchanged"
    assert repeated["save_progress"]["already_saved_indices"] == [0]
    assert h.wf.workspace.checkpoint_path.read_bytes() == complete_bytes
    snapshot = h.wf.workspace.snapshot()
    expected = [h.advice(index)["建议内容"] for index in range(4)]
    assert snapshot["draft"]["病害明细"] == seed
    assert [row["建议内容"] for row in snapshot["draft"]["维修建议"]] == expected

    # Omitting the optional checks reaches the real targeted workflow error;
    # the script supplies them in the next call, without re-saving any row.
    missing = h.submit("complete_review", {"basis": "脚本审核尚未填写两类结论"})
    assert missing["error"]["code"] == "GAP_CHECK_REQUIRED"
    assert not h.wf._gap_state()["closed"]
    reviewed = h.submit("complete_review", {"basis": "脚本核对已有病害与四条措施",
        "gap_checks": {"病害明细": "保留原先成功保存的甲处渗水",
                       "维修建议": "四条原文明示措施均已按成功回执保存"}})
    assert reviewed["ok"] and len(reviewed["actual_word_pages"]) == 3
    assert h.wf._gap_state()["closed"]
    assert h.wf.model_self_check is None  # Export/readback is not model review.
    # Complete every real current-file page, even if a view needs continuation.
    h.read_all()
    assert h.wf._artifact_read_coverage(h.wf.workspace.snapshot())
    h.wf._check_budget()
    finished = h.submit("finish_report", {"basis": "脚本模型已读取实际Word三表并核对四条措施"})
    assert finished["ok"] and h.wf._finish_is_current()
    assert h.wf.model_self_check["verified_file_binding"]["workspace_version"] == snapshot["version"]
    assert h.wf.deadline == deadline and h.wf.settings.max_model_calls == 16
    assert len(h.client.calls) <= 16

    # Independently read the actual DOCX bytes, not just the workspace JSON.
    doc = Document(h.wf.artifact_delivery["docx_path"])
    advice = next(table for table in doc.tables if
        "建议内容" in [cell.text for cell in table.rows[0].cells])
    headers = [cell.text for cell in advice.rows[0].cells]
    content_index, sequence_index = headers.index("建议内容"), headers.index("序号")
    assert [row.cells[content_index].text for row in advice.rows[1:]] == expected
    assert [row.cells[sequence_index].text for row in advice.rows[1:]] == ["1", "2", "3", "4"]
    defects = next(table for table in doc.tables if
        "病害描述" in [cell.text for cell in table.rows[0].cells])
    assert len(defects.rows) == 2
    assert any(cell.text == seed[0]["病害描述"] for cell in defects.rows[1].cells)
