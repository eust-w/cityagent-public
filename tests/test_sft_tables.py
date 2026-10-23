from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import threading

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest


SCRIPT = Path(__file__).resolve().parents[1] / "scripts/build_sft_dataset.py"
SPEC = importlib.util.spec_from_file_location("sft_table_builder", SCRIPT)
builder = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(builder)

ADVICE = ["序号", "建议类别", "建议内容", "病害部位"]
DISEASE = ["序号", "病害部位", "病害类型", "病害描述", "是否新增", "上一次定检状态", "发展程度"]


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    return table


def sample_document():
    doc = Document()
    add_table(doc, ["字段", "内容"], [["桥梁名称", "合成示例桥"], ["总体评分", "90"]])
    add_table(doc, ADVICE, [["1", "预防性养护", "定期检查", "无"]])
    add_table(doc, DISEASE, [["1", "桥墩", "裂缝", "宽0.1mm", "否", "无", "无"]])
    return doc


def parse(doc, tmp_path):
    path = tmp_path / "label.docx"
    doc.save(path)
    errors = []
    answer = builder.label_answer(path, tmp_path / "cache", diagnostics=errors)
    return answer, errors


def test_independent_equal_cells_keep_their_columns(tmp_path):
    doc = sample_document()
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["病害明细"] == [dict(zip(DISEASE, ["1", "桥墩", "裂缝", "宽0.1mm", "否", "无", "无"]))]
    assert builder.read_docx_tables(tmp_path / "label.docx")[2][1][-2:] == ["无", "无"]


def test_report_text_keeps_independent_equal_values_and_document_order(tmp_path):
    doc = Document()
    doc.add_paragraph("第一段：表格之前")
    add_table(doc, ["比重Wij", "MDPi", "权重Wi"], [["0.00", "0.00", "0.15"]])
    doc.add_paragraph("第二段：两表之间")
    add_table(doc, ["前次等级", "本次等级"], [["A级", "A级"]])
    doc.add_paragraph("第三段：表格之后")
    path = tmp_path / "report.docx"
    doc.save(path)
    assert builder.report_text(path, tmp_path / "cache").splitlines() == [
        "第一段：表格之前", "比重Wij MDPi 权重Wi", "0.00 0.00 0.15",
        "第二段：两表之间", "前次等级 本次等级", "A级 A级", "第三段：表格之后",
    ]


def test_report_text_does_not_repeat_one_horizontally_merged_xml_cell(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "同一个合并格"
    table.cell(0, 2).text = "独立格"
    path = tmp_path / "report.docx"
    doc.save(path)
    assert builder.docx_text(path) == "同一个合并格 独立格"


def test_legacy_report_uses_libreoffice_txt_and_preserves_utf8_text(tmp_path, monkeypatch):
    expected = "\ufeff第一行\n0.00\t0.00\n\n"
    converted = tmp_path / "converted.txt"
    converted.write_text(expected, encoding="utf-8")
    calls = []

    def fake_convert(path, target, cache):
        calls.append((path, target, cache))
        return converted

    monkeypatch.setattr(builder, "convert_doc", fake_convert)
    path, cache = tmp_path / "report.doc", tmp_path / "cache"
    assert builder.report_text(path, cache) == expected
    assert calls == [(path, "txt", cache)]


def test_legacy_report_invalid_utf8_is_rejected(tmp_path, monkeypatch):
    converted = tmp_path / "converted.txt"
    converted.write_bytes(b"broken\xffutf8")
    monkeypatch.setattr(builder, "convert_doc", lambda *args: converted)
    assert builder.report_text(tmp_path / "report.doc", tmp_path / "cache") is None


def test_builder_standalone_can_import_formal_extractor_outside_repository(tmp_path):
    result = subprocess.run([sys.executable, "-I", "-B", str(SCRIPT), "--help"],
                            cwd=tmp_path, capture_output=True, text=True, check=True)
    assert "--train-root" in result.stdout


def test_basic_information_allows_merged_layout_columns(tmp_path):
    doc = sample_document()
    original = doc.tables[0]
    table = doc.add_table(rows=3, cols=3)
    for row, key, value in zip(table.rows, ["字段", "桥梁名称", "总体评分"], ["内容", "合成示例桥", "90"]):
        row.cells[0].merge(row.cells[1]).text = key
        row.cells[2].text = value
    original._tbl.addprevious(table._tbl)
    original._tbl.getparent().remove(original._tbl)
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["基本信息"] == {"桥梁名称": "合成示例桥", "总体评分": "90"}


def test_basic_equal_field_and_value_are_not_deduplicated(tmp_path):
    doc = sample_document()
    doc.tables[0].rows[1].cells[1].text = "桥梁名称"
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["基本信息"]["桥梁名称"] == "桥梁名称"


@pytest.mark.parametrize("value_header", ["内容", "示例", "提取内容"])
def test_official_three_column_basic_table_ignores_explanation(tmp_path, value_header):
    doc = sample_document()
    original = doc.tables[0]
    table = add_table(doc, ["字段", value_header, "说明"],
                      [["桥梁名称", "合成示例桥", "桥梁全称"], ["总体评分", "90", "说明不是答案"]])
    original._tbl.addprevious(table._tbl)
    original._tbl.getparent().remove(original._tbl)
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["基本信息"] == {"桥梁名称": "合成示例桥", "总体评分": "90"}


def test_headerless_basic_table_keeps_first_bridge_name_row(tmp_path):
    doc = sample_document()
    table = doc.tables[0]
    table._tbl.remove(table.rows[0]._tr)
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["基本信息"]["桥梁名称"] == "合成示例桥"


def test_verified_header_aliases_map_to_canonical_fields(tmp_path):
    doc = sample_document()
    doc.tables[1].rows[0].cells[2].text = "建议内容（原文复刻）"
    for index, alias in [(1, "位 置"), (2, "病害种类"), (3, "病 害 情 况"), (5, "上一次定检状态（2014年）")]:
        doc.tables[2].rows[0].cells[index].text = alias
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert list(answer["维修建议"][0]) == ADVICE
    assert list(answer["病害明细"][0]) == DISEASE


def test_aliases_cannot_hide_duplicate_fields(tmp_path):
    doc = sample_document()
    doc.tables[2].rows[0].cells[2].text = "位置"
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("重复字段" in error for error in errors)


def test_specific_location_is_description_only_in_verified_full_template(tmp_path):
    doc = sample_document()
    table = doc.tables[2]
    for cell, text in zip(table.rows[0].cells, ["序号", "位置", "病害种类", "具体位置", "是否新增", "上一次定检状态", "发展程度"]):
        cell.text = text
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["病害明细"][0]["病害描述"] == "宽0.1mm"
    table.rows[0].cells[1].text = "病害部位"
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("缺少字段" in error for error in errors)


def test_expanded_grid_with_one_merged_header_field_is_preserved(tmp_path):
    doc = sample_document()
    original = doc.tables[2]
    table = doc.add_table(rows=2, cols=8)
    for row, values in zip(table.rows, [DISEASE, ["1", "桥墩", "裂缝", "宽0.1mm", "否", "无", "无"]]):
        row.cells[6].merge(row.cells[7])
        for index, value in enumerate(values):
            row.cells[index].text = value
    original._tbl.addprevious(table._tbl)
    original._tbl.getparent().remove(original._tbl)
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["病害明细"][0]["发展程度"] == "无"
    assert len(answer["病害明细"][0]) == 7


def test_grid_span_can_change_between_rows_without_changing_logical_fields(tmp_path):
    doc = sample_document()
    original = doc.tables[2]
    table = doc.add_table(rows=2, cols=8)
    table.rows[0].cells[6].merge(table.rows[0].cells[7])
    for index, value in enumerate(DISEASE):
        table.rows[0].cells[index].text = value
    table.rows[1].cells[5].merge(table.rows[1].cells[6])
    for index, value in zip([0, 1, 2, 3, 4, 5, 7], ["1", "桥墩", "裂缝", "宽0.1mm", "否", "原有裂缝", "轻微"]):
        table.rows[1].cells[index].text = value
    original._tbl.addprevious(table._tbl)
    original._tbl.getparent().remove(original._tbl)
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["病害明细"][0]["上一次定检状态"] == "原有裂缝"
    assert answer["病害明细"][0]["发展程度"] == "轻微"


def test_extra_independent_cell_under_merged_header_is_rejected(tmp_path):
    doc = sample_document()
    original = doc.tables[2]
    table = doc.add_table(rows=2, cols=8)
    table.rows[0].cells[6].merge(table.rows[0].cells[7])
    for index, value in enumerate(DISEASE):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(["1", "桥墩", "裂缝", "宽0.1mm", "否", "无", "无", "额外字段"]):
        table.rows[1].cells[index].text = value
    original._tbl.addprevious(table._tbl)
    original._tbl.getparent().remove(original._tbl)
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("多列" in error for error in errors)


def test_explanation_only_basic_row_is_skipped_but_unlabelled_value_is_rejected(tmp_path):
    doc = sample_document()
    original = doc.tables[0]
    table = add_table(doc, ["字段", "内容", "说明"],
                      [["桥梁名称", "合成示例桥", "桥梁全称"], ["", "", "模板说明"]])
    original._tbl.addprevious(table._tbl)
    original._tbl.getparent().remove(original._tbl)
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert answer["基本信息"] == {"桥梁名称": "合成示例桥"}
    table.rows[2].cells[1].text = "没有字段名的内容"
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("空字段名" in error for error in errors)


def test_vertical_merge_inherits_value_without_collapsing_rows(tmp_path):
    doc = sample_document()
    table = doc.tables[2]
    values = ["2", "桥墩", "露筋", "钢筋外露", "否", "无", "无"]
    for cell, value in zip(table.add_row().cells, values):
        cell.text = value
    table.cell(1, 1).merge(table.cell(2, 1)).text = "桥墩"
    answer, errors = parse(doc, tmp_path)
    assert errors == []
    assert len(answer["病害明细"]) == 2
    assert [row["病害部位"] for row in answer["病害明细"]] == ["桥墩", "桥墩"]
    assert [row["病害类型"] for row in answer["病害明细"]] == ["裂缝", "露筋"]


def test_horizontal_merge_across_independent_fields_is_rejected(tmp_path):
    doc = sample_document()
    doc.tables[2].cell(1, 5).merge(doc.tables[2].cell(1, 6)).text = "无"
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("跨字段合并" in error for error in errors)


@pytest.mark.parametrize("table_index", [1, 2])
def test_missing_required_header_is_rejected(tmp_path, table_index):
    doc = sample_document()
    doc.tables[table_index].rows[0].cells[-1].text = "未定义字段"
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("缺少字段" in error for error in errors)


def test_duplicate_header_names_are_rejected_even_when_cells_are_distinct(tmp_path):
    doc = sample_document()
    doc.tables[2].rows[0].cells[-1].text = "上一次定检状态"
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("重复字段" in error for error in errors)


def test_duplicate_basic_information_keys_are_rejected(tmp_path):
    doc = sample_document()
    doc.tables[0].rows[2].cells[0].text = "桥梁名称"
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("重复字段" in error for error in errors)


def test_omitted_grid_cell_is_rejected_instead_of_shifting_values(tmp_path):
    doc = sample_document()
    row = doc.tables[2].rows[1]
    row._tr.remove(row._tr.tc_lst[0])
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    row._tr.get_or_add_trPr().append(before)
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("缺列" in error for error in errors)


def test_blank_header_is_rejected(tmp_path):
    doc = sample_document()
    doc.tables[1].rows[0].cells[0].text = ""
    answer, errors = parse(doc, tmp_path)
    assert answer is None
    assert any("空字段" in error for error in errors)


@pytest.mark.parametrize("name, expected", [
    ("示例桥-有对比年度的信息提取报告.docx", "含有对比年度"),
    ("示例桥-含对比年度的信息提取报告.docx", "含有对比年度"),
    ("示例桥-含有对比年度的信息提取报告.docx", "含有对比年度"),
    ("示例桥-无对比年度的信息提取报告.docx", "无对比年度"),
    ("示例桥.docx", "未知"),
])
def test_comparison_kind_recognizes_filename_variants(name, expected):
    assert builder.compare_kind(name) == expected


def test_main_records_parser_version_and_rejection_reason(tmp_path, monkeypatch):
    reports = tmp_path / "reports" / "2018年"
    labels = tmp_path / "labels" / "2018年"
    reports.mkdir(parents=True)
    labels.mkdir(parents=True)
    (reports / "示例桥.docx").write_bytes(b"synthetic report fixture")
    document = sample_document()
    document.tables[2].rows[0].cells[-1].text = "上一次定检状态"
    document.save(labels / "示例桥-有对比年度的信息提取报告.docx")
    monkeypatch.setattr(builder, "YEARS", ["2018年"])
    monkeypatch.setattr(builder, "report_text", lambda *args: "合成报告" * 200)
    out = tmp_path / "invalid-build"
    monkeypatch.setattr(builder.sys, "argv", [str(SCRIPT), "--train-root", str(reports.parent), "--label-root", str(labels.parent), "--out-dir", str(out)])
    builder.main()
    assert (out / "sft_train.jsonl").read_text() == ""
    stats = json.loads((out / "stats.json").read_text())
    assert stats["label_parser_version"] == builder.LABEL_PARSER_VERSION
    assert stats["report_parser_version"] == builder.REPORT_PARSER_VERSION
    assert stats["conversion_tool_versions"] == {}
    assert stats["label_parse_failures"] == 1
    assert "重复字段" in (out / "anomalies.md").read_text()

    document = sample_document()
    document.save(labels / "示例桥-有对比年度的信息提取报告.docx")
    (reports / "示例桥.docx").rename(reports / "示例桥.doc")
    monkeypatch.setattr(builder, "soffice_version", lambda: "LibreOffice synthetic-test")
    out = tmp_path / "valid-build"
    monkeypatch.setattr(builder.sys, "argv", [str(SCRIPT), "--train-root", str(reports.parent), "--label-root", str(labels.parent), "--out-dir", str(out)])
    builder.main()
    record = json.loads((out / "sft_train.jsonl").read_text())
    assert record["meta"]["label_parser_version"] == builder.LABEL_PARSER_VERSION
    assert record["meta"]["report_parser_version"] == builder.REPORT_PARSER_VERSION
    assert record["meta"]["conversion_tool_versions"] == {"soffice": "LibreOffice synthetic-test"}
    stats = json.loads((out / "stats.json").read_text())
    assert stats["conversion_tool_versions"] == record["meta"]["conversion_tool_versions"]
    assert record["meta"]["compare_kind"] == "含有对比年度"
    assert json.loads(record["output"])["病害明细"][0]["发展程度"] == "无"


def test_doc_conversion_uses_separate_profiles_for_concurrent_requests(tmp_path, monkeypatch):
    profiles = []

    def fake_convert(command, **kwargs):
        profiles.append(command[1])
        cache = Path(command[command.index("--outdir") + 1])
        source = Path(command[-1])
        (cache / f"{source.stem}.docx").write_bytes(b"synthetic converted document")

    monkeypatch.setattr(builder.subprocess, "run", fake_convert)
    cache = tmp_path / "cache with spaces"
    assert builder.convert_doc(tmp_path / "first.doc", "docx", cache) == cache / "first.docx"
    assert builder.convert_doc(tmp_path / "second.doc", "docx", cache) == cache / "second.docx"
    assert profiles[0] != profiles[1]
    assert all(profile.startswith("-env:UserInstallation=file:///") for profile in profiles)
    assert all("cache%20with%20spaces" in profile for profile in profiles)
    assert not list(cache.glob(".soffice-*"))


@pytest.mark.parametrize("empty_file", [False, True])
def test_successful_converter_exit_without_nonempty_output_is_rejected(tmp_path, monkeypatch, empty_file):
    def fake_convert(command, **kwargs):
        if empty_file:
            (tmp_path / "cache" / "label.docx").touch()

    monkeypatch.setattr(builder.subprocess, "run", fake_convert)
    errors = []
    answer = builder.label_answer(tmp_path / "label.doc", tmp_path / "cache", diagnostics=errors)
    assert answer is None
    assert errors == ["标签 .doc 转换失败"]


def test_parallel_extraction_keeps_serial_output_order_and_isolates_errors(tmp_path, monkeypatch):
    reports, labels = tmp_path / "reports" / "2018年", tmp_path / "labels" / "2018年"
    reports.mkdir(parents=True)
    labels.mkdir(parents=True)
    label_source = tmp_path / "label-source.docx"
    sample_document().save(label_source)
    for stem in ["A桥", "B桥", "C桥", "D桥"]:
        # 内容相同的文件对仍需使用不同缓存目录。
        (reports / f"{stem}.docx").write_bytes(b"identical synthetic report contents")
        (labels / f"{stem}-无对比年度的信息提取报告.docx").write_bytes(label_source.read_bytes())
    monkeypatch.setattr(builder, "YEARS", ["2018年"])
    outputs, stats, anomalies = [], [], []
    for workers in (1, 2):
        second_started = threading.Event()
        completion_order, caches = [], []

        def read_report(path, cache):
            caches.append(cache)
            if workers == 2 and path.stem == "A桥":
                assert second_started.wait(timeout=5), "Second worker never started"
            if path.stem == "C桥":
                raise OSError("synthetic single-pair failure")
            completion_order.append(path.stem)
            if path.stem == "B桥":
                second_started.set()
            return f"{path.stem}的合成训练正文。" * 100

        monkeypatch.setattr(builder, "report_text", read_report)
        out = tmp_path / f"workers-{workers}"
        monkeypatch.setattr(builder.sys, "argv", [str(SCRIPT), "--train-root", str(reports.parent),
                            "--label-root", str(labels.parent), "--out-dir", str(out), "--workers", str(workers)])
        builder.main()
        outputs.append((out / "sft_train.jsonl").read_bytes())
        run_stats = json.loads((out / "stats.json").read_text())
        assert run_stats.pop("workers") == workers
        stats.append(run_stats)
        anomalies.append((out / "anomalies.md").read_text())
        assert len(caches) == len(set(caches)) == 4
        if workers == 2:
            assert completion_order.index("B桥") < completion_order.index("A桥")
    assert outputs[0] == outputs[1]
    assert stats[0] == stats[1]
    assert anomalies[0] == anomalies[1]
    assert stats[0]["extraction_failures"] == 1
    assert "OSError" in anomalies[0]
    assert [json.loads(line)["meta"]["report"] for line in outputs[1].splitlines()] == ["A桥.docx", "B桥.docx", "D桥.docx"]


@pytest.mark.parametrize("workers", [0, 5])
def test_worker_count_is_bounded(tmp_path, monkeypatch, workers):
    monkeypatch.setattr(builder.sys, "argv", [str(SCRIPT), "--train-root", str(tmp_path),
                        "--label-root", str(tmp_path), "--out-dir", str(tmp_path / "out"), "--workers", str(workers)])
    with pytest.raises(SystemExit) as exc:
        builder.main()
    assert exc.value.code == 2
    assert not (tmp_path / "out").exists()
