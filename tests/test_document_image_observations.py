"""Local synthetic DOCX/backend contracts. No model, HTTP, labels or real reports."""
from copy import deepcopy
import json
from pathlib import Path
import sys
from unittest.mock import patch

import pytest
from docx import Document
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from bridge_evidence.contest.document_workspace import DocumentWorkspace, WorkspaceError, WorkspaceIdentityError, bind_image_observation
from bridge_evidence.contest.document_tools import DocumentTools
from bridge_evidence.contest.evidence_manifest import build_manifest
from bridge_evidence.contest import tong_react_adapter as adapter


class SyntheticOwnBackend:
    cache_key = "SYNTHETIC-OWN-VISION-NO-MODEL"
    model = "SYNTHETIC-OWN-MODEL"
    requires_observation_id = True

    def __init__(self):
        self.calls = []

    def __call__(self, blob, **kwargs):
        self.calls.append(kwargs)
        return {"kind": "synthetic observation", "text": "仅合成图片内容", "region_seen": kwargs["region"],
                "question_seen": kwargs["question"], "call_id": "synthetic-image-call-" + str(len(self.calls))}


@pytest.fixture
def env(tmp_path):
    class Env:
        def __init__(self):
            png = tmp_path / "synthetic.png"
            Image.new("RGB", (12, 12), "blue").save(png)
            doc = Document()
            doc.add_paragraph("合成桥甲，当前年度。")
            doc.add_picture(str(png))
            doc.add_paragraph("合成桥乙，另一年度。")
            doc.add_picture(str(png))
            self.source = tmp_path / "synthetic.docx"
            doc.save(self.source)
            self.manifest = build_manifest(self.source)
            self.identity = {"model": "synthetic", "purpose": "synthetic tests only"}
            self.directory = tmp_path / "workspace"
            self.workspace = DocumentWorkspace(self.manifest, self.directory, self.identity)
            self.backend = SyntheticOwnBackend()
            self.tools = self.make_tools()
            self.images = [row["id"] for row in self.manifest["images"]]

        def make_tools(self, *, backend=None, budget=20000, on_observation=None, on_read=None):
            return DocumentTools(self.manifest, count_tokens=len, max_tool_tokens=budget, source_path=self.source,
                                 vision_backend=self.backend if backend is None else backend,
                                 native_response_budget=True,
                                 on_read=self.workspace.record_read if on_read is None else on_read,
                                 on_image_observation=self.workspace.record_image_observation if on_observation is None else on_observation)

        def image(self, index=0, *, question="这幅合成图显示什么？", region="{}", tools=None):
            return (tools or self.tools).read_image(self.images[index], question, region, [])

        def update(self, **kwargs):
            return self.workspace.apply_update({"base_version": self.workspace.snapshot()["version"], **kwargs})

        def fact(self, observation, *, image=None, specific=True):
            ref = {"image_id": image or observation["image_id"]}
            if specific:
                ref["observation_id"] = observation["observation_id"]
            return {"id": "image_fact", "value": "合成桥", "source_refs": [ref], "status": "uncertain"}

        def ready(self):
            for unit in self.manifest["units"]:
                self.workspace.record_read(unit["id"], 0, len(unit["text"]))
            images = [self.image(i) for i in range(2)]
            assert all(row["ok"] for row in images), images
            self.update(facts=[self.fact(images[0])], draft={"基本信息": {"桥梁名称": "合成桥"}, "维修建议": [], "病害明细": []},
                        draft_fact_ids=["image_fact"], draft_bindings={"基本信息.桥梁名称": ["image_fact"]})
            self.workspace.record_review(self.review())
            assert self.workspace.delivery_status()["ready"]
            return images

        def review(self):
            snapshot = self.workspace.snapshot()
            return {"workspace_version": snapshot["version"], "draft_sha256": snapshot["draft_sha256"],
                    "checked_evidence_ids": snapshot["reading_progress"]["review_required_evidence_ids"],
                    "checked_image_observation_ids": sorted(snapshot["image_observations"]),
                    "passed": True, "model_call_ids": ["synthetic-independent-review"],
                    "unsupported_claims": [], "missing_items": [], "conflicts": [], "unresolved_items": []}

    return Env()


def session(tools):
    return {"handlers": tools.handlers(), "tool_calls": [], "max_calls": 1000,
            "tool_budget_exceeded": False, "count_tokens": len, "max_tool_tokens": tools.max_tool_tokens}


def invoke(tools, image_id):
    return adapter._invoke_active(session(tools), "read_image", {"image_id": image_id, "question": "这幅合成图显示什么？", "region_json": "{}", "context_refs": []})


def test_new_crop_or_question_invalidates_review_and_same_cached_observation_is_noop(env):
    images = env.ready()
    before = env.workspace.snapshot()
    duplicate = env.image()
    assert duplicate["observation_id"] == images[0]["observation_id"]
    assert duplicate["cache"]["judgment_hit"] is True
    assert env.workspace.snapshot() == before
    changed_question = env.image(question="这个合成图中的局部内容？")
    after = env.workspace.snapshot()
    assert changed_question["observation_id"] != images[0]["observation_id"]
    assert after["version"] == before["version"] + 1
    assert after["review"] is None
    assert not env.workspace.delivery_status()["ready"]
    changed_crop = env.image(region='{"x":0,"y":0,"width":4,"height":4}')
    assert changed_crop["observation_id"] != changed_question["observation_id"]
    assert len(env.workspace.snapshot()["image_observations"]) == 4


def test_repeat_bytes_in_other_occurrence_get_different_identity_and_cannot_share_ref(env):
    first, second = env.image(0), env.image(1)
    assert first["image_sha256"] == second["image_sha256"]
    assert first["observation_id"] != second["observation_id"]
    with pytest.raises(WorkspaceError, match="for this occurrence"):
        env.update(facts=[env.fact(first, image=env.images[1])])
    with pytest.raises(WorkspaceError, match="for this image occurrence"):
        env.workspace.load_image_observation(env.images[1], first["observation_id"])


def test_success_persists_complete_immutable_sidecar_but_snapshot_has_only_bounded_index(env):
    output = env.image()
    snapshot = env.workspace.snapshot()
    assert snapshot["image_reads"] == [env.images[0]]
    index = snapshot["image_observations"][output["observation_id"]]
    assert set(index) == {"image_id", "observation_sha256", "sidecar_sha256", "requires_observation_id"}
    assert "judgment" not in json.dumps(snapshot, ensure_ascii=False)
    restored = DocumentWorkspace(env.manifest, env.directory, env.identity)
    body = restored.load_image_observation(env.images[0], output["observation_id"])
    assert body["judgment"] == output["judgment"]
    assert body["question"] == output["question"]
    assert "cache" not in body
    body["judgment"]["text"] = "mutating returned object"
    assert restored.load_image_observation(env.images[0], output["observation_id"])["judgment"] == output["judgment"]


def test_own_model_fact_requires_specific_delivered_observation(env):
    output = env.image()
    with pytest.raises(WorkspaceError, match="specific observation_id"):
        env.update(facts=[env.fact(output, specific=False)])
    env.update(facts=[env.fact(output)])
    assert env.workspace.snapshot()["facts"]["image_fact"]["source_refs"][0]["observation_id"] == output["observation_id"]
    missing = deepcopy(output)
    missing["observation_id"] = "io_" + "0" * 64
    with pytest.raises(WorkspaceError, match="actually delivered"):
        env.update(facts=[env.fact(missing)])


def test_legacy_ocr_bare_ref_freezes_current_observations_and_later_crops_do_not_expand_it(env):
    class OCR:
        cache_key = "SYNTHETIC OCR"
        def recognize(self, blob, *, region):
            return {"text": "合成OCR文字", "region": region}
    tools = env.make_tools(backend=OCR())
    first = env.image(tools=tools)
    env.update(facts=[env.fact(first, specific=False)])
    oldref = env.workspace.snapshot()["facts"]["image_fact"]["source_refs"][0]
    assert oldref["observation_ids"] == [first["observation_id"]]
    assert "Legacy" in oldref["observation_limitation"]
    second = env.image(question="另一个问题", tools=tools)
    assert second["observation_id"] != first["observation_id"]
    assert env.workspace.snapshot()["facts"]["image_fact"]["source_refs"][0] == oldref
    env.update(facts=[{"id": "image_fact", "value": "updated same evidence"}])
    assert env.workspace.snapshot()["facts"]["image_fact"]["source_refs"][0] == oldref
    with pytest.raises(WorkspaceError, match="managed by the workspace"):
        env.update(facts=[env.fact(first) | {"source_refs": [oldref]}])


def test_old_image_read_only_checkpoint_stays_compatible_with_explicit_unknown_content_scope(env):
    state = env.workspace.snapshot()
    state.pop("image_observations")
    state.pop("image_observation_index_sha256")
    state.pop("workspace_version")
    state.pop("reading_progress")
    env.workspace.checkpoint_path.write_text(json.dumps(state))
    workspace = DocumentWorkspace(env.manifest, env.directory, env.identity)
    workspace.record_read(env.images[0], 0, 1)
    env.workspace = workspace
    env.update(facts=[{"id": "legacy", "value": "old uncertain evidence", "status": "uncertain", "source_refs": [{"image_id": env.images[0]}]}])
    reference = workspace.snapshot()["facts"]["legacy"]["source_refs"][0]
    assert reference["observation_ids"] == []
    assert "no persisted observation content" in reference["observation_limitation"]


def test_own_model_review_must_name_all_actual_observations_not_just_same_image_id(env):
    env.ready()
    env.image(question="新增局部观察")
    review = env.review()
    review["checked_image_observation_ids"].pop()
    with pytest.raises(WorkspaceError, match="all delivered"):
        env.workspace.record_review(review)
    valid = env.review()
    valid["checked_image_observation_ids"].append("io_" + "0" * 64)
    with pytest.raises(WorkspaceError, match="undelivered"):
        env.workspace.record_review(valid)
    env.workspace.record_review(env.review())
    assert env.workspace.delivery_status()["ready"]


def test_complete_native_boundary_is_counted_before_any_observation_callback(env):
    prototype = invoke(env.tools, env.images[0])
    assert prototype["ok"]
    exact = len(adapter._json_bytes(prototype).decode())
    # Use independent workspaces so the rejected call cannot inherit coverage.
    workspace = DocumentWorkspace(env.manifest, env.directory / "isolated", env.identity)
    rejected = env.make_tools(budget=exact-1, on_observation=workspace.record_image_observation, on_read=workspace.record_read)
    result = invoke(rejected, env.images[0])
    assert not result["ok"]
    assert result["result"]["status"] == "budget_exceeded"
    assert workspace.snapshot()["image_reads"] == []
    assert workspace.snapshot()["image_observations"] == {}
    assert rejected.image_observations == []
    assert len(adapter._json_bytes(result).decode()) <= exact-1
    fits = env.make_tools(budget=exact, on_observation=workspace.record_image_observation, on_read=workspace.record_read)
    # New synthetic backend call numbers have the same character width.
    delivered = invoke(fits, env.images[0])
    assert delivered["ok"]
    assert workspace.snapshot()["image_reads"] == [env.images[0]]
    assert len(adapter._json_bytes(delivered).decode()) == exact


def test_persistence_failure_returns_error_grants_no_read_and_retry_reuses_valid_upstream_result(env):
    with patch.object(env.workspace, "_save", side_effect=OSError("synthetic disk failure")):
        failed = env.image()
    assert not failed["ok"]
    assert env.workspace.snapshot()["image_reads"] == []
    assert env.workspace.snapshot()["image_observations"] == {}
    assert env.tools.image_observations == []
    assert len(list((env.directory / "image-observations").glob("*.json"))) == 1  # orphan is not delivery
    assert len(env.backend.calls) == 1
    retry = env.image()
    assert retry["ok"]
    assert len(env.backend.calls) == 1  # actual successful inference isn't needlessly repeated
    assert env.workspace.snapshot()["image_reads"] == [env.images[0]]


def test_new_callback_atomically_owns_image_read_and_does_not_invoke_second_legacy_callback(env):
    def must_not_run(*_args):
        raise RuntimeError("second callback would fail after persistence")
    tools = env.make_tools(on_read=must_not_run)
    assert env.image(tools=tools)["ok"]
    assert env.workspace.snapshot()["image_reads"] == [env.images[0]]
    assert env.workspace.snapshot()["version"] == 1


def test_default_on_read_only_remains_compatible(env):
    reads = []
    tools = DocumentTools(env.manifest, count_tokens=len, max_tool_tokens=20000, source_path=env.source,
                          vision_backend=env.backend, on_read=lambda *args: reads.append(args))
    output = env.image(tools=tools)
    assert output["ok"]
    assert reads == [(env.images[0], 0, 1)]
    assert tools.image_observations[0]["observation_id"] == output["observation_id"]


@pytest.mark.parametrize("failure", ["body", "missing", "wrong_image", "index_hash"])
def test_sidecar_or_index_tamper_fails_restore_instead_of_using_old_review(env, failure):
    env.ready()
    snapshot = env.workspace.snapshot()
    oid = next(iter(snapshot["image_observations"]))
    path = env.directory / "image-observations" / (oid + ".json")
    if failure == "body":
        data = json.loads(path.read_text())
        data["observation"]["judgment"]["text"] = "tampered observation"
        path.write_text(json.dumps(data))
    elif failure == "missing":
        path.unlink()
    else:
        state = json.loads(env.workspace.checkpoint_path.read_text())
        if failure == "wrong_image":
            current_image = state["image_observations"][oid]["image_id"]
            state["image_observations"][oid]["image_id"] = next(value for value in env.images if value != current_image)
        else:
            state["image_observations"] = {}
        env.workspace.checkpoint_path.write_text(json.dumps(state))
    with pytest.raises(WorkspaceIdentityError):
        DocumentWorkspace(env.manifest, env.directory, env.identity)


def test_callback_rejects_rebound_image_source_and_stale_observation_hash(env):
    output = env.image()
    different = deepcopy(output)
    different["question"] = "changed after token budgeting"
    with pytest.raises(WorkspaceIdentityError, match="identity/hash"):
        env.workspace.record_image_observation(env.images[0], different)
    with pytest.raises(WorkspaceIdentityError, match="occurrence differs"):
        env.workspace.record_image_observation(env.images[1], output)
    assert env.workspace.snapshot()["version"] == 1


def test_observation_index_and_body_limits_fail_closed_without_truncating(env):
    workspace = DocumentWorkspace(env.manifest, env.directory / "bounded", env.identity, max_image_observations=1)
    tools = env.make_tools(on_observation=workspace.record_image_observation, on_read=workspace.record_read)
    first = env.image(tools=tools)
    second = env.image(1, tools=tools)
    assert first["ok"] and not second["ok"]
    assert workspace.snapshot()["image_reads"] == [env.images[0]]
    assert len(workspace.snapshot()["image_observations"]) == 1
    tiny = DocumentWorkspace(env.manifest, env.directory / "tiny", env.identity, max_image_observation_bytes=10)
    with pytest.raises(WorkspaceError, match="persistence bound"):
        tiny.record_image_observation(env.images[0], first)
    assert tiny.snapshot()["image_reads"] == []


def test_structured_vision_error_is_recoverable_bounded_and_never_cached_as_success(env):
    class VisionObservationError(Exception):
        def as_tool_error(self):
            return {"ok": False, "status": "vision_budget_exceeded", "error": {"type": "VisionBudget", "message": "Use smaller crop", "recoverable": True},
                    "recovery": {"tool": "read_image", "region": {"x":0,"y":0,"width":4,"height":4}}}
    class Backend(SyntheticOwnBackend):
        def __call__(self, blob, **kwargs):
            self.calls.append(kwargs)
            raise VisionObservationError("private full response must not be shown")
    backend = Backend()
    tools = env.make_tools(backend=backend, budget=1000)
    first = env.image(tools=tools)
    second = env.image(tools=tools)
    assert first["status"] == "vision_budget_exceeded"
    assert first["error"]["recoverable"] is True
    assert first["recovery"]["tool"] == "read_image"
    assert "private full response" not in json.dumps(first)
    assert len(adapter._json_bytes(adapter.wrap_tool_response("read_image", first)).decode()) <= 1000
    assert not second["ok"] and len(backend.calls) == 2
    assert tools._judgment_cache == {} and tools.image_observations == []
    assert env.workspace.snapshot()["image_reads"] == []


def test_oversized_structured_error_uses_bounded_error_without_partial_recovery(env):
    class VisionObservationError(Exception):
        def as_tool_error(self):
            return {"ok": False, "status": "error", "error": {"type": "VisionBudget", "message": "x"*10000, "recoverable": True}}
    def backend(*args, **kwargs):
        raise VisionObservationError()
    tools = env.make_tools(backend=backend, budget=1000)
    result = env.image(tools=tools)
    assert result["status"] == "budget_exceeded"
    assert tools.image_observations == []
    assert env.workspace.snapshot()["image_reads"] == []


def test_content_hash_is_not_just_question_crop_or_model_name(env):
    result = env.image()
    changed = {k:v for k,v in result.items() if k not in {"observation_id", "observation_sha256"}}
    changed["judgment"] = {"text": "Different actual response, identical request"}
    assert bind_image_observation(changed)["observation_id"] != result["observation_id"]
    cached = deepcopy(result)
    cached["cache"] = {"base_hit": True, "judgment_hit": True}
    assert bind_image_observation(cached)["observation_id"] == result["observation_id"]


def test_persisted_delivery_cannot_enable_only_bare_payload_budget(env):
    with pytest.raises(ValueError, match="native_response_budget=True"):
        DocumentTools(env.manifest, count_tokens=len, source_path=env.source, vision_backend=env.backend,
                      on_image_observation=env.workspace.record_image_observation)


def test_new_path_repeated_calls_retain_small_index_and_individual_delivery_receipts(env):
    first = env.image()
    second = env.image()
    assert first["observation_id"] == second["observation_id"]
    assert env.tools.image_observations == [{"image_id": env.images[0], "observation_id": first["observation_id"],
                                            "observation_sha256": first["observation_sha256"], "storage": "workspace_sidecar"}]
    assert len(env.tools.receipts) == 2
    assert all(row["observation_ids"] == [first["observation_id"]] for row in env.tools.receipts)
    assert env.tools.receipts[0]["result_sha256"] != env.tools.receipts[1]["result_sha256"]  # cache-hit telemetry differs


def test_delivery_review_is_bound_to_observation_index_hash_besides_workspace_version(env):
    env.ready()
    state = json.loads(env.workspace.checkpoint_path.read_text())
    state["review"]["image_observation_index_sha256"] = "0" * 64
    env.workspace.checkpoint_path.write_text(json.dumps(state))
    status = env.workspace.delivery_status()
    assert not status["ready"]
    assert "semantic_review_stale" in status["blockers"]
